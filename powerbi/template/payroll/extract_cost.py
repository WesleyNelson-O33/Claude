#!/usr/bin/env python3
"""
Extracts actual employment cost per employee per job from Employment Hero
Pay Run Audit Report exports.

Point it at a folder of PayRunAudit-*.xlsx files:
    python3 extract_cost.py <folder> [output_folder]

Cost basis: Gross Earnings + SG Super + Employer Liabilities.
Payroll tax and workers comp are NOT in the report - they are levied on the
total wage bill, not per person. Add them as an uplift (ON_COST_UPLIFT below).
"""
import openpyxl, os, sys, csv, io, glob, collections, datetime

ON_COST_UPLIFT = 0.00   # e.g. 0.065 for 6.5% payroll tax + workers comp

def is_total_row(r):
    """Every audit sheet ends with a Total row. Summing it doubles the figures -
    and because that row carries no Units, hours still reconcile while dollars
    quietly double. Guard every loop with this."""
    return r is None or r[0] is None or str(r[0]).strip().lower().startswith("total")

def sheet_rows(ws, first_col="Employee Id"):
    """Audit sheets sometimes carry a title row above the header."""
    it = ws.iter_rows(values_only=True)
    hdr = None
    for r in it:
        if r and r[0] == first_col:
            hdr = [str(x) if x is not None else "" for x in r]
            break
    if hdr is None:
        return [], {}
    ix = {h: i for i, h in enumerate(hdr) if h}
    return it, ix

def period_of(path):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["Summary"]
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "Report Period" and r[1]:
            txt = str(r[1])
            end = txt.split("-")[-1].strip()
            try:
                d = datetime.datetime.strptime(end, "%d/%m/%Y").date()
                return txt, d.strftime("%Y-%m-%d"), d.strftime("%Y-%m")
            except ValueError:
                return txt, "", ""
    return "", "", ""

def main(src, out):
    files = sorted(glob.glob(os.path.join(src, "PayRunAudit*.xlsx"))) or \
            sorted(glob.glob(os.path.join(src, "*udit*.xlsx")))
    if not files:
        print("no PayRunAudit*.xlsx found in", src); return
    lines, totals = [], collections.defaultdict(lambda: collections.Counter())

    for f in files:
        label, end_date, month = period_of(f)
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)

        # ---- per employee per job, from Earnings Details ----
        it, ix = sheet_rows(wb["Earnings Details"])
        for r in it:
            if is_total_row(r): continue
            def g(col, default=0):
                i = ix.get(col)
                v = r[i] if i is not None and i < len(r) else None
                return v if v is not None else default
            units = g("Units") or 0
            gross = g("Gross Earnings") or 0
            sup   = g("SG Super") or 0
            lines.append({
                "Pay Period End": end_date, "Month": month,
                "Employee Id": str(g("Employee Id", "")).strip(),
                "Employee External Id": str(g("Employee External Id", "")).strip(),
                "Employee Name": str(g("Employee Name", "")).strip(),
                "Job No": str(g("Location External Id", "")).strip(),
                "Job Name (payroll)": str(g("Location Name", "")).strip(),
                "Pay Category": str(g("Pay Category Name", "")).strip(),
                "Unit Type": str(g("Unit Type", "")).strip(),
                "Hours": round(units, 4) if isinstance(units, (int, float)) else 0,
                "Rate": g("Rate", ""),
                "Gross": round(gross, 2) if isinstance(gross, (int, float)) else 0,
                "SG Super": round(sup, 2) if isinstance(sup, (int, float)) else 0,
            })

        # ---- per employee totals, for the cost rate ----
        it, ix = sheet_rows(wb["Pay Run Totals"])
        for r in it:
            if is_total_row(r): continue
            def g(col, default=0):
                i = ix.get(col)
                v = r[i] if i is not None and i < len(r) else None
                return v if v is not None else default
            k = (end_date, str(g("Employee Id", "")).strip())
            t = totals[k]
            t["hours"] += g("Total Hours") or 0
            t["gross"] += g("Gross Earnings") or 0
            t["super"] += (g("SG Super") or 0) + (g("Employer Contribution Super") or 0)
            t["net"]   += g("Net Earnings") or 0
            t["ext"] = str(g("Employee External Id", "")).strip()
            t["name"] = (str(g("Employee First Name", "")).strip() + " " +
                         str(g("Employee Surname", "")).strip()).strip()
            t["month"] = month

        # ---- employer liabilities, if any ----
        if "Employer Liabilities" in wb.sheetnames:
            it, ix = sheet_rows(wb["Employer Liabilities"])
            for r in it:
                if is_total_row(r): continue
                k = (end_date, str(r[0]).strip())
                i = ix.get("Amount")
                if i is not None and isinstance(r[i], (int, float)):
                    totals[k]["liab"] += r[i]

    os.makedirs(out, exist_ok=True)
    p1 = os.path.join(out, "cost_by_person_job.csv")
    with io.open(p1, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(lines[0].keys())); w.writeheader(); w.writerows(lines)

    rate = []
    for (end, eid), t in sorted(totals.items()):
        cost = (t["gross"] + t["super"] + t["liab"]) * (1 + ON_COST_UPLIFT)
        rate.append({"Pay Period End": end, "Month": t["month"], "Employee Id": eid,
                     "Employee External Id": t["ext"], "Employee Name": t["name"],
                     "Hours Paid": round(t["hours"], 2), "Gross": round(t["gross"], 2),
                     "Super": round(t["super"], 2), "Employer Liabilities": round(t["liab"], 2),
                     "On-cost Uplift %": ON_COST_UPLIFT,
                     "Total Cost": round(cost, 2),
                     "Cost Rate $/hr": round(cost / t["hours"], 2) if t["hours"] else ""})
    p2 = os.path.join(out, "cost_rate_by_employee.csv")
    with io.open(p2, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rate[0].keys())); w.writeheader(); w.writerows(rate)

    # reconcile the line detail against the per-employee totals
    det = sum(l["Gross"] for l in lines)
    tot = sum(t["gross"] for t in totals.values())
    ok = abs(det - tot) < max(1.0, tot * 0.001)
    print("RECONCILIATION  Earnings Details $%s  vs  Pay Run Totals $%s   %s"
          % (format(det, ",.2f"), format(tot, ",.2f"), "OK" if ok else "*** MISMATCH ***"))
    if not ok:
        print("    Do not use these files until this reconciles.")
    print("files read:        %d" % len(files))
    print("cost_by_person_job.csv     %d rows" % len(lines))
    print("cost_rate_by_employee.csv  %d rows" % len(rate))
    print("total gross  $%s" % format(sum(l["Gross"] for l in lines), ",.2f"))
    print("total super  $%s" % format(sum(l["SG Super"] for l in lines), ",.2f"))
    if ON_COST_UPLIFT == 0:
        print("\nNOTE: ON_COST_UPLIFT is 0. Set it to your payroll tax + workers comp"
              "\n      rate (as a decimal) at the top of this script and re-run.")

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "."
    out = sys.argv[2] if len(sys.argv) > 2 else "."
    main(src, out)
