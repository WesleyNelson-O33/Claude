"""
Render ranges of the Financial Controller Pack to PNGs for the deck.

The deck is built from snips of the pack, so each picture here is the pack's own
cells and its own charts, with the pack's own formatting. Nothing is redrawn.

How it works, per snip: take a copy of the pack, hide every sheet except the one
we want so LibreOffice exports only that page, set the print area to the range,
convert to PDF, then trim the white margin off the render.

Sheets are hidden rather than deleted. Deleting them would turn every
cross-sheet formula into #REF and the render would come out full of errors.
"""
import os, shutil, subprocess, sys, json
import openpyxl
from openpyxl.worksheet.properties import PageSetupProperties

HERE = os.path.dirname(os.path.abspath(__file__))
PACK = os.path.join(HERE, '..', '..', 'pack', 'CTS Financial Controller Pack - v20.xlsx')
WORK = os.path.join(HERE, 'snipwork')
OUTDIR = os.path.join(HERE, '..', 'snips')
SOFFICE = '/root/.claude/skills/synced/e9aa5ee9-a99d-40c9-9bb1-cc5bc5010cf8_723c92f1-3883-471f-8787-0ed0739f050b/pptx/scripts/office/soffice.py'

# id -> (sheet, print range, collapse the outline first, columns to hide,
#        cells to set first)
# Hidden columns let a snip sit a table next to a chart without the working
# columns that happen to live between them.
SNIPS = {
    'pl_income':   ('Summary',        'A5:I58',     True),
    'pl_result':   ('Summary',        'A59:I181',   True),
    'dept_result': ('Overhead Split', 'A140:H148',  False),
    # The same result rolled up to the three departments the P&L is built on.
    'dept_three':  ('Overhead Split', 'A152:H158',  False),
    'rev_by_dept': ('Cost Analysis',  'F190:AA212', False),
    'rev_month':   ('Cost Analysis',  'H229:AA251', False),
    'rev_ytd':     ('Cost Analysis',  'H259:T282',  False),
    'rev_years':   ('Cost Analysis',  'H339:T362',  False),
    'rev_q1':      ('Cost Analysis',  'H305:X328',  False),
    'xero_hist':   ('Xero History',   'F8:AA31',    False),
    'xero_cum':    ('Xero History',   'F33:AA58',   False),
    'xero_qtr':    ('Xero History',   'F59:AA84',   False),
    'xero_roll':   ('Xero History',   'F85:AA110',  False),
    'cost_dist':   ('Cost Analysis',  'I33:AC59',    False),
    'other_opex':  ('Cost Analysis',  'I4:AC30',     False),
    'labour_mix':  ('Cost Analysis',  'I61:AC87',    False),
    'income_month':('Cost Analysis',  'I117:AC140',  False),  # I142 starts the next chart
    'staff_split': ('Overheads MoM',  'A160:P178',  False),
    'top10':       ('Top 10 Clients', 'G4:AA27',     False),
    'dso_calc':    ('DSO',            'A5:B15',     False),
    'dso_who':     ('DSO',            'J23:T46',    False),
    'util_tbl':    ('Utilisation',    'A66:D73',    False),
    # The utilisation calculation itself: hours, utilisation per cent and FTE
    # for each department, for the three closed months. Split in two, because
    # thirty rows on one slide shrinks the figures to nothing.
    # Chargeable per cent on her own report's basis, the three person groups.
    # The month on one line per group. The three month history is on the tab
    # at rows 120 to 164 for anyone who wants it.
    'util_month':  ('Utilisation',    'A170:H176',  False),
    'util_pct_a':  ('Utilisation',    'A4:F19',     False),
    'util_pct_b':  ('Utilisation',    'A20:F37',    False),
    # One slide per department, off the Dept P&L tab with its selector set.
    'dept_Support':     ('Dept P&L', 'A7:G25', False, '', {'B5': 'Support'}),
    'dept_Events':      ('Dept P&L', 'A7:G25', False, '', {'B5': 'Events'}),
    'dept_Video':       ('Dept P&L', 'A7:G25', False, '', {'B5': 'Video'}),
    'dept_Integration': ('Dept P&L', 'A7:G25', False, '', {'B5': 'Integration'}),
    'dept_Consulting':  ('Dept P&L', 'A7:G25', False, '', {'B5': 'Consulting'}),
    'util_chart':  ('Utilisation',    'F65:AA90',    False),
    # The same hours picture rolled up to the three departments.
    'util_three':  ('Utilisation',    'F92:AA117',   False),
}


# On the Summary tab departments sit at outline level 2, sub-departments at 3
# and the account detail at 4. The screenshot shows departments and
# sub-departments but not the account categories, which also sit at level 3.
SUBDEPTS = {'events', 'video', 'integration', 'consulting'}


# Where a snip starts below the sheet's own header row, the header is printed
# again at the top of the snip so the columns can be read.
TITLE_ROWS = {'pl_result': '5:5', 'util_pct_b': '8:8'}


def collapse(ws):
    """Reproduce the on-screen view: departments and sub-departments, no
    account detail. Level 3 is kept only where the row is a sub-department."""
    for r in range(1, ws.max_row + 1):
        d = ws.row_dimensions.get(r)
        if d is None or not d.outlineLevel:
            continue
        label = str(ws.cell(r, 1).value or '').strip().lower()
        if d.outlineLevel >= 4 or (d.outlineLevel == 3 and label not in SUBDEPTS):
            d.hidden = True


def build(snip_id, sheet, rng, do_collapse, hide_cols='', set_cells=None):
    wb = openpyxl.load_workbook(PACK)
    for ws in wb.worksheets:
        ws.sheet_state = 'visible' if ws.title == sheet else 'hidden'
    wb.active = wb.index(wb[sheet])      # LibreOffice prints from the active sheet
    ws = wb[sheet]
    if do_collapse:
        collapse(ws)
    for col in hide_cols:
        ws.column_dimensions[col].hidden = True
    for ref, value in (set_cells or {}).items():
        ws[ref] = value
    ws.print_area = rng
    if snip_id in TITLE_ROWS:
        ws.print_title_rows = TITLE_ROWS[snip_id]
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.gridLines = False
    ws.print_options.horizontalCentered = False
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(ws.page_margins, side, 0.15)
    ws.page_margins.header = 0
    ws.page_margins.footer = 0
    # The pack prints a company name and a page number. Neither belongs in a snip.
    for hf in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter,
               ws.firstHeader, ws.firstFooter):
        for part in (hf.left, hf.center, hf.right):
            part.text = None
    ws.HeaderFooter.differentFirst = False
    ws.HeaderFooter.differentOddEven = False
    path = os.path.join(WORK, f'{snip_id}.xlsx')
    wb.save(path)
    return path


def trim(png):
    """Cut the white page margin off, leaving a small even border."""
    from PIL import Image, ImageChops
    im = Image.open(png).convert('RGB')
    bg = Image.new('RGB', im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).convert('L').point(lambda v: 255 if v > 12 else 0).getbbox()
    if not box:
        return
    pad = 10
    box = (max(box[0] - pad, 0), max(box[1] - pad, 0),
           min(box[2] + pad, im.width), min(box[3] + pad, im.height))
    im.crop(box).save(png)


def main(only=None):
    os.makedirs(WORK, exist_ok=True)
    os.makedirs(OUTDIR, exist_ok=True)
    made = {}
    for snip_id, spec in SNIPS.items():
        if only and snip_id not in only:
            continue
        sheet, rng, coll = spec[0], spec[1], spec[2]
        hide = spec[3] if len(spec) > 3 else ''
        sets = spec[4] if len(spec) > 4 else None
        x = build(snip_id, sheet, rng, coll, hide, sets)
        subprocess.run([sys.executable, SOFFICE, '--headless', '--convert-to', 'pdf',
                        '--outdir', WORK, x], check=True, capture_output=True)
        pdf = os.path.join(WORK, f'{snip_id}.pdf')
        stem = os.path.join(WORK, snip_id)
        subprocess.run(['pdftoppm', '-png', '-r', '200', '-f', '1', '-l', '1', pdf, stem],
                       check=True, capture_output=True)
        src = next(p for p in sorted(os.listdir(WORK))
                   if p.startswith(snip_id + '-') and p.endswith('.png'))
        dst = os.path.join(OUTDIR, f'{snip_id}.png')
        shutil.move(os.path.join(WORK, src), dst)
        trim(dst)
        from PIL import Image
        made[snip_id] = Image.open(dst).size
        print(f'  {snip_id:<13} {sheet:<16} {rng:<12} {made[snip_id][0]}x{made[snip_id][1]}')
    sizes_path = os.path.join(OUTDIR, 'sizes.json')
    old = json.load(open(sizes_path)) if os.path.exists(sizes_path) else {}
    old.update({k: list(v) for k, v in made.items()})
    json.dump(old, open(sizes_path, 'w'), indent=1)
    return made


if __name__ == '__main__':
    main(sys.argv[1:] or None)
