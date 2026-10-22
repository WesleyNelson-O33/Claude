"""Pages 7 and 8 become one page, and page 10 picks up the new dot chart.

Both pages answered the same question with one comparison each: revenue against
budget, and revenue against the same month last year. Side by side they read as
one point, so page 7 now carries both charts and page 8 goes.

Page 8 is removed rather than page 7 kept empty. Nothing is added to the deck in
this pass, only removed, because adding and removing a slide in the same pass
makes python-pptx hand two parts the same name and the file comes out corrupt.
"""
import sys
sys.path.insert(0, '/home/user/Claude/deck/edits')
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from addslide import balance, need, write, TOP, FLOOR, LEFT, SIZE

SNIP = '/home/user/Claude/deck/snips'
GAP = 0.20

BULLETS = [
    'September billed $562k against a $414k budget, $148k ahead of it, and '
    'against $356k last September.',
    'Year to date is $1,624k against $1,996k, short by $372k or 19%.',
    'Almost all of that shortfall is August, which billed $603k against a $973k '
    'budget.',
    'Against FY25 rather than the weak FY26, September is 28% ahead. That is the '
    'steadier read on growth.',
]
NOTES = """Revenue against budget month by month, and revenue for the month across three years. Both taken from the Cost Analysis tab.

Against budget, month by month. July billed $459,049 against $609,277, short by $150,228. August billed $603,467 against $973,168, short by $369,702. September billed $561,945 against $413,665, ahead by $148,280. Year to date $1,624,460 against $1,996,110, short by $371,650 or 18.6%.

August is where the year to date shortfall sits. It was budgeted as the second largest month of the FY27 budget and came in $369,702 under. September went the other way and recovered $148,280 of it.

Against the same month in earlier years. September 2026 $561,945, September 2025 $356,323, September 2024 $440,243. September 2026 is $205,621 or 57.7% ahead of last year and $121,701 or 27.6% ahead of the year before. September 2025 was itself 19.1% below September 2024.

The 57.7% figure flatters the result because FY26 was the weakest of the three years. Against FY25 the same month is 27.6% ahead, which is about 13% a year over two years, and that is the number to carry into the reforecast.

The two charts sit together because they answer the same question from two sides. The month is ahead of where it was a year ago and ahead of its own budget. The year to date is behind budget, and that is an August problem rather than a September one.

Xero holds the last three financial years, so September 2024 is the earliest September available for the comparison."""


def relay2(slide, snips, bullets, notes, names, title):
    by = {sh.name: sh for sh in slide.shapes}
    tree = slide.shapes._spTree
    by['Text 1'].text_frame.paragraphs[0].runs[0].text = title

    cols = list(balance(bullets, 4.05))
    ph_panel = max(need(c, 4.05) for c in cols if c) + 0.30
    strip_top = FLOOR - ph_panel
    band = strip_top - 0.12 - TOP

    half = (8.90 - GAP) / 2
    aspects = [Image.open(s).width / Image.open(s).height for s in snips]
    pic_h = min([band] + [half / a for a in aspects])

    for sh in list(slide.shapes):
        if sh.name in names:
            tree.remove(sh._element)
    for k, (snip, name, aspect) in enumerate(zip(snips, names, aspects)):
        w = pic_h * aspect
        x = LEFT + k * (half + GAP) + (half - w) / 2
        pic = slide.shapes.add_picture(snip, Inches(x), Inches(TOP + (band - pic_h) / 2),
                                       width=Inches(w), height=Inches(pic_h))
        pic.name = name
        tree.remove(pic._element)
        tree.insert(2, pic._element)

    panel, label = by['commentary panel'], by['Text 4']
    panel.left, panel.top = Inches(LEFT), Inches(strip_top)
    panel.width, panel.height = Inches(8.90), Inches(ph_panel)
    label.left, label.top, label.width = Inches(0.75), Inches(strip_top + 0.08), Inches(3.0)
    for key, x, items in (('Text 5', 0.75, cols[0]), ('Text 5b', 5.10, cols[1])):
        box = by[key]
        box.left, box.top = Inches(x), Inches(strip_top + 0.27)
        box.width, box.height = Inches(4.05), Inches(ph_panel - 0.30)
        write(box, items)
        for para in box.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(SIZE)

    slide.notes_slide.notes_text_frame.text = notes
    return pic_h


def swap_picture(slide, snip, name):
    """Same geometry, new image. The snip is the same shape as the one it
    replaces, so nothing has to be laid out again."""
    old = next(sh for sh in slide.shapes if sh.name == name)
    geom = (old.left, old.top, old.width, old.height)
    tree = slide.shapes._spTree
    tree.remove(old._element)
    pic = slide.shapes.add_picture(snip, geom[0], geom[1], width=geom[2], height=geom[3])
    pic.name = name
    tree.remove(pic._element)
    tree.insert(2, pic._element)


if __name__ == '__main__':
    p = Presentation(sys.argv[1])

    # page 10 takes the dot version of the thirty six month chart
    swap_picture(p.slides[9], f'{SNIP}/xero_hist.png', 'snip xero_hist')
    print('page 10: dot chart swapped in')

    h = relay2(p.slides[6], [f'{SNIP}/rev_ytd.png', f'{SNIP}/rev_years.png'],
               BULLETS, NOTES, ['snip rev_month_budget', 'snip rev_years'],
               'Revenue, against budget and year on year')
    print(f'page 7: two charts at {h:.2f} high')

    ids = p.slides._sldIdLst
    el = list(ids)[7]
    p.part.drop_rel(el.rId)
    ids.remove(el)
    print('page 8 removed')

    p.save(sys.argv[2])
