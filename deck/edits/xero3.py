"""Page 11 becomes the cumulative chart, not the rolling twelve months.

The rolling line was the better statistic and the worse slide. Cumulative year
to date against budget says more on one picture: the year is ahead of both
completed years, it is behind the plan set for it, and the plan itself is below
what the business billed two years ago.

The rolling twelve month block and chart stay on the Xero History tab. They are
useful working, they are just not a board slide.
"""
import sys
sys.path.insert(0, '/home/user/Claude/deck/edits')
from pptx import Presentation
from addslide import relay

SNIP = '/home/user/Claude/deck/snips'

BULLETS = [
    'Year to date revenue is $1,624k, ahead of both completed years at this '
    'point in the year.',
    'It is $372k or 19% behind budget. Almost all of that opened in August, and '
    'September closed $148k of it back.',
    'FY26 lost its year in December and January, $1,045k behind FY25 across '
    'those two months.',
    'The full year budget of $6,305k is itself 10% below what FY25 actually '
    'billed.',
]
NOTES = """Revenue year to date on a cumulative basis, three years against the FY27 budget, taken from the Xero History tab.

At September: FY27 $1,624,460, FY26 $1,228,582, FY25 $1,380,102, and the FY27 budget $1,996,110.

The year is ahead of both completed years at this point. It is $371,650 or 18.6% behind the budget set for it.

Where that gap came from, month by month. July was $150,228 behind budget. August was $369,702 behind, which is where almost all of it opened. September was $148,280 ahead of budget, which closed part of it back. August was budgeted at $973,168, the second largest month of the FY27 budget, and billed $603,467.

What the two completed lines show is where a year gets lost. FY25 and FY26 track each other almost exactly from July to November. They separate in December and January. December was $395,542 behind and January $649,658 behind, $1,045,200 across the two months, against a full year gap of $887,171. Every other month of FY26 combined was ahead of FY25.

That matters now because FY27 is about to enter the same two months, and the budget has December at $344,970 and January at $353,969, both well below what FY25 billed in them.

The last point is the budget itself. The full year FY27 budget is $6,304,601. FY25 actually billed $7,000,357. The plan for this year is $695,756 or 9.9% below what the business delivered two years ago, and $191,415 or 3.1% above FY26. Being 18.6% behind that budget is the gap against a plan that was already set below the FY25 result.

The FY27 line stops at September because the months beyond it hold budget rather than actuals. The budget line runs the full year because a budget is set for the whole year."""

p = Presentation(sys.argv[1])
slide = p.slides[10]
w, h, strip = relay(slide, f'{SNIP}/xero_cum.png', BULLETS, NOTES,
                    'snip xero_roll', 'Revenue year to date, against budget')
for sh in slide.shapes:
    if sh.name == 'snip xero_roll':
        sh.name = 'snip xero_cum'
print(f'page 11: picture {w:.2f} x {h:.2f}, strip {strip:.2f}')
p.save(sys.argv[2])
