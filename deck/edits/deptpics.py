"""Put the re-rendered department snips back on their pages.

Pages 17, 18, 19, 25 and 26 carry the snip whole. Pages 20 to 24 carry it with
the seven overhead detail rows cut out, so the page shows the trading lines and
the total rather than the breakdown. Each picture goes back at the same place
and the same size, so nothing on the page moves.
"""
import io
import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from relayout import drop_rows

DECK = 'deck/CTS Monthly Business Review - 30092026.pptx'
SNIPS = 'deck/snips'
WORK = '/tmp/claude-0/-home-user-Claude/a00ffe59-433c-507b-9006-917d8c60083a/scratchpad/pics'

WHOLE = {17: 'dept_result', 18: 'dept_mix', 19: 'dept_three',
         25: 'dept_pl', 26: 'dept_line'}
CROPPED = {20: 'dept_Support', 21: 'dept_Events', 22: 'dept_Video',
           23: 'dept_Integration', 24: 'dept_Consulting'}
ROWS, DETAIL = 18, {9, 10, 11, 12, 13, 14, 15}

os.makedirs(WORK, exist_ok=True)
prs = Presentation(DECK)

for page, tag in list(WHOLE.items()) + list(CROPPED.items()):
    slide = prs.slides[page - 1]
    pic = next(sh for sh in slide.shapes
               if sh.shape_type == 13 and sh.name != 'cts logo')
    was = Image.open(io.BytesIO(pic.image.blob)).size
    im = Image.open(f'{SNIPS}/{tag}.png').convert('RGB')
    if page in CROPPED:
        im = drop_rows(im, ROWS, DETAIL)
    if im.size != was:
        raise SystemExit(f'page {page} {tag}: {im.size} but the page had {was}')
    path = f'{WORK}/{tag}.png'
    im.save(path)

    left, top, width, height = pic.left, pic.top, pic.width, pic.height
    name = pic.name
    tree = slide.shapes._spTree
    idx = list(tree).index(pic._element)
    tree.remove(pic._element)
    new = slide.shapes.add_picture(path, left, top, width=width, height=height)
    new.name = name
    tree.remove(new._element)
    tree.insert(idx, new._element)
    print(f'  page {page:<4}{tag:<20}{im.size[0]}x{im.size[1]}  '
          f'{width / 914400:.2f}x{height / 914400:.2f}in')

prs.save(DECK)
print('deck saved')
