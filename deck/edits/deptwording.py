"""Two wording corrections she asked for, both wrong before any of this work.

Page 3 said Support was the only department covering its allocated overheads.
Integration covers its own too, at 23,756 of gross profit against 13,193 of
overhead, which is what page 18 already says. Page 19 is left alone: it reads
the month as three groups, and on that basis Support really is the only one.

Page 25's note called Events the second largest loss. At 46,496 it is the
largest of the five.
"""
from pptx import Presentation

DECK = 'deck/CTS Monthly Business Review - 30092026.pptx'

OLD_ONLY = 'Support is the only department covering its allocated overheads.'
NEW_ONLY = ('Support and Integration are the only departments covering their '
            'allocated overheads.')
OLD_LOSS = 'the three large departments and the second largest loss'
NEW_LOSS = 'the three large departments and the largest loss'

BODY = {3: [(OLD_ONLY, NEW_ONLY)]}
NOTES = {3: [(OLD_ONLY, NEW_ONLY)], 25: [(OLD_LOSS, NEW_LOSS)]}


def swap(tf, old, new):
    hit = False
    for para in tf.paragraphs:
        if old not in para.text:
            continue
        runs = para.runs
        if runs:
            runs[0].text = para.text.replace(old, new)
            for r in runs[1:]:
                r.text = ''
            hit = True
    return hit


if __name__ == '__main__':
    prs = Presentation(DECK)
    missed = []
    for page, pairs in BODY.items():
        slide = prs.slides[page - 1]
        for old, new in pairs:
            hit = any(swap(sh.text_frame, old, new)
                      for sh in slide.shapes if sh.has_text_frame)
            print(f'  page {page} body   {"ok" if hit else "MISSED"}  {old[:56]}')
            if not hit:
                missed.append((page, 'body', old))
    for page, pairs in NOTES.items():
        slide = prs.slides[page - 1]
        for old, new in pairs:
            hit = swap(slide.notes_slide.notes_text_frame, old, new)
            print(f'  page {page} notes  {"ok" if hit else "MISSED"}  {old[:56]}')
            if not hit:
                missed.append((page, 'notes', old))
    if missed:
        raise SystemExit('some phrases were not found, nothing saved')
    prs.save(DECK)
    print('deck saved')
