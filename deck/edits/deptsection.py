"""A section divider for the departmental pages, in the department colours.

The departmental pages sat inside Financial performance with no break in front of
them. They now have their own section, and its dots are the five department
colours rather than the three from the palette: Support red, Events dark teal,
Video bright teal, Integration blue and Consulting purple, in P&L order.

Adding a section pushes the numbering on, so the dividers after it move up one,
the eyebrow on the departmental pages changes to match, and the agenda gains a
line. The agenda rows are respaced to fit six without running into the footer.
"""
import copy, json, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

D = json.load(open('/home/user/Claude/pack/data/dept_colours.json'))['departments']
DEPTS = ['Support', 'Events', 'Video', 'Integration', 'Consulting']

# left, top, diameter. Five circles down the right, overlapping the way the other
# dividers do, and all of them inside the page.
DOTS = [(6.45, 0.10, 2.45), (7.80, 0.95, 2.05), (7.25, 2.25, 1.70),
        (7.70, 2.70, 2.25), (6.45, 3.45, 1.50)]

AT = 15                       # the new divider becomes page 16
HEADLINE = 'Departmental performance'
STRAPLINE = "Every department's result for September 2026 and how it was made"
NOTES = ("Departmental performance for September 2026.\n\n"
         "The pages that follow give the company result by department, the same "
         "result rolled up the way the P&L groups the business, then a page for "
         "each of the five departments, and finally departmental revenue against "
         "budget and against the same month last year.")

AGENDA = [
    ('01', 'Financial performance',
     'Profit and loss summary, revenue against budget, and the three year comparisons'),
    ('02', 'Departmental performance',
     "Every department's result for September 2026 and how it was made"),
    ('03', 'Operating expenditure',
     'Where the money goes: cost distribution, operating expense, labour and employment'),
    ('04', 'Customers and receivables',
     'Top ten customers, receivables and days sales outstanding'),
    ('05', 'People and utilisation',
     'Chargeable hours by department and the profit they produce'),
    ('06', 'Budget and outlook',
     'Revenue by month against budget, and what we do about it'),
]
ROW_TOP, ROW_STEP = 1.55, 0.60

DEPT_PAGES = range(16, 24)    # after the insert: the eight departmental pages


def set_text(shape, text):
    para = shape.text_frame.paragraphs[0]
    para.runs[0].text = text
    for r in list(para.runs)[1:]:
        r._r.getparent().remove(r._r)


p = Presentation(sys.argv[1])

# ---- the divider itself, cloned off the first one ---------------------------
src = p.slides[3]
new = p.slides.add_slide(src.slide_layout)
for sh in list(new.shapes):
    sh._element.getparent().remove(sh._element)
for sh in src.shapes:
    new.shapes._spTree.append(copy.deepcopy(sh._element))

tree = new.shapes._spTree
for sh in list(new.shapes):
    if sh.name.startswith('dot '):
        tree.remove(sh._element)
for i, (name, (left, top, dia)) in enumerate(zip(DEPTS, DOTS)):
    dot = new.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top),
                               Inches(dia), Inches(dia))
    dot.name = f'dot {i}'
    dot.fill.solid()
    dot.fill.fore_color.rgb = RGBColor.from_string(D[name])
    dot.line.fill.background()
    dot.shadow.inherit = False
    tree.remove(dot._element)
    tree.insert(2, dot._element)            # behind the text, like the others

by = {sh.name: sh for sh in new.shapes}
set_text(by['Text 2'], '02')
set_text(by['Text 4'], HEADLINE)
set_text(by['Text 5'], STRAPLINE)
ns = new.notes_slide
if ns.notes_text_frame is None:        # a cloned slide starts without one
    ns.shapes._spTree.append(
        copy.deepcopy(src.notes_slide.notes_placeholder._element))
ns.notes_text_frame.text = NOTES

ids = p.slides._sldIdLst
el = ids[-1]
ids.remove(el)
ids.insert(AT, el)
print(f'divider added as page {AT + 1}, dots in the five department colours')

# ---- the dividers after it move up one --------------------------------------
moved = []
for i, s in enumerate(p.slides, start=1):
    if i <= AT + 1:
        continue
    num = next((sh for sh in s.shapes if sh.name == 'Text 2'), None)
    if num is None or not num.text_frame.text.strip().isdigit():
        continue
    set_text(num, f'{int(num.text_frame.text.strip()) + 1:02d}')
    moved.append(i)
print(f'dividers renumbered on pages {moved}')

# ---- the departmental pages get their own eyebrow ---------------------------
for i in DEPT_PAGES:
    eb = next((sh for sh in p.slides[i].shapes if sh.name == 'Text 0'), None)
    if eb is not None:
        set_text(eb, 'DEPARTMENTAL PERFORMANCE')
print(f'eyebrow set on pages {DEPT_PAGES.start + 1} to {DEPT_PAGES.stop}')

# ---- the agenda gains a line and is respaced --------------------------------
ag = p.slides[1]
rows = [(f'Shape {2 + 4 * k}', f'Text {3 + 4 * k}', f'Text {4 + 4 * k}',
         f'Text {5 + 4 * k}') for k in range(5)]
have = {sh.name: sh for sh in ag.shapes}
tree = ag.shapes._spTree
last = rows[-1]
clone = []
for name in last:
    el = copy.deepcopy(have[name]._element)
    tree.append(el)
    sh = next(x for x in ag.shapes if x._element is el)
    clone.append(sh)
for sh, name in zip(clone, (f'Shape {2 + 4 * 5}', f'Text {3 + 4 * 5}',
                            f'Text {4 + 4 * 5}', f'Text {5 + 4 * 5}')):
    sh.name = name
rows.append(tuple(sh.name for sh in clone))

have = {sh.name: sh for sh in ag.shapes}
for k, (names, (num, head, sub)) in enumerate(zip(rows, AGENDA)):
    top = ROW_TOP + k * ROW_STEP
    box, n_, h_, s_ = (have[x] for x in names)
    box.top = Inches(top)
    n_.top = Inches(top)
    h_.top = Inches(top - 0.02)
    s_.top = Inches(top + 0.22)
    set_text(n_, num)
    set_text(h_, head)
    set_text(s_, sub)
print(f'agenda rebuilt with {len(AGENDA)} lines')

p.save(sys.argv[2])
