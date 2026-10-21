"""Slide 30: the table reads as recommendations rather than instructions."""
import sys
from pptx import Presentation

ROWS = {
 0: ('Area', 'Recommendation', 'Owner', 'Suggested by'),
 1: ('P&L', 'That FY27 is reforecast on gross margin rather than revenue. At a 38.9% '
            'gross margin, break even requires $549k a month against the $541k currently '
            'being delivered.'),
 2: ('Overhead keys', 'Events carries $111k of overhead on $73k of revenue and cannot return a '
                      'profit on that basis. I would suggest the keys are tested for FY28 '
                      'and left unchanged within the current year.'),
 3: ('Consulting', '$20k of running costs against $4k of fees. A decision on whether the '
                   'department is rebuilt, combined with Integration or closed would be '
                   'timely this quarter.'),
 4: ('Customers', 'The top three customers represent 48.7% of revenue. A named retention '
                  'plan for each, with the growth target set on customers four to twenty, '
                  'would reduce that concentration.'),
 5: ('Cash', 'Days sales outstanding is 47.2 days. Reaching 30 days would release $304k. '
             'Aware Super and Deloitte Services P/L for Deloitte Services Trust hold most '
             'of the balance.'),
 6: ('Utilisation', "Chargeable per cent is 72% against the report's own 81% target. Five "
                    'additional points across Support and Production would be worth '
                    'approximately $16k a month at no additional cost.'),
 7: ('Leave', 'Production recorded 333 hours of leave on 882 available. Planning leave '
              'against the booking calendar would protect delivery capacity.'),
}

p = Presentation(sys.argv[1])
for sh in p.slides[29].shapes:
    if not sh.has_table:
        continue
    t = sh.table
    for i, row in enumerate(t.rows):
        if i not in ROWS:
            continue
        vals = ROWS[i]
        for j, text in enumerate(vals):
            cell = t.cell(i, j)
            para = cell.text_frame.paragraphs[0]
            if not para.runs:
                continue
            para.runs[0].text = text
            for r in list(para.runs)[1:]:
                r._r.getparent().remove(r._r)
p.save(sys.argv[2])
print('actions table reworded as recommendations')
