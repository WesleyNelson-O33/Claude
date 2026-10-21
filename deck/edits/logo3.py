"""The CTS logo, top right on every light page.

Replaces the text wordmark that stood in while there was no logo file. Danica
supplied the artwork, which is trimmed to its ink and held in deck/assets.
"""
import sys
from pptx import Presentation
from pptx.util import Inches

LOGO = '/home/user/Claude/deck/assets/cts-logo.png'
W = 1.05                       # inches wide; the artwork is 1.779:1
RIGHT, TOP = 9.45, 0.20        # the page's right margin, clear of the title

def place(slide):
    tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if sh.name in ('cts mark', 'cts mark sub', 'cts logo'):
            tree.remove(sh._element)
    h = W / 1.779
    pic = slide.shapes.add_picture(LOGO, Inches(RIGHT - W), Inches(TOP),
                                   width=Inches(W), height=Inches(h))
    pic.name = 'cts logo'

p = Presentation(sys.argv[1])
n = 0
for s in p.slides:
    if s.slide_layout.name != 'CTS_LIGHT':
        continue
    place(s)
    n += 1
p.save(sys.argv[2])
print(f'logo placed on {n} light pages')
