"""Slide 5: the table full width, the spreadsheet heading cropped off, and the
commentary moved to a strip underneath."""
import copy, io, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]
WORK = '/tmp/claude-0/-home-user-Claude/a00ffe59-433c-507b-9006-917d8c60083a/scratchpad/work'

p = Presentation(SRC)
s = p.slides[4]
by = {sh.name: sh for sh in s.shapes}
pic = by['snip pl_income']

# crop everything above the navy header row: the "Profit & Loss" heading, the
# month selector and the Actual labels are spreadsheet furniture, and the
# header row already carries the months
im = Image.open(io.BytesIO(pic.image.blob)).convert('RGB')
TOP = 172
im.crop((0, TOP, im.width, im.height)).save(f'{WORK}/pl_income_cropped.png')
w, h = im.width, im.height - TOP

tree = s.shapes._spTree
tree.remove(pic._element)
X, Y, W = 0.55, 1.38, 8.90
new = s.shapes.add_picture(f'{WORK}/pl_income_cropped.png', Inches(X), Inches(Y),
                           width=Inches(W), height=Inches(W * h / w))
new.name = 'snip pl_income'
tree.remove(new._element)
tree.insert(2, new._element)

# the commentary becomes a strip under the table, in two columns
PY, PH = 3.84, 1.22
panel, label, body = by['commentary panel'], by['Text 4'], by['Text 5']
panel.left, panel.top, panel.width, panel.height = Inches(0.55), Inches(PY), Inches(8.90), Inches(PH)
label.left, label.top, label.width = Inches(0.75), Inches(PY + 0.09), Inches(3.0)

bullets = [para.text for para in body.text_frame.paragraphs if para.text.strip()]
half = 2                       # the first two run longer than the last three
cols = [bullets[:half], bullets[half:]]

second = copy.deepcopy(body._element)
tree.append(second)
body2 = [sh for sh in s.shapes if sh._element is second][0]
body2.name = 'Text 5b'

for box, items in ((body, cols[0]), (body2, cols[1])):
    box.left = Inches(0.75 if items is cols[0] else 5.10)
    box.top, box.width, box.height = Inches(PY + 0.29), Inches(4.05), Inches(PH - 0.33)
    tf = box.text_frame
    keep = tf.paragraphs[0]
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    template_run = keep.runs[0]
    for r in list(keep.runs)[1:]:
        r._r.getparent().remove(r._r)
    template_run.text = items[0]
    for extra in items[1:]:
        np = copy.deepcopy(keep._p)
        keep._p.getparent().append(np)
        from pptx.text.text import _Paragraph
        para = _Paragraph(np, tf)
        para.runs[0].text = extra
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size = Pt(8)

p.save(OUT)
print('slide 5 rebuilt: table full width, commentary underneath')
