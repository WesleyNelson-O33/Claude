"""A chart for the departmental result, with the table kept underneath it.

The chart goes in as page 17 and the table it came from moves to page 18, so the
board reads the picture and the numbers are on the next page if anyone asks.

Of the four cuts drawn, this is the one that explains the table rather than
restating it. What each department earned sits above the line, the overhead
charged to it sits below, and the diamond is what was left. A department can be
on a good margin and still finish below the line, and that is visible here and
nowhere else.

This is a monthly page.
"""
import sys
sys.path.insert(0, '/home/user/Claude/deck/edits')
from pptx import Presentation
from addslide import add

SNIP = '/home/user/Claude/deck/snips'

BULLETS = [
    'Gross profit of $188k did not cover overheads of $229k. That gap is the '
    '$41k loss.',
    'Only Support and Integration covered the overhead charged to them.',
    'Events carries 48% of all overhead on 13% of revenue, so a 78% margin still '
    'ends in a $54k loss.',
    'Consulting earned $74 of gross profit on $4k of revenue and carries $20k of '
    'overhead.',
]
NOTES = """Gross profit against the overhead charged to each department for September 2026, taken from section 7 of the Overhead Split tab. The bar above the line is what the department earned, the bar below is the overhead allocated to it, and the diamond is the result.

Support earned $101,290 and carries $72,283, leaving $29,007.
Events earned $56,881 and carries $111,108, leaving a loss of $54,227.
Video earned $5,813 and carries $12,659, leaving a loss of $6,846.
Integration earned $23,756 and carries $13,285, leaving $10,471.
Consulting earned $74 and carries $19,944, leaving a loss of $19,870.

The company earned $187,814 of gross profit against $229,278 of overhead. The $41,464 difference is the loss for the month. Put plainly, the business did not earn enough gross profit to carry its own overhead base.

Two things come out of the shape of this chart.

The first is Events. Its gross margin was 77.6%, the second highest of the five, and it still produced the largest loss on the page. It carries $111,108 of overhead, which is 48.5% of the company total, on $73,304 of revenue, which is 13.0% of the company total. No improvement in margin fixes that. Either Events carries more revenue or it carries less overhead, and the overhead keys are worth testing before the next budget.

The second is Integration. It brought in $234,858 of revenue, the second largest of the five, and converted it into $23,756 of gross profit at a 10.1% margin. It covers its overhead and returns $10,471, so it is not a loss maker, but it is doing the most work for the least return.

Support is the one department behaving as it should. A 41.7% margin on contract revenue, covering its allocated overhead, and returning $29,007.

The full table is on the next page."""

p = Presentation(sys.argv[1])
w, h, strip = add(p, 8, 16, 'DEPARTMENTAL PERFORMANCE',
                  'Departmental result, September 2026',
                  f'{SNIP}/dept_mix.png', BULLETS, NOTES, 'snip dept_mix')
print(f'page 17: the chart | picture {w:.2f} x {h:.2f}, strip {strip:.2f}')

table = p.slides[17]
title = next(sh for sh in table.shapes if sh.name == 'Text 1')
para = title.text_frame.paragraphs[0]
para.runs[0].text = 'Departmental result in full, September 2026'
for r in list(para.runs)[1:]:
    r._r.getparent().remove(r._r)
print('page 18: the table, retitled')

p.save(sys.argv[2])
