"""Remove one slide by its number. Nothing is added in this pass, because adding
and removing together makes python-pptx hand two parts the same name."""
import sys
from pptx import Presentation

p = Presentation(sys.argv[1])
at = int(sys.argv[3]) - 1
title = next((sh.text_frame.text.strip() for sh in p.slides[at].shapes
              if sh.name == 'Text 1'), '(no title)')
ids = p.slides._sldIdLst
el = list(ids)[at]
p.part.drop_rel(el.rId)
ids.remove(el)
print(f'removed page {at + 1}: {title}')
p.save(sys.argv[2])
