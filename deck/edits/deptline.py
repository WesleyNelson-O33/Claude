"""The departmental result above and below the line.

The gross profit page put all five departments above the line, because all five
do make a gross profit. Read quickly it says everybody made money, and three of
them did not. This page says who did.

It goes in front of the gross profit page, and that page is retitled to say
before overhead, so the order reads result first and the component second.

This is a monthly page.
"""
import sys
sys.path.insert(0, '/home/user/Claude/deck/edits')
from pptx import Presentation
from addslide import add

SNIP = '/home/user/Claude/deck/snips'

BULLETS = [
    'Two of the five made a profit. Support kept $29k and Integration $10k.',
    'Events lost $54k, Consulting $20k and Video $7k.',
    'All five made a gross profit. It is the overhead charged to them that puts '
    'three below the line.',
    'The five together are the $41k loss for the month.',
]
NOTES = """The result by department for September 2026 after the overhead charged to each one, taken from section 7 of the Overhead Split tab. Teal is above the line, red is below.

Support $29,007. Integration $10,471. Video a loss of $6,846. Consulting a loss of $19,870. Events a loss of $54,227. The five add to the $41,464 loss for the month.

Two of the five made a profit.

The reason this page exists is that the gross profit page behind it shows all five departments above the line, and that is true: every department earned more than its direct cost. Support made $101,290 of gross profit, Events $56,881, Integration $23,756, Video $5,813 and even Consulting made $74. On gross profit nobody lost money.

What puts three of them below the line is the overhead charged to them. $229,278 of it across the month, against $187,814 of gross profit earned. The company cannot cover its overhead base out of the margin it is making, and the departments where that bites hardest are the ones carrying overhead out of proportion to their revenue.

Events is the clearest case. It carries $111,108 of overhead, 48.5% of the company total, on 13.0% of the revenue. Its margin is the second best of the five at 77.6% and it still posts the largest loss on the page.

Consulting is the other kind of problem. It earned $74 of gross profit on $3,531 of revenue and carries $19,944 of overhead, so almost the whole of its loss is overhead on a department that is barely trading.

Video's loss of $6,846 is the smallest and the least structural. It billed $7,303 against $20,025 in the same month last year, so it is a volume problem on a department with a 79.6% margin.

Two points for the board. First, improving margin does not fix Events or Consulting, because their losses are overhead, not margin. Second, the overhead keys were set for FY27 and have not been tested since, which is the action already on the recommendations page."""

p = Presentation(sys.argv[1])
w, h, strip = add(p, 8, 25, 'DEPARTMENTAL PERFORMANCE',
                  'Result by department, above and below the line',
                  f'{SNIP}/dept_line.png', BULLETS, NOTES, 'snip dept_line')
print(f'page 26: above and below the line | picture {w:.2f} x {h:.2f}, strip {strip:.2f}')

title = next(sh for sh in p.slides[26].shapes if sh.name == 'Text 1')
para = title.text_frame.paragraphs[0]
para.runs[0].text = 'Gross profit before overhead, year on year'
for r in list(para.runs)[1:]:
    r._r.getparent().remove(r._r)
print('page 27 retitled so it cannot read as the net result')

p.save(sys.argv[2])
