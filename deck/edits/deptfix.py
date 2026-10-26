"""Put the corrected September departmental figures through the deck.

The overhead allocation now follows the analysis's tab 5 and the department
pages, which use the 49/18/3/1/1 Staff key and leave other income out.  So the
September results move:

    Support      29,007 -> 20,832   overhead  72,283 -> 80,458
    Events      -54,227 -> -46,496  overhead 111,108 -> 103,377
    Video        -6,846 -> -6,626   overhead  12,659 -> 12,440
    Integration  10,471 -> 10,563   overhead  13,285 -> 13,193
    Consulting  -19,870 -> -20,692  overhead  19,944 -> 20,766

The five now come to a 42,419 loss rather than 41,464, because the 955 of other
income sits outside the departments.  Company pages still read 41,464, which is
the net after it, so only the sentences that add the departments up change.

Her wording is kept.  Only the figures move.
"""
from pptx import Presentation

DECK = 'deck/CTS Monthly Business Review - 30092026.pptx'

EDITS = {
    3: [('+$29k after overhead', '+$21k after overhead'),
        ('+$10k after overhead', '+$11k after overhead'),
        ('$(54)k after overhead', '$(46)k after overhead'),
        ('$(20)k after overhead', '$(21)k after overhead'),
        ('$(41)k after overhead', '$(42)k after overhead')],
    17: [('a $29k profit on $243k revenue', 'a $21k profit on $243k revenue'),
         ('a $54k loss after $111k in allocated overheads',
          'a $46k loss after $103k in allocated overheads'),
         ('limited the final profit to $10k', 'limited the final profit to $11k'),
         ('against approximately $20k of overheads',
          'against approximately $21k of overheads')],
    18: [('overheads of $229k. That gap is the $41k loss.',
          'overheads of $230k. That gap is the $42k loss.'),
         ('Events carries 48% of all overhead', 'Events carries 45% of all overhead'),
         ('still ends in a $54k loss', 'still ends in a $46k loss'),
         ('and carries $20k of overhead', 'and carries $21k of overhead')],
    19: [('$29k profit on $243k revenue', '$21k profit on $243k revenue'),
         ('a $61k loss after $124k of overheads', 'a $53k loss after $116k of overheads'),
         ('contributed to a $9k loss', 'contributed to a $10k loss')],
    20: [('After $72k of allocated overheads, Support recorded a $29k profit.',
          'After $80k of allocated overheads, Support recorded a $21k profit.'),
         ('Its allocated office wages were $42k', 'Its allocated office wages were $43k')],
    21: [('a $54k loss after $111k of allocated overheads',
          'a $46k loss after $103k of allocated overheads'),
         ('Approximately $72k relates to shared overhead allocations.',
          'Approximately $64k relates to shared overhead allocations.')],
    22: [('against approximately $13k of overheads', 'against approximately $12k of overheads')],
    23: [('After $13k of overheads, Integration reported a $10k profit.',
          'After $13k of overheads, Integration reported an $11k profit.')],
    24: [('a $20k loss after approximately $20k of overheads',
          'a $21k loss after approximately $21k of overheads')],
    25: [('Support earned $243k, cost $214k and kept $29k.',
          'Support earned $243k, cost $222k and kept $21k.'),
         ('Events earned $73k, cost $128k and lost $54k.',
          'Events earned $73k, cost $120k and lost $46k.'),
         ('cost $224k and kept $10k', 'cost $224k and kept $11k'),
         ('against $603k of cost. That is the $41k loss.',
          'against $604k of cost. That is the $42k loss.')],
    26: [('Support kept $29k and Integration $10k.',
          'Support kept $21k and Integration $11k.'),
         ('Events lost $54k, Consulting $20k and Video $7k. Together the five are the $41k loss.',
          'Events lost $46k, Consulting $21k and Video $7k. Together the five are the $42k loss.'),
         ('left only $10k after overhead', 'left only $11k after overhead')],
}


def swap(shape, old, new):
    """Replace inside a run where possible, so the formatting survives."""
    done = 0
    for para in shape.text_frame.paragraphs:
        if old not in para.text:
            continue
        for run in para.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                done += 1
                break
        else:
            # the phrase is split across runs, so rewrite the first and clear the rest
            runs = para.runs
            if runs:
                runs[0].text = para.text.replace(old, new)
                for r in runs[1:]:
                    r.text = ''
                done += 1
    return done


if __name__ == '__main__':
    prs = Presentation(DECK)
    missed = []
    for page, pairs in EDITS.items():
        slide = prs.slides[page - 1]
        for old, new in pairs:
            hits = 0
            for sh in slide.shapes:
                if sh.has_text_frame:
                    hits += swap(sh, old, new)
            if hits:
                print(f'  page {page:<4}{old[:58]}')
            else:
                missed.append((page, old))
    if missed:
        for page, old in missed:
            print(f'  MISSED page {page}: {old}')
        raise SystemExit('some phrases were not found, nothing saved')
    prs.save(DECK)
    print('deck saved')
