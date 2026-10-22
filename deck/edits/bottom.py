"""The last nine pages move their commentary to the bottom, like every other page.

These pages had the chart on the left and a tall commentary panel down the right.
They are the only ones left that did not match, so they now use the same strip as
the rest: two balanced columns across the foot of the page, and the picture in the
standard box above it.

There is a cost and it is worth being plain about it. The commentary on these
pages runs to five full bullets, which were written to fill a tall narrow column.
Laid across the foot of the page that still needs a deep strip, and a deep strip
leaves less room above it, so these charts come out smaller than they were. The
only way to have both is shorter commentary on these pages.
"""
import copy, sys
sys.path.insert(0, '/home/user/Claude/deck/edits')
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from addslide import balance, need, write, FLOOR, LEFT, SIZE

SNIP = '/home/user/Claude/deck/snips'

TOP, CLEAR = 1.10, 0.10          # the same box the rest of the deck uses
WIDTH = 8.90


def to_bottom(slide):
    by = {sh.name: sh for sh in slide.shapes}
    tree = slide.shapes._spTree
    body = by['Text 5']
    bullets = [p.text for p in body.text_frame.paragraphs if p.text.strip()]
    cols = list(balance(bullets, 4.05))

    ph = max(need(c, 4.05) for c in cols if c) + 0.30
    strip_top = FLOOR - ph

    extra = by.get('Text 5b')
    if extra is None:
        el = copy.deepcopy(body._element)
        tree.append(el)
        extra = next(sh for sh in slide.shapes if sh._element is el)
        extra.name = 'Text 5b'

    panel, label = by['commentary panel'], by['Text 4']
    panel.left, panel.top = Inches(LEFT), Inches(strip_top)
    panel.width, panel.height = Inches(WIDTH), Inches(ph)
    label.left, label.top, label.width = Inches(0.75), Inches(strip_top + 0.08), Inches(3.0)

    for box, x, items in ((body, 0.75, cols[0]), (extra, 5.10, cols[1])):
        box.left, box.top = Inches(x), Inches(strip_top + 0.27)
        box.width, box.height = Inches(4.05), Inches(ph - 0.30)
        write(box, items)
        for para in box.text_frame.paragraphs:
            para.alignment = None
            for r in para.runs:
                r.font.size = Pt(SIZE)

    # the picture goes back from file, because the charts on these pages were
    # redrawn in the pack at the shape of the box they now have to fill
    old = next(sh for sh in slide.shapes
               if 'PICTURE' in str(sh.shape_type) and sh.name != 'cts logo')
    name = old.name
    path = f"{SNIP}/{name.replace('snip ', '')}.png"
    box_h = strip_top - CLEAR - TOP
    im = Image.open(path)
    aspect = im.width / im.height
    h = min(box_h, WIDTH / aspect)
    w = h * aspect
    tree.remove(old._element)
    pic = slide.shapes.add_picture(path, Inches(LEFT + (WIDTH - w) / 2),
                                   Inches(TOP + (box_h - h) / 2),
                                   width=Inches(w), height=Inches(h))
    pic.name = name
    tree.remove(pic._element)
    tree.insert(2, pic._element)
    return name, w, h, ph


p = Presentation(sys.argv[1])
for i, slide in enumerate(p.slides):
    panel = next((sh for sh in slide.shapes if sh.name == 'commentary panel'), None)
    if panel is None or panel.top / 914400 >= 3.0:
        continue
    name, w, h, ph = to_bottom(slide)
    print(f'p{i+1:>2} {name:20} picture {w:.2f}x{h:.2f}, strip {ph:.2f}')
p.save(sys.argv[2])
