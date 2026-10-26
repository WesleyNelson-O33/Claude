"""Bring the speaker notes onto the corrected September figures.

The notes carry the detail to the dollar, so every one of them moved when the
allocation changed. Her wording is kept. Only the figures move, plus two things
that were already wrong before this: page 18 still described a diamond marker
that was taken off the chart, and the department totals are now stated as the
result before other income, because the 955 of other income sits outside the
departments and the company net profit is after it.
"""
from pptx import Presentation

DECK = 'deck/CTS Monthly Business Review - 30092026.pptx'

EDITS = {
    3: [('and $29k after allocated overheads', 'and $21k after allocated overheads'),
        ('Events generated $73k at 77.6% and a loss of $54k.',
         'Events generated $73k at 77.6% and a loss of $46k.'),
        ('Integration generated $235k at 10.1% and $10k.',
         'Integration generated $235k at 10.1% and $11k.'),
        ('Consulting generated $4k at 2.1% and a loss of $20k.',
         'Consulting generated $4k at 2.1% and a loss of $21k.')],
    17: [('Overhead $72,283, of which $59,182 is allocated admin. Result $29,007.',
          'Overhead $80,458, of which $67,357 is allocated admin. Result $20,832.'),
         ('Overhead $111,108, of which $71,672 is allocated admin. Result a loss of $54,227.',
          'Overhead $103,377, of which $63,941 is allocated admin. Result a loss of $46,496.'),
         ('Overhead $12,659, of which $9,782 is allocated admin. Result a loss of $6,846.',
          'Overhead $12,440, of which $9,562 is allocated admin. Result a loss of $6,626.'),
         ('Overhead $13,285, of which $4,148 is allocated admin. Result $10,471.',
          'Overhead $13,193, of which $4,056 is allocated admin. Result $10,563.'),
         ('Overhead $19,944, of which $13,979 is allocated admin. Result a loss of $19,870.',
          'Overhead $20,766, of which $14,801 is allocated admin. Result a loss of $20,692.'),
         ('The five departments total a loss of $41,464, which is net profit for the month.',
          'The five departments total a loss of $42,419. Net profit for the month is a loss '
          'of $41,464, after $955 of other income that sits outside the departments.'),
         ('And Events carries $71,672 of office costs on $73,304 of sales.',
          'And Events carries $63,941 of office costs on $73,304 of sales.')],
    18: [('and the diamond is the result', 'and the third bar is the result'),
         ('Support earned $101,290 and carries $72,283, leaving $29,007.',
          'Support earned $101,290 and carries $80,458, leaving $20,832.'),
         ('Events earned $56,881 and carries $111,108, leaving a loss of $54,227.',
          'Events earned $56,881 and carries $103,377, leaving a loss of $46,496.'),
         ('Video earned $5,813 and carries $12,659, leaving a loss of $6,846.',
          'Video earned $5,813 and carries $12,440, leaving a loss of $6,626.'),
         ('Integration earned $23,756 and carries $13,285, leaving $10,471.',
          'Integration earned $23,756 and carries $13,193, leaving $10,563.'),
         ('Consulting earned $74 and carries $19,944, leaving a loss of $19,870.',
          'Consulting earned $74 and carries $20,766, leaving a loss of $20,692.'),
         ('against $229,278 of overhead. The $41,464 difference is the loss for the month.',
          'against $230,233 of overhead. The $42,419 difference is the loss for the month '
          'before other income.'),
         ('It carries $111,108 of overhead, which is 48.5% of the company total',
          'It carries $103,377 of overhead, which is 44.9% of the company total'),
         ('It covers its overhead and returns $10,471',
          'It covers its overhead and returns $10,563'),
         ('covering its allocated overhead, and returning $29,007.',
          'covering its allocated overhead, and returning $20,832.')],
    19: [('All overhead $72,283. Result $29,007.', 'All overhead $80,458. Result $20,832.'),
         ('All overhead $123,767. Result a loss of $61,073.',
          'All overhead $115,817. Result a loss of $53,122.'),
         ('All overhead $33,229. Result a loss of $9,398.',
          'All overhead $33,959. Result a loss of $10,129.'),
         ('The three total a loss of $41,464, which is net profit for the month, and the '
          'check row on the tab reconciles it to the five department view.',
          'The three total a loss of $42,419, the company result before other income, and '
          'the check row on the tab reconciles it to the five department view.'),
         ('it carries $123,767 of running costs against $80,607 of sales',
          'it carries $115,817 of running costs against $80,607 of sales')],
    20: [('were $72,283, of which $13,101 is its own spending and $59,182 is its share of the office',
          'were $80,458, of which $13,101 is its own spending and $67,357 is its share of the office'),
         ('That leaves $29,007 for the month.', 'That leaves $20,832 for the month.'),
         ('employment expense of $51,242, of which $8,980 is its own staff and $42,262 its share',
          'employment expense of $51,785, of which $8,980 is its own staff and $42,805 its share')],
    21: [('were $111,108, of which $39,436 is its own spending and $71,672 is its share of the office',
          'were $103,377, of which $39,436 is its own spending and $63,941 is its share of the office'),
         ('That leaves a loss of $54,227 for the month.', 'That leaves a loss of $46,496 for the month.'),
         ('employment expense of $50,959, of which $35,665 is its own staff and $15,294 its share',
          'employment expense of $51,550, of which $35,665 is its own staff and $15,885 its share')],
    22: [('were $12,659, of which $2,878 is its own spending and $9,782 is its share of the office',
          'were $12,440, of which $2,878 is its own spending and $9,562 is its share of the office'),
         ('That leaves a loss of $6,846 for the month.', 'That leaves a loss of $6,626 for the month.'),
         ('employment expense of $3,961, of which $2,660 is its own staff and $1,301 its share',
          'employment expense of $5,294, of which $2,660 is its own staff and $2,634 its share')],
    23: [('were $13,285, of which $9,136 is its own spending and $4,148 is its share of the office',
          'were $13,193, of which $9,136 is its own spending and $4,056 is its share of the office'),
         ('That leaves $10,471 for the month.', 'That leaves $10,563 for the month.'),
         ('employment expense of $8,832, of which $8,563 is its own staff and $269 its share',
          'employment expense of $9,440, of which $8,563 is its own staff and $877 its share')],
    24: [('were $19,944, of which $5,965 is its own spending and $13,979 is its share of the office',
          'were $20,766, of which $5,965 is its own spending and $14,801 is its share of the office'),
         ('That leaves a loss of $19,870 for the month.', 'That leaves a loss of $20,692 for the month.'),
         ('employment expense of $9,496, of which $5,545 is its own staff and $3,951 its share',
          'employment expense of $6,422, of which $5,545 is its own staff and $877 its share')],
    25: [('Support earned $242,948, cost $213,942 and kept $29,007.',
          'Support earned $242,948, cost $222,117 and kept $20,832.'),
         ('Events earned $73,304, cost $127,531 and lost $54,227.',
          'Events earned $73,304, cost $119,800 and lost $46,496.'),
         ('Video earned $7,303, cost $14,149 and lost $6,846.',
          'Video earned $7,303, cost $13,929 and lost $6,626.'),
         ('Integration earned $234,858, cost $224,387 and kept $10,471.',
          'Integration earned $234,858, cost $224,295 and kept $10,563.'),
         ('Consulting earned $3,531, cost $23,401 and lost $19,870.',
          'Consulting earned $3,531, cost $24,223 and lost $20,692.'),
         ('Across the five, $561,945 of revenue against $603,408 of cost, which is the '
          '$41,464 loss for the month.',
          'Across the five, $561,945 of revenue against $604,364 of cost, which is the '
          '$42,419 loss for the month before other income.'),
         ('plus $72,283 of overhead, giving $213,942.', 'plus $80,458 of overhead, giving $222,117.'),
         ('the five results add to the $41,464 net loss on the P&L.',
          'the five results add to the $42,419 loss before other income, which is the '
          '$41,464 net loss on the P&L once the $955 of other income is added.'),
         ('and almost the same cost, $213,942 and $224,387. The results are $29,007 and $10,471.',
          'and almost the same cost, $222,117 and $224,295. The results are $20,832 and $10,563.'),
         ('because $111,108 of the $127,531 it costs is allocated overhead',
          'because $103,377 of the $119,800 it costs is allocated overhead'),
         ('Between them they earned $10,834 and cost $37,550.',
          'Between them they earned $10,834 and cost $38,152.')],
    26: [('Support $29,007. Integration $10,471. Video a loss of $6,846. Consulting a loss of '
          '$19,870. Events a loss of $54,227. The five add to the $41,464 loss for the month.',
          'Support $20,832. Integration $10,563. Video a loss of $6,626. Consulting a loss of '
          '$20,692. Events a loss of $46,496. The five add to the $42,419 loss for the month '
          'before other income.'),
         ('$229,278 across the month against $187,814 of gross profit earned.',
          '$230,233 across the month against $187,814 of gross profit earned.'),
         ('It carries $111,108 of overhead, 48.5% of the company total, on 13.0% of the revenue.',
          'It carries $103,377 of overhead, 44.9% of the company total, on 13.0% of the revenue.'),
         ('and carries $19,944 of overhead, so almost the whole of its loss',
          'and carries $20,766 of overhead, so almost the whole of its loss'),
         ('so after $13,285 of overhead only $10,471 reached the line',
          'so after $13,193 of overhead only $10,563 reached the line'),
         ("Video's loss of $6,846 is the smallest", "Video's loss of $6,626 is the smallest")],
}


if __name__ == '__main__':
    prs = Presentation(DECK)
    missed = []
    for page, pairs in EDITS.items():
        slide = prs.slides[page - 1]
        if not slide.has_notes_slide:
            missed.append((page, 'no notes slide'))
            continue
        tf = slide.notes_slide.notes_text_frame
        for old, new in pairs:
            hit = False
            for para in tf.paragraphs:
                if old not in para.text:
                    continue
                runs = para.runs
                if runs:
                    runs[0].text = para.text.replace(old, new)
                    for r in runs[1:]:
                        r.text = ''
                    hit = True
            if hit:
                print(f'  page {page:<4}{old[:62]}')
            else:
                missed.append((page, old))
    if missed:
        for page, old in missed:
            print(f'  MISSED page {page}: {old}')
        raise SystemExit('some phrases were not found, nothing saved')
    prs.save(DECK)
    print('notes saved')
