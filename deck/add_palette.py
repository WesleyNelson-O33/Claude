"""Put the CTS palette on the side of every slide, off the canvas.

Two ways to reach it while she works:
  1. The theme colour scheme, so the nine show up in PowerPoint's own colour
     picker under Theme Colors. That is what fix_theme.py sets.
  2. A strip of swatches parked to the right of each slide, outside the slide
     area. She can see them while editing and eyedropper off them. They do not
     appear in the show, in a print, or in a PDF.
"""
import os, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

D = os.path.dirname(os.path.abspath(__file__))
DECK = f'{D}/CTS Monthly Business Review - September 2026.pptx'
B = json.load(open(f'{D}/../pack/data/brand.json'))

SWATCH = [('Bright Teal',  B['bright_teal'],  'Logo green'),
          ('Dark Teal',    B['dark_teal'],    'Production'),
          ('Primary Blue', B['primary_blue'], 'Logo blue'),
          ('Deep Blue',    B['deep_blue'],    ''),
          ('Accent Red',   B['accent_red'],   'Logo red, Support'),
          ('Deep Red',     B['deep_red'],     ''),
          ('Accent Purple',B['purple'],       'Consulting'),
          ('Charcoal',     B['charcoal'],     'Text'),
          ('White',        B['white'],        '')]

TAG = 'CTS palette'
prs = Presentation(DECK)
SW, SH = prs.slide_width, prs.slide_height
X = SW + Inches(0.5)                 # parked off the right edge of the slide
W, H, GAP = Inches(2.3), Inches(0.42), Inches(0.08)

def dark(hexv):
    r, g, b = (int(hexv[i:i+2], 16) for i in (0, 2, 4))
    return (r*299 + g*587 + b*114) / 1000 < 150

n = 0
for slide in prs.slides:
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith(TAG):
            sh._element.getparent().remove(sh._element)
    top = Inches(0.4)
    lab = slide.shapes.add_textbox(X, top, W, Inches(0.3))
    lab.text_frame.text = f'{TAG} — not on the slide, eyedropper from here'
    p = lab.text_frame.paragraphs[0]
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(B['charcoal'])
    top += Inches(0.34)
    for name, hexv, note in SWATCH:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X, top, W, H)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(hexv)
        box.line.color.rgb = RGBColor.from_string(B['grey'])
        box.line.width = Pt(0.75)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        tf.text = f'{name}   #{hexv}' + (f'   {note}' if note else '')
        r = tf.paragraphs[0].runs[0]
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(
            B['white'] if dark(hexv) else B['charcoal'])
        top += H + GAP
    n += 1

prs.save(DECK)
print(f'palette parked beside {n} slides, {len(SWATCH)} swatches each')
print('slide width', SW / 914400, 'in; strip sits at', X / 914400, 'in')
