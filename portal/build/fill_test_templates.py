"""Fill copies of the templates from the seed, for an end to end test of Build.

Writes to a scratch folder, never over the real templates. The P&L, budget
and ledger come straight from the seed data files, so a Build from these
copies should reproduce the seed to the cent. Earnings lines are synthesised
from the staff rows in the Earnings Details shape.
"""
import json, re, shutil, sys
from datetime import date
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path("/home/user/Claude/portal")
SRC = ROOT / "templates"
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/claude-0/-home-user-Claude/fe27351c-b17a-525d-8bba-ade44e966c97/scratchpad/filled")
OUT.mkdir(parents=True, exist_ok=True)


def js(name, var):
    txt = (ROOT / "data" / name).read_text()
    m = re.search(r"window\.%s = (.*);\s*$" % var, txt, re.S)
    return json.loads(m.group(1))


fin = js("CTS_fin_data.js", "CTS_FIN")
gl = js("CTS_gl_data.js", "CTS_GL")
staff = js("CTS_staff_data.js", "CTS_STAFF")
pipe = js("CTS_pipeline_data.js", "CTS_PIPELINE")
MONTHS = fin["months"]


def month_label(k):
    y, m = int(k[:4]), int(k[5:])
    return date(y, m, 1).strftime("%b-%y")


# ---- 01 P&L and 07 Budget: accounts down, the template's own months across
for src, key in (("01 Xero P&L.xlsx", "actual"), ("07 Budget.xlsx", "budget")):
    wb = load_workbook(SRC / src)
    ws = wb["P&L" if key == "actual" else "Budget"]
    heads = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_of = {h: i + 1 for i, h in enumerate(heads)}
    grid = fin[key]
    r = 2
    for name in sorted(grid):
        ws.cell(row=r, column=1, value=name)
        for mk, cents in grid[name].items():
            lbl = month_label(mk)
            if lbl in col_of:
                # Xero shows costs positive; the seed stores them negative
                ws.cell(row=r, column=col_of[lbl], value=abs(cents) / 100.0)
        r += 1
    # a total row, as a real Xero paste carries them
    ws.cell(row=r, column=1, value="Total Income")
    wb.save(OUT / src)
    print("filled", src, r - 2, "accounts")

# ---- 02 GL: the thirteen Xero columns from the seed ledger
wb = load_workbook(SRC / "02 Xero GL Transactions.xlsx")
ws = wb["GL"]
ci = {c: i for i, c in enumerate(gl["cols"])}
for n, row in enumerate(gl["rows"], start=2):
    d = row[ci["date"]]
    ws.cell(row=n, column=1, value=n - 1)
    ws.cell(row=n, column=2, value="")
    ws.cell(row=n, column=3, value=row[ci["account"]])
    ws.cell(row=n, column=4, value=row[ci["source"]])
    ws.cell(row=n, column=5, value=date(int(d[:4]), int(d[5:7]), int(d[8:10])))
    ws.cell(row=n, column=6, value=row[ci["contact"]])
    ws.cell(row=n, column=7, value=row[ci["debit"]] or None)
    ws.cell(row=n, column=8, value=row[ci["credit"]] or None)
    ws.cell(row=n, column=9, value=row[ci["jobNo"]])
    ws.cell(row=n, column=10, value=row[ci["invoiceNo"]])
    ws.cell(row=n, column=11, value="")
    ws.cell(row=n, column=12, value=row[ci["description"]])
    ws.cell(row=n, column=13, value=row[ci["costCentre"]])
wb.save(OUT / "02 Xero GL Transactions.xlsx")
print("filled 02 GL", len(gl["rows"]), "lines")

# ---- 03 Earnings: one line per staff month per bucket, in the report shape
wb = load_workbook(SRC / "03 Employment Hero Earnings.xlsx")
ws = wb["Earnings"]
LOC = {"ONSITE": "PWC IT Support - Brisbane [ONS]", "PRODUCTION": "Ashfield AGM Event [PRD]",
       "VIDEO": "Studio B [VID]", "INTEGRATION": "Trentham Site [INT]",
       "CONSULTING": "Marchmont Advisory [CONS]", "ADMIN": "CTS Head Office [CTS]"}
sc = {c: i for i, c in enumerate(staff["cols"])}
n = 3
for row in staff["rows"]:
    mk = row[sc["month"]]
    if mk not in ("2026-07", "2026-08"):
        continue
    dept, name, emp = row[sc["dept"]], row[sc["name"]], row[sc["employment"]] or "Permanent"
    total_h = row[sc["chargeable"]] + row[sc["nonChargeable"]] + row[sc["leave"]] + row[sc["publicHoliday"]]
    rate = (row[sc["grossCents"]] / 100.0) / total_h if total_h else 0
    for cat, hrs in ((emp + " Ordinary Hours", row[sc["chargeable"]] + row[sc["nonChargeable"]]),
                     ("Annual Leave", row[sc["leave"]]), ("Public Holiday", row[sc["publicHoliday"]])):
        if not hrs:
            continue
        gross = round(hrs * rate, 2)
        ws.cell(row=n, column=1, value=n); ws.cell(row=n, column=2, value=row[sc["empId"]])
        ws.cell(row=n, column=3, value=name); ws.cell(row=n, column=6, value=cat)
        ws.cell(row=n, column=7, value=round(hrs, 2)); ws.cell(row=n, column=8, value="Hours")
        ws.cell(row=n, column=11, value=LOC[dept])
        ws.cell(row=n, column=12, value="15/%s/%s - shift" % (mk[5:7], mk[:4]))
        ws.cell(row=n, column=13, value=round(rate, 4)); ws.cell(row=n, column=14, value="per Hour")
        ws.cell(row=n, column=15, value=gross); ws.cell(row=n, column=16, value=gross)
        ws.cell(row=n, column=17, value=round(gross * 0.12, 2))
        n += 1
# one back pay with no date, reallocated on the Adjustments tab
ws.cell(row=n, column=3, value="Tessa Moreau"); ws.cell(row=n, column=6, value="Back Payment")
ws.cell(row=n, column=7, value=0); ws.cell(row=n, column=8, value="")
ws.cell(row=n, column=11, value=LOC["ONSITE"]); ws.cell(row=n, column=12, value="correction to PR14 rate")
ws.cell(row=n, column=15, value=240.0); ws.cell(row=n, column=17, value=28.8)
loc = wb["Locations"]
for i, (d, l) in enumerate(LOC.items(), start=2):
    loc.cell(row=i, column=1, value=l); loc.cell(row=i, column=2, value=d)
    loc.cell(row=i, column=3, value="Non-chargeable" if d == "ADMIN" else "Chargeable")
adj = wb["Adjustments"]
adj.cell(row=2, column=1, value="Tessa Moreau"); adj.cell(row=2, column=2, value="Back Payment")
adj.cell(row=2, column=3, value="PR14"); adj.cell(row=2, column=4, value=6)
adj.cell(row=2, column=5, value=LOC["ONSITE"]); adj.cell(row=2, column=6, value="Worked time")
wb.save(OUT / "03 Employment Hero Earnings.xlsx")
print("filled 03 Earnings", n - 3, "lines plus one back pay")

# ---- 04 05 06: pipeline placeholders from the seed
def iso(s):
    return date(int(s[:4]), int(s[5:7]), int(s[8:10])) if s else None
wb = load_workbook(SRC / "04 Zoho Deals.xlsx"); ws = wb["Deals"]
for n, d in enumerate(pipe["deals"], start=2):
    for c, v in enumerate([d["name"], d["account"], d["stage"], d["amount"] / 100, iso(d["close"]), d["owner"],
                           d["prob"] * 100, d["expected"] / 100, d["type"], d["source"], "", iso(d["created"]), None, d["dept"]], start=1):
        ws.cell(row=n, column=c, value=v)
wb.save(OUT / "04 Zoho Deals.xlsx")
wb = load_workbook(SRC / "05 OnRent Orders.xlsx"); ws = wb["Orders"]
for n, o in enumerate(pipe["orders"], start=2):
    for c, v in enumerate([o["no"], o["client"], o["title"], iso(o["start"]), iso(o["end"]), o["dept"], o["status"],
                           o["value"] / 100, o["owner"], iso(o["created"])], start=1):
        ws.cell(row=n, column=c, value=v)
wb.save(OUT / "05 OnRent Orders.xlsx")
wb = load_workbook(SRC / "06 Qwilr Quotes.xlsx"); ws = wb["Quotes"]
for n, q in enumerate(pipe["quotes"], start=2):
    for c, v in enumerate([q["ref"], q["client"], q["title"], iso(q["sent"]), q["status"], iso(q["accepted"]),
                           q["value"] / 100, q["owner"], q["dept"]], start=1):
        ws.cell(row=n, column=c, value=v)
wb.save(OUT / "06 Qwilr Quotes.xlsx")
print("filled pipeline:", len(pipe["deals"]), "deals", len(pipe["orders"]), "orders", len(pipe["quotes"]), "quotes")

# ---- 08 Config: as shipped, with one split base confirmed to exercise that path
wb = load_workbook(SRC / "08 Config.xlsx")
wb["Split bases"]["E3"] = "Confirmed"
wb.save(OUT / "08 Config.xlsx")
print("copied 08 Config with Staff basis set to Confirmed")
print("filled templates in", OUT)
