"""Paste the year so far into the new split file, so it opens with real numbers."""
import os, openpyxl
from openpyxl.styles import Font

SRC = "/home/user/Claude/ptax/out/EH Payroll Split FY2027.xlsx"
OLD = "/home/user/Claude/ptax/EH Payroll Split File FY2027.xlsx"
F = Font(name="Aptos Narrow", size=11)
D, M = "dd/mm/yyyy", "#,##0.00"


def copy(dst, src_rows, ncols, skip, datecols, moneycols, key):
    n = 0
    for row in src_rows:
        if row[key] is None:
            continue
        n += 1
        for j in range(ncols):
            v = row[skip + j]
            if v is None:
                continue
            c = dst.cell(row=2 + n, column=1 + j, value=v)
            c.font = F
            if j in datecols:
                c.number_format = D
            elif j in moneycols:
                c.number_format = M
    return n


def main():
    wb = openpyxl.load_workbook(SRC)
    old = openpyxl.load_workbook(OLD, data_only=True, read_only=True)
    n = copy(wb["1 Pay run totals"],
             old["EH data"].iter_rows(min_row=3, values_only=True),
             19, 1, {1, 2}, set(range(7, 19)), 4)
    print("pay run total lines:", n)
    n = copy(wb["2 Earnings"],
             old["EH data (2)"].iter_rows(min_row=3, values_only=True),
             20, 1, {1, 2}, {17, 18, 19, 15}, 4)
    print("earnings lines:", n)
    n = copy(wb["3 Employees"],
             old["Employee Details Report"].iter_rows(min_row=8, values_only=True),
             23, 1, {4}, set(), 1)
    print("people:", n)
    wb.save(SRC)
    print("wrote", SRC, "%.1f KB" % (os.path.getsize(SRC) / 1024))


if __name__ == "__main__":
    main()
