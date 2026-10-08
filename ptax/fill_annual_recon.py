"""Paste the year to date into the annual reconciliation workbook, so it opens
with real numbers in it rather than an empty shell."""
import os, shutil, openpyxl
from openpyxl.styles import Font

SRC = "/home/user/Claude/ptax/out/Payroll Tax Annual Reconciliation FY2027.xlsx"
SPLIT = "/home/user/Claude/ptax/EH Payroll Split File FY2027.xlsx"
FONT = Font(name="Aptos Narrow", size=11)


def main():
    wb = openpyxl.load_workbook(SRC)
    src = openpyxl.load_workbook(SPLIT, data_only=True, read_only=True)

    # 1 Earnings: the earnings report is columns B to U of her EH data (2) tab
    ws = wb["1 Earnings"]
    it = src["EH data (2)"].iter_rows(min_row=3, values_only=True)
    n = 0
    for row in it:
        if row[4] is None and row[18] is None:
            continue
        n += 1
        for j in range(20):                      # her B..U -> my A..T
            v = row[1 + j]
            if v is not None:
                c = ws.cell(row=2 + n, column=1 + j, value=v)
                c.font = FONT
                if j == 1 or j == 2:
                    c.number_format = "dd/mm/yyyy"
                elif j in (17, 18, 19, 15):
                    c.number_format = "#,##0.00"
    print("earnings rows pasted:", n)

    # 2 Employees: her Employee Details Report is columns B onwards from row 8
    we = wb["2 Employees"]
    m = 0
    for row in src["Employee Details Report"].iter_rows(min_row=8, values_only=True):
        if row[1] is None:
            continue
        m += 1
        for j in range(23):                      # her B..X -> my A..W
            v = row[1 + j]
            if v is not None:
                c = we.cell(row=2 + m, column=1 + j, value=v)
                c.font = FONT
                if j == 4:
                    c.number_format = "dd/mm/yyyy"
    print("employee rows pasted:", m)

    out = os.path.join(os.path.dirname(SRC),
                       "Payroll Tax Annual Reconciliation FY2027.xlsx")
    wb.save(out)
    print("wrote", out, "%.1f KB" % (os.path.getsize(out) / 1024))


if __name__ == "__main__":
    main()
