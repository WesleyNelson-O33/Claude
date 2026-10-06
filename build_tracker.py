import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = "/home/user/Claude/CTS_Bonus_Leave_Tracker.xlsx"
NROWS = 150
R0, R1 = 5, 5 + NROWS - 1
EH_MAX = 3000

F = "Arial"
f_norm = Font(name=F, size=10); f_bold = Font(name=F, size=10, bold=True); f_title = Font(name=F, size=14, bold=True)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF"); f_input = Font(name=F, size=10, color="0000FF")
f_note = Font(name=F, size=9, italic=True, color="555555")
fill_in = PatternFill("solid", fgColor="FFF2CC"); fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_hin = PatternFill("solid", fgColor="2E75B6"); fill_hout = PatternFill("solid", fgColor="1F4E78"); fill_heh = PatternFill("solid", fgColor="C00000")
fill_hp = PatternFill("solid", fgColor="7F7F7F"); fill_ex = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE = "dd/mm/yyyy"; NUM = 'General;-General;"-"'

wb = Workbook()
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def hdr(ws, cell, text, fill):
    ws[cell] = text; ws[cell].font = f_hdr; ws[cell].fill = fill; ws[cell].alignment = center; ws[cell].border = border

# ============================ STAFF ============================
ws = wb.active; ws.title = "Staff"; ws.sheet_properties.tabColor = "70AD47"
# (key, header, width, kind, note)
spec = [
 ("name", "Name", 12, "in", "type here"), ("sur", "Surname", 12, "in", "type here"),
 ("type", "EmploymentType", 13, "in", "current: Full-Time / Part-Time / Casual"),
 ("dob", "DateOfBirth", 11, "in", "type here"), ("bday", "Birthday", 11, "out", "next birthday"),
 ("start", "StartDate", 11, "in", "type here"),
 ("cas", "Casual Start", 11, "in", "date casual began"), ("pte", "PTE Start", 11, "in", "date part-time began"), ("fte", "FTE Start", 11, "in", "date full-time began"),
 ("anniv_in", "AnniversaryDate", 12, "in", "optional: HR-adjusted service start (unpaid leave)"), ("anniv", "Anniversary", 11, "out", "next anniversary of commencement (FT/PT start)"),
 ("ten", "Tenure", 8, "out", "years since commencement (FT/PT start)"), ("tenymd", "Tenure (Y/M/D)", 12, "out", "qualifying service counted"),
 ("hrs_due", "Bonus leave Accrued", 11, "out", "HOURS due per policy to date"), ("days_due", "Days", 7, "out", "days due per policy to date"),
 ("pthrs", "PT Hrs/Week", 8, "in", "part-timers only; blank = 24+"),
 ("chg", "Status Change", 40, "out", "from the three start dates"), ("elig", "Eligible?", 24, "out", ""), ("next", "Next Accrual", 11, "out", "date of next day"),
 ("open", "Opening Balance (hrs)", 10, "eh", "EH go-live load (only if commenced before the report start)"), ("acc", "Leave Accrued (hrs)", 10, "eh", "accrued in EH since commencement"),
 ("taken", "Leave Taken (hrs)", 10, "eh", "taken since commencement"), ("close", "Closing Balance (hrs)", 10, "eh", "EH balance now"),
 ("due", "Due per policy (hrs)", 10, "eh", "policy accrual from report start to as-of"), ("toacc", "Hours to Accrue", 10, "eh", "due - accrued in EH"),
 ("accq", "Add to EH now?", 9, "eh", "Yes = hours owed"), ("action", "Action", 36, "eh", ""),
 ("gap", "Opening gap vs policy (hrs)", 10, "eh", "0 when the report covers full history"), ("rows", "Rows in EH report", 8, "eh", "Bonus Leave rows found"),
 ("check", "Data Check", 30, "out", ""),
 ("mdays", "Earned this month (days)", 10, "eh", "days whose milestone falls in the accrual month"), ("mdate", "Milestone this month", 11, "eh", "date of that milestone"),
 # helpers
 ("h_act", "h active", 6, "hp", ""), ("h_S", "h S", 10, "hp", ""), ("h_n", "h n dates", 6, "hp", ""),
 ("h_D1", "h D1", 10, "hp", ""), ("h_D2", "h D2", 10, "hp", ""), ("h_D3", "h D3", 10, "hp", ""),
 ("h_T1", "h T1", 10, "hp", ""), ("h_T2", "h T2", 10, "hp", ""), ("h_T3", "h T3", 10, "hp", ""),
 ("h_P1", "h P1", 6, "hp", ""), ("h_P2", "h P2", 6, "hp", ""), ("h_P3", "h P3", 6, "hp", ""), ("h_Pc", "h Pc", 6, "hp", ""),
 ("h_M0", "h M0", 6, "hp", ""), ("h_B1", "h B1", 6, "hp", ""), ("h_B2", "h B2", 6, "hp", ""),
 ("h_MP", "h MP", 6, "hp", ""), ("h_B1P", "h B1P", 6, "hp", ""), ("h_B2P", "h B2P", 6, "hp", ""),
 ("h_dn", "h daysNow", 6, "hp", ""), ("h_dp", "h daysPFY", 6, "hp", ""), ("h_no", "h nextoff", 6, "hp", ""),
 ("h_rows", "h rows", 6, "hp", ""), ("h_last", "h lastrow", 6, "hp", ""),
 ("h_MM", "h MM", 6, "hp", ""), ("h_B1M", "h B1M", 6, "hp", ""), ("h_B2M", "h B2M", 6, "hp", ""), ("h_dm", "h daysLastMonth", 6, "hp", ""), ("h_lastoff", "h lastoff", 6, "hp", ""),
]
C = {k: get_column_letter(i + 1) for i, (k, *_) in enumerate(spec)}
kind = {k: kd for k, _, _, kd, _ in spec}
first_help = C["h_act"]; last_help = C["h_lastoff"]

ws["A1"] = "CTS BONUS LEAVE TRACKER"; ws["A1"].font = f_title
ws["A2"] = (f'="ACCRUING FOR "&UPPER(TEXT(AsOfDate,"mmmm yyyy"))&"  |  As at "&TEXT(AsOfDate,"dd/mm/yyyy")&"  |  Report from "&TEXT(PFYDate+1,"dd/mm/yyyy")&"  |  Staff: "&COUNTIF(${C["h_act"]}$5:${C["h_act"]}${R1},1)'
            f'&"  |  Accruing: "&COUNTIF(${C["elig"]}$5:${C["elig"]}${R1},"Yes*")&"  |  To accrue now: "&ROUND(SUMIF(${C["toacc"]}$5:${C["toacc"]}${R1},">0"),2)&""'
            f'&" hrs for "&COUNTIF(${C["accq"]}$5:${C["accq"]}${R1},"Yes")&" people  |  Earned this month: "&ROUND(SUM(${C["mdays"]}$5:${C["mdays"]}${R1}),2)&" days"')
ws["D1"] = "Accrual month = the month of the As-of date on Settings.  EARNED THIS MONTH = days whose milestone falls in that month.  HOURS TO ACCRUE = everything still missing in EH (this month + any earlier months not yet posted).  Post Hours to Accrue; the Action says how much is this month and how much is catch-up."
ws["D1"].font = f_note
ws["A2"].font = f_bold
for k, h, w, kd, note in spec:
    L = C[k]; ws.column_dimensions[L].width = w
    hdr(ws, f"{L}4", h, {"in": fill_hin, "out": fill_hout, "eh": fill_heh, "hp": fill_hp}[kd])
    c = ws[f"{L}3"]; c.value = note; c.font = f_note; c.alignment = Alignment(wrap_text=True, vertical="bottom")
ws.row_dimensions[3].height = 36; ws.row_dimensions[4].height = 42
ws.freeze_panes = "C5"
ws.column_dimensions.group(first_help, last_help, hidden=True, outline_level=1)

def Fn(p, s):
    return (f'IF({p}="FT",IF({s}<FT_First,0,IF({s}<FT_Y5,1+INT(({s}-FT_First)/FT_Int1),'
            f'FT_PreY5+FT_PerYear*INT(({s}-FT_Y5)/12)+FT_AnnivDays+INT(MOD({s}-FT_Y5,12)/FT_Int2))),'
            f'IF({p}="PT",IF({s}<PT_First,0,1+INT(({s}-PT_First)/PT_Int)),0))')
def months(S, D):
    return f'IF({D}<{S},-1,12*(YEAR({D})-YEAR({S}))+MONTH({D})-MONTH({S})-IF(DAY({D})<MIN(DAY({S}),DAY(EOMONTH({D},0))),1,0))'
def isFT(t): return f'OR(ISNUMBER(SEARCH("full",{t})),UPPER(TRIM({t}))="FT",UPPER(TRIM({t}))="FTE")'
def isPT(t): return f'OR(ISNUMBER(SEARCH("part",{t})),UPPER(TRIM({t}))="PT",UPPER(TRIM({t}))="PTE")'
def isCAS(t): return f'OR(ISNUMBER(SEARCH("cas",{t})),UPPER(TRIM({t}))="CAS")'
def norm(t):
    return f'IF({t}="","",IF({isFT(t)},"Full-Time",IF({isPT(t)},"Part-Time",IF({isCAS(t)},"Casual","Unknown"))))'
def pathway(t, h):
    return f'IF({isFT(t)},"FT",IF(AND({isPT(t)},OR({h}="",N({h})>=PT_MinHours)),"PT","None"))'
def eh(col, r):
    return f'SUMIFS(INDEX(LH_Data,0,LH_Col{col}),INDEX(LH_Data,0,LH_ColFirst),$A{r},INDEX(LH_Data,0,LH_ColSur),$B{r},INDEX(LH_Data,0,LH_ColCat),BonusCat)*UnitFactor'
def nextdate(d, asof):
    return f'DATE(YEAR({asof})+IF(DATE(YEAR({asof}),MONTH({d}),DAY({d}))<={asof},1,0),MONTH({d}),DAY({d}))'

def formulas(r):
    c = lambda k: f'{C[k]}{r}'
    v = lambda k: f'${C[k]}{r}'
    dm1 = lambda k: f'({C[k]}{r}-1)'
    A = f'{c("h_act")}=0'
    S = c("h_S"); dates = f'${C["cas"]}{r}:${C["fte"]}{r}'
    def typeof(dcell):
        return f'IF({dcell}="","",IF({dcell}=${C["pte"]}{r},"Part-Time",IF({dcell}=${C["fte"]}{r},"Full-Time","Casual")))'
    def days_at(M, B1, B2):
        return (f'{Fn(c("h_P1"), B1)}+IF(AND({c("h_n")}>=2,{c("h_D2")}<={{D}}),{Fn(c("h_P2"), B2)}-{Fn(c("h_P2"), B1)},0)'
                f'+IF(AND({c("h_n")}>=3,{c("h_D3")}<={{D}}),{Fn(c("h_P3"), M)}-{Fn(c("h_P3"), B2)},0)')
    f = {}
    f["h_act"] = f'=IF(OR(LEN($A{r}&$B{r})=0,${C["start"]}{r}=""),0,1)'
    f["h_S"] = f'=IF({A},"",IF(${C["anniv_in"]}{r}<>"",${C["anniv_in"]}{r},IF(OR(CasualCounts="Yes",COUNT(${C["cas"]}{r}:${C["fte"]}{r})=0,COUNT(${C["pte"]}{r}:${C["fte"]}{r})=0),${C["start"]}{r},MIN(${C["pte"]}{r}:${C["fte"]}{r}))))'
    f["h_n"] = f'=IF({A},"",COUNT({dates}))'
    for k, i in (("h_D1", 1), ("h_D2", 2), ("h_D3", 3)):
        f[k] = f'=IF({A},"",IFERROR(SMALL({dates},{i}),""))'
    for k, d in (("h_T1", "h_D1"), ("h_T2", "h_D2"), ("h_T3", "h_D3")):
        f[k] = f'=IF({A},"",{typeof(c(d))})'
    # pathway of the first type; if no dates entered at all, use the current EmploymentType
    f["h_P1"] = f'=IF({A},"",IF({c("h_n")}=0,{pathway(v("type"), v("pthrs"))},{pathway(c("h_T1"), v("pthrs"))}))'
    f["h_P2"] = f'=IF({A},"",IF({c("h_n")}<2,"",{pathway(c("h_T2"), v("pthrs"))}))'
    f["h_P3"] = f'=IF({A},"",IF({c("h_n")}<3,"",{pathway(c("h_T3"), v("pthrs"))}))'
    f["h_Pc"] = f'=IF({A},"",IF(AND({c("h_n")}>=3,{c("h_D3")}<=AsOfDate),{c("h_P3")},IF(AND({c("h_n")}>=2,{c("h_D2")}<=AsOfDate),{c("h_P2")},{c("h_P1")})))'
    f["h_M0"] = f'=IF({A},"",{months(S, "AsOfDate")})'
    f["h_B1"] = f'=IF({A},"",IF(OR({c("h_n")}<2,{c("h_D2")}>AsOfDate),{c("h_M0")},{months(S, dm1("h_D2"))}))'
    f["h_B2"] = f'=IF({A},"",IF(OR({c("h_n")}<3,{c("h_D3")}>AsOfDate),{c("h_M0")},{months(S, dm1("h_D3"))}))'
    f["h_MP"] = f'=IF({A},"",{months(S, "PFYDate")})'
    f["h_B1P"] = f'=IF({A},"",IF(OR({c("h_n")}<2,{c("h_D2")}>PFYDate),{c("h_MP")},{months(S, dm1("h_D2"))}))'
    f["h_B2P"] = f'=IF({A},"",IF(OR({c("h_n")}<3,{c("h_D3")}>PFYDate),{c("h_MP")},{months(S, dm1("h_D3"))}))'
    f["h_dn"] = f'=IF({A},"",' + days_at(c("h_M0"), c("h_B1"), c("h_B2")).replace("{D}", "AsOfDate") + ')'
    f["h_dp"] = f'=IF({A},"",' + days_at(c("h_MP"), c("h_B1P"), c("h_B2P")).replace("{D}", "PFYDate") + ')'
    M0 = c("h_M0"); Pc = c("h_Pc")
    f["h_no"] = (f'=IF({A},"",IF({Pc}="FT",IF({M0}<FT_First,FT_First,IF({M0}<FT_Y5,MIN(FT_Y5,FT_First+FT_Int1*(INT(({M0}-FT_First)/FT_Int1)+1)),'
                 f'FT_Y5+12*INT(({M0}-FT_Y5)/12)+FT_Int2*(INT(MOD({M0}-FT_Y5,12)/FT_Int2)+1))),'
                 f'IF({Pc}="PT",IF({M0}<PT_First,PT_First,PT_First+PT_Int*(INT(({M0}-PT_First)/PT_Int)+1)),"")))')
    Dm = 'EOMONTH(AsOfDate,-1)'
    f["h_MM"] = f'=IF({A},"",{months(S, Dm)})'
    f["h_B1M"] = f'=IF({A},"",IF(OR({c("h_n")}<2,{c("h_D2")}>{Dm}),{c("h_MM")},{months(S, dm1("h_D2"))}))'
    f["h_B2M"] = f'=IF({A},"",IF(OR({c("h_n")}<3,{c("h_D3")}>{Dm}),{c("h_MM")},{months(S, dm1("h_D3"))}))'
    f["h_dm"] = f'=IF({A},"",' + days_at(c("h_MM"), c("h_B1M"), c("h_B2M")).replace("{D}", Dm) + ')'
    f["h_lastoff"] = (f'=IF({A},"",IF({c("h_Pc")}="FT",IF({c("h_M0")}<FT_First,"",IF({c("h_M0")}<FT_Y5,FT_First+FT_Int1*INT(({c("h_M0")}-FT_First)/FT_Int1),FT_Y5+FT_Int2*INT(({c("h_M0")}-FT_Y5)/FT_Int2))),'
                      f'IF({c("h_Pc")}="PT",IF({c("h_M0")}<PT_First,"",PT_First+PT_Int*INT(({c("h_M0")}-PT_First)/PT_Int)),"")))')
    f["mdays"] = f'=IF({A},"",{c("days_due")}-{c("h_dm")})'
    f["mdate"] = f'=IF(OR({A},{c("mdays")}="",N({c("mdays")})=0,{c("h_lastoff")}=""),"",EDATE({S},{c("h_lastoff")}))'
    f["h_rows"] = f'=IF({A},"",IF(NOT(LH_Ready),0,COUNTIFS(INDEX(LH_Data,0,LH_ColFirst),$A{r},INDEX(LH_Data,0,LH_ColSur),$B{r},INDEX(LH_Data,0,LH_ColCat),BonusCat)))'
    f["h_last"] = f'=IF({A},"",IF(OR(NOT(LH_Ready),{c("h_rows")}=0),0,SUMPRODUCT(MAX((INDEX(LH_Data,0,LH_ColFirst)=$A{r})*(INDEX(LH_Data,0,LH_ColSur)=$B{r})*(INDEX(LH_Data,0,LH_ColCat)=BonusCat)*ROW(LH_Data)))))'
    # visible
    f["bday"] = f'=IF(OR({A},${C["dob"]}{r}=""),"",{nextdate(v("dob"), "AsOfDate")})'
    f["anniv"] = f'=IF({A},"",{nextdate(S, "AsOfDate")})'
    f["ten"] = f'=IF({A},"",ROUND((AsOfDate-{S})/365.25,1))'
    f["tenymd"] = f'=IF({A},"",IF(AsOfDate<{S},"-",DATEDIF({S},AsOfDate,"y")&"y "&DATEDIF({S},AsOfDate,"ym")&"m "&DATEDIF({S},AsOfDate,"md")&"d"))'
    f["days_due"] = f'=IF({A},"",{c("h_dn")})'
    f["hrs_due"] = f'=IF({A},"",{c("days_due")}*HoursPerDay)'
    def seg(tfrom, tto, d, p_from, p_to):
        return (f'{tfrom}&" to "&{tto}&" on "&TEXT({d},"dd/mm/yyyy")&IF({d}>AsOfDate," (future)",IF({p_from}={p_to},"",IF({p_to}="None"," - STOPPED accruing",IF({p_from}="None"," - commencement for Bonus Leave"," - pathway changed"))))')
    f["chg"] = (f'=IF({A},"",IF({c("h_n")}<2,"",{seg(c("h_T1"), c("h_T2"), c("h_D2"), c("h_P1"), c("h_P2"))}'
                f'&IF({c("h_n")}<3,"","; "&{seg(c("h_T2"), c("h_T3"), c("h_D3"), c("h_P2"), c("h_P3"))})))')
    f["elig"] = (f'=IF({A},"",IF({Pc}="None",IF({c("days_due")}>0,"No longer accruing","Not eligible - "&IF(AND({isPT(v("type"))},${C["pthrs"]}{r}<>"",N(${C["pthrs"]}{r})<PT_MinHours),"under "&PT_MinHours&" hrs",IF({norm(v("type"))}="Unknown","type not recognised",${C["type"]}{r}))),'
                 f'IF({c("days_due")}>0,"Yes - Accrual Required","Not yet - from "&TEXT({c("next")},"dd/mm/yyyy"))))')
    f["next"] = f'=IF({A},"",IF({Pc}="None","",EDATE({S},{c("h_no")})))'
    load = f'SUMIFS(INDEX(LH_Data,0,LH_ColAcc),INDEX(LH_Data,0,LH_ColFirst),$A{r},INDEX(LH_Data,0,LH_ColSur),$B{r},INDEX(LH_Data,0,LH_ColCat),BonusCat,INDEX(LH_Data,0,LH_ColPeriod),LH_OpenLabel)'
    since = f'INDEX(LH_Data,0,LH_ColStart),">="&{S}'
    f["acc"] = f'=IF({A},"",IF(NOT(LH_Ready),"",SUMIFS(INDEX(LH_Data,0,LH_ColAcc),INDEX(LH_Data,0,LH_ColFirst),$A{r},INDEX(LH_Data,0,LH_ColSur),$B{r},INDEX(LH_Data,0,LH_ColCat),BonusCat,{since})*UnitFactor))'
    f["taken"] = f'=IF({A},"",IF(NOT(LH_Ready),"",SUMIFS(INDEX(LH_Data,0,LH_ColTaken),INDEX(LH_Data,0,LH_ColFirst),$A{r},INDEX(LH_Data,0,LH_ColSur),$B{r},INDEX(LH_Data,0,LH_ColCat),BonusCat,{since})*UnitFactor))'
    f["close"] = f'=IF({A},"",IF(NOT(LH_Ready),"",IF({c("h_last")}=0,0,INDEX(INDEX(LH_Data,0,LH_ColClose),{c("h_last")}-1)*UnitFactor)))'
    f["open"] = f'=IF({A},"",IF(NOT(LH_Ready),"",IF({S}>PFYDate,0,IF(LH_ColPeriod=0,{c("close")}-{c("acc")}+{c("taken")},{load}*UnitFactor))))'
    f["due"] = f'=IF({A},"",({c("h_dn")}-{c("h_dp")})*HoursPerDay)'
    f["toacc"] = f'=IF({A},"",IF(NOT(LH_Ready),"",IF({c("h_rows")}=0,{c("hrs_due")},IF(PFYDate<${C["start"]}{r},{c("hrs_due")}-N({c("open")})-N({c("acc")}),{c("due")}-N({c("acc")})))))'
    f["accq"] = f'=IF(OR({A},{c("toacc")}=""),"",IF({c("toacc")}>0,"Yes","No"))'
    f["gap"] = f'=IF(OR({A},{c("open")}=""),"",{c("h_dp")}*HoursPerDay-{c("open")})'
    f["rows"] = f'=IF({A},"",{c("h_rows")})'
    X = c("toacc"); N_ = c("days_due"); Rw = c("h_rows"); Mh = c("hrs_due")
    f["action"] = (f'=IF({A},IF(LEN($A{r}&$B{r})=0,"","FIX INPUT - StartDate missing"),'
                   f'IF({c("check")}<>"","FIX INPUT - "&{c("check")},'
                   f'IF(NOT(LH_Ready),"Paste Leave History export (Settings must show YES)",'
                   f'IF({X}>0,"ACCRUE "&ROUND({X},2)&" hrs ("&ROUND({X}/HoursPerDay,2)&" days): "&IF(N({c("mdays")})>0,{c("mdays")}&" day(s) for "&UPPER(TEXT(AsOfDate,"mmm yyyy"))&" (milestone "&UPPER(TEXT({c("mdate")},"dd mmm yy"))&")"&IF({X}/HoursPerDay>{c("mdays")}+0.001," + "&ROUND({X}/HoursPerDay-{c("mdays")},2)&" day(s) catch-up from earlier months",""),"all catch-up from earlier months, nothing new this month")&IF({Rw}=0,"; no Bonus Leave in EH yet",""),'
                   f'IF({X}<0,"CHECK - EH accrued "&ROUND(-{X},2)&" hrs more than policy since commencement",'
                   f'IF(N({c("close")})>{N_}*HoursPerDay+0.01,"CHECK - EH balance "&ROUND({c("close")}-{N_}*HoursPerDay,2)&" hrs more than policy entitlement",'
                   f'IF({N_}>0,"OK - nothing to add"&IF(N({c("mdays")})>0," ("&{c("mdays")}&" day(s) for "&UPPER(TEXT(AsOfDate,"mmm yyyy"))&" already in EH)",""),'
                   f'IF({Pc}="None",{c("elig")},'
                   f'"Not yet - first day on "&TEXT({c("next")},"dd/mm/yyyy")))))))))')
    # data check
    latest_type = f'IF({c("h_n")}>=3,{c("h_T3")},IF({c("h_n")}=2,{c("h_T2")},{c("h_T1")}))'
    f["check"] = (f'=IF({A},"",TRIM('
                  f'IF(OR(ISTEXT(${C["start"]}{r}),ISTEXT(${C["pte"]}{r}),ISTEXT(${C["fte"]}{r}),ISTEXT(${C["cas"]}{r}),ISTEXT(${C["anniv_in"]}{r}),ISTEXT(${C["dob"]}{r})),"A date is typed as text - retype it as dd/mm/yyyy. ","")&'
                  f'IF({norm(v("type"))}="Unknown","EmploymentType not recognised - use Full-Time, Part-Time or Casual. ","")&'
                  f'IF(AND({c("h_n")}>=1,{latest_type}<>{norm(v("type"))}),"EmploymentType does not match latest start date ("&{latest_type}&"). ","")&'
                  f'IF(AND({c("h_n")}>=1,{c("h_D1")}<>${C["start"]}{r}),"Earliest type date should equal StartDate. ","")&'
                  f'IF(AND({c("h_n")}>=2,{c("h_D1")}={c("h_D2")}),"Two types start on the same date. ","")&'
                  f'IF(AND({isPT(v("type"))},${C["pthrs"]}{r}=""),"Part-time: enter PT Hrs/Week (blank = 24+). ","")))')
    return f

for r in range(R0, R1 + 1):
    fr = formulas(r)
    for k, _, _, kd, _ in spec:
        cell = ws[f"{C[k]}{r}"]; cell.border = border
        if kd == "in":
            cell.fill = fill_in; cell.font = f_input
        else:
            cell.fill = fill_calc; cell.font = f_norm; cell.value = fr[k]
    for k in ("dob", "bday", "start", "pte", "fte", "cas", "anniv_in", "anniv", "next", "mdate", "h_S", "h_D1", "h_D2", "h_D3"): ws[f"{C[k]}{r}"].number_format = DATE
    for k in ("hrs_due", "days_due", "pthrs", "open", "acc", "taken", "close", "due", "toacc", "gap", "mdays"): ws[f"{C[k]}{r}"].number_format = NUM
    ws[f"{C['ten']}{r}"].number_format = "0.0"

dv = DataValidation(type="list", formula1='"Full-Time,Part-Time,Casual"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{C['type']}{R0}:{C['type']}{R1}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="20000", allow_blank=True, error="Enter a date"); ws.add_data_validation(dvd)
for k in ("dob", "start", "pte", "fte", "cas", "anniv_in"): dvd.add(f"{C[k]}{R0}:{C[k]}{R1}")
ws.auto_filter.ref = f"A4:{C['check']}{R1}"
ac = C["action"]
ws.conditional_formatting.add(f"{ac}{R0}:{ac}{R1}", FormulaRule(formula=[f'LEFT({ac}{R0},6)="ACCRUE"'], fill=PatternFill("solid", fgColor="FCE4D6"), font=Font(name=F, size=10, bold=True, color="C00000")))
ws.conditional_formatting.add(f"{ac}{R0}:{ac}{R1}", FormulaRule(formula=[f'LEFT({ac}{R0},5)="CHECK"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, size=10, bold=True, color="9C0006")))
ws.conditional_formatting.add(f"{ac}{R0}:{ac}{R1}", FormulaRule(formula=[f'LEFT({ac}{R0},3)="FIX"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, size=10, bold=True, color="9C0006")))
ws.conditional_formatting.add(f"{ac}{R0}:{ac}{R1}", FormulaRule(formula=[f'LEFT({ac}{R0},2)="OK"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, size=10, color="006100")))
for k, test, color in (("accq", '="Yes"', "FCE4D6"), ("mdays", '>0', "FCE4D6"), ("elig", 'LEFT(X,3)="Yes"', "C6EFCE"), ("chg", '<>""', "FFEB9C"), ("check", '<>""', "FFC7CE")):
    L = C[k]
    formula = f'LEFT({L}{R0},3)="Yes"' if test.startswith("LEFT") else f'{L}{R0}{test}'
    ws.conditional_formatting.add(f"{L}{R0}:{L}{R1}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))

D = dt.date
# name, sur, type, dob, start, pte, fte, cas, anniv_in, pthrs
examples = [
 ("Jane", "Citizen", "Full-Time", D(1990, 4, 12), D(2021, 2, 15), None, D(2021, 2, 15), None, None, None),
 ("Tom", "Nguyen", "Part-Time", D(1985, 11, 30), D(2020, 6, 1), D(2020, 6, 1), None, None, None, 24),
 ("Sam", "Lee", "Casual", D(2000, 1, 20), D(2019, 3, 4), None, None, D(2019, 3, 4), None, None),
 ("Priya", "Patel", "Part-Time", D(1992, 7, 8), D(2022, 3, 1), D(2025, 7, 1), D(2022, 3, 1), None, None, 20),
 ("Liam", "Brown", "Full-Time", D(1988, 2, 29), D(2020, 1, 10), D(2020, 1, 10), D(2024, 1, 1), None, None, None),
 ("Olivia", "Smith", "Full-Time", D(1995, 9, 3), D(2022, 5, 5), None, D(2022, 5, 5), None, D(2022, 11, 2), None),
 ("Ava", "Jones", "Full-Time", D(2001, 12, 25), D(2024, 12, 1), None, D(2024, 12, 1), None, None, None),
 ("Chris", "Taylor", "Full-Time", D(1993, 6, 6), D(2019, 9, 2), None, D(2021, 3, 1), D(2019, 9, 2), None, None),
]
keys = ["name", "sur", "type", "dob", "start", "pte", "fte", "cas", "anniv_in", "pthrs"]
import os
for i, ex in (enumerate(examples) if not os.environ.get("CLEAN") else []):
    for k, v in zip(keys, ex):
        cell = ws[f"{C[k]}{R0 + i}"]; cell.value = v; cell.fill = fill_ex; cell.font = f_input; cell.border = border
        if isinstance(v, dt.date): cell.number_format = DATE

# ============================ LEAVE HISTORY ============================
eh_hdrs = ["Employee Id", "External Id", "First Name", "Surname", "Leave Category", "Opening Balance", "Closing Balance", "Pay Period", "Notes", "Leave Accrued", "Leave Taken", "Unit Type"]
periods = ["01/07/2026 - 14/07/2026", "15/07/2026 - 28/07/2026", "29/07/2026 - 11/08/2026", "12/08/2026 - 25/08/2026", "26/08/2026 - 08/09/2026", "09/09/2026 - 22/09/2026"]
def hist(eid, fn, sn, cat, opening, events):
    rows = []; bal = 0
    for i, p in enumerate(periods):
        acc, tk = events.get(p, (0, 0))
        if i == 0: acc += opening
        rows.append((eid, "", fn, sn, cat, bal, bal + acc - tk, p, "", acc, tk, "Hours")); bal = bal + acc - tk
    return rows
lh_rows = (hist(101, "Jane", "Citizen", "Bonus Leave", 64, {periods[3]: (0, 8)}) + hist(101, "Jane", "Citizen", "Annual Leave", 120, {periods[1]: (5.84, 0)})
           + hist(102, "Tom", "Nguyen", "Bonus Leave", 32, {}) + hist(104, "Priya", "Patel", "Bonus Leave", 8, {periods[2]: (0, 8)})
           + hist(105, "Liam", "Brown", "Bonus Leave", 72, {periods[4]: (8, 0)}) + hist(108, "Chris", "Taylor", "Bonus Leave", 16, {}))
e = wb.create_sheet("Leave History"); e.sheet_properties.tabColor = "C00000"
for j, h in enumerate(eh_hdrs, 1): e.cell(row=1, column=j, value=h).font = f_bold
for i, row in (enumerate(lh_rows, 2) if not os.environ.get("CLEAN") else []):
    for j, v in enumerate(row, 1): e.cell(row=i, column=j, value=v).font = f_norm
e["N1"] = ("SAMPLE - clear this sheet, then " if not os.environ.get("CLEAN") else "") + "Paste the EH Leave History Report at A1 (headers in row 1, paste as values). Run it from the first day of business to today."; e["N1"].font = Font(name=F, size=10, bold=True, color="C00000")
e["N2"] = "Rows must be in date order within each employee (EH's default). The latest row's Closing Balance is used as the balance now."; e["N2"].font = f_note
for j, w in enumerate([11, 10, 12, 12, 16, 12, 12, 24, 10, 12, 12, 10], 1): e.column_dimensions[get_column_letter(j)].width = w
e.freeze_panes = "A2"
name("LH_Hdr", "'Leave History'!$A$1:$AZ$1"); name("LH_Data", f"'Leave History'!$A$2:$AZ${EH_MAX}")
e["N1"] = "Period start text (auto)"; e["O1"] = "Period start date (auto - do not delete)"
for cc in ("N1", "O1"): e[cc].font = f_note; e[cc].fill = fill_calc
for rr in range(2, EH_MAX + 1):
    e[f"N{rr}"] = f'=IF(H{rr}="","",IFERROR(LEFT(H{rr},FIND(" ",H{rr})-1),H{rr}))'
    e[f"O{rr}"] = (f'=IFERROR(DATE(VALUE(MID(N{rr},FIND("/",N{rr},FIND("/",N{rr})+1)+1,4)),'
                   f'VALUE(MID(N{rr},FIND("/",N{rr})+1,FIND("/",N{rr},FIND("/",N{rr})+1)-FIND("/",N{rr})-1)),'
                   f'VALUE(LEFT(N{rr},FIND("/",N{rr})-1))),0)')
    e[f"O{rr}"].number_format = DATE
e.column_dimensions["N"].width = 14; e.column_dimensions["O"].width = 16
e["N2"].value = e["N2"].value
name("LH_ColStart", "15")
e["Q2"] = "Columns N and O are helpers that read the pay period start date. Paste your export into A:L only and never clear N:O."; e["Q2"].font = f_note

# ============================ SETTINGS ============================
st = wb.create_sheet("Settings"); st.sheet_properties.tabColor = "FFC000"
for col, w in zip("ABCDEFGH", [3, 44, 22, 10, 10, 3, 3, 3]): st.column_dimensions[col].width = w
st.column_dimensions["F"].width = 70
st["B2"] = "SETTINGS"; st["B2"].font = f_title
def setting(row, label, value, nm, note="", fmt=None, calc=False):
    st.cell(row=row, column=2, value=label).font = f_norm
    c = st.cell(row=row, column=3, value=value); c.font = f_norm if calc else f_input; c.fill = fill_calc if calc else fill_in; c.border = border
    if fmt: c.number_format = fmt
    if note: st.cell(row=row, column=6, value=note).font = f_note
    name(nm, f"Settings!$C${row}")
hdr(st, "B4", "General", fill_hout); hdr(st, "C4", "Value", fill_hout)
setting(5, "As-of date", "=TODAY()", "AsOfDate", "Type a fixed date for a month-end run.", DATE)
setting(6, "Day BEFORE the Leave History report starts", dt.date(2019, 6, 30), "PFYDate", "Run the report from the first day of business and enter the day before its start date here. Then Hours to Accrue = lifetime policy entitlement minus everything EH has ever credited.", DATE)
setting(7, "Hours per Bonus Leave day", 8, "HoursPerDay", "Policy p.3")
setting(8, "Bonus Leave category name in EH", "Bonus Leave", "BonusCat", "Must match the Leave Category text in the export exactly.")
setting(9, "Unit Type in the report", "Hours", "LHUnit", "Hours or Days (the Unit Type column of the export).")
st["G9"] = '=IF(LHUnit="Days",HoursPerDay,1)'; st["G9"].font = f_note; name("UnitFactor", "Settings!$G$9")

st["B10"] = "Pay period column (to find 'Opening Balances' rows)"; st["B10"].font = f_norm
c = st["C10"]; c.value = "Pay Period"; c.font = f_input; c.fill = fill_in; c.border = border
d = st["D10"]; d.value = '=IF(ISNUMBER(MATCH($C10,LH_Hdr,0)),"Yes","NO")'; d.font = f_bold; d.fill = fill_calc; d.border = border; d.alignment = center
st["G10"] = '=IFERROR(MATCH($C10,LH_Hdr,0),0)'; st["G10"].font = f_note
name("LH_ColPeriod", "Settings!$G$10")
st["B19"] = "Text in Pay Period that marks an opening-balance load"; st["B19"].font = f_norm
c = st["C19"]; c.value = "Opening Balances"; c.font = f_input; c.fill = fill_in; c.border = border
st["F19"] = "EH shows the balance loaded at go-live as a row with this Pay Period text. It is treated as the opening balance, not as an accrual."; st["F19"].font = f_note
name("LH_OpenLabel", "Settings!$C$19")
hdr(st, "B11", "Leave History export column headers", fill_hout); hdr(st, "C11", "Header text", fill_hout); hdr(st, "D11", "Found?", fill_hout)
hmap = [(12, "First name", "First Name", "First"), (13, "Surname", "Surname", "Sur"), (14, "Leave category", "Leave Category", "Cat"),
        (15, "Leave accrued", "Leave Accrued", "Acc"), (16, "Leave taken", "Leave Taken", "Taken"), (17, "Closing balance", "Closing Balance", "Close")]
for row, label, default, key in hmap:
    st.cell(row=row, column=2, value=label).font = f_norm
    c = st.cell(row=row, column=3, value=default); c.font = f_input; c.fill = fill_in; c.border = border
    d = st[f"D{row}"]; d.value = f'=IF(ISNUMBER(MATCH($C{row},LH_Hdr,0)),"Yes","NO")'; d.font = f_bold; d.fill = fill_calc; d.border = border; d.alignment = center
    st[f"G{row}"] = f'=IFERROR(MATCH($C{row},LH_Hdr,0),0)'; st[f"G{row}"].font = f_note
    name(f"LH_Col{key}", f"Settings!$G${row}")
st["F12"] = "Taken from your Leave History export. Change only if EH renames a column. All six must say Yes."; st["F12"].font = f_note
st["B18"] = "Export ready?"; st["B18"].font = f_bold
st["D18"] = '=IF(LH_Ready,"YES","NO")'; st["D18"].font = f_bold; st["D18"].fill = fill_calc; st["D18"].border = border; st["D18"].alignment = center
st["G18"] = "=AND(LH_ColFirst>0,LH_ColSur>0,LH_ColCat>0,LH_ColAcc>0,LH_ColTaken>0,LH_ColClose>0)"; st["G18"].font = f_note
name("LH_Ready", "Settings!$G$18")
st.conditional_formatting.add("D12:D18", FormulaRule(formula=['D12="NO"'], fill=PatternFill("solid", fgColor="FFC7CE")))
st.conditional_formatting.add("D12:D18", FormulaRule(formula=['OR(D12="Yes",D12="YES")'], fill=PatternFill("solid", fgColor="C6EFCE")))
dvu = DataValidation(type="list", formula1='"Hours,Days"', allow_blank=True); st.add_data_validation(dvu); dvu.add("C9")

hdr(st, "B20", "Policy rules (CTS Bonus Leave Policy 2026 v1.0)", fill_hout); hdr(st, "C20", "Value", fill_hout)
setting(21, "Full-time: first day at (months)", 36, "FT_First", "3 years - p.5")
setting(22, "Full-time: months between days, years 3-5", 4, "FT_Int1", "1 day every 4 months - p.5")
setting(23, "Full-time: 5-year milestone (months)", 60, "FT_Y5", "p.5")
setting(24, "Full-time: days on 5-yr and later anniversaries", 2, "FT_AnnivDays", "p.5")
setting(25, "Full-time: months between days after year 5", 3, "FT_Int2", "1 day every 3 months - p.5")
setting(26, "Part-time: first day at (months)", 60, "PT_First", "5 years - p.5-6")
setting(27, "Part-time: months between days", 4, "PT_Int", "1 day every 4 months - p.6")
setting(28, "Part-time: minimum hours per week", 24, "PT_MinHours", "Under 24 hrs = not eligible - p.4")
setting(29, "Full-time days before 5-yr milestone (calc)", "=1+INT((FT_Y5-1-FT_First)/FT_Int1)", "FT_PreY5", "", None, True)
setting(30, "Full-time days per year from year 6 (calc)", "=FT_AnnivDays+INT(11/FT_Int2)", "FT_PerYear", "", None, True)
setting(31, "Does casual service count towards qualifying service?", "No", "CasualCounts", "No = service counts from the first full-time or part-time start date (CTS decision; policy p.3 defines commencement as the original start date). Yes = service counts from StartDate.")
dvc = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); st.add_data_validation(dvc); dvc.add("C31")

notes = [
 "HOW TO USE",
 "1. Staff tab: type Name, Surname, EmploymentType (current type), DateOfBirth, StartDate.",
 "2. PTE Start / FTE Start / Casual Start: the date each employment type began. Always fill the one for the type they started on (same as StartDate). If they later changed type, fill the date the new type began as well. The latest date is the current type. Up to two changes are handled (e.g. Casual > Full-Time > Part-Time). Column Q writes the change out and says whether accrual stopped, started or changed pathway.",
 "3. AnniversaryDate: only if HR has pushed the qualifying-service start out for unpaid or parental leave (policy p.4-5). Otherwise leave blank and StartDate is used.",
 "4. Bonus leave Accrued = HOURS the policy says the person should have received in total to date. Days = the same in days. Eligible? and Next Accrual show who is in the scheme and when the next day lands.",
 "5. Leave History tab: run the EH Leave History Report from the first day of business to today and paste it at A1 (headers in row 1). On this tab set the day BEFORE the report start date. A full-history report is what makes the reconciliation clean: leave taken is then fully accounted for.",
 "6. The red columns on Staff then read straight from the report: Opening Balance (a loaded balance, normally 0), Leave Accrued (everything EH has credited), Leave Taken, Closing Balance (balance now), Due per policy (policy accrual from the report start to today), Hours to Accrue (= policy entitlement - opening balance - EH accrued), Add to EH now? and an Action.",
 "7. Filter the Action column: FIX INPUT = a problem with what was typed. ACCRUE = add those hours in EH now. CHECK = EH has credited more than the policy allows (an HR decision, not a formula problem). OK = nothing to add. Not eligible / Not yet = no accrual due and why. If a person has no Bonus Leave rows in EH at all, Hours to Accrue is their full policy entitlement.",
 "8b. Milestones tab: pick one person and see every milestone date, the type they were on that date and the days earned. All Staff Breakdown tab: the same for everyone, one line per person, so you can show how any number was reached.",
 "8. Opening gap vs policy only matters if you run a report that starts mid-way (e.g. from 1 July). It is then policy hours at the report start minus the opening balance, which mixes leave taken before the start with accruals never posted. With a full-history report it is 0.",
 "",
 "POLICY IN ONE LINE EACH",
 "Full-time: 1 day at 3 years, then 1 day every 4 months. At 5 years 2 days, then 1 day every 3 months (5 a year). Policy p.5, Schedule 1 p.9.",
 "Part-time 24+ hrs: nothing until 5 years, then 1 day every 4 months (3 a year). Policy p.5-6.",
 "Casual, contractor, part-time under 24 hrs: not eligible (p.3-4). No pro-rata (p.7). Status on the accrual date decides the pathway (p.7-8). 1 day = 8 hours (p.3). By CTS decision, casual service does not count: qualifying service runs from the first full-time or part-time start date (Settings switch 'Does casual service count'). The policy itself defines commencement as the original start date (p.3), so keep HR's written decision on file.",
 "",
 "ASSUMPTIONS: PT Hrs/Week blank means 24+ (eligible). Staff are matched to EH on First Name + Surname, so spelling must match the export. Rows in the report are assumed to be in date order within each employee. Terminated staff forfeit the balance (p.8) - delete them from the Staff tab and zero the EH balance.",
]
for i, t in enumerate(notes, 34):
    c = st.cell(row=i, column=2, value=t); c.font = f_bold if (t.isupper() or t.startswith("ASSUMPTIONS")) else f_norm
    st.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if len(t) > 120: st.row_dimensions[i].height = 30 if len(t) < 240 else 44

import sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from milestones import add_milestones, add_all_staff
add_milestones(wb, R0, R1); add_all_staff(wb, R0, R1)
wb._sheets = [wb[n] for n in ["Staff", "Leave History", "Milestones", "All Staff Breakdown", "Settings"]]
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "| columns:", {k: C[k] for k in ("cas", "anniv_in", "chg", "action", "check", "h_act", "h_last")})
