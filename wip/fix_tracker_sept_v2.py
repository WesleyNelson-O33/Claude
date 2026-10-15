"""Apply the "Fix on tracker" lines to Danica's own copy of the FY27 tracker (her 15/10 version,
which already has DAFF, the Aware amount and the ten pre-production lines as debits).
Same XML editing as fill_fy27_sept_wip.py so tables, comments and formulas survive."""
import re, sys, zipfile
import openpyxl
exec(open(__file__.replace("fix_tracker_sept_v2.py", "fill_fy27_sept_wip.py")).read().split("z = zipfile.ZipFile(SRC)")[0]
     .replace("SRC, WIPSCHED, OUT = sys.argv[1], sys.argv[2], sys.argv[3]", "SRC, OUT = sys.argv[1], sys.argv[2]"))

z = zipfile.ZipFile(SRC)
sh = {k: Sheet(z.read(v)) for k, v in PATHS.items()}
ro = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
wip = sh["WIP Movements"]
rows = list(ro["WIP Movements"].iter_rows(min_row=1, max_row=804, values_only=True))
def row(r): return rows[r - 1]
assert str(row(21)[1]).strip() == "2506208" and row(21)[9] == 6042.5
assert str(row(52)[1]).strip() == "2604704" and row(52)[9] == -4585
assert str(row(173)[1]).strip() == "26092101" and row(173)[9] == -950
assert str(row(186)[1]).strip() == "2609803" and row(187)[1] is None

# three existing rows
wip.set("J21", 6402.5)
wip.set("R21", "Based on labour hours. Sep 26: 49.25 hrs x $130 (WIP Schedule row 195).")
wip.set("H52", "Reversal of prior accrual")
wip.set("J173", -550)
wip.set("R173", "Separate journal; Auto-reverse. Labour 400 and equipment 150. Video labour 400 is on 26092101V below.")

R70539 = "Auto-reversed 01/09/2026 (Xero journal 70539)"
NEW = [
    ("26092101V", "VIDEO", "INV-10645", "Bankwest - BCEC WGEA Report Launch 13.10.2026 | Video Labour Revenue (INV-10645)", "Deferral - invoiced in advance", "Revenue", -400, "42150", "Separate journal; Auto-reverse"),
    ("26073107", "ONSITE", None, "DTTL | SkyBus for KK for DTTL HOB Tech & Coverage 6-9 Sep 26", "Release of deferred cost", "Cost", -40, "58110", R70539),
    ("26073107", "ONSITE", None, "DTTL | Return flights for KK ADL-HOB for DTTL HOB Tech & Coverage 6-9 Sep 26", "Release of deferred cost", "Cost", -531.82, "58110", R70539),
    ("26073107", "ONSITE", None, "DTTL | Accommodation for KK for DTTL HOB Tech & Coverage 6-9 Sep 26", "Release of deferred cost", "Cost", -1589.98, "58110", R70539),
    ("26071701", "PRODUCTION", None, "ASX Live - 9.9.26 [PRD] Goget Travel Expense", "Release of deferred cost", "Cost", -69.75, "58115", R70539),
    ("26061602", "PRODUCTION", None, "ASX CEO Connect (Extended Results) 1.9.26 [PRD] Pre-prod labour", "Reversal of prior accrual", "Revenue", -1852.5, "42100", R70539),
    ("26071701", "PRODUCTION", None, "ASX Live - 9.9.26 [PRD] Pre-prod labour", "Reversal of prior accrual", "Revenue", -1885, "42100", R70539),
    ("26072102", "PRODUCTION", None, "Ventia - RUOK? Day - 10.09.2026 [PRD] pre-prd labour", "Reversal of prior accrual", "Revenue", -585, "42100", R70539),
    ("26082502", "PRODUCTION", None, "ICC - Presenter & Speaker Preparation - September 2026 [PRD] pre-prod labour", "Reversal of prior accrual", "Revenue", -292.5, "42100", R70539),
    ("26032302", "PRODUCTION", None, "RSNSW - OGM - 2.9.2026 [PRD] pre-prod labour", "Reversal of prior accrual", "Revenue", -260, "42100", R70539),
    ("26061602V", "VIDEO", None, "ASX CEO Connect (Extended Results) 1.9.26 [VID] pre-vid labour", "Reversal of prior accrual", "Revenue", -190, "42150", R70539),
    ("2591012C", "PRODUCTION", None, "St Luke Equipment Sale | DL to advise to invoice in SEPT26 | Indicative Revenue", "Release of deferral", "Revenue", 4227.27, "42600", R70539),
    ("26082502", "PRODUCTION", "INV-10674", "ICC - September Presenter Support Presentations - ComBio - 1/10/2026 | Labour Revenue (INV-10674)", "Deferral - invoiced in advance", "Revenue", -762.5, "42100", "Separate journal; Auto-reverse"),
    ("2506208", "PRODUCTION", None, "CBA AGM 2026 - Goget hire | Travel Expense (INV10498230926)", "Deferred cost - paid in advance", "Cost", 72, "58115", "Separate journal; Auto-reverse"),
    ("26062901", "INTEGRATION", None, "Aware | SYD 28.06 Services [INT] Labour Revenue Pre work", "Deferral - invoiced in advance", "Revenue", -148, "42800",
     "Based on labour hours; Reverse when invoiced. Same treatment as August (row 162)."),
]
for i, (job, cc, inv, desc, typ, rc, amt, gl, note) in enumerate(NEW):
    r = 187 + i
    for c, v in (("A", SEP), ("B", job), ("E", cc), ("F", inv), ("G", desc), ("H", typ), ("I", rc), ("J", amt), ("K", gl), ("R", note)):
        wip.set("%s%d" % (c, r), v)
    wip.cell("A%d" % r).set("s", wip.cell("A186").get("s"))

# Deferrals: Aware Appspace on Xero's 12-month schedule
d = sh["Deferrals"]
rv = list(ro["Deferrals"].iter_rows(min_row=57, max_row=58, values_only=True))
assert rv[0][1] == "INV-10635" and rv[1][1] == "WIP-474"
d.set("J57", None)
d.set("G57", 19430.92)
d.set("K57", rv[0][10] + ". Xero journals 71180/71192: 12 months Aug 26 - Jul 27 at 1,766.45, Aug and Sep both released in Sep. "
      "Amount here is the invoice less August's month, so Sep releases two months as Xero did.")
d.set("G58", 17487.85)
d.set("J58", 1589.81)
d.set("K58", rv[1][10] + ". Xero journals 71180/71192: 12 months Aug 26 - Jul 27 at 1,589.81, Aug and Sep both released in Sep. "
      "Amount here is the bill less August's month, so Sep releases two months as Xero did.")

# Lists: job 2609803 is in Xero's job list but not the tracker's
lists = [x[0] for x in ro["Lists"].iter_rows(min_row=5, max_row=1500, min_col=16, max_col=16, values_only=True)]
if "2609803" not in [str(j).strip() for j in lists if j is not None]:
    nxt = 5 + max(i for i, j in enumerate(lists) if j is not None) + 1
    sh["Lists"].set("P%d" % nxt, "2609803")
    sh["Lists"].set("Q%d" % nxt, "Automic - Noumi LTD AGM - 5.11.26")
    sh["Lists"].set("R%d" % nxt, "Production")

wbxml = z.read("xl/workbook.xml").decode()
if "fullCalcOnLoad" not in wbxml:
    wbxml = re.sub(r"<calcPr([^>]*)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbxml)
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as o:
    for item in z.infolist():
        name = item.filename
        if name in PATHS.values():
            data = sh[[k for k, v in PATHS.items() if v == name][0]].bytes()
        elif name == "xl/workbook.xml":
            data = wbxml.encode()
        else:
            data = z.read(name)
        o.writestr(item, data)

# check: September closing from the constants
sep = sum(r[9] for r in rows[4:] if hasattr(r[0], "month") and r[0].year == 2026 and r[0].month == 9 and isinstance(r[9], (int, float)))
sep += (6402.5 - 6042.5) + 400 + sum(x[6] for x in NEW)
deferral_sep = 2521.70 + (-17664.47 + 19270.34) + (15898.04 - 17343.33)
print("written", OUT, "closing", round(-128544.93 + sep + deferral_sep, 2))
