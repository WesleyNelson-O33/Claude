"""Schedule of the changes made to the FY27 Revenue Tracker by fill_fy27_sept_wip.py."""
import datetime, sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

SRC, NEW, WIPS, OUT = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
from openpyxl.utils import column_index_from_string as CI
_ws = openpyxl.load_workbook(WIPS, data_only=True)["Schedule Updated"]
SCHED = []
for _r in range(6, _ws.max_row + 1):
    _v = _ws.cell(_r, CI("BY")).value
    if isinstance(_v, (int, float)) and abs(_v) > 0.004:
        SCHED.append((_r, str(_ws.cell(_r, 1).value).split("\n")[0].strip(), round(_v, 2)))
_used = set()
def sched_row(job, amt):
    for r, j, v in SCHED:
        if r not in _used and j == job and abs(v - amt) < 0.006:
            _used.add(r); return r
    return None
a = openpyxl.load_workbook(SRC, data_only=True)
b = openpyxl.load_workbook(NEW, data_only=True)
wa, wb_ = a["WIP Movements"], b["WIP Movements"]

rows = []  # tab, cell/row, job, description, what changed, before, after, effect on 11300, source
def wiprow(r, what, before, after, effect, src):
    rows.append(("WIP Movements", "Row %d" % r, str(wb_["B%d" % r].value), wb_["G%d" % r].value, what, before, after, effect, src))

wiprow(11, "Amount", wa["J11"].value, wb_["J11"].value, wb_["J11"].value - wa["J11"].value,
       "WIP Schedule row 174 and Xero journal 71188: 15.25 hrs x $165")
wiprow(21, "Amount", wa["J21"].value, wb_["J21"].value, wb_["J21"].value - wa["J21"].value,
       "WIP Schedule row 195 and Xero journal 71188: 49.25 hrs x $130")
wiprow(52, "Type and amount", "%s %s" % (wa["H52"].value, wa["J52"].value), "%s %s" % (wb_["H52"].value, wb_["J52"].value),
       wb_["J52"].value - wa["J52"].value, "WIP Schedule row 377 and Xero journal 71188: invoiced INV-10596, accrued labour reversed")
wiprow(173, "Amount", wa["J173"].value, wb_["J173"].value, wb_["J173"].value - wa["J173"].value,
       "Xero journal 71186: video labour 400 moved to 26092101V (new row 177)")
r = 177
while wb_["B%d" % r].value:
    note = wb_["R%d" % r].value or ""
    if "70539" in note: src = "Xero journal 70539 (01/09/2026 reversal of August's auto-reversing lines)"
    elif "71188" in note: src = "Xero journal 71188 (WIP Journal September 2026)"
    elif str(wb_["B%d" % r].value) == "26062901": src = "Xero journal 71188 (WIP Journal September 2026)"
    else: src = "Xero journal 71186 (Pre-PROD & INT accrual for October 2026)"
    rows.append(("WIP Movements", "Row %d (new)" % r, str(wb_["B%d" % r].value), wb_["G%d" % r].value,
                 "New line: " + str(wb_["H%d" % r].value), None, wb_["J%d" % r].value, wb_["J%d" % r].value, src))
    r += 1
da, db = a["Deferrals"], b["Deferrals"]
rows.append(("Deferrals", "Row 57", "26091401", "Aware Super Appspace (Revenue) INV-10635", "Amount Override",
             da["G57"].value, db["G57"].value, 1605.87,
             "Xero journals 71180 and 71192: 12 months Aug 26 - Jul 27, Aug and Sep released in Sep"))
rows.append(("Deferrals", "Row 58", "26091401", "Aware Super Appspace (Cost) WIP-474",
             "Amount Override and Per Month Override", "%s / %s" % (da["G58"].value, da["J58"].value),
             "%s / %s" % (db["G58"].value, db["J58"].value), -1445.29,
             "Xero journals 71180 and 71192: 12 months Aug 26 - Jul 27, Aug and Sep released in Sep"))
rows.append(("Lists", "Row 466 (new)", "2609803", "Automic - Noumi LTD AGM - 5.11.26", "Job added to job list",
             None, "Production", 0, "Job is in Xero's Job Numbers list but was missing from the tracker"))
rows.append(("WIP Summary", "Column G", "51 jobs", "Per WIP Schedule (paste in)", "WIP Schedule closing balance per job pasted",
             None, "51 jobs", 0, "WIP Schedule, Schedule Updated tab, September balance column BZ"))
for cell, label in (("B23", "Closing balance per the WIP Schedule file"), ("B25", "Closing balance per Xero GL 11300")):
    rows.append(("Month-End", cell, "", label, "Typed in", a["Month-End"][cell].value, b["Month-End"][cell].value, 0,
                 "WIP Schedule total and Xero Account Transactions 11300 at 30/09/2026"))

out = openpyxl.Workbook()
ws = out.active
ws.title = "Changes"
ws.append(["Revenue Tracker changes - September 2026 WIP"])
ws.append(["Source file: FY27_Revenue_Tracker__29092026.xlsx. Updated file: FY27_Revenue_Tracker_30092026_WIP.xlsx."])
ws.append([])
ws.append(["Tracker WIP closing before changes", -136622.23])
ws.append(["Total effect of changes", None])
ws.append(["Tracker WIP closing after changes", None])
ws.append(["Xero 11300 at 30/09/2026", -146188.18])
ws.append(["Difference", None])
ws.append([])
hdr = ["#", "Tab", "Where", "Job", "Description", "What changed", "Before", "After", "Effect on 11300", "Source", "WIP Schedule row"]
ws.append(hdr)
h = ws.max_row
FIXED = {"Row 11": "174", "Row 21": "195", "Row 52": "377", "Row 173": "473, 475", "Row 57": "471", "Row 58": "472",
         "Column G": "all rows, column BZ", "B23": "column BZ total"}
for i, x in enumerate(rows, 1):
    ref = FIXED.get(x[1])
    if ref is None and x[1].endswith("(new)") and x[0] == "WIP Movements":
        ref = sched_row(x[2], x[6]); ref = str(ref) if ref else ""
    ws.append([i] + list(x) + [ref or ""])
last = ws.max_row
ws["B5"] = "=SUM(I%d:I%d)" % (h + 1, last)
ws["B6"] = "=B4+B5"
ws["B8"] = "=ROUND(B6-B7,2)"
ws["A1"].font = Font(bold=True, size=13)
for c in ws[h]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
for rr in range(4, 9):
    ws["A%d" % rr].font = Font(bold=True); ws["B%d" % rr].number_format = "#,##0.00"
for rr in range(h + 1, last + 1):
    for col in "GHI":
        ws["%s%d" % (col, rr)].number_format = "#,##0.00"
    for col in "EJ":
        ws["%s%d" % (col, rr)].alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEFGHIJK", (5, 15, 14, 12, 50, 30, 16, 16, 15, 60, 16)):
    ws.column_dimensions[col].width = w
ws.column_dimensions["A"].width = 36
ws.freeze_panes = "A%d" % (h + 1)
out.save(OUT)
print(len(rows), "changes; effect", round(sum(x[7] for x in rows), 2))
