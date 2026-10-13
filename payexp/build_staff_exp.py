"""Staff expense reimbursement journals for Xero, one per pay run.

Copies the layout of August's journal 69541 (EH Expense Reimbursement):
expense lines tracked to the claim's job, GST on the expense, and one credit
to 21410 Payroll Accruals Payable tracked to CTS / 9000.
"""
import csv, json, sys
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

UP = "/root/.claude/uploads/a91f56e6-b7da-5dad-8e0b-6d5f516f9360/"
# (pay run audit, EH pay run number, journal date) as Danica set them
RUNS = [
    ("07fe6af9-PayRunAudit-CTS_-_Fortnight-20260904.xlsx", "37005", "30/09/2026"),
    ("f91e975d-PayRunAudit-CTS_-_Fortnight-20260918.xlsx", "37006", "30/09/2026"),
    ("c9578b8f-PayRunAudit-CTS_-_Fortnight-20261002.xlsx", "37007", "02/10/2026"),
]
TRACKING = sys.argv[1]  # saved Xero tracking categories JSON
HEAD = ["*Narration", "*Date", "Description", "*AccountCode", "*TaxRate", "*Amount",
        "TrackingName1", "TrackingOption1", "TrackingName2", "TrackingOption2"]

# Expense category -> (account by department tag, tax rate)
ACCT = {
    "Direct Travel (chargeable)": ({"[PRD]": "58115", "[ONS]": "58110"}, "GST on Expenses"),
    "Direct Freight (chargeable)": ({"[PRD]": "58215", "[ONS]": "58210"}, "GST on Expenses"),
    "Taxis/Parking (non-chargeable)": ("61903", "GST on Expenses"),
}
CC = {"[CTS]": "CTS", "[PRD]": "PRODUCTION", "[ONS]": "ONSITE", "[VID]": "VIDEO",
      "[INT]": "INTEGRATION", "[CONS]": "CONSULTING"}
# 37007, as Danica set it: no GST on any line, and the ASX confetti claims
# left out because they go in as bills paid from 21410.
NO_GST = {"37007"}
SKIP = {"37007": {"ASX - Lighting and Confetti Test - 29.9.26 [PRD]"}}
D = lambda x: Decimal(str(x)).quantize(Decimal("0.01"), ROUND_HALF_UP)

jobs = {}
for c in json.load(open(TRACKING))["tracking_categories"]:
    if c["name"] == "Job Numbers":
        for o in c["options"]:
            jobs.setdefault(o["name"].split(" - ", 1)[1], o["name"])

for f, run_no, DATE in RUNS:
    wb = openpyxl.load_workbook(UP + f, read_only=True)
    s = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)}
    a, b = s["Report Period"].split(" - ")
    narr = "EH Expense Reimbursement %s (%s-%s)" % (run_no, a.replace("/", "."), b.replace("/", "."))
    rows, total = [], Decimal("0")
    for r in list(wb["Expenses"].iter_rows(values_only=True))[1:]:
        if r[0] == "Total Expenses":
            assert run_no in SKIP or D(r[7]) == total, (f, r[7], total)
            continue
        cat, loc, amt = r[4], r[5], D(r[7])
        if loc in SKIP.get(run_no, ()):
            continue
        tag = loc[loc.rindex("["):]
        accts, tax = ACCT[cat]
        if run_no in NO_GST:
            tax = "BAS Excluded"
        acct = accts if isinstance(accts, str) else accts[tag]
        # GST-inclusive, as claimed: August's 60.00 split 54.55 + 5.45 is the
        # tax-inclusive split, so the draft is set to Tax Inclusive after import.
        rows.append([narr, DATE, loc, acct, tax, str(amt),
                     "Cost Centres", CC[tag], "Job Numbers", jobs[loc]])
        total += amt
    rows.append([narr, DATE, narr, "21410", "BAS Excluded", str(-total),
                 "Cost Centres", "CTS", "Job Numbers", "9000 - OFFICE / ADMIN [CTS]"])
    out = "/home/user/Claude/payexp/out/Staff Expenses %s-%s Xero Journal.csv" % (
        a.replace("/", "."), b.replace("/", "."))
    with open(out, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(HEAD)
        w.writerows(rows)
    print(out, len(rows), total)
