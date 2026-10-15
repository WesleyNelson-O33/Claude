"""COGS and Expenses read as positive figures on the statements.

The numbers underneath do not change. A cost is still negative, so gross
profit, EBITDA, EBIT and net profit all keep calculating exactly as they do
now. Only what the cell shows changes.

A normal cost prints as a plain positive figure. A cost line that has gone
the other way, a credit or a reversal, prints red in brackets, because that
is the one that needs looking at.

Variance columns are left alone. A variance has to keep its sign or it stops
meaning anything.
"""
import os, shutil
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = ('/tmp/claude-0/-home-user-Claude/a00ffe59-433c-507b-9006-917d8c60083a/'
       'scratchpad/v83/her9.xlsx')
OUT = f'{HERE}/CTS Financial Controller Pack - v10.xlsx'

R0, R1 = 59, 172          # Cost of Sales down to the last Other Expenses line

# The columns that hold a level, not a variance. Only these flip.
LEVELS = {
    'Summary':     ['B', 'C', 'D', 'E', 'G', 'K', 'L'],
    'P&L FY27':    [chr(c) for c in range(ord('B'), ord('M') + 1)]
                   + ['N', 'O', 'Q', 'U', 'V'],
    'Budget FY27': [chr(c) for c in range(ord('B'), ord('M') + 1)] + ['N', 'O'],
    'Actual FY26': [chr(c) for c in range(ord('B'), ord('M') + 1)] + ['N', 'O'],
    'Budget FY26': [chr(c) for c in range(ord('B'), ord('M') + 1)] + ['N', 'O'],
}

COST_0 = '[Red]"($"#,##0\\);\\$#,##0;\\-'
COST_2 = '[Red]"($"#,##0.00\\);\\$#,##0.00;\\-'

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

flipped = 0
for name, cols in LEVELS.items():
    ws = wb[name]
    # Prove the block is where it should be before touching a single cell.
    assert str(ws.cell(R0, 1).value).strip() == 'Cost of Sales', (name, ws.cell(R0, 1).value)
    assert str(ws.cell(164, 1).value).strip() == 'Expenses Overheds', name
    assert str(ws.cell(171, 1).value).strip() == 'Other Expenses', name
    assert str(ws.cell(173, 1).value).strip() == 'Gross profit', name
    for r in range(R0, R1 + 1):
        for col in cols:
            c = ws[f'{col}{r}']
            fmt = (c.number_format or '').upper()
            if '%' in fmt or 'MMM' in fmt:
                continue
            if '.00' in fmt:
                c.number_format = COST_2
            elif '#,##0' in fmt:
                c.number_format = COST_0
            else:
                continue
            flipped += 1

print('cost cells now showing positive:', flipped)
wb.save(OUT)
print('saved', OUT)
