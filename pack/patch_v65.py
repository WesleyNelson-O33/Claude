"""Three stronger cuts of the same three years of revenue.

Three bars for one month, and three bars for one quarter, are tables drawn as
charts. They carry three numbers each. These three blocks carry the same data in
forms that say something:

  cumulative year to date, which answers whether the year is ahead and by how
    much at every point in it, not just at the reporting month
  all four quarters across the three years, which shows what is still to come
    rather than only the quarter that has closed
  the rolling twelve months, which strips the seasonality out and leaves the
    trend, and for a business with an October and a January like this one that
    is the only honest trend line

Everything reads the Xero History table above it, so nothing is a second copy.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE = 'FF242C39', 'FFFFFFFF'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
BODY = Font(name='Calibri', sz=11, color=NAVY)

FY25, FY26, FY27 = 10, 22, 34        # first row of each year in the table above
YEARS = [('FY25', FY25), ('FY26', FY26), ('FY27', FY27)]

CUM_BAND, CUM_HEAD, CUM_FIRST = 56, 57, 58          # cumulative, twelve months
Q_BAND, Q_HEAD, Q_FIRST = 72, 73, 74                # four quarters, three years
R_BAND, R_HEAD, R_FIRST = 80, 81, 82                # rolling twelve months

wb = openpyxl.load_workbook(WB)
ws = wb['Xero History']
assert ws['A9'].value == 'Month'

def band(row, cols, text):
    ws.cell(row, 1, text)
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = BAND
        cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

def header(row, labels):
    for c, text in enumerate(labels, start=1):
        cell = ws.cell(row, c, text)
        cell.fill = BAND
        cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
        cell.border = Border(bottom=RULE)
        cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

def money(row, col, formula, bold=False):
    cell = ws.cell(row, col, formula)
    cell.number_format = MONEY
    cell.font = Font(name='Calibri', sz=11, b=bold, color=NAVY)

# ---- cumulative revenue year to date, one line per year --------------------
band(CUM_BAND, 4, 'REVENUE YEAR TO DATE, CUMULATIVE')
header(CUM_HEAD, ['Month', 'FY25', 'FY26', 'FY27'])
for i in range(12):
    r = CUM_FIRST + i
    ws.cell(r, 1, f'=LEFT(A{FY27 + i},3)').font = BODY      # Jul, Aug, Sep ...
    for j, (_, base) in enumerate(YEARS):
        last = base + i
        money(r, 2 + j, f'=IF(C{last}="","",SUM($C${base}:$C${last}))')

# ---- all four quarters, three years ----------------------------------------
band(Q_BAND, 4, 'REVENUE BY QUARTER, YEAR ON YEAR')
header(Q_HEAD, ['Quarter', 'FY25', 'FY26', 'FY27'])
for q in range(4):
    r = Q_FIRST + q
    ws.cell(r, 1, f'Q{q + 1}').font = Font(name='Calibri', sz=11, b=True, color=NAVY)
    for j, (_, base) in enumerate(YEARS):
        a, b = base + q * 3, base + q * 3 + 2
        money(r, 2 + j, f'=IF(COUNT(C{a}:C{b})<3,"",SUM(C{a}:C{b}))', bold=True)

# ---- the rolling twelve months ---------------------------------------------
band(R_BAND, 4, 'REVENUE, ROLLING TWELVE MONTHS')
header(R_HEAD, ['Twelve months to', 'Revenue'])
for k in range(12, 37):                       # Jun 2025 is the first full twelve
    r, src = R_FIRST + k - 12, 9 + k
    ws.cell(r, 1, f'=A{src}').font = BODY
    money(r, 2, f'=IF(C{src}="","",SUM(C{src - 11}:C{src}))')

wb.save(WB)
print(f'Xero History rows {CUM_BAND} to {R_FIRST + 24} added: cumulative year to '
      f'date, four quarters year on year, and the rolling twelve months')
