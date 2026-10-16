"""Take the staff split back off the Summary.

It belongs on P&L by Dept only. The input block on Overheads MoM stays,
because that is where the office staff figures are typed and P&L by Dept
reads them from there.
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Border

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
wb = openpyxl.load_workbook(OUT)
s = wb['Summary']
n = 0
for r in range(194, 199):
    for c in range(1, 7):
        cell = s.cell(r, c)
        if cell.value is not None or cell.has_style:
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()
            cell.border = Border()
            n += 1
print(f'Summary: split block cleared, {n} cells')
wb.save(OUT)
print('saved')
