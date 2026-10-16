"""P&L FY27: a month block in columns U to X, and everything from PY moves right.

She asked for the month beside the year to date, so the sheet now reads:

    N   FY total
    O   YTD
    Q R S      YTD against budget, variance, variance %
    U V W X    the month, the month's budget, variance, variance %   <- new
    Z AA AB AC PY, PY budget, variance, variance %                   (was U to X)
    AE..AH     YOY, YOY budget, variance, variance %                 (was Z to AC)
    AI         Risk                                                  (was AD)

Five columns go in, four of them used and a blank at Y, which is the same
spacing every other group on this sheet already has. Nothing outside the sheet
pointed at those columns, so the shift is contained.

The month column is the YTD column with its "<=" taken off, so a money line
picks up the one month and a margin line recomputes its ratio, exactly as the
year to date block does.
"""
import os, re, copy
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as gl, column_index_from_string as ci

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
SHIFT = 5
FIRST = ci('U')                       # the first column that moves

wb = openpyxl.load_workbook(OUT)
s = wb['P&L FY27']
LAST = s.max_column

# where each moving column ends up
MAP = {gl(c): gl(c + SHIFT) for c in range(FIRST, LAST + 1)}
PAT = re.compile(r'\b(' + '|'.join(sorted(MAP, key=len, reverse=True)) + r')(\$?\d+)\b')
move = lambda f: PAT.sub(lambda m: MAP[m.group(1)] + m.group(2), f)

# 1. lift the old block out, then put it back five columns along
grab = {}
for c in range(FIRST, LAST + 1):
    for r in range(1, s.max_row + 1):
        cell = s.cell(r, c)
        if cell.value is not None or cell.has_style:
            grab[(r, c)] = (cell.value, copy.copy(cell._style), cell.number_format)
for (r, c) in grab:
    s.cell(r, c).value = None
n = 0
for (r, c), (v, st, nf) in grab.items():
    t = s.cell(r, c + SHIFT)
    t.value = move(v) if isinstance(v, str) and v.startswith('=') else v
    t._style = st
    t.number_format = nf
    n += 1
print(f'columns U..{gl(LAST)} moved to {gl(FIRST+SHIFT)}..{gl(LAST+SHIFT)}, {n} cells')

# 2. the new month block, built off the year to date block beside it
YTD, BUD, VAR, VPC = ci('O'), ci('Q'), ci('R'), ci('S')
NEW = {'U': YTD, 'V': BUD, 'W': VAR, 'X': VPC}
for lab, src in (('U', 'Month'), ('V', 'Month budget'), ('W', 'Var $'), ('X', 'Var %')):
    c = s.cell(5, ci(lab), src)
    c._style = copy.copy(s.cell(5, NEW[lab])._style)
    c.alignment = Alignment(horizontal='right')

ONE = 'Setup!$B$5'
built = 0
for r in range(6, s.max_row + 1):
    o = s.cell(r, YTD).value
    if not isinstance(o, str) or not o.startswith('='):
        continue
    q = s.cell(r, BUD).value
    # a money line sums one month instead of everything to date; a margin line
    # recomputes its own ratio, so it only needs its column letter swapped
    u = o.replace('"<="&' + ONE, ONE).replace('O', 'U') if '"<="' not in o \
        else o.replace('"<="&' + ONE, ONE)
    u = o.replace('"<="&' + ONE, ONE) if '"<="' in o else re.sub(r'\bO(\$?\d+)', r'U\1', o)
    v = q.replace('"<="&' + ONE, ONE) if isinstance(q, str) and '"<="' in q \
        else (re.sub(r'\bQ(\$?\d+)', r'V\1', q) if isinstance(q, str) else None)
    s.cell(r, ci('U'), u)
    if v:
        s.cell(r, ci('V'), v)
    s.cell(r, ci("W"), f"=IFERROR(U{r}-V{r},\"\")")
    s.cell(r, ci('X'), f'=IFERROR(W{r}/ABS(V{r}),0)')
    for lab in ('U', 'V', 'W', 'X'):
        t = s.cell(r, ci(lab))
        t._style = copy.copy(s.cell(r, NEW[lab])._style)
        t.number_format = s.cell(r, NEW[lab]).number_format
    built += 1
print('month block written on', built, 'rows')

for lab in ('U', 'V', 'W', 'X'):
    s.column_dimensions[lab].width = s.column_dimensions[gl(NEW[lab])].width or 13
for c in range(FIRST + SHIFT, LAST + SHIFT + 1):
    s.column_dimensions[gl(c)].width = s.column_dimensions[gl(c - SHIFT)].width or 13

wb.save(OUT)
print('saved')
