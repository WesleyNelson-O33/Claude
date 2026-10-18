"""
Chargeable per cent for the month, on one line per group.

The deck had four slides on this: the three month history split over two, and
the hours and FTE detail split over another two. Danica asked for one slide.

Thirty six rows cannot be read on a slide, so this is the month on Setup laid
out across instead of down. One row per person group, plus the company, with
the billable, non-billable and leave hours, the hours available, the chargeable
per cent, the target and the over or under reading.

It reads off the three month block above it with INDEX and MATCH on the month
row, so it follows the reporting month on Setup and nothing has to be retyped.
The history stays on the tab for anyone who wants it.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)
HRS, PCT = '#,##0.00;;', '0.0%'
SIGNED = '0.0%;[RED]-0.0%;\\-'

HEAD = ['Person Group', 'Billable', 'Non-billable', 'Leave', 'Available',
        'Chargeable %', 'Target', 'Over / under']
#  column -> (row offset in the source block, number format)
COLS = [(2, 0, HRS), (3, 1, HRS), (4, 2, HRS), (5, 4, HRS),
        (6, 5, PCT), (7, 6, PCT), (8, 7, SIGNED)]
GROUPS = [('Production', 123), ('Consulting', 134), ('Support', 145),
          ('Company', 156)]

wb = openpyxl.load_workbook(WB)
ws = wb['Utilisation']
assert ws['A120'].value.startswith('UTILISATION, AS THE MONTHLY REPORT'), ws['A120'].value

ws['A170'] = 'CHARGEABLE PER CENT, THE MONTH ON SETUP'
for c in range(1, 9):
    cell = ws.cell(170, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(HEAD, start=1):
    cell = ws.cell(172, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

for i, (name, src) in enumerate(GROUPS):
    r = 173 + i
    last = name == 'Company'
    ws.cell(r, 1, name).font = Font(name='Calibri', sz=11, b=last, color=NAVY)
    for c, off, fmt in COLS:
        cell = ws.cell(r, c)
        cell.value = (f'=INDEX($D{src + off}:$O{src + off},'
                      f'MATCH(Setup!$B$5,$D$122:$O$122,0))')
        cell.number_format = fmt
        cell.font = Font(name='Calibri', sz=11, b=last or c in (6, 8), color=NAVY)
    if last:
        for c in range(1, 9):
            ws.cell(r, c).fill = TOTAL

wb.save(WB)
print('Utilisation rows 170 to 176 added: the month on one line per group')
