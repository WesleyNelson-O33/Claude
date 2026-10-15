"""The CTS palette across the whole book, and the charts she asked for.

Charts are native Excel charts, not pictures, so they move when the month
moves or a new GL is pasted.
"""
import os, json, collections, copy
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import DataPoint

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
R = json.load(open(f'{HERE}/data/newsheet_rows.json'))
P = json.load(open(f'{HERE}/data/pldept.json'))
D = json.load(open(f'{HERE}/data/dept_colours.json'))

TEAL, AQUA, BLUE, NAVY = B['dark_teal'], B['bright_teal'], B['primary_blue'], B['deep_blue']
RED, DRED, PURP, INK = B['accent_red'], B['deep_red'], B['purple'], B['charcoal']
ALLOWED = {v.upper() for k, v in B.items() if k != '_note'}

wb = openpyxl.load_workbook(OUT)

# ------------------------------------------------- 1. the exact palette ----
# Anything that is not one of her nine, or one of the two banding neutrals,
# is mapped onto the nearest colour she gave.
FONT_MAP = {'1A1A1A': INK, '5F6670': INK, '000000': INK, 'FF0000': RED,
            'C00000': DRED, '0000FF': TEAL, '808080': INK, '595959': INK}
FILL_MAP = {'C9CED6': B['grey'], 'F2F2F2': B['grey_lt'], 'FBE8EC': B['grey_lt'],
            'FFFF00': B['grey'], 'DDEBF7': B['grey'], 'FFF2CC': B['grey']}
n = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f = c.fill
            if f is not None and f.fill_type == 'solid':
                rgb = getattr(f.start_color, 'rgb', None)
                if isinstance(rgb, str):
                    h = rgb[-6:].upper()
                    if h not in ALLOWED:
                        c.fill = PatternFill('solid', fgColor=FILL_MAP.get(h, B['grey_lt']))
                        n += 1
            fc = getattr(c.font, 'color', None)
            rgb = getattr(fc, 'rgb', None) if fc is not None else None
            if isinstance(rgb, str):
                h = rgb[-6:].upper()
                if h not in ALLOWED:
                    nf = copy.copy(c.font)
                    nf.color = openpyxl.styles.Color(rgb='FF' + FONT_MAP.get(h, INK))
                    c.font = nf
                    n += 1
print('cells moved onto the CTS palette:', n)

# ------------------------------------------------------------ 2. charts ----
def style(ch, title, w, h, labels=True):
    ch.title = title
    ch.width, ch.height = w, h
    ch.style = None
    ch.y_axis.majorGridlines = None
    ch.legend = None
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.numFmt = '#,##0'

def paint(ch, colours):
    for s, col in zip(ch.series, colours):
        s.graphicalProperties = GraphicalProperties(solidFill=col)
        s.graphicalProperties.line = LineProperties(noFill=True)

def paint_depts(ser, names):
    """One colour a department, her convention: teals Production, blue and
    purple Consulting, red Support."""
    ser.data_points = [
        DataPoint(idx=i,
                  spPr=GraphicalProperties(solidFill=D['departments'].get(n, INK),
                                           ln=LineProperties(noFill=True)))
        for i, n in enumerate(names)]

# Top 10 clients -------------------------------------------------------------
ws = wb['Top 10 Clients']
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.add_data(Reference(ws, min_col=4, min_row=6, max_row=16), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=7, max_row=16))
style(ch, 'Top 10 customers, year to date', 20, 11)
paint(ch, [TEAL])
ws.add_chart(ch, 'G5')

# P&L by Dept ----------------------------------------------------------------
ws = wb['P&L by Dept']
m = P['month_rows']
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
HDR_M = m['Total Income'] - 1
ch.add_data(Reference(ws, min_col=1, min_row=m['Total Income'], max_col=7,
                      max_row=m['Total Income']), from_rows=True, titles_from_data=True)
ch.add_data(Reference(ws, min_col=1, min_row=m['Result after all overhead'], max_col=7,
                      max_row=m['Result after all overhead']), from_rows=True,
            titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=HDR_M, max_col=7))
style(ch, 'Revenue and result after overhead, this month', 22, 10)
ch.legend = openpyxl.chart.legend.Legend(); ch.legend.position = 'b'
DEPTS = [ws.cell(HDR_M, c).value for c in range(2, 8)]
paint(ch, [TEAL, INK])
paint_depts(ch.series[0], DEPTS)
ws.add_chart(ch, 'J5')

y = P['ytd_rows']
ch2 = BarChart(); ch2.type = 'col'; ch2.grouping = 'clustered'
HDR_Y = y['Total Income'] - 1
ch2.add_data(Reference(ws, min_col=1, min_row=y['Total Income'], max_col=7,
                       max_row=y['Total Income']), from_rows=True, titles_from_data=True)
ch2.add_data(Reference(ws, min_col=1, min_row=y['Result after all overhead'], max_col=7,
                       max_row=y['Result after all overhead']), from_rows=True,
             titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=2, min_row=HDR_Y, max_col=7))
style(ch2, 'Revenue and result after overhead, year to date', 22, 10)
ch2.legend = openpyxl.chart.legend.Legend(); ch2.legend.position = 'b'
paint(ch2, [TEAL, INK])
paint_depts(ch2.series[0], [ws.cell(HDR_Y, c).value for c in range(2, 8)])
ws.add_chart(ch2, 'J27')

# DSO ------------------------------------------------------------------------
ws = wb['DSO']
ch = BarChart(); ch.type = 'col'
ch.add_data(Reference(ws, min_col=2, min_row=18, max_row=24), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=19, max_row=24))
style(ch, 'Aged receivables', 17, 9)
paint(ch, [TEAL])
ws.add_chart(ch, 'J5')

ch = BarChart(); ch.type = 'bar'
ch.add_data(Reference(ws, min_col=3, min_row=30, max_row=40), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=31, max_row=40))
style(ch, 'Who owes the most', 17, 11)
paint(ch, [BLUE])
ws.add_chart(ch, 'J24')

T = R['dso_trend']
ch = LineChart()
ch.add_data(Reference(ws, min_col=1, min_row=T + 5, max_col=13, max_row=T + 5),
            from_rows=True, titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=T + 1, max_col=13))
style(ch, 'DSO month by month', 22, 9, labels=False)
for s in ch.series:
    s.graphicalProperties = GraphicalProperties()
    s.graphicalProperties.line = LineProperties(solidFill=TEAL, w=28000)
    s.smooth = False
ws.add_chart(ch, f'A{T + 9}')

# Overhead Split -------------------------------------------------------------
ws = wb['Overhead Split']
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
ch.add_data(Reference(ws, min_col=1, min_row=16, max_col=6, max_row=18),
            from_rows=True, titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=15, max_col=6))
style(ch, 'The three split keys', 20, 9, labels=False)
ch.legend = openpyxl.chart.legend.Legend(); ch.legend.position = 'b'
ch.y_axis.numFmt = '0%'
paint(ch, [INK, B['grey'], NAVY])
ws.add_chart(ch, 'V5')
print('charts added: 7')

wb.save(OUT)
print('saved', OUT)
