import openpyxl, warnings
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook("CTS Financial Controller Pack.xlsx")
pl, eng, bud, summ = wb["P&L FY27"], wb["Engine"], wb["Budget_Paste"], wb["Summary"]

# 1. every P&L subcategory cell must point at the Engine row holding that dept+subcat
bad = 0; checked = 0
cur_dept = None
for r in range(6, pl.max_row + 1):
    label = pl.cell(row=r, column=1).value
    if not isinstance(label, str):
        continue
    if not label.startswith(" "):
        cur_dept = label.strip(); continue
    if not label.startswith("        "):
        continue
    sub = label.strip()
    f = pl.cell(row=r, column=2).value
    if not isinstance(f, str) or "Engine" not in f:
        continue
    checked += 1
    er = int(f.split("Engine!$E")[1].split(",")[0])
    key = eng.cell(row=er, column=1).value
    if key != f"{cur_dept}|{sub}":
        bad += 1
        if bad < 4: print("MISMATCH", r, cur_dept, sub, "->", key)
print(f"P&L subcategory links checked {checked}, mismatched {bad}")

# 2. the last month column must land on Engine column P
r = 8
print("P&L last month cell:", pl.cell(row=r, column=13).value)
print("Engine header E4..P4:", [eng.cell(row=4, column=c).value for c in (5, 16)])
print("Budget header E4..P4:", [bud.cell(row=4, column=c).value for c in (5, 16)])

# 3. formula lengths must stay under Excel's 8192 limit
mx = 0; where = None
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and len(c.value) > mx:
                mx, where = len(c.value), f"{ws.title}!{c.coordinate}"
print("longest formula:", mx, "chars at", where)
print("Summary B6:", summ["B6"].value[:200])
print("Summary total row:", summ.max_row, summ.cell(row=summ.max_row, column=1).value)
