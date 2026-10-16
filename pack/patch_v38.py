"""Profitability, month on month and year on year.

  month on month   gross profit, EBITDA and net profit across the twelve months
  year on year     net profit and EBITDA this year against the same months last
                   year, off the Actual FY26 tab

From next month on the P&L carries budget rather than actual, so the heading
says so rather than letting it read as achieved.
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
CC = json.load(open(f'{HERE}/data/costcharts.json'))
INK, WHITE, GREY, GREYL, AQUA = (B['charcoal'], B['white'], B['grey'],
                                 B['grey_lt'], B['bright_teal'])
MONEY = '\\$#,##0;[Red]"($"#,##0\\);\\-'
PCT = '0.0%;[Red](0.0%);\\-'
PL, F26 = "'P&L FY27'", "'Actual FY26'"
ROWS = {'Gross profit': 173, 'EBITDA': 176, 'Net profit': 180, 'Total Income': 6}

wb = openpyxl.load_workbook(OUT)
ws = wb['Cost Analysis']
thin = Side(style='thin', color=GREY)
hair = Side(style='medium', color=AQUA)
S = CC['bvm_last'] + 4

def banner(r, text, n=7):
    ws.cell(r, 1, text)
    for c in range(1, n + 1):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).font = Font(name='Calibri', size=12, bold=True, color=WHITE)

def header(r, labels):
    for i, t in enumerate(labels, start=1):
        c = ws.cell(r, i, t)
        c.font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(horizontal='right' if i > 1 else 'left', wrap_text=True)
        c.border = Border(bottom=hair)

banner(S, 'PROFITABILITY BY MONTH, ACTUAL TO DATE THEN BUDGET, AGAINST LAST YEAR')
header(S + 1, ['Month', 'Gross profit', 'EBITDA', 'Net profit',
               'Net profit last year', 'EBITDA last year', 'Net profit %'])
for i in range(12):
    r, c = S + 2 + i, gl(2 + i)
    ws.cell(r, 1, f'={PL}!{c}$5').number_format = 'mmm\\-yy'
    ws.cell(r, 2, f'={PL}!{c}${ROWS["Gross profit"]}').number_format = MONEY
    ws.cell(r, 3, f'={PL}!{c}${ROWS["EBITDA"]}').number_format = MONEY
    ws.cell(r, 4, f'={PL}!{c}${ROWS["Net profit"]}').number_format = MONEY
    ws.cell(r, 5, f'={F26}!{c}${ROWS["Net profit"]}').number_format = MONEY
    ws.cell(r, 6, f'={F26}!{c}${ROWS["EBITDA"]}').number_format = MONEY
    ws.cell(r, 7, f'=IFERROR({PL}!{c}${ROWS["Net profit"]}/'
                  f'{PL}!{c}${ROWS["Total Income"]},0)').number_format = PCT
    for cc in range(1, 8):
        ws.cell(r, cc).font = Font(name='Calibri', size=11, color=INK)
        ws.cell(r, cc).border = Border(bottom=thin)
        if cc > 1:
            ws.cell(r, cc).alignment = Alignment(horizontal='right')
LAST = S + 13
ws.cell(LAST + 1, 1, 'Check, net profit to date against the P&L YTD')
ws.cell(LAST + 1, 4, f'=ROUND(SUMIF($A${S+2}:$A${LAST},"<="&Setup!$B$5,'
                     f'$D${S+2}:$D${LAST})-{PL}!$O${ROWS["Net profit"]},2)')
ws.cell(LAST + 1, 4).number_format = MONEY
for cc in range(1, 8):
    ws.cell(LAST + 1, cc).font = Font(name='Calibri', size=11, bold=True, color=INK)
    ws.cell(LAST + 1, cc).border = Border(bottom=thin)

json.dump({**CC, 'prof_hdr': S + 1, 'prof_first': S + 2, 'prof_last': LAST},
          open(f'{HERE}/data/costcharts.json', 'w'), indent=1)
print(f'profitability block: rows {S} to {LAST}')
wb.save(OUT)
print('saved')
