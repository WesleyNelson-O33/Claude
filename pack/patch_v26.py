"""A short label for the customer charts.

The full Xero contact name stays in column B and is what everything matches
on. The chart axis reads a shortened version instead, so the names stop
crowding the bars. Two formula columns so the rule is readable:

  Name tidied   company suffixes come off, and a trading name like
                "X P/L for X Trust" is cut back to X
  Chart label   text inside brackets wins, so CBA (Place South) becomes
                Place South; anything still over 26 characters is cut
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
INK, WHITE = B['charcoal'], B['white']
SUFFIX = ('" Pty Limited"', '" Pty Ltd"', '" Limited"', '" Ltd"', '" P/L"')

def tidy(ref):
    s = ref
    for token in SUFFIX:
        s = f'SUBSTITUTE({s},{token},"")'
    s = f'TRIM({s})'
    return (f'=IF({ref}="","",IF(ISNUMBER(SEARCH(" for ",{s})),'
            f'TRIM(LEFT({s},SEARCH(" for ",{s})-1)),{s}))')

def label(full, tid):
    return (f'=IF({full}="","",IF(ISNUMBER(SEARCH("(",{full})),'
            f'MID({full},SEARCH("(",{full})+1,SEARCH(")",{full})-SEARCH("(",{full})-1),'
            f'IF(LEN({tid})>26,LEFT({tid},25)&"…",{tid})))')

wb = openpyxl.load_workbook(OUT)

def add(sheet, c1, hdr, first, last):
    ws = wb[sheet]
    for col, text in ((c1, 'Name tidied'), (c1 + 1, 'Chart label')):
        c = ws.cell(hdr, col, text)
        c.font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(horizontal='left')
        ws.column_dimensions[gl(col)].width = 28
    for r in range(first, last + 1):
        ws.cell(r, c1, tidy(f'$B{r}'))
        ws.cell(r, c1 + 1, label(f'$B{r}', f'${gl(c1)}{r}'))
        for col in (c1, c1 + 1):
            ws.cell(r, col).font = Font(name='Calibri', size=11, color=INK)
    print(f'{sheet}: {gl(c1)} tidied, {gl(c1+1)} chart label, rows {first} to {last}')

add('Top 10 Clients', 6, 6, 7, 16)
add('Top 10 YoY', 9, 4, 5, 14)
wb.save(OUT)
print('saved')
