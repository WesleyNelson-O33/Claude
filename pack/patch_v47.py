"""
The income by month chart printed every label twice.

Column B on that block is "actual to date then budget", so for the months after
the reporting month it holds the same figure as the budget column. Two bars of
identical height, two identical labels, printed on top of each other.

Column D now carries the actual for closed months only and is blank after
that. A blank charts as a zero height bar, which is invisible, and the label
format hides zeros, so the closed months show two bars and the rest show one.
"""
import openpyxl, sys
from openpyxl.styles import Font, Alignment

WB = sys.argv[1] if len(sys.argv) > 1 else 'CTS Financial Controller Pack - v20.xlsx'
FIRST, LAST = 112, 123

wb = openpyxl.load_workbook(WB)
c = wb['Cost Analysis']
c['D111'] = 'Actual, closed months'
c['D111'].font = Font(bold=True, size=10)
c['D111'].alignment = Alignment(horizontal='right')
for r in range(FIRST, LAST + 1):
    cell = c.cell(r, 4)
    cell.value = f'=IF($A{r}<=Setup!$B$5,$B{r},"")'
    cell.number_format = c.cell(r, 2).number_format
wb.save(WB)
print(f'Cost Analysis D{FIRST}:D{LAST} added, actual for closed months only')
