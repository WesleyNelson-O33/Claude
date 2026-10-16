"""Take my notes off the P&L tabs.

The sheets keep their titles, the month, and the headings that label a real
block of data. Everything that was me explaining the sheet to her comes off.

Her own note on Summary U4 stays. So does the heading over the departmental
budget expenses on Budget FY27, because that labels data rather than
explaining it.
"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
MONTH = '=TEXT(Setup!$B$5,"mmm yyyy")'

CLEAR = {
    'P&L FY27':    ['A2', 'C3', 'A221'],
    'Budget FY27': ['A2', 'C3'],
    'Actual FY26': ['A2', 'C3'],
    'Budget FY26': ['A2', 'C3'],
    'Summary':     ['C3'],
    'P&L by Dept': ['A2'],
    'Dept P&L':    ['A2', 'C5', 'A27'],
}
# the month line keeps the month and loses the instruction glued onto it
KEEPMONTH = {'P&L by Dept': 'A3'}

wb = openpyxl.load_workbook(OUT)
n = 0
for sheet, cells in CLEAR.items():
    ws = wb[sheet]
    for ref in cells:
        old = ws[ref].value
        if old is None:
            continue
        ws[ref] = None
        n += 1
        print(f'  {sheet}!{ref}  removed: {str(old)[:76]}')
for sheet, ref in KEEPMONTH.items():
    ws = wb[sheet]
    print(f'  {sheet}!{ref}  reduced to the month')
    ws[ref] = MONTH
    n += 1

print(f'\nnotes removed: {n}')
print('kept: Summary!U4 (her note), Budget FY27!A195 (heading over real data)')
wb.save(OUT)
print('saved')
