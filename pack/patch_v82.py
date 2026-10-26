"""Clear two stray rows on Xero History.

The monthly net profit block runs rows 136 to 138 and the quarterly one 142 to
144. Rows 139 and 140 carry a second copy of the first two quarter rows, written
over the gap and over the quarterly band heading, so the sheet shows a Q1 2024
line sitting inside the monthly block and a figure printed beside the heading.
The charts read 136:138 and 142:144, so nothing downstream was wrong, but the
tab reads as if the monthly block has four periods.
"""
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'

wb = openpyxl.load_workbook(WB)
ws = wb['Xero History']
assert str(ws['A134'].value).startswith('NET PROFIT AND NET LOSS, THE MONTH')
assert str(ws['A140'].value).startswith('NET PROFIT AND NET LOSS, THE QUARTER')

for coord in ('A139', 'B139', 'C139', 'D139', 'B140', 'C140', 'D140'):
    if ws[coord].value is not None:
        print(f'  cleared {coord}: {str(ws[coord].value)[:40]}')
        ws[coord] = None

wb.save(WB)
print('saved')
