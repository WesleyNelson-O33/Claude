"""Employment expenses split into office staff and departmental staff.

From Gi, in the reporting video at 1:32:24 and in the manual:

  departmental staff  the employment cost that genuinely belongs to that
                      department
  office staff        admin salaries pushed into that department, either by
                      the overhead split or because someone clocked into one
                      of its cost centres

Departmental staff is the employment expense sitting on that department's own
cost centre. Checked against her August P&L and all five tie to the cent:
Support 7,827, Events 24,841, Video 6,414, Integration 5,308, Consulting 1,381,
and the year to date figures too.

Office staff is the admin pot, 85,081.57 in August, which also ties to her CTS
office staff exactly. How that pot splits across the five does NOT come out of
anything in this pack. In Gi's own file that cell is grey, not linked, and she
updates it each month off the labour hours report. Her August split works out
at 67.85, 25.19, 4.18, 1.39 and 1.39 per cent, which is not the headcount key,
not total hours, not chargeable hours and not FTE. So it is an input here, with
a check that the five add back to the pot.
"""
import os, json, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
INK, WHITE, GREY, GREYL, AQUA = (B['charcoal'], B['white'], B['grey'],
                                 B['grey_lt'], B['bright_teal'])
COST = '[Red]"($"#,##0\\);\\$#,##0;\\-'
PCT = '0.0%;[Red](0.0%);\\-'
DEPTS = [('Support', 'ONSITE'), ('Events', 'PRODUCTION'), ('Video', 'VIDEO'),
         ('Integration', 'INTEGRATION'), ('Consulting', 'CONSULTING')]
CA, CB, CD = ('Cleanup!$A$5:$A$20004', 'Cleanup!$B$5:$B$20004', 'Cleanup!$D$5:$D$20004')
CF = 'Cleanup!$F$5:$F$20004'

wb = openpyxl.load_workbook(OUT)
ws = wb['Overheads MoM']
HR = 4                                   # the month heading row on that sheet
S = 130
thin = Side(style='thin', color=GREY)
hair = Side(style='medium', color=AQUA)

def banner(r, text):
    ws.cell(r, 1, text)
    for c in range(1, 15):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).font = Font(name='Calibri', size=12, bold=True, color=WHITE)

def head(r, first):
    ws.cell(r, 1, first)
    for i in range(12):
        ws.cell(r, 2 + i, f'={gl(2+i)}${HR}').number_format = 'mmm\\-yy'
    ws.cell(r, 14, 'Year to date')
    for c in range(1, 15):
        ws.cell(r, c).font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).alignment = Alignment(horizontal='right' if c > 1 else 'left')
        ws.cell(r, c).border = Border(bottom=hair)

def line(r, label, body, bold=False, fill=None, input_=False):
    ws.cell(r, 1, label)
    for i in range(12):
        c = 2 + i
        ws.cell(r, c, body(gl(c))) if body else None
        ws.cell(r, c).number_format = COST
        if input_:
            ws.cell(r, c).fill = PatternFill('solid', fgColor=AQUA)
    ws.cell(r, 14, f'=SUMIF($B${HR}:$M${HR},"<="&Setup!$B$5,$B{r}:$M{r})')
    ws.cell(r, 14).number_format = COST
    for c in range(1, 15):
        ws.cell(r, c).font = Font(name='Calibri', size=11, bold=bold, color=INK)
        ws.cell(r, c).border = Border(bottom=thin)
        if fill and not (input_ and 1 < c < 14):
            ws.cell(r, c).fill = PatternFill('solid', fgColor=fill)
        if c > 1:
            ws.cell(r, c).alignment = Alignment(horizontal='right')

banner(S, 'EMPLOYMENT EXPENSES, OFFICE STAFF AND DEPARTMENTAL STAFF')
ws.cell(S + 1, 1, 'Departmental staff is calculated. Office staff is typed, the teal cells, '
                  'because the way the admin pot splits is not in this pack.')
ws.cell(S + 1, 1).font = Font(name='Calibri', size=9, italic=True, color=INK)
head(S + 2, 'Department')

r = S + 3
DEPT_ROWS, OFF_ROWS, TOT_ROWS = {}, {}, {}
for label, tag in DEPTS:
    dr = r
    line(r, f'{label}, departmental staff',
         lambda c, t=tag: f'=SUMIFS({CA},{CD},"{t}",{CF},"Employment Expenses",'
                          f'{CB},{c}${HR})')
    DEPT_ROWS[label] = r; r += 1
    line(r, f'{label}, office staff', None, input_=True)
    OFF_ROWS[label] = r; r += 1
    line(r, f'{label}, total employment expense',
         lambda c, a=dr, b=dr + 1: f'={c}{a}+{c}{b}', bold=True, fill=GREYL)
    TOT_ROWS[label] = r; r += 2

POT = r
line(POT, 'The admin pot, employment expense on the CTS cost centre',
     lambda c: f'=SUMIFS({CA},{CD},"ADMIN",{CF},"Employment Expenses",{CB},{c}${HR})',
     bold=True, fill=GREY)
r += 1
ENT = r
line(ENT, 'Office staff entered above',
     lambda c: '=' + '+'.join(f'{c}{OFF_ROWS[d]}' for d, _ in DEPTS), bold=True)
r += 1
line(r, 'Check, the pot less what is entered',
     lambda c: f'=ROUND({c}{POT}-{c}{ENT},2)', bold=True)
CHK = r
r += 2
line(r, 'Check, departmental staff against the P&L employment expense line',
     lambda c: '=ROUND(' + '+'.join(f'{c}{DEPT_ROWS[d]}' for d, _ in DEPTS)
               + f'+{c}{POT}-{c}{57},2)', bold=True)
print(f'block written: rows {S} to {r}')
print('  departmental staff rows:', DEPT_ROWS)
print('  office staff input rows:', OFF_ROWS)

json.dump({'start': S, 'dept': DEPT_ROWS, 'office': OFF_ROWS, 'total': TOT_ROWS,
           'pot': POT, 'check': CHK}, open(f'{HERE}/data/staffsplit.json', 'w'), indent=1)
wb.save(OUT)
print('saved')
