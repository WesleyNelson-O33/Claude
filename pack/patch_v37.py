"""Budget against actual, as a block and two charts.

  by line        the six lines that matter, year to date, actual against budget
  by month       total income month by month, actual against budget

The monthly one is honest about itself: from next month on the P&L carries
budget rather than actual, so the two series converge. The heading says so.
"""
import os, json, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
CC = json.load(open(f'{HERE}/data/costcharts.json'))
INK, WHITE, GREY, GREYL, AQUA = (B['charcoal'], B['white'], B['grey'],
                                 B['grey_lt'], B['bright_teal'])
MONEY = '\\$#,##0;[Red]"($"#,##0\\);\\-'
PCT = '0.0%;[Red](0.0%);\\-'
PL = "'P&L FY27'"
BU = "'Budget FY27'"
LINES = [('Total Income', 6), ('Cost of Sales', 59), ('Gross profit', 173),
         ('Expenses Overheds', 164), ('EBITDA', 176), ('Net profit', 180)]

wb = openpyxl.load_workbook(OUT)
ws = wb['Cost Analysis']
thin = Side(style='thin', color=GREY)
hair = Side(style='medium', color=AQUA)
S = 100

def banner(r, text, n=6):
    ws.cell(r, 1, text)
    for c in range(1, n + 1):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).font = Font(name='Calibri', size=12, bold=True, color=WHITE)

def header(r, labels):
    for i, t in enumerate(labels, start=1):
        c = ws.cell(r, i, t)
        c.font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(horizontal='right' if i > 1 else 'left', wrap_text=True)
        c.border = Border(bottom=hair)

def plain(r, n, bold=False, fill=None):
    for c in range(1, n + 1):
        ws.cell(r, c).font = Font(name='Calibri', size=11, bold=bold, color=INK)
        ws.cell(r, c).border = Border(bottom=thin)
        if fill:
            ws.cell(r, c).fill = PatternFill('solid', fgColor=fill)
        if c > 1:
            ws.cell(r, c).alignment = Alignment(horizontal='right')

banner(S, 'BUDGET AGAINST ACTUAL, YEAR TO DATE')
header(S + 1, ['Line', 'Actual', 'Budget', 'Var $', 'Var %', ''])
for i, (nm, pr) in enumerate(LINES):
    r = S + 2 + i
    ws.cell(r, 1, nm)
    ws.cell(r, 2, f'={PL}!$O${pr}').number_format = MONEY
    ws.cell(r, 3, f'={PL}!$Q${pr}').number_format = MONEY
    ws.cell(r, 4, f'=$B{r}-$C{r}').number_format = MONEY
    ws.cell(r, 5, f'=IFERROR($D{r}/ABS($C{r}),0)').number_format = PCT
    plain(r, 5)
LAST = S + 1 + len(LINES)

M0 = LAST + 3
banner(M0, 'TOTAL INCOME BY MONTH, ACTUAL TO DATE THEN BUDGET, AGAINST BUDGET', 4)
header(M0 + 1, ['Month', 'Actual', 'Budget', ''])
for i in range(12):
    r = M0 + 2 + i
    ws.cell(r, 1, f'={PL}!{gl(2+i)}$5').number_format = 'mmm\\-yy'
    ws.cell(r, 2, f'={PL}!{gl(2+i)}$6').number_format = MONEY
    ws.cell(r, 3, f'={BU}!{gl(2+i)}$6').number_format = MONEY
    plain(r, 3)
MLAST = M0 + 13
ws.cell(MLAST + 1, 1, 'Check, the actual column to date against the P&L YTD')
ws.cell(MLAST + 1, 2, f'=ROUND(SUMIF($A${M0+2}:$A${MLAST},"<="&Setup!$B$5,'
                      f'$B${M0+2}:$B${MLAST})-{PL}!$O$6,2)').number_format = MONEY
plain(MLAST + 1, 3, bold=True)

json.dump({**CC, 'bva_first': S + 2, 'bva_last': LAST, 'bva_hdr': S + 1,
           'bvm_first': M0 + 2, 'bvm_last': MLAST, 'bvm_hdr': M0 + 1},
          open(f'{HERE}/data/costcharts.json', 'w'), indent=1)
print(f'budget against actual: by line {S+2} to {LAST}, by month {M0+2} to {MLAST}')
wb.save(OUT)
print('saved')
