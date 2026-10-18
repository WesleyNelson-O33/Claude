"""
The hours picture, grouped into the three departments.

Danica asked for the hours chart, not a revenue chart, rolled up the way the
P&L groups the business: Support on its own, Production as Events and Video,
Consulting as Integration and Consulting.

This adds the block the chart reads from, under the five department version, so
the two sit together on the tab. It references the five department rows rather
than the pastes, so there is one source of truth and the two cannot drift.

It also removes the revenue chart data added in patch_v52, which is not what was
asked for.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
NOFILL = PatternFill(fill_type=None)
RULE = Side(style='medium', color=NAVY)
HRS = '#,##0.00;;'

wb = openpyxl.load_workbook(WB)

# --- the three department hours block, Utilisation ----------------------
ws = wb['Utilisation']
assert ws['A66'].value.startswith('HOURS BY DEPARTMENT'), ws['A66'].value
ws['A76'] = 'HOURS BY THE THREE DEPARTMENTS, THE MONTH ON SETUP'
for c in range(1, 5):
    cell = ws.cell(76, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(['Department', 'Billable hours', 'Non-billable hours',
                          'Leave'], start=1):
    cell = ws.cell(77, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

# five department rows: 68 Support, 69 Events, 70 Video, 71 Integration,
# 72 Consulting. Production is Events plus Video, Consulting is Integration
# plus Consulting.
GROUPS = [('Support', [68]), ('Production', [69, 70]), ('Consulting', [71, 72])]
for i, (name, rows) in enumerate(GROUPS):
    r = 78 + i
    ws.cell(r, 1, name).font = Font(name='Calibri', sz=11, color=NAVY)
    for c in (2, 3, 4):
        col = chr(64 + c)
        ws.cell(r, c, '=' + '+'.join(f'{col}{s}' for s in rows)).number_format = HRS
        ws.cell(r, c).font = Font(name='Calibri', sz=11, color=NAVY)

ws.cell(81, 1, 'Total')
for c in (2, 3, 4):
    col = chr(64 + c)
    ws.cell(81, c, f'=SUM({col}78:{col}80)').number_format = HRS
for c in range(1, 5):
    cell = ws.cell(81, c)
    cell.fill = TOTAL
    cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)

# --- take the revenue chart data back out, Overhead Split ---------------
ws = wb['Overhead Split']
for r in range(163, 168):
    for c in range(1, 5):
        cell = ws.cell(r, c)
        cell.value = None
        cell.fill = NOFILL
        cell.border = Border()

wb.save(WB)
print('Utilisation rows 76 to 81 added: hours for the three departments')
print('Overhead Split rows 163 to 167 cleared: the revenue chart data is gone')
