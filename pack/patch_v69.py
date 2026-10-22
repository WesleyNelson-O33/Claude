"""Helper columns for the departmental profit and loss chart.

Revenue, total cost and the result, one column each, so the chart reads three
plain series rather than the build script doing arithmetic.

Total cost is cost of sales plus overhead. Cost of sales is not a line on the
table, it is revenue less gross profit, and overhead sits on its own row, which
is why a figure like Support's 213,942 does not appear anywhere on the printed
table. It is 141,659 of cost of sales plus 72,283 of overhead.

The columns sit at P, Q and R on rows 142 to 147, above every snip range on the
tab, so they never show up in a picture.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
NOTE = Font(name='Calibri', sz=10, color='FF6B7585')
HEAD, FIRST, LAST = 142, 143, 147

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws[f'A{HEAD}'].value == 'Department'

ws.cell(HEAD - 1, 16, 'Working for the profit and loss chart. Total cost is cost '
                      'of sales plus overhead.').font = NOTE
for c, name in ((16, 'Revenue'), (17, 'Total cost'), (18, 'Result')):
    cell = ws.cell(HEAD, c, name)
    cell.font = Font(name='Calibri', sz=10, b=True, color='FF242C39')

for r in range(FIRST, LAST + 1):
    for c, f in ((16, f'=$B{r}'),
                 (17, f'=-($B{r}-$E{r})+$G{r}'),      # cost of sales plus overhead
                 (18, f'=$H{r}')):
        cell = ws.cell(r, c, f)
        cell.number_format = MONEY
        cell.font = NOTE
    ws.column_dimensions[ws.cell(HEAD, c).column_letter].width = 12

wb.save(WB)
print('Overhead Split columns P, Q and R added: revenue, total cost and result')
