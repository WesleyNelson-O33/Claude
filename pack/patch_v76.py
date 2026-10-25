"""Pick up the legal fees that were falling into the catch-all.

The GL codes 800 a month to an account called plainly "Legal Fees".  The
overhead split names two legal rows, "Legal Fees - Easton Belle" (3 Way) and
"Legal Fees - Other (Employsure)" (Staff), so neither matched and the 800 fell
into "Everything else coded to admin", where it was spread on the 3 Way key.

The P&L analysis puts the whole 800 on Legal Fees - Other (Employsure), on the
Staff key, in both July and August.  So this points that row at the GL name.
"""
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
ROW = 59
GL_NAME = 'Legal Fees'

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws.cell(ROW, 1).value == 'Legal Fees - Other (Employsure)', ws.cell(ROW, 1).value
assert ws.cell(ROW, 2).value == 'Staff'

for c in range(3, 15):                      # C..N, the twelve months
    cell = ws.cell(ROW, c)
    old = cell.value
    assert f'$A{ROW}' in old, old
    cell.value = old.replace(f'$A{ROW}', f'"{GL_NAME}"')

print(f'row {ROW} now reads the GL account "{GL_NAME}"')
print('  ', ws.cell(ROW, 3).value)
wb.save(WB)
