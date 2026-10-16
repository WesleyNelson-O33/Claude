"""Two small fixes after the month block went in.

W errored on the margin rows. A margin line returns an empty string when the
department has no income, and an empty string less an empty string is a
#VALUE. Wrap it, the same way the Var % column beside it already is.

And the subcategory index on the Dept P&L workings was headed "#", which reads
like an error in the cell. Call it No. instead.
"""
import os
import openpyxl
from openpyxl.utils import column_index_from_string as ci

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
wb = openpyxl.load_workbook(OUT)
s = wb['P&L FY27']
n = 0
for r in range(6, s.max_row + 1):
    if s.cell(r, ci('W')).value == f'=U{r}-V{r}':
        s.cell(r, ci('W'), f'=IFERROR(U{r}-V{r},"")')
        n += 1
print('Var $ cells wrapped:', n)

d = wb['Dept P&L']
for r in range(1, d.max_row + 1):
    if d.cell(r, 1).value == '#':
        d.cell(r, 1, 'No.')
        print(f'Dept P&L A{r} header renamed to No.')
wb.save(OUT)
print('saved')
