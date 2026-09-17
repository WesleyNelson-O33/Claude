import openpyxl, warnings, re
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook("CTS Financial Controller Pack.xlsx")
print("sheets:", wb.sheetnames)
print("defined names:", sorted(wb.defined_names.keys()))
print()
for name in ("Cleanup", "Engine", "P&L FY27", "Summary", "PL_Check", "Utilisation"):
    ws = wb[name]
    print(f"--- {name}  dims={ws.calculate_dimension()}")
for ref in ["Cleanup!A5", "Cleanup!C5", "Cleanup!D5", "Cleanup!F5",
            "Engine!A5", "Engine!E5", "Engine!Q5",
            "P&L FY27!A6", "P&L FY27!B8", "P&L FY27!B7",
            "Summary!B6", "Utilisation!B6", "PL_Check!B5"]:
    sh, cell = ref.split("!")
    print(f"{ref:<18} {str(wb[sh][cell].value)[:190]}")
# formula sanity: every formula should reference a sheet that exists
bad = []
names = set(wb.sheetnames)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                for m in re.finditer(r"([A-Za-z0-9_& ]+)!", v):
                    t = m.group(1).strip().strip("'")
                    if t and t not in names and not t[0].isdigit():
                        bad.append((ws.title, c.coordinate, t))
print("\nbroken sheet references:", len(bad), bad[:5])
