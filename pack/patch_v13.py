"""Two controls were reading the wrong thing. The data was never wrong.

PL_Check
  1. Xero calls the income total "Total Trading Income", not "Total Income",
     so the lookup never found it and the income line read nil.
  2. Column C held August figures, typed in by hand when August was the
     reporting month. The GL moved to September, the typed column did not,
     so the control compared two different months. Column C now reads the
     pasted P&L, so it moves on its own. She can still type over it.

Balance_Sheet
  Her Xero balance sheet export puts the account name in column B, not A,
  because Xero indents. Column A is empty. The lookup was matching against
  the empty column A and reading the names out of column B as if they were
  figures, so every line came back nil. It now matches column B, and Setup
  points at column C, which is where the 30 Sept 2026 balances are. Her
  export carries no comparative columns, so last month and start of year are
  set to nil rather than being pointed at something that is not there.
"""
import os, re, shutil
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = f'{HERE}/CTS Financial Controller Pack - v12.xlsx'
OUT = f'{HERE}/CTS Financial Controller Pack - v13.xlsx'

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

# ------------------------------------------------------------- PL_Check ----
ws = wb['PL_Check']
assert ws['G5'].value == 'Total Income', ws['G5'].value
ws['G5'] = 'Total Trading Income'
for r in range(5, 9):
    ws[f'C{r}'] = f'=$I{r}'
ws['A2'] = ('Paste the Xero profit and loss on PL_Paste, whole export, at A1. '
            'Column C reads it. If there is no paste this month, type the '
            'totals straight into column C. Every status should read OK '
            'before you report.')

# -------------------------------------------------------- Balance_Sheet ----
bsh = wb['Balance_Sheet']
n = 0
for row in bsh.iter_rows():
    for c in row:
        if isinstance(c.value, str) and 'BS_Paste!$A$1:$A$400' in c.value:
            c.value = c.value.replace('BS_Paste!$A$1:$A$400', 'BS_Paste!$B$1:$B$400')
            n += 1
print('Balance_Sheet lookups repointed at the name column:', n)
bsh['A2'] = ('Paste the Xero balance sheet on BS_Paste, whole export, at A1. '
             'This reads it by matching the account name in column B, which '
             'is where Xero puts it. Tell Setup rows 30 to 32 which columns '
             'hold the balances.')

# ------------------------------------------------------------------ Setup --
st = wb['Setup']
st['B30'] = 3        # column C holds this month's balance
st['B31'] = 0        # her export has no last month column
st['B32'] = 0        # her export has no start of year column
st['C30'] = 'Column B is 2, C is 3, and so on. The Xero export used in September had it in C.'
st['C31'] = 'Leave at 0 if your export has no comparative column.'
st['C32'] = 'Leave at 0 if your export has no comparative column.'
st['B35'] = '=COUNTA(BS_Paste!$B$1:$B$400)'

wb.save(OUT)
print('saved', OUT)
