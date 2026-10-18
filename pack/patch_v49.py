"""
July utilisation was on the wrong basis.

August and September were regrouped in patch_v19 by the department the work was
FOR, taken from the Raw data tab of her own utilisation report and filtered to
the people on her own pivot. That is the basis the P&L splits revenue on, and it
is why Video and Integration carry hours in those two months.

July was never redone. patch_v39 typed it off the Person Group totals on the
"July" tab, which groups by the department the person BELONGS to. Nobody is
assigned to Video or Integration, so both read nil for July while the work they
did was sitting inside Production and Consulting. The same tab also left the
admin cost centre out altogether.

So the one tab held two different bases, and the worked hours block further down
it was already on the job basis. July is now on the job basis too.

Source: July 2026 Utilisation Report, Raw data tab, month Jul-26, filtered to
the twenty three people on her own "July" pivot and grouped by Job Group. The
filter ties to her Grand Total to the cent: 2,863.59 chargeable, 688.54
non-chargeable, 398.50 leave.

What moves, July, chargeable / non-chargeable / leave:

  Support      2,541.09 / 34.04 / 270.50  ->  2,556.59 / 37.64 / 270.50
  Production     258.00 / 550.00 / 112.00  ->    233.42 / 529.00 / 112.00
  Video            0.00 / 0.00 / 0.00      ->     10.83 /  20.17 /   0.00
  Integration      0.00 / 0.00 / 0.00      ->     40.50 /   0.00 /   0.00
  Consulting      64.50 / 104.50 / 16.00   ->     22.25 / 100.90 /  16.00
  CTS             blank                    ->      0.00 /   0.83 /   0.00

Nothing in the P&L moves. The office staff split runs off the worked hours by
cost centre further down the tab, not off this block.
"""
import json, os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
H = json.load(open(f'{HERE}/../deck/data/util_job_basis.json'))['Jul-26']

# pack row block -> the job group it holds. Column D is July 2026.
ROWS = {'Onsite': (9, 'Support'), 'Production': (15, 'Production'),
        'Video': (21, 'Video'), 'Integration': (27, 'Integration'),
        'Consulting': (33, 'Consulting'), 'CTS': (39, 'CTS')}

wb = openpyxl.load_workbook(WB)
ws = wb['Utilisation']
n = 0
for dept, (r0, job) in ROWS.items():
    assert str(ws.cell(r0, 2).value).strip() == dept, (dept, ws.cell(r0, 2).value)
    assert str(ws.cell(r0, 3).value).strip() in ('Chargeable', 'Billable hours'), dept
    for off, key in ((0, 'ch'), (1, 'nc'), (2, 'lv')):
        ws.cell(r0 + off, 4, round(H[job][key], 2))
        n += 1

# The chart block below was renamed to plain English in patch_v48. The main
# block still said Chargeable, so the deck showed one wording on the table and
# another on the chart. Same words everywhere now.
for r0, _ in ROWS.values():
    ws.cell(r0, 3, 'Billable hours')
    ws.cell(r0 + 1, 3, 'Non-billable hours')

wb.save(WB)
print(f'July utilisation repasted on the job basis: {n} cells')
