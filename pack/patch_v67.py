"""Helper columns for the dots on the thirty six month chart.

The chart puts the bars behind in grey and the line and its dots in front, with
each year's dots in its own colour. LibreOffice ignores a marker colour set per
data point, so each year has to be its own series, fed from a column that holds
only that year and is blank everywhere else. These are those columns.

They sit out at AC, AD and AE, clear of the snip range, so they never show up in
the picture the deck uses.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NOTE = Font(name='Calibri', sz=10, color='FF6B7585')
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
FIRST = 10
SPANS = [('FY25', 29, 0, 12), ('FY26', 30, 12, 24), ('FY27', 31, 24, 36)]

wb = openpyxl.load_workbook(WB)
ws = wb['Xero History']
assert ws['A9'].value == 'Month'

ws.cell(8, 29, 'Dot series, one column per year. Working for the chart, not a report.')
ws.cell(8, 29).font = NOTE
for name, col, lo, hi in SPANS:
    ws.cell(9, col, name).font = NOTE
    for i in range(lo, hi):
        r = FIRST + i
        cell = ws.cell(r, col, f'=IF($C${r}="","",$C${r})')
        cell.number_format = MONEY
        cell.font = NOTE
    ws.column_dimensions[ws.cell(9, col).column_letter].width = 11

wb.save(WB)
print('Xero History columns AC, AD and AE added: one dot series per year')
