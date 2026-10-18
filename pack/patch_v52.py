"""
Two fixes on Overhead Split.

The column headers ran into each other. "Gross profit" and "Gross margin %" sat
side by side in two narrow columns with no gap, so on a slide they read as one
word. Columns D to H are now wide enough for their own headings.

And a chart data block for section 8, so the three departments can be shown as
a picture as well as a table. Revenue, gross profit and the result after all
overhead, for Support, Production and Consulting. It references section 8, so
there is still one source of truth.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE = 'FF242C39', 'FFFFFFFF'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']

for col, w in (('D', 13), ('E', 14), ('F', 15), ('G', 14), ('H', 13)):
    ws.column_dimensions[col].width = w

ws['A163'] = 'CHART DATA, THE THREE DEPARTMENTS'
for c in range(1, 5):
    cell = ws.cell(163, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(['Department', 'Revenue', 'Gross profit', 'Result'], start=1):
    cell = ws.cell(164, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

NAMES = {165: 'Support', 166: 'Production', 167: 'Consulting'}
for i, (r, name) in enumerate(NAMES.items()):
    src = 155 + i
    ws.cell(r, 1, name).font = Font(name='Calibri', sz=11, color=NAVY)
    for c, col in ((2, 'B'), (3, 'E'), (4, 'H')):
        ws.cell(r, c, f'={col}{src}').number_format = MONEY
        ws.cell(r, c).font = Font(name='Calibri', sz=11, color=NAVY)

wb.save(WB)
print('Overhead Split: columns widened, chart data for the three departments at rows 163 to 167')
