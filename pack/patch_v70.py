"""Gross profit by department, this September against last September.

What was asked for was profit and loss by department year on year. Only part of
that exists.

Revenue and cost of sales are held by department on both the Actual FY26 and the
P&L FY27 tabs, so gross profit by department can be compared for the two years.
Overhead was not split by department before this financial year, which is why the
department pages print "not split last year" against Total Expenses and Net
Profit/Loss. So a net result by department cannot go back a year, and September
2024 is not held by department anywhere in the pack.

The row of each department's income and cost of sales total is the lookup table
on Dept P&L rows 43 to 47, and both tabs use the same layout, so the same two
rows read either year. The column comes from the period number on Setup, so the
block follows the reporting month.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
BODY = Font(name='Calibri', sz=11, color=NAVY)
BOLD = Font(name='Calibri', sz=11, b=True, color=NAVY)
NOTE = Font(name='Calibri', sz=10, color='FF6B7585')

# department, income total row, cost of sales total row. Dept P&L rows 43 to 47.
DEPTS = [('Support', 41, 130), ('Events', 8, 61), ('Video', 16, 78),
         ('Integration', 25, 96), ('Consulting', 33, 113)]
BANDROW, HEAD, FIRST = 170, 171, 172
TOT = FIRST + len(DEPTS)
YR = 'YEAR(Setup!$B$5)'

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']

ws.cell(BANDROW, 1, 'GROSS PROFIT BY DEPARTMENT, YEAR ON YEAR')
for c in range(1, 4):
    ws.cell(BANDROW, c).fill = BAND
    ws.cell(BANDROW, c).font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, label in enumerate(['Department',
                           f'=TEXT(Setup!$B$5,"mmm")&" "&({YR}-1)',
                           f'=TEXT(Setup!$B$5,"mmm")&" "&{YR}'], start=1):
    cell = ws.cell(HEAD, c, label)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

for i, (name, inc, cos) in enumerate(DEPTS):
    r = FIRST + i
    ws.cell(r, 1, name).font = BODY
    for c, tab in ((2, 'Actual FY26'), (3, 'P&L FY27')):
        cell = ws.cell(r, c, f"=INDEX('{tab}'!$B${inc}:$M${inc},Setup!$B$8)"
                             f"+INDEX('{tab}'!$B${cos}:$M${cos},Setup!$B$8)")
        cell.number_format = MONEY
        cell.font = BODY

ws.cell(TOT, 1, 'Total').font = BOLD
for c in (2, 3):
    cell = ws.cell(TOT, c, f'=SUM({chr(64+c)}{FIRST}:{chr(64+c)}{TOT-1})')
    cell.number_format = MONEY
    cell.font = BOLD
for c in range(1, 4):
    ws.cell(TOT, c).fill = TOTAL

ws.cell(TOT + 2, 1, 'Gross profit only. Overhead was not split by department '
                    'before FY27, so a net result cannot go back a year.').font = NOTE

wb.save(WB)
print(f'Overhead Split rows {BANDROW} to {TOT} added: gross profit by '
      f'department, two years')
