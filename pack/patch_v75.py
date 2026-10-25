"""Allocate other income and other expenses the way the P&L analysis does.

The admin pot carries interest received, income other and interest expense.
'Dept P&L' swept them into its "Other admin overhead" plug, so they sat inside
each department's overhead as a credit.  The analysis allocates them on the
Staff key but shows them below the line, as other income.

So: 'P&L by Dept' picks them up on its own "Other income" and "Other Expenses"
rows, and 'Dept P&L' stops carrying them in the plug.  The two tabs then show
the same overhead for the same department and month.
"""
import re
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
OS = "'Overhead Split'"
SUB = f'{OS}!$U$23:$U$86'
AMT = f'{OS}!$C$23:$N$86'
HDR = f'{OS}!$C$22:$N$22'

# name, date row, other income row, other expenses row, key columns
BLOCKS = [
    ('Production',    6,  24,  69, ['$Q$', '$R$']),
    ('Consulting',   83, 101, 146, ['$S$', '$T$']),
    ('Support',     160, 169, 196, ['$P$']),
    ('General Admin', 210, 219, 246, []),
]
COLS = [openpyxl.utils.get_column_letter(c) for c in range(2, 14)]   # B..M


def alloc(col, date_row, keys, sub):
    pick = f'INDEX({AMT},0,MATCH({col}${date_row},{HDR},0))'
    test = f'({SUB}="{sub}")'
    if not keys:
        return f'-SUMPRODUCT({test}*{pick})'
    key = '+'.join(f'{OS}!{k}23:{k}86' for k in keys)
    key = key if len(keys) == 1 else f'({key})'
    return f'SUMPRODUCT({test}*{key}*{pick})'


wb = openpyxl.load_workbook(WB)
ws = wb['P&L by Dept']
changed = 0
for name, date_row, inc_row, exp_row in [(b[0], b[1], b[2], b[3]) for b in BLOCKS]:
    keys = dict((b[0], b[4]) for b in BLOCKS)[name]
    for row, sub in ((inc_row, 'Other Income'), (exp_row, 'Other Expenses')):
        for col in COLS:
            cell = ws[f'{col}{row}']
            old = cell.value
            m = re.fullmatch(r'=IF\((.+?)<=Setup!\$B\$5,(.+),(.+)\)', old)
            if not m:
                raise SystemExit(f'{name} {col}{row}: unexpected shape: {old}')
            when, direct, forecast = m.groups()
            cell.value = (f'=IF({when}<=Setup!$B$5,{direct}'
                          f'+{alloc(col, date_row, keys, sub)},{forecast})')
            changed += 1
    print(f'  {name:<14} other income row {inc_row}, other expenses row {exp_row}')

# 'Dept P&L' plug: the department's whole admin share less the six named
# categories.  Narrow it to the 'Other' subcategory so interest and other
# income no longer sit inside overhead.
dp = wb['Dept P&L']
old = dp['B22'].value
new = ('=INDEX(OFFSET(Engine!$A$1,INDEX($F$43:$F$48,MATCH($B$5,$A$43:$A$48,0))+22,'
       'MATCH(Setup!$B$5,Engine!$E$4:$P$4,0)+3,6,1),1)*0'
       f'+IF(INDEX($G$43:$G$48,MATCH($B$5,$A$43:$A$48,0))=0,'
       f'-SUMPRODUCT(({SUB}="Other")*INDEX({AMT},0,MATCH(Setup!$B$5,{HDR},0))),'
       f'SUMPRODUCT(({SUB}="Other")*INDEX({OS}!$P$23:$T$86,0,'
       'INDEX($G$43:$G$48,MATCH($B$5,$A$43:$A$48,0)))'
       f'*INDEX({AMT},0,MATCH(Setup!$B$5,{HDR},0))))')
dp['B22'] = new
print('  Dept P&L       B22 narrowed to the Other subcategory')

wb.save(WB)
print(f'{changed + 1} cells rewritten, saved')
