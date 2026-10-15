"""Push data through the repaired pack and check every figure ties."""
import json, random, datetime as dt, shutil
import openpyxl
HERE = __file__.rsplit('/', 1)[0]
SRC, TST = f'{HERE}/CTS Financial Controller Pack - fixed.xlsx', f'{HERE}/_t.xlsx'
MASTER = json.load(open(f'{HERE}/data/lists_master.json'))
BSACC = json.load(open(f'{HERE}/data/bs_accounts.json'))
random.seed(11)
CC = {'ONSITE': 'ONSITE', 'PRODUCTION': 'PRODUCTION', 'VIDEO': 'VIDEO',
      'INTEGRATION': 'INTEGRATION', 'CONSULTING': 'CONSULTING', 'ADMIN': 'CTS'}
MONTHS = [dt.date(2026, 7, 15), dt.date(2026, 8, 15), dt.date(2026, 9, 15)]
pl_accts = [m for m in MASTER
            if m['cat'] in ('Income', 'Cost of Sales', 'Expenses', 'Other Income', 'Other Expenses')]
rows, exp_sub, exp_add = [], {}, {}
def add(d, code, name, dr, cr, cc, job=''):
    rows.append([len(rows) + 1, code, name, 'Invoice', d, 'Test', dr or None, cr or None,
                 job, '', '', 'test', cc])
def book(d, p, dept, m, dr, cr, cc, job=''):
    add(d, m['code'], m['name'], dr, cr, cc, job)
    sgn = cr - dr
    exp_sub[(dept, m['cat'], m['sub'], p)] = round(exp_sub.get((dept, m['cat'], m['sub'], p), 0) + sgn, 2)
    if m['addback']:
        exp_add[(dept, m['addback'], p)] = round(exp_add.get((dept, m['addback'], p), 0) + sgn, 2)
for p, d in enumerate(MONTHS, 1):
    for dept, tag in CC.items():
        for m in random.sample(pl_accts, 24):
            amt = round(random.uniform(300, 8000), 2)
            income = m['cat'] in ('Income', 'Other Income')
            book(d, p, dept, m, 0 if income else amt, amt if income else 0, tag)
byc = {m['code']: m for m in MASTER}
book(MONTHS[2], 3, 'ADMIN', byc['61300'], 12000, 0, 'CTS')     # depreciation
book(MONTHS[2], 3, 'ADMIN', byc['61730'], 71153, 0, 'CTS')     # instant asset write off
book(MONTHS[2], 3, 'ADMIN', byc['92000'], 30000, 0, 'CTS')     # income tax
book(MONTHS[2], 3, 'PRODUCTION', byc['42100'], 0, 5000, '', 'J9 [PRD]')   # bracket fallback
add(dt.date(2026, 6, 20), '41100', 'Contract Support Staff', 0, 999999, 'ONSITE')  # prior FY
add(MONTHS[2], '99999', 'Nope', 0, 55555, 'ONSITE')                                # unknown
add(MONTHS[2], '41100', 'Contract Support Staff', 0, 777, 'NOWHERE')               # bad tag

shutil.copy(SRC, TST)
wb = openpyxl.load_workbook(TST)
gp = wb['GL_Paste']
for i, row in enumerate(rows):                      # her paste zone starts at A2
    for c, v in enumerate(row, 1):
        gp.cell(2 + i, c, v)
CATS = ['Income', 'Cost of Sales', 'Expenses', 'Other Income', 'Other Expenses']
sep = {c: round(sum(v for (d, cat, s, p), v in exp_sub.items() if cat == c and p == 3), 2)
       for c in CATS}
pp = wb['PL_Paste']
pp.cell(1, 1, 'Account'); pp.cell(1, 2, 'Sep 2026')
for i, (c, lbl) in enumerate(zip(CATS, ['Total Income', 'Total Cost of Sales',
                                        'Total Operating Expenses', 'Total Other Income',
                                        'Total Other Expenses'])):
    pp.cell(3 + i, 1, lbl); pp.cell(3 + i, 2, sep[c])
pc = wb['PL_Check']                                  # her typed column, same figures
for i, c in enumerate(CATS): pc.cell(5 + i, 3, sep[c])
bp = wb['BS_Paste']
bp.cell(1, 1, 'Account'); bp.cell(1, 2, '30 Sep 2026')
bp.cell(1, 3, '31 Aug 2026'); bp.cell(1, 4, '1 Jul 2026')
bal, r = {}, 3
for a in BSACC:
    if a['group'] == 'Equity': continue
    v = round(random.uniform(1000, 90000), 2); bal[a['code']] = v
    bp.cell(r, 1, a['name'])
    for c, k in ((2, 1.0), (3, 0.97), (4, 0.9)): bp.cell(r, c, round(v * k, 2))
    r += 1
grp = {a['code']: a['group'] for a in BSACC}
assets = round(sum(v for c, v in bal.items() if grp[c].endswith('Assets')), 2)
liabs = round(sum(v for c, v in bal.items() if 'Liabilities' in grp[c]), 2)
eq = round(assets - liabs, 2)
bp.cell(r, 1, 'Retained Earnings')
for c, k in ((2, 1.0), (3, 0.97), (4, 0.9)): bp.cell(r, c, round(eq * k, 2))
wb.save(TST)
json.dump({'sub': {'|'.join(map(str, k)): v for k, v in exp_sub.items()},
           'add': {'|'.join(map(str, k)): v for k, v in exp_add.items()}, 'sep': sep,
           'assets': assets, 'liabs': liabs, 'eq': eq}, open(f'{HERE}/_exp.json', 'w'))
print(f'{len(rows)} GL lines; Sep per the GL:', {k: round(v, 2) for k, v in sep.items()})
