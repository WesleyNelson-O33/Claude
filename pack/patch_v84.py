"""Label the FY26 tabs with FY26 dates.

Actual FY26 and Budget FY26 hold the prior year, July 2025 to June 2026, but
their month row was written a year forward, July 2026 to June 2027, and the
Month cell at the top read the current reporting month rather than the
comparative one. So a tab of September 2025 figures was headed Sep-26.

The formulas that read those tabs compensated for it, in two different ways.
Summary added twelve months to its own date before matching. Xero History took
a year off the label it printed. Everything else matched the current reporting
month against a row that was a year out and so landed on the right column by
accident.

This puts the dates right and takes the compensation out: the month row becomes
July 2025 to June 2026, the Month cell reads the comparative month from Setup
B10, anything matching a date against those rows now matches B10 rather than B5,
Summary drops its twelve month shift, and Xero History stops subtracting a year
from a label that is now correct.

Not one figure in the pack changes. Only the labels.
"""
import os
import re
import sys
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

FY26 = ('Actual FY26', 'Budget FY26')
# July 2025 to June 2026, the year those tabs actually hold
MONTHS = [datetime(2025 + (0 if m >= 7 else 1), m, 1)
          for m in list(range(7, 13)) + list(range(1, 7))]
ROW5 = re.compile(r"'(Actual FY26|Budget FY26)'!\$B\$5:\$M\$5")
EDATE12 = re.compile(r"EDATE\(([A-Z]{1,2})\$5,12\)")
YEARLESS1 = re.compile(r"\(YEAR\('Actual FY26'!([A-Z]{1,2})\$5\)-1\)")

wb = openpyxl.load_workbook(WB)

# 1. the month row and the Month cell on the two FY26 tabs
for t in FY26:
    ws = wb[t]
    assert ws['A5'].value == 'Line' and ws['A3'].value == 'Month', t
    for i, d in enumerate(MONTHS):
        c = ws.cell(5, 2 + i, d)
        c.number_format = 'mmm-yy'
    ws['B3'] = '=Setup!$B$10'
    print(f'  {t}: month row set to {MONTHS[0]:%b %Y} .. {MONTHS[-1]:%b %Y}, '
          f'Month cell now reads Setup B10')

# 2. anything that matches a date against those rows moves to the comparative month
n_match = 0
for ws in wb.worksheets:
    if ws.title in ('Cleanup', 'P&L Analysis'):
        continue
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or not ROW5.search(c.value):
                continue
            if 'Setup!$B$5' in c.value:
                c.value = c.value.replace('Setup!$B$5', 'Setup!$B$10')
                n_match += 1
print(f'  {n_match} formulas now match Setup B10 instead of B5')

# 3. Summary drops the twelve month shift
n_sum = 0
ws = wb['Summary']
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and EDATE12.search(c.value) and ROW5.search(c.value):
            c.value = EDATE12.sub(r'\1$5', c.value)
            n_sum += 1
print(f'  {n_sum} Summary formulas no longer shift the date forward a year')

# 4. Xero History stops taking a year off the label
n_lab = 0
ws = wb['Xero History']
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and YEARLESS1.search(c.value):
            c.value = YEARLESS1.sub(r"YEAR('Actual FY26'!\1$5)", c.value)
            n_lab += 1
print(f'  {n_lab} Xero History labels no longer subtract a year')

assert n_match and n_sum and n_lab, 'one of the three rewrites found nothing'
wb.save(WB)
print('saved')
