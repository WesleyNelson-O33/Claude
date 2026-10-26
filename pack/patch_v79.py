"""Allocate the admin depreciation so EBITDA and the EBITDA margin are right.

The company's depreciation is the Instant Asset Write Off account.  Most of it
is coded to admin, so it sits in the admin pot and is spread across the
departments on the Office Dept key, inside the General & Administrative Exp
line.  The depreciation row on each block, though, reads the Engine, which only
ever gives what is coded direct to that department.

Since the blocks now carry allocated overhead, those two no longer matched.
Support's overhead held 175 of allocated write off in September with nothing to
add back, Production's held 8,573 with only its own 184 added back, and General
Admin added back 9,798 that its contra had already taken out, so it showed a
phantom EBITDA on no revenue.

Each block's depreciation row now adds its share of the admin write off, and
General Admin carries the contra.  EBITDA, EBIT and both their margins read
correctly on every block, and the company depreciation is unchanged because the
five keys add to one.
"""
import re
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
OS = "'Overhead Split'"
IAWO = 52                      # Overhead Split row: Instant Asset Write Off, admin
HDR = f'{OS}!$C$22:$N$22'

# row, date row, key columns on the overhead split ([] = the contra)
BLOCKS = [(72, 6, ['$Q$', '$R$']), (149, 83, ['$S$', '$T$']),
          (199, 160, ['$P$']), (249, 210, [])]
COLS = [openpyxl.utils.get_column_letter(c) for c in range(2, 14)]

wb = openpyxl.load_workbook(WB)
ws = wb['P&L by Dept']
osheet = wb['Overhead Split']
assert osheet.cell(IAWO, 1).value == 'Instant Asset Write Off', osheet.cell(IAWO, 1).value
assert osheet.cell(IAWO, 2).value == 'Office Dept'

changed = 0
for row, date_row, keys in BLOCKS:
    for col in COLS:
        cell = ws[f'{col}{row}']
        old = cell.value
        m = re.fullmatch(r'=IF\((.+?)<=Setup!\$B\$5,(.+),0\)', old)
        if not m:
            raise SystemExit(f'{col}{row}: unexpected shape: {old}')
        when, direct = m.groups()
        amount = f'INDEX({OS}!$C{IAWO}:$N{IAWO},1,MATCH({col}${date_row},{HDR},0))'
        if keys:
            key = '+'.join(f'{OS}!{k}{IAWO}' for k in keys)
            share = f'({key})*{amount}' if len(keys) > 1 else f'{key}*{amount}'
            share = f'+IF(Setup!$B$28="Yes",{share},0)'
        else:
            share = f'-IF(Setup!$B$28="Yes",{amount},0)'
        cell.value = f'=IF({when}<=Setup!$B$5,{direct}{share},0)'
        changed += 1
    print(f'  row {row:<5}{"contra" if not keys else "+ " + " ".join(keys)}')

wb.save(WB)
print(f'{changed} cells rewritten, saved')
