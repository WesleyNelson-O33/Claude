"""Push data through the new layout and check every figure ties."""
import json, random, datetime as dt, shutil, subprocess, sys
import openpyxl
HERE = __file__.rsplit('/', 1)[0]
SRC, TST = f'{HERE}/CTS Financial Controller Pack - v4.xlsx', f'{HERE}/_v4.xlsx'
MASTER = json.load(open(f'{HERE}/data/lists_master.json'))
BSACC = json.load(open(f'{HERE}/data/bs_accounts.json'))
random.seed(23)
CC = {'ONSITE': 'ONSITE', 'PRODUCTION': 'PRODUCTION', 'VIDEO': 'VIDEO',
      'INTEGRATION': 'INTEGRATION', 'CONSULTING': 'CONSULTING', 'ADMIN': 'CTS'}
MONTHS = [dt.date(2026, 7, 15), dt.date(2026, 8, 15), dt.date(2026, 9, 15)]
pl_accts = [m for m in MASTER
            if m['cat'] in ('Income', 'Cost of Sales', 'Expenses', 'Other Income', 'Other Expenses')]
rows, exp = [], {}
def gl(d, code, name, dr, cr, cc, job=''):
    rows.append([len(rows) + 1, code, name, 'Invoice', d, 'Test', dr or None, cr or None,
                 job, '', '', 'test', cc])
def book(d, p, dept, m, dr, cr, cc, job=''):
    gl(d, m['code'], m['name'], dr, cr, cc, job)
    k = (dept, m['cat'], m['sub'], p)
    exp[k] = round(exp.get(k, 0) + cr - dr, 2)
    if m['addback']:
        a = (dept, m['addback'], p); exp[a] = round(exp.get(a, 0) + cr - dr, 2)
for p, d in enumerate(MONTHS, 1):
    for dept, tag in CC.items():
        for m in random.sample(pl_accts, 24):
            amt = round(random.uniform(300, 8000), 2)
            inc = m['cat'] in ('Income', 'Other Income')
            book(d, p, dept, m, 0 if inc else amt, amt if inc else 0, tag)
byc = {m['code']: m for m in MASTER}
book(MONTHS[2], 3, 'ADMIN', byc['61300'], 12000, 0, 'CTS')
book(MONTHS[2], 3, 'ADMIN', byc['61730'], 71153, 0, 'CTS')
book(MONTHS[2], 3, 'ADMIN', byc['92000'], 30000, 0, 'CTS')
book(MONTHS[2], 3, 'PRODUCTION', byc['42100'], 0, 5000, '', 'J9 [PRD]')

shutil.copy(SRC, TST)
wb = openpyxl.load_workbook(TST)
gp = wb['GL_Paste']
gp.cell(1, 1, 'Corporate Technology Services Pty Ltd')
gp.cell(2, 1, 'Account Transactions for P&L analysis')
for i, h in enumerate(['Index', 'Account Code', 'Account Name', 'Source', 'Date', 'Contact',
                       'Debit', 'Credit', 'Job Numbers', 'Invoice Number', 'Reference',
                       'Description', 'Cost Centres']):
    gp.cell(3, 1 + i, h)
for i, row in enumerate(rows):
    for c, v in enumerate(row, 1): gp.cell(4 + i, c, v)
CATS = ['Income', 'Cost of Sales', 'Expenses', 'Other Income', 'Other Expenses']
sep = {c: round(sum(v for k, v in exp.items() if len(k) == 4 and k[1] == c and k[3] == 3), 2)
       for c in CATS}
pp = wb['PL_Paste']; pp.cell(1, 1, 'Account'); pp.cell(1, 2, 'Sep 2026')
for i, (c, lbl) in enumerate(zip(CATS, ['Total Income', 'Total Cost of Sales',
                                        'Total Operating Expenses', 'Total Other Income',
                                        'Total Other Expenses'])):
    pp.cell(3 + i, 1, lbl); pp.cell(3 + i, 2, sep[c])
bp = wb['BS_Paste']; bp.cell(1, 1, 'Account')
for c, h in ((2, '30 Sep 2026'), (3, '31 Aug 2026'), (4, '1 Jul 2026')): bp.cell(1, c, h)
bal, r = {}, 3
for a in BSACC:
    if a['group'] == 'Equity': continue
    v = round(random.uniform(1000, 90000), 2); bal[a['code']] = v
    bp.cell(r, 1, a['name'])
    for c, k in ((2, 1.0), (3, 0.97), (4, 0.9)): bp.cell(r, c, round(v * k, 2))
    r += 1
grp = {a['code']: a['group'] for a in BSACC}
eqp = round(sum(v for c, v in bal.items() if grp[c].endswith('Assets'))
            - sum(v for c, v in bal.items() if 'Liabilities' in grp[c]), 2)
bp.cell(r, 1, 'Retained Earnings')
for c, k in ((2, 1.0), (3, 0.97), (4, 0.9)): bp.cell(r, c, round(eqp * k, 2))
wb.save(TST)
print(f'{len(rows)} GL lines; Sep per the GL:', sep)
j = json.loads(subprocess.run([sys.executable, '/mnt/skills/public/xlsx/scripts/recalc.py',
                               TST, '1200'], capture_output=True, text=True).stdout)
print('recalc:', j['status'], j['total_errors'], 'errors')

v = openpyxl.load_workbook(TST, data_only=True)
pl = v['P&L FY27']
LBL = {pl.cell(r, 1).value: r for r in range(6, 178) if pl.cell(r, 1).value}
def R(name): return LBL[name]
ok = True
def check(lab, got, want):
    global ok
    good = abs(round(got or 0, 2) - round(want, 2)) < 0.02
    ok = ok and good
    print(f'  {lab:<34}{(got or 0):>15,.2f} vs {want:>15,.2f}  {"OK" if good else "MISMATCH"}')

YTD = 15
inc = pl.cell(R('Total Income '), YTD).value
cos = pl.cell(R('    Cost of Sales'), YTD).value
ovh = pl.cell(R('    Expenses Overheds '), YTD).value
oi = pl.cell(R('Other income '), YTD).value
oe = pl.cell(R('    Other Expenses'), YTD).value
dep = pl.cell(R('    Depreciation & amortisation'), YTD).value
print('=== statement ties, year to date ===')
check('Total Income = GL income', inc, sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Income'))
check('Cost of Sales = GL cost of sales', cos, sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Cost of Sales'))
check('Overheads = GL expenses', ovh, sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Expenses'))
check('Other income = GL other income', oi, sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Other Income'))
check('Other expenses = GL other exp', oe, sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Other Expenses'))
check('Gross profit', pl.cell(R('    Gross profit'), YTD).value, inc + cos)
check('EBITDA', pl.cell(R('    EBITDA'), YTD).value, inc + cos + ovh - dep)
check('EBIT', pl.cell(R('    EBIT'), YTD).value, inc + cos + ovh)
check('Net profit', pl.cell(6, YTD).value, inc + cos + ovh + oi + oe)
check('September net vs the Xero P&L', pl.cell(6, 4).value, sum(sep.values()))
print('=== income by department adds back to the total ===')
dsum = 0
for d in ['Production', 'Video', 'Onsite', 'Integration', 'Consulting', 'General Admin']:
    rr = [r for r in range(6, 178) if pl.cell(r, 1).value == '    ' + d]
    want = sum(v_ for k, v_ in exp.items() if len(k) == 4 and k[1] == 'Income'
               and k[0] == {'Production': 'PRODUCTION', 'Video': 'VIDEO', 'Onsite': 'ONSITE',
                            'Integration': 'INTEGRATION', 'Consulting': 'CONSULTING',
                            'General Admin': 'ADMIN'}[d])
    got = pl.cell(rr[0], YTD).value or 0
    dsum += got
    check(f'  {d} income', got, want)
check('six departments = Total Income', dsum, inc)
print('=== every tab has the same rows ===')
for nm in ['Summary', 'Budget FY27', 'Actual FY26', 'Budget FY26']:
    w = v[nm]
    same = all(w.cell(r, 1).value == pl.cell(r, 1).value for r in range(6, 178))
    print(f'  {nm:<14}{"identical" if same else "DIFFERENT"}')
    ok = ok and same
print('=== profit margin per department, each month ===')
NAMES = {'Production': 'PRODUCTION', 'Video': 'VIDEO', 'Onsite': 'ONSITE',
         'Integration': 'INTEGRATION', 'Consulting': 'CONSULTING', 'General Admin': 'ADMIN'}
mhead = [r for r in range(6, 200) if pl.cell(r, 1).value == 'Profit Margin Per Department'][0]
for i, (disp, code) in enumerate(NAMES.items()):
    mr = mhead + 1 + i
    for mcol, p_ in ((2, 1), (3, 2), (4, 3)):
        inc_ = sum(x for k, x in exp.items() if len(k) == 4 and k[1] == 'Income'
                   and k[0] == code and k[3] == p_)
        cos_ = sum(x for k, x in exp.items() if len(k) == 4 and k[1] == 'Cost of Sales'
                   and k[0] == code and k[3] == p_)
        want = (inc_ + cos_) / inc_ if inc_ else 0
        got = pl.cell(mr, mcol).value or 0
        check(f'  {disp} margin month {p_}', got * 100, want * 100)
    # and the year to date column
    inc_ = sum(x for k, x in exp.items() if len(k) == 4 and k[1] == 'Income' and k[0] == code)
    cos_ = sum(x for k, x in exp.items() if len(k) == 4 and k[1] == 'Cost of Sales' and k[0] == code)
    check(f'  {disp} margin YTD', (pl.cell(mr, 15).value or 0) * 100,
          ((inc_ + cos_) / inc_ if inc_ else 0) * 100)
print('=== the margin block is on every tab ===')
for nm in ['Summary', 'Budget FY27', 'Actual FY26', 'Budget FY26']:
    w = v[nm]
    same = all(w.cell(r, 1).value == pl.cell(r, 1).value for r in range(6, mhead + 7))
    print(f'  {nm:<14}{"identical" if same else "DIFFERENT"}')
    ok = ok and same

b = v['Budget FY27']
print('=== her budget figures survive ===')
check('Budget income FY', b.cell(R('Total Income '), 14).value, 6304600.65)
check('Budget cost of sales FY', b.cell(R('    Cost of Sales'), 14).value, -3379146.73)
check('Budget overheads FY', b.cell(R('    Expenses Overheds '), 14).value, -2528203.17)
check('Budget net FY', b.cell(6, 14).value, 397250.75)
print('=== controls ===')
pc = v['PL_Check']
print('  PL_Check:', [pc.cell(5 + i, 10).value for i in range(5)], '| net', pc.cell(10, 10).value)
print('  Balance sheet:', v['Balance_Sheet'].cell(100, 6).value)
print()
print('ALL CHECKS PASSED' if ok else 'SOMETHING IS WRONG')
