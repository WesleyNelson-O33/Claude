"""
Two corrections to the year on year revenue block.

The P&L FY27 tab carries budget in the months that are not closed yet, so a
straight reference to it put October at 950,416 and charted a budget figure as
though it were actual. FY27 now stops at the reporting month.

With FY27 a part year, a "Full year" row comparing it to full years of FY23 to
FY26 would be misleading, so that row is gone. The first quarter block below is
the like for like total.
"""
import os, sys
import openpyxl
from openpyxl.styles import Border, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
FIRST, LAST, TOTROW = 282, 293, 294

wb = openpyxl.load_workbook(WB)
ws = wb['Cost Analysis']
assert ws.cell(281, 6).value == 'FY27'

for i, r in enumerate(range(FIRST, LAST + 1)):
    col = chr(ord('B') + i)
    ws.cell(r, 6, f"=IF('P&L FY27'!{col}$5<=Setup!$B$5,'P&L FY27'!{col}6,\"\")")

for c in range(1, 7):
    cell = ws.cell(TOTROW, c)
    cell.value = None
    cell.fill = PatternFill(fill_type=None)
    cell.border = Border()

wb.save(WB)
print('FY27 now stops at the reporting month; the full year row is removed')
