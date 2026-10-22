"""A cumulative budget line beside the three cumulative years.

The cumulative chart showed FY27 ahead of both completed years and said nothing
about plan. Both facts belong on it: the year is ahead of where it was and
behind where it was asked to be.

Budget is the Total Income row of the Budget FY27 tab, cumulated the same way
as the three actual years. It runs the full twelve months, because a budget is
set for the whole year.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE = 'FF242C39', 'FFFFFFFF'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'

CUM_HEAD, CUM_FIRST = 57, 58

wb = openpyxl.load_workbook(WB)
ws = wb['Xero History']
assert ws[f'A{CUM_HEAD - 1}'].value == 'REVENUE YEAR TO DATE, CUMULATIVE'

cell = ws.cell(CUM_HEAD, 5, 'FY27 budget')
cell.fill = BAND
cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
cell.border = Border(bottom=RULE)
cell.alignment = Alignment(horizontal='right')

for i in range(12):
    col = chr(ord('B') + i)                       # B..M on Budget FY27
    cell = ws.cell(CUM_FIRST + i, 5, f"=SUM('Budget FY27'!$B$6:{col}$6)")
    cell.number_format = MONEY
    cell.font = Font(name='Calibri', sz=11, color=NAVY)

# the band above now needs to cover the extra column
for c in range(1, 6):
    b = ws.cell(CUM_HEAD - 1, c)
    b.fill = BAND
    b.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

wb.save(WB)
print('Xero History column E added: FY27 budget, cumulated')
