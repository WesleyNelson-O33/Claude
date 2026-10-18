"""
The three departments, the way the P&L groups them.

The Summary tab has three trading departments, not five. Production is the
parent of Events and Video. Consulting is the parent of Integration and
Consulting. Support stands on its own. The pack only ever showed the five
children, so nobody could see the three numbers the P&L is actually built on.

Section 8 on Overhead Split rolls section 7 up to those three. It references
section 7 rather than repeating its formulas, so there is one source of truth
and the two sections cannot drift. Gross margin is recalculated on each rolled
up line, because a margin cannot be added together.

The check at the bottom is section 8's result against section 7's. It is zero.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
PCT = '0.0%'

HEAD = ['Department', 'Revenue', 'Budget', 'Var', 'Gross profit',
        'Gross margin %', 'All overhead', 'Result']
# label -> the section 7 rows it rolls up
GROUPS = [('Support', [143]),
          ('Production (Events and Video)', [144, 145]),
          ('Consulting (Integration and Consulting)', [146, 147])]

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws['A140'].value.startswith('7.'), ws['A140'].value
assert ws['A148'].value == 'Total', ws['A148'].value

ws['A152'] = '8. THE THREE DEPARTMENTS, AS THE P&L GROUPS THEM'
for c in range(1, 9):
    cell = ws.cell(152, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(HEAD, start=1):
    cell = ws.cell(154, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

for i, (label, rows) in enumerate(GROUPS):
    r = 155 + i
    ws.cell(r, 1, label).font = Font(name='Calibri', sz=11, color=NAVY)
    for c in (2, 3, 4, 5, 7, 8):
        col = chr(64 + c)
        ws.cell(r, c, '=' + '+'.join(f'{col}{s}' for s in rows)).number_format = MONEY
        ws.cell(r, c).font = Font(name='Calibri', sz=11, color=NAVY)
    ws.cell(r, 6, f'=IFERROR(E{r}/B{r},0)').number_format = PCT
    ws.cell(r, 6).font = Font(name='Calibri', sz=11, color=NAVY)

ws.cell(158, 1, 'Total')
for c in (2, 3, 4, 5, 7, 8):
    col = chr(64 + c)
    ws.cell(158, c, f'=SUM({col}155:{col}157)').number_format = MONEY
ws.cell(158, 6, '=IFERROR(E158/B158,0)').number_format = PCT
for c in range(1, 9):
    cell = ws.cell(158, c)
    cell.fill = TOTAL
    cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)

ws.cell(160, 1, 'Check, against the five department result above')
ws.cell(160, 1).font = Font(name='Calibri', sz=11, b=True, color=NAVY)
ws.cell(160, 2, '=H158-H148').number_format = MONEY
ws.cell(160, 2).font = Font(name='Calibri', sz=11, b=True, color=NAVY)

wb.save(WB)
print('Overhead Split section 8 added: the three departments, rows 152 to 160')
