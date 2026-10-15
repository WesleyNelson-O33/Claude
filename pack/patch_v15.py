"""Give Actual FY26 and Budget FY26 the YTD column the Summary already asks for.

The Summary has P6 = E6 - 'Actual FY26'!O6 and U6 = G6 - 'Budget FY26'!O6.
Neither of those tabs has a column O, so both formulas were subtracting a
blank cell. They looked like they worked and did not: P was simply repeating
this year's YTD actual and U was repeating this year's YTD budget.

Column O is added to both tabs, built exactly like the YTD column that
already exists on P&L FY27. Nothing else is touched. Columns Q, R and S on
the Summary are still empty, because they are her headings and I do not know
what she wants in them.
"""
import os, shutil, copy
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = f'{HERE}/CTS Financial Controller Pack - v14.xlsx'
OUT = f'{HERE}/CTS Financial Controller Pack - v15.xlsx'

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

src = wb['P&L FY27']
rows = [r for r in range(6, src.max_row + 1)
        if isinstance(src.cell(r, 15).value, str) and src.cell(r, 15).value.startswith('=')]
print('rows that carry a YTD figure on P&L FY27:', len(rows), rows[0], '..', rows[-1])

for name in ('Actual FY26', 'Budget FY26'):
    ws = wb[name]
    assert ws.cell(5, 14).value == 'FY total', (name, ws.cell(5, 14).value)
    assert ws.cell(5, 15).value is None, (name, 'column O is already used')
    hdr = ws.cell(5, 15)
    hdr.value = 'YTD'
    hdr._style = copy.copy(ws.cell(5, 14)._style)
    for r in rows:
        if not isinstance(ws.cell(r, 14).value, str):
            continue                      # a spacer or a heading, not a figure row
        c = ws.cell(r, 15)
        c.value = (f"=SUMIF('{name}'!$B$5:$M$5,\"<=\"&Setup!$B$5,'{name}'!$B{r}:$M{r})")
        c._style = copy.copy(ws.cell(r, 14)._style)
    ws.column_dimensions['O'].width = ws.column_dimensions['N'].width
    print(f'  {name}: YTD column written')

wb.save(OUT)
print('saved', OUT)
