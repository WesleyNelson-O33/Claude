"""Fill the utilisation gaps from the raw timesheet in her own report.

The September and August utilisation reports carry a Raw data tab, 49,827
lines, with every timesheet entry: who worked, which job, whether it was
chargeable, and how many hours. Her September sheet is a pivot of that,
grouped by the department the PERSON belongs to. That is why Video and
Integration were blank: nobody is assigned to them, their work is done by
Production and Consulting people.

Grouping the same lines by the department the work was FOR fills every row,
and it is the basis the P&L already uses, because revenue is split by the
job's cost centre. Totals are unchanged, only the split.

Checked against her own sheets on her own list of people. August ties to the
cent. September is one hour apart on the chargeable and non-chargeable split,
and identical in total.
"""
import os, json, shutil
import openpyxl
from openpyxl.styles import Font, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = f'{HERE}/CTS Financial Controller Pack - v18.xlsx'
OUT = f'{HERE}/CTS Financial Controller Pack - v19.xlsx'
H = json.load(open('/home/user/Claude/deck/data/util_job_basis.json'))
MID = '5F6670'

# pack row block -> her job group
ROWS = {'Onsite': 9, 'Production': 15, 'Video': 21, 'Integration': 27,
        'Consulting': 33, 'CTS': 39}
JOB = {'Onsite': 'Support', 'Production': 'Production', 'Video': 'Video',
       'Integration': 'Integration', 'Consulting': 'Consulting', 'CTS': 'CTS'}
COL = {'Aug-26': 5, 'Sep-26': 6}

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb['Utilisation']

n = 0
for dept, r0 in ROWS.items():
    assert str(ws.cell(r0, 2).value).strip() == dept, (dept, ws.cell(r0, 2).value)
    assert str(ws.cell(r0, 3).value).strip() == 'Chargeable', dept
    for mon, c in COL.items():
        x = H[mon][JOB[dept]]
        for off, key in ((0, 'ch'), (1, 'nc'), (2, 'lv')):
            ws.cell(r0 + off, c, round(x[key], 2))
            n += 1
print('hour cells written:', n)

ws['A45'] = ('Hours come from the Raw data tab of the monthly utilisation report, grouped by '
             'the department the work was for, which is the same basis the P&L splits revenue '
             'on. Her own sheet groups by the department the person belongs to, so the two '
             'will not agree by department. The totals are the same either way.')
ws['A45'].font = Font(name='Calibri', sz=9, italic=True, color=MID)
ws['A45'].alignment = Alignment(horizontal='left', vertical='center')
wb.save(OUT)
print('saved', OUT)
