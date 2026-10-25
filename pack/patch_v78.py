"""Put the 'Dept P&L' plug back to the whole unnamed share.

patch_v75 narrowed it to the 'Other' subcategory so that page's overhead would
match the overhead line on 'P&L by Dept'.  That broke something more important:
'Overhead Split' section 7, which drives the deck's departmental result chart
and the result table, carries the whole admin share, so its Result and this
page's Net Profit/Loss no longer agreed.

Section 7's Result is also the figure the analysis calls Net Profit/(Loss): it
carries the allocated interest.  So the plug goes back and the two tabs agree
on the department result.  They present the interest differently, which is the
point: 'P&L by Dept' shows it below the line as the analysis does, this page
keeps it in the plug.
"""
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
ORIGINAL = (
    "=SUM(OFFSET(Engine!$A$1,INDEX($F$43:$F$48,MATCH($B$5,$A$43:$A$48,0))+22,"
    "MATCH(Setup!$B$5,Engine!$E$4:$P$4,0)+3,6,1))"
    "+IF(INDEX($G$43:$G$48,MATCH($B$5,$A$43:$A$48,0))=0,"
    "-INDEX('Overhead Split'!$C$95:$N$95,1,"
    "MATCH(Setup!$B$5,'Overhead Split'!$C$22:$N$22,0)),"
    "INDEX('Overhead Split'!$C$90:$N$94,"
    "INDEX($G$43:$G$48,MATCH($B$5,$A$43:$A$48,0)),"
    "MATCH(Setup!$B$5,'Overhead Split'!$C$22:$N$22,0)))"
    "-SUM($B16,$B17,$B18,$B19,$B20,$B21)")

wb = openpyxl.load_workbook(WB)
ws = wb['Dept P&L']
assert ws['A22'].value == 'Other admin overhead', ws['A22'].value
ws['B22'] = ORIGINAL
wb.save(WB)
print('Dept P&L B22 restored')
