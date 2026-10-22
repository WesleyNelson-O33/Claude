"""Profit and loss split for the departmental result chart.

The gross profit chart shows all five departments above the line, because every
one of them does make a gross profit. It is the overhead charged to them that
puts three below the line. A chart of the net result is needed so that is not
read the wrong way.

The result goes into two columns, a profit column and a loss column, so the chart
colours itself without the build script having to read the figures. A freshly
patched cell has no cached value until recalc, so reading the sign at build time
fails on a first run.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
NOTE = Font(name='Calibri', sz=10, color='FF6B7585')
HEAD, FIRST, LAST = 142, 143, 147

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws[f'A{HEAD}'].value == 'Department'

for c, name in ((19, 'Profit'), (20, 'Loss')):
    ws.cell(HEAD, c, name).font = Font(name='Calibri', sz=10, b=True,
                                       color='FF242C39')
for r in range(FIRST, LAST + 1):
    for c, f in ((19, f'=IF(N($H{r})>0,$H{r},"")'),
                 (20, f'=IF(N($H{r})<=0,$H{r},"")')):
        cell = ws.cell(r, c, f)
        cell.number_format = MONEY
        cell.font = NOTE

wb.save(WB)
print('Overhead Split columns S and T added: the result split into profit and loss')
