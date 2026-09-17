import openpyxl, warnings
warnings.filterwarnings('ignore')
SRC = "/root/.claude/uploads/a91f56e6-b7da-5dad-8e0b-6d5f516f9360/71b3078d-PL_Analysis_-_Aug_2026_-_Xero.xlsm"
src = openpyxl.load_workbook(SRC, data_only=True, read_only=True)["1. Xero GL Transactions"]
wb = openpyxl.load_workbook("CTS Financial Controller Pack.xlsx")
gl = wb["GL_Paste"]
from openpyxl.styles import Font
BODY = Font(name="Arial", size=10)
n = 0
for row in src.iter_rows(min_row=2, max_row=20001, max_col=13, values_only=True):
    if row[1] is None:
        break
    n += 1
    for i, v in enumerate(row):
        c = gl.cell(row=1 + n, column=1 + i, value=v)
        c.font = BODY
        if i == 4:
            c.number_format = "dd mmm yyyy"
        elif i in (6, 7):
            c.number_format = '#,##0.00'
print("GL lines loaded:", n)

# PL_Check: the August category totals off the Xero P&L, as the worked example.
pl = wb["PL_Check"]
for r, v in ((5, 603466.54), (6, -335516.00), (7, -208889.82), (8, 861.35), (9, 0.00)):
    pl.cell(row=r, column=3, value=v)
pl["A13"] = "Example above: the August 2026 Xero P&L, entered so you can see the check working."
pl["A13"].font = Font(name="Arial", size=9, italic=True, color="595959")
wb.save("CTS Financial Controller Pack.xlsx")
print("saved")
