"""Undo patch_v76.

patch_v76 pointed the "Legal Fees - Other (Employsure)" row at the GL account
"Legal Fees" so the 800 a month stopped falling into the catch-all, and spread
it on the Staff key, which is the key the P&L analysis labels that account with.

The analysis does not actually compute it that way.  Its formula for that row
reads the percentage from the row above (Legal Fees - Easton Belle, 3 Way), so
the 800 is spread on 3 Way, not Staff.  Leaving the row unmatched puts the 800
in the catch-all, which is also on 3 Way, so the pack reproduces the analysis.

Whether the account should be on Staff or 3 Way is Danica's call.
"""
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
ROW = 59

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws.cell(ROW, 1).value == 'Legal Fees - Other (Employsure)'

for c in range(3, 15):
    cell = ws.cell(ROW, c)
    assert '"Legal Fees"' in cell.value, cell.value
    cell.value = cell.value.replace('"Legal Fees"', f'$A{ROW}')

print('  ', ws.cell(ROW, 3).value)
wb.save(WB)
