"""Align the Staff key with the one the P&L analysis uses.

The pack and the P&L analysis allocate admin overhead the same way: the amount
coded to admin, times a key, plus whatever is coded to the department directly.
They carry the same three keys, and two of them already agree to the fourth
decimal. 3 Way is 0.33, 0.26, 0.08, 0.05, 0.28 in both. Office Dept is 0.0179,
0.875, 0.0713, 0.0179, 0.0179 in both.

The Staff key is the one that differs, and it differs because of the headcount
behind it. The pack had 39 Support, 24 Events, 3 Video, 1 Integration and nothing
for Consulting, over a base of 67. The analysis works on 49, 18, 3, 1 and 1 over
a base of 72, which is where its 0.680556, 0.25, 0.041667, 0.013889, 0.013889
come from.

This sets the headcount to the analysis's, so the key computes to the same five
figures rather than being hard coded over the top of the formula.

The headcount itself is a fact about the business that I cannot check from the
files. It is flagged for Danica to confirm.
"""
import os, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

# Support, Events, Video, Integration, Consulting, then Admin
HEADCOUNT = {'Support': 49, 'Events': 18, 'Video': 3, 'Integration': 1,
             'Consulting': 1}

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws['A5'].value == 'Department' and ws['B5'].value == 'Headcount'

before = {}
for r in range(6, 11):
    name = ws.cell(r, 1).value
    before[name] = ws.cell(r, 2).value
    ws.cell(r, 2, HEADCOUNT[name])

wb.save(WB)
print('Staff key headcount changed to match the P&L analysis')
for name, old in before.items():
    print(f'  {name:12} {old} -> {HEADCOUNT[name]}')
print(f'  base {sum(before.values())} -> {sum(HEADCOUNT.values())}')
