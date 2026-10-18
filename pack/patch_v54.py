"""
Utilisation per cent, calculated the way her own report calculates it.

The pack divided billable hours by hours worked, billable plus non-billable,
with leave left out of the sum. Her report divides billable hours by Workable
Business Hours, the hours the person was available for in the month, so leave
and public holidays stay in the bottom of the fraction. The two give very
different answers. Support reads 99.5 per cent my way and 91 per cent hers.

Hers is the source, so hers is what the pack and the deck now show.

Workable Business Hours cannot be worked out inside the pack. It is each
person's available hours for the month, which needs the headcount of each
group and who is part time, and neither is in here. So it is a paste, like the
hours, taken from the Person Group total rows of her own monthly report.

Her report has three person groups, because nobody is assigned to Video or
Integration. That is why this block has three groups and not five. The five
department hours block above it stays as it is: it is grouped by the job the
work was for, which is what ties to revenue, and it is what the chart reads.

To stop the two being confused, the per cent on the five department blocks is
renamed. It is the billable share of the hours actually worked, which is a
different and still useful number. "Chargeable %" now means only one thing in
this pack, and it means what her report means by it.

Seeded from the July, August and September 2026 Utilisation Reports, tabs
"July", "August" and "September", the Person Group total rows. Each month ties
to her Grand Total: July 72.57 per cent, August 80.79, September 71.81.
"""
import json, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
H = json.load(open(f'{HERE}/../deck/data/util_her_basis.json'))

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)
HRS, PCT = '#,##0.00;;', '0.0%'

LINES = ['Billable hours', 'Non-billable hours', 'Leave', 'Public holiday',
         'Workable business hours', 'Chargeable %']
KEYS = ['ch', 'nc', 'lv', 'ph', 'workable']
GROUPS = [('Production', 123), ('Consulting', 130), ('Support', 137)]
MONTHCOL = {'Jul-26': 4, 'Aug-26': 5, 'Sep-26': 6}      # D, E, F

wb = openpyxl.load_workbook(WB)
ws = wb['Utilisation']

# --- the five department blocks: say what their per cent actually is -----
for r in (12, 18, 24, 30, 36, 42):
    assert ws.cell(r, 3).value in ('Utilisation %', 'Billable share of hours worked')
    ws.cell(r, 3, 'Billable share of hours worked')

# --- her basis ----------------------------------------------------------
ws['A120'] = 'UTILISATION, AS HER REPORT CALCULATES IT'
for c in range(1, 16):
    cell = ws.cell(120, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(['FY (year it ends)', 'Person Group', 'Hours'], start=1):
    cell = ws.cell(122, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
for c in range(4, 16):
    cell = ws.cell(122, c, f'={gl(c)}$4')
    cell.number_format = 'mmm-yy'
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='right')

def block(r0, name, values):
    """values: key -> {month -> number}, or None for the company total."""
    for i, line in enumerate(LINES):
        r = r0 + i
        ws.cell(r, 1, '=Setup!$B$7' if i == 0 else None)
        if i == 0:
            ws.cell(r, 2, name).font = Font(name='Calibri', sz=11, b=True, color=NAVY)
        ws.cell(r, 3, line).font = Font(name='Calibri', sz=11,
                                        b=(line == 'Chargeable %'), color=NAVY)
        for c in range(4, 16):
            cell = ws.cell(r, c)
            if line == 'Chargeable %':
                cell.value = f'=IFERROR({gl(c)}{r0}/{gl(c)}{r0 + 4},0)'
                cell.number_format = PCT
                cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)
            else:
                if values is None:
                    cell.value = ('=' + '+'.join(f'{gl(c)}{g[1] + i}' for g in GROUPS))
                else:
                    mon = next((m for m, col in MONTHCOL.items() if col == c), None)
                    cell.value = values[KEYS[i]][mon] if mon else None
                cell.number_format = HRS
                cell.font = Font(name='Calibri', sz=11, color=NAVY)

for name, r0 in GROUPS:
    vals = {k: {m: H[m][name][k] for m in MONTHCOL} for k in KEYS}
    block(r0, name, vals)

block(144, 'Company', None)
for i in range(6):
    for c in range(1, 16):
        cell = ws.cell(144 + i, c)
        cell.fill = TOTAL
        f = cell.font
        cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)

# Check against the Grand Total row of her own report, each month.
ws.cell(151, 3, 'Check, her Grand Total').font = Font(name='Calibri', sz=11, b=True, color=NAVY)
for m, c in MONTHCOL.items():
    cell = ws.cell(151, c, round(H[m]['Grand']['pct'], 6))
    cell.number_format = PCT
    cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)

wb.save(WB)
print('Utilisation rows 120 to 151 added: her Chargeable % basis, three person groups')
print('the five department per cent renamed to "Billable share of hours worked"')
