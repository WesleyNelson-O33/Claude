"""
Section 6 of Overhead Split: the office staff split now runs off a weighting
table instead of worked hours.

Why. Her Graph file does not calculate the office staff split. Rows 6 to 16 of
its PL tab are typed values, and "Office Staff" on row 42 is just row 41 less
row 43, where row 43 is also typed. The allocated employment figure on row 41
(65,554.35 for Onsite in August) appears nowhere in the P&L Analysis workbook
and nowhere in the Labour Hours Reports. Worked hours do not reproduce it:
August hours by the person's own department give Support 68.44 per cent against
her 67.85, and Video 1.71 against her 4.18.

Backing the weights out of her own two months gives Support 48.6, Events 18.3,
Video 3.0, Integration 1.0, Consulting 1.0, and that set barely moves between
July and August. So the driver is a weighting she keeps, not hours.

The weights below are taken from her July and August decks, so the pack now
reproduces her figures. August onward carries the August weights forward until
she replaces them. The table is typed, visible and labelled, on the same row
layout as before, so rows 120 to 136 are untouched.
"""
import openpyxl, sys

WB = sys.argv[1] if len(sys.argv) > 1 else 'CTS Financial Controller Pack - v20.xlsx'

# Her own office staff split, per month, straight off the Graph files.
HER = {
    'C': (34377.91, 12935.82, 2129.60, 707.79, 707.79),      # July 2026
    'D': (57727.36, 21435.90, 3552.91, 1182.75, 1182.75),    # August 2026
}
ROWS = (113, 114, 115, 116, 117)          # Support, Events, Video, Integration, Consulting
COLS = [chr(ord('C') + i) for i in range(12)]   # C..N, July to June

# Weights scale to Consulting = 1 so the column reads like a headcount.
weights = {c: tuple(v / split[4] for v in split) for c, split in HER.items()}
carry = weights['D']                      # August carries forward to the rest of the year

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']

ws['A111'] = ("Allocated on a weighting per department, typed below. Her own split, backed out of "
              "the July and August decks. Replace the weights when the make up changes.")
ws['A112'] = 'Weighting'

fmt = '#,##0.0000'
for col in COLS:
    w = weights.get(col, carry)
    for row, value in zip(ROWS, w):
        cell = ws[f'{col}{row}']
        cell.value = value
        cell.number_format = fmt
ws['A118'] = '    Total weighting'
for col in COLS:
    ws[f'{col}118'].number_format = fmt

wb.save(WB)
print(f'section 6 rewired: weights typed into {len(COLS) * len(ROWS)} cells')
for col, lab in (('C', 'July'), ('D', 'August')):
    print(f'  {lab:<8}' + '  '.join(f'{v:0.4f}' for v in weights[col]))
