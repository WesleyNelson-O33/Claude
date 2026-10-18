"""
Borders and wording, every tab.

Three faults ran through the pack.

Section headings. Most are a navy band with white text at twelve point. Nine
were plain bold text on white, so the tab read as one long list with no breaks
in it. The numbering also mixed ALL CAPS with sentence case on the same tab.

Table headers. The house style is a navy row with white text and a medium rule
under it. Six tables had bold black text and no rule, which is why those blocks
look unfinished on a slide: Utilisation hours, the department result, the
employment expense split, revenue by department, the Dept P&L header, and the
Budget FY27 expense table, which had the band but no rule.

Wording. Acronyms the board does not read (GPM, NPM, COGS, LY, GM %), trailing
spaces in the Summary headers, and double spaces in the numbering.

Left alone on purpose: account names that come from Xero, such as "General &
Administrative Exp" and "Work in Progess". They look like typos and they are
hers, so they stay. Row indentation on the P&L tabs stays too; it is the
hierarchy.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)

# sheet -> heading rows, header rows, total rows
WORK = {
    'Setup':          ([4, 13, 19, 27], [], []),
    'Utilisation':    ([47, 66], [67], [73]),
    'Overhead Split': ([4, 14, 20, 88, 99, 110, 140], [142], [148]),
    'Overheads MoM':  ([88, 130, 160], [161], [167]),
    'Cost Analysis':  ([4, 19, 47, 62, 100, 110, 127, 190], [191], [197]),
    'Dept P&L':       ([], [7], []),
    'Budget FY27':    ([], [196], []),
    'Top 10 Clients': ([5, 22, 105], [], []),
    'Top 10 YoY':     ([17, 33], [], []),
    'DSO':            ([5, 17, 29, 76], [], []),
    'P&L by Dept':    ([258], [], []),
}

# Plain English, one wording per thing, across every tab.
WORDS = {
    'Setup': {'A4': '1. REPORTING MONTH', 'A13': '2. RISK FLAG THRESHOLDS',
              'A19': '3. DATA STATUS (CALCULATED)',
              'A27': '4. PASTED REPORTS AND THE EBITDA BRIDGE'},
    'Overhead Split': {
        'A4': '1. HEADCOUNT, WHICH DRIVES THE STAFF KEY',
        'A14': '2. THE THREE KEYS',
        'A20': '3. ADMIN OVERHEAD BY ACCOUNT AND KEY',
        'A88': '4. ADMIN OVERHEAD SPREAD ACROSS THE DEPARTMENTS',
        'A99': "5. EACH DEPARTMENT'S OWN OVERHEAD",
        'A110': '6. OFFICE STAFF SPLIT, THE ADMIN EMPLOYMENT EXPENSE SPREAD ACROSS THE DEPARTMENTS',
        'A140': '7. DEPARTMENT RESULT, THE MONTH ON SETUP',
        'F142': 'Gross margin %'},
    'Utilisation': {'C8': 'Hours',
                    'B9': 'Support', 'B10': 'Support', 'B11': 'Support',
                    'B15': 'Events', 'B16': 'Events', 'B17': 'Events',
                    'B39': 'Admin', 'B40': 'Admin', 'B41': 'Admin',
                    'A47': 'WORKED HOURS BY COST CENTRE',
                    'A49': 'Prior year (paste once a year)',
                    'A57': 'This year (paste each month)'},
    'Dept P&L': {'A13': 'Total Cost of Sales', 'A15': 'Gross profit margin',
                 'A25': 'Net profit margin', 'D7': 'Var vs last year'},
    'Summary': {'K5': 'MoM', 'L5': 'MoM vs Budget', 'M5': 'Var', 'N5': 'Var %',
                'U5': 'YOY Actual', 'V5': 'YOY Budget'},
    'P&L by Dept': {'R6': 'MoM vs Budget', 'R83': 'MoM vs Budget',
                    'R210': 'MoM vs Budget'},
    'PL_Check': {'A7': 'Expenses (including interest, tax and currency)'},
}

wb = openpyxl.load_workbook(WB)

n_words = 0
for sheet, cells in WORDS.items():
    ws = wb[sheet]
    for ref, text in cells.items():
        if ws[ref].value != text:
            ws[ref] = text
            n_words += 1

# "no LY split" reads as code. Say it in words.
ws = wb['Dept P&L']
for r in range(16, 25):
    if ws.cell(r, 3).value == '="no LY split"':
        ws.cell(r, 3, '="not split last year"')
        n_words += 1

def width(ws, row, floor=4):
    last = floor
    for c in range(1, min(ws.max_column, 30) + 1):
        if ws.cell(row, c).value is not None:
            last = c
    return last

n_band = n_head = n_tot = 0
for sheet, (headings, headers, totals) in WORK.items():
    ws = wb[sheet]
    for r in headings:
        # as wide as the table under it, so the band squares off the block
        nxt = next((h for h in headers if h > r), None)
        cols = width(ws, nxt) if nxt and nxt - r <= 3 else max(width(ws, r + 2), 4)
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.fill = BAND
            cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)
        n_band += 1
    for r in headers:
        for c in range(1, width(ws, r) + 1):
            cell = ws.cell(r, c)
            cell.fill = BAND
            cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
            cell.border = Border(bottom=RULE)
        n_head += 1
    for r in totals:
        for c in range(1, width(ws, r) + 1):
            cell = ws.cell(r, c)
            cell.fill = TOTAL
            cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)
        n_tot += 1

wb.save(WB)
print(f'wording fixed: {n_words} cells')
print(f'section bands: {n_band}   table headers: {n_head}   total rows: {n_tot}')
