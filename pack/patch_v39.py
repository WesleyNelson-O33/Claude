"""
July utilisation hours were never pasted, so the July office staff pot of
50,858.92 had nothing to allocate on and section 6 of Overhead Split left the
whole amount sitting in admin. Column D on Utilisation was empty for all
fifteen rows.

Figures copied from July 2026 Utilisation Report, tab "July", the Person Group
totals: Support 2,541.09 / 34.04 / 270.50, Production 258.00 / 550.00 / 112.00,
Consulting 64.50 / 104.50 / 16.00. That tab has no Video or Integration group,
so both are nil for July. Public holiday hours are nil for every group in July,
so nothing had to be decided about them.
"""
import openpyxl, sys

WB = sys.argv[1] if len(sys.argv) > 1 else 'CTS Financial Controller Pack - v20.xlsx'

# Utilisation row -> July hours. Column D is the July 2026 month.
JULY = {
     9: 2541.09,  10:  34.04,  11: 270.50,   # Onsite      chargeable / non-ch / leave
    15:  258.00,  16: 550.00,  17: 112.00,   # Production
    21:    0.00,  22:   0.00,  23:   0.00,   # Video        no group in the report
    27:    0.00,  28:   0.00,  29:   0.00,   # Integration  no group in the report
    33:   64.50,  34: 104.50,  35:  16.00,   # Consulting
}

wb = openpyxl.load_workbook(WB)
u = wb['Utilisation']
for row, hours in JULY.items():
    u.cell(row, 4).value = hours
    u.cell(row, 4).number_format = u.cell(row, 5).number_format
wb.save(WB)
print(f'July utilisation pasted: {len(JULY)} cells in column D')
