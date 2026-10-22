"""
Revenue against budget, year to date, by department.

The deck compares revenue to budget for the month and by month. What it has not
had is the same comparison for the year so far, which is the figure the board
works to.

Year to date is summed across the closed months rather than read from a stored
total, so it follows the reporting month on Setup like everything else.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
WB = sys.argv[1] if len(sys.argv) > 1 else f'{HERE}/CTS Financial Controller Pack - v20.xlsx'

NAVY, WHITE, GREYL = 'FF242C39', 'FFFFFFFF', 'FFF4F4F4'
BAND = PatternFill('solid', start_color=NAVY, end_color=NAVY)
TOTAL = PatternFill('solid', start_color=GREYL, end_color=GREYL)
RULE = Side(style='medium', color=NAVY)
MONEY = '#,##0;[RED]\\(#,##0\\);\\-'
PCT = '0.0%;[RED]-0.0%;\\-'

DEPTS = [('Support', 41), ('Events', 8), ('Video', 16),
         ('Integration', 25), ('Consulting', 33)]
BANDROW, HEAD, FIRST = 250, 251, 252
TOTROW = FIRST + len(DEPTS)

def ytd(tab, row):
    return (f"=SUMPRODUCT(('{tab}'!$B$5:$M$5<=Setup!$B$5)*"
            f"'{tab}'!$B{row}:$M{row})")

wb = openpyxl.load_workbook(WB)
ws = wb['Cost Analysis']
assert ws.cell(BANDROW - 20, 1).value == 'REVENUE BY MONTH AND DEPARTMENT'

ws.cell(BANDROW, 1, 'REVENUE AGAINST BUDGET, YEAR TO DATE')
for c in range(1, 6):
    cell = ws.cell(BANDROW, c)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=12, b=True, color=WHITE)

for c, text in enumerate(['Department', 'Actual', 'Budget', 'Var $', 'Var %'], start=1):
    cell = ws.cell(HEAD, c, text)
    cell.fill = BAND
    cell.font = Font(name='Calibri', sz=10, b=True, color=WHITE)
    cell.border = Border(bottom=RULE)
    cell.alignment = Alignment(horizontal='left' if c == 1 else 'right')

for i, (name, row) in enumerate(DEPTS):
    r = FIRST + i
    ws.cell(r, 1, name).font = Font(name='Calibri', sz=11, color=NAVY)
    ws.cell(r, 2, ytd('P&L FY27', row)).number_format = MONEY
    ws.cell(r, 3, ytd('Budget FY27', row)).number_format = MONEY
    ws.cell(r, 4, f'=B{r}-C{r}').number_format = MONEY
    ws.cell(r, 5, f'=IFERROR(D{r}/C{r},0)').number_format = PCT
    for c in range(2, 6):
        ws.cell(r, c).font = Font(name='Calibri', sz=11, color=NAVY)

ws.cell(TOTROW, 1, 'Total')
for c, col in ((2, 'B'), (3, 'C'), (4, 'D')):
    ws.cell(TOTROW, c, f'=SUM({col}{FIRST}:{col}{TOTROW - 1})').number_format = MONEY
ws.cell(TOTROW, 5, f'=IFERROR(D{TOTROW}/C{TOTROW},0)').number_format = PCT
for c in range(1, 6):
    cell = ws.cell(TOTROW, c)
    cell.fill = TOTAL
    cell.font = Font(name='Calibri', sz=11, b=True, color=NAVY)

wb.save(WB)
print(f'Cost Analysis rows {BANDROW} to {TOTROW} added: revenue against budget, year to date')
