"""
CTS Financial Controller Pack  -  builder.

Build order, and it matters:
    python3.13 build_fcp_v2.py
    python3.13 /mnt/skills/public/xlsx/scripts/recalc.py "CTS Financial Controller Pack v2.xlsx" 900
    python3.13 ../reporting/tools/fix_outline.py "CTS Financial Controller Pack v2.xlsx" "P&L FY27,Summary"

Step three is not optional. LibreOffice strips the row-outline settings when it
recalculates, and without them the collapse buttons sit against the wrong rows.

Three paste zones, each bare at A1 because Excel refuses a whole-sheet paste
anywhere else:  GL_Paste, PL_Paste, BS_Paste.  Everything else is formulas.

Sign convention, carried over from the pack this replaces:
    Amount = Credit less Debit.  Income positive, costs negative.
    So gross profit is Income PLUS Cost of Sales, not minus.
"""
import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import Outline
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, Reference

HERE = __file__.rsplit('/', 1)[0]
MASTER = json.load(open(f'{HERE}/data/lists_master.json'))
BSACC  = json.load(open(f'{HERE}/data/bs_accounts.json'))
BUDGET = json.load(open(f'{HERE}/data/budget_fy27.json'))

OUT = f'{HERE}/CTS Financial Controller Pack v2.xlsx'

# ----------------------------------------------------------------- look ----
AR = 'Arial'
NAVY = PatternFill('solid', fgColor='1F3864')
YELL = PatternFill('solid', fgColor='FFFF00')
GREY = PatternFill('solid', fgColor='EDEDED')
BAND = PatternFill('solid', fgColor='DCE6F1')
NEWF = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOPB = Border(top=Side(style='thin', color='1F3864'))
MONEY = '$#,##0;($#,##0);"-"'
MON2  = '$#,##0.00;($#,##0.00);"-"'
PCT   = '0.0%;(0.0%);"-"'
NUM1  = '#,##0.0;(#,##0.0);"-"'
RATIO = '0.00x;(0.00x);"-"'

def f(sz=10, b=False, col=None):
    return Font(name=AR, size=sz, bold=b, color=col)

def title(ws, t, sub=None):
    ws['A1'] = t; ws['A1'].font = f(14, True, '1F3864')
    if sub:
        ws['A2'] = sub; ws['A2'].font = f(9, False, '595959')

def hdr(ws, row, labels, start=1, fmt=None):
    for i, t in enumerate(labels):
        c = ws.cell(row, start + i, t)
        c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
        c.alignment = Alignment(horizontal='center', wrap_text=True, vertical='center')
        if fmt: c.number_format = fmt

def inp(c, nf=MONEY):
    c.font = f(10, False, '0000FF'); c.fill = YELL; c.border = BOX; c.number_format = nf

def widths(ws, spec):
    for col, w in spec.items(): ws.column_dimensions[col].width = w

# ------------------------------------------------------- the taxonomy ------
DEPTS = ['ONSITE', 'PRODUCTION', 'VIDEO', 'INTEGRATION', 'CONSULTING', 'ADMIN']
INC = ['Contract Income', 'Non Contract Income', 'Work in Progess',
       'Freight Collected & Travel', 'Credit Card Fees Collected',
       'Meal & Travel Allowances', 'Miscellaneous Income']
COS = ['Equipment - Purchase', 'Equipment - Hires', 'Sub-Contract Labour',
       'Service - hires', 'Subscriptions & Licences', 'Equipment -Maintenance Support',
       'Stamp Duty', 'Damage Waiver', 'Supplier Per Diems', 'Direct Wages & Salaries',
       'Direct Salaries', 'Direct Wages', 'Direct Superannuation',
       'Direct Workers Comp', 'Other Direct Expenses', 'Credit Card fees paid - COGS']
EXP = ['General & Administrative Exp', 'Business Insurances',
       'Advertising & Promotion Exp', 'Operating Expenses',
       'Employment Expenses', 'Credit Card fees paid']
SUBS = ([('Income', s) for s in INC] + [('Cost of Sales', s) for s in COS]
        + [('Expenses', s) for s in EXP]
        + [('Other Income', 'Other Income'), ('Other Expenses', 'Other Expenses')])
ADDBACKS = ['D&A', 'IAWO', 'INTEREST_EXP', 'INTEREST_INC', 'TAX']
ADDLABEL = {'D&A': 'Depreciation & amortisation', 'IAWO': 'Instant asset write off',
            'INTEREST_EXP': 'Interest expense', 'INTEREST_INC': 'Interest income',
            'TAX': 'Income tax'}
# labour subcategories, used for the cost-structure block on the dashboard
LABOUR = ['Direct Wages & Salaries', 'Direct Salaries', 'Direct Wages',
          'Direct Superannuation', 'Direct Workers Comp']

# ----------------------------- the P&L block, one department, 48 rows ------
# (offset, label, level, kind)   kind drives how the row is written
BLOCK = (
    [(0, '{D}', 0, 'dept')]
    + [(1, 'Income', 1, 'sum_inc')]
    + [(2 + i, '        ' + s, 2, ('eng', 'Income', s)) for i, s in enumerate(INC)]
    + [(9, 'Cost of Sales', 1, 'sum_cos')]
    + [(10 + i, '        ' + s, 2, ('eng', 'Cost of Sales', s)) for i, s in enumerate(COS)]
    + [(26, '    Gross profit', 1, 'gp'),
       (27, '    Gross margin %', 1, 'gp_pct'),
       (28, 'Overheads', 1, 'sum_exp')]
    + [(29 + i, '        ' + s, 2, ('eng', 'Expenses', s)) for i, s in enumerate(EXP)]
    + [(35, '    Other income (excluding interest)', 1, 'oi_ex'),
       (36, '    Other costs (excluding interest and tax)', 1, 'oe_ex'),
       (37, '    EBITDA', 1, 'ebitda'),
       (38, '    EBITDA %', 1, 'ebitda_pct'),
       (39, '    Depreciation & amortisation', 1, 'dep'),
       (40, '    EBIT', 1, 'ebit'),
       (41, '    EBIT %', 1, 'ebit_pct'),
       (42, '    Interest income', 1, 'int_inc'),
       (43, '    Interest expense', 1, 'int_exp'),
       (44, '    Net profit before tax', 1, 'npbt'),
       (45, '    Income tax', 1, 'tax'),
       (46, '    Net profit after tax', 1, 'npat'),
       (47, '    Net margin %', 1, 'net_pct')]
)
BLEN = 48
STARTS = [6 + i * (BLEN + 1) for i in range(6)]      # 6, 55, 104, 153, 202, 251
TOTAL_START = STARTS[-1] + BLEN + 1                  # 300
PL_LAST = TOTAL_START + BLEN - 1                     # 347
MARGIN_KINDS = {'gp_pct', 'ebitda_pct', 'ebit_pct', 'net_pct'}
# rows a user types into on the budget / prior-year sheets
INPUT_KINDS = {'eng', 'oi_ex', 'oe_ex', 'dep', 'int_inc', 'int_exp', 'tax'}

def block_rows(start):
    """(absolute row, label, level, kind) for one department block."""
    return [(start + off, lab, lvl, kind) for off, lab, lvl, kind in BLOCK]

def off_of(kind):
    for off, lab, lvl, k in BLOCK:
        if k == kind: return off
    raise KeyError(kind)

def off_eng(cat, sub):
    for off, lab, lvl, k in BLOCK:
        if isinstance(k, tuple) and k[1] == cat and k[2] == sub: return off
    raise KeyError((cat, sub))

# ------------------------------------------------- engine row numbers ------
ENG_R0 = 5
SUB_ROW = {}                       # (dept, cat, sub) -> engine row
r = ENG_R0
for d in DEPTS:
    for cat, sub in SUBS:
        SUB_ROW[(d, cat, sub)] = r; r += 1
ADD_ROW = {}                       # (dept, addback) -> engine row
for d in DEPTS:
    for a in ADDBACKS:
        ADD_ROW[(d, a)] = r; r += 1
ENG_R1 = r - 1

CL_R0, CL_R1 = 5, 20004
GL_R0, GL_R1 = 2, 20001            # row 1 of GL_Paste is the first pasted row too
GL_ROWS = 20000

wb = Workbook()

# =========================================================== 1. README =====
ws = wb.active; ws.title = 'README'
title(ws, 'CTS Financial Controller Pack', 'Three pastes a month. Everything else is formulas.')
widths(ws, {'A': 4, 'B': 26, 'C': 108})
rows = [
    ('h', 'EACH MONTH', ''),
    ('r', '1.  Setup B5', 'Pick the reporting month. The financial year, the period and every comparative follow from it.'),
    ('r', '2.  GL_Paste', 'Xero: Account Transactions for P&L analysis, 1 July to the end of the reporting month. Select the whole export and paste at A1. Title rows do not matter, the pack ignores any row with no real date in the Date column.'),
    ('r', '3.  PL_Paste', 'Xero: Profit and Loss for the reporting month. Paste the whole export at A1. This replaces typing the five category totals by hand.'),
    ('r', '4.  BS_Paste', 'Xero: Balance Sheet as at the end of the reporting month, with the two comparison columns turned on. Paste the whole export at A1.'),
    ('r', '5.  Controls', 'Every line must read OK before you report anything.'),
    ('r', '6.  Read', 'Dashboard first, then Summary, then P&L FY27 for the detail.'),
    ('h', 'ONCE A YEAR', ''),
    ('r', 'Budget FY27', "This year's budget or rolling forecast. Revise when the forecast moves."),
    ('r', 'Actual FY26', "Last year's actuals. Drives the PY and YOY columns. Empty at the moment, so those columns read nil."),
    ('r', 'Budget FY26', "Last year's budget. Also empty."),
    ('r', 'Utilisation', 'Hours per department per month. Working days and working hours calculate themselves.'),
    ('r', 'Lists', 'Public holidays, and any new Xero account.'),
    ('h', 'WHERE EACH NUMBER COMES FROM', ''),
    ('r', 'The P&L', 'The GL paste, always. Not the P&L paste. The P&L paste exists only to prove the GL paste is complete.'),
    ('r', 'The balance sheet', 'The BS paste. The GL export is a P&L report, so it carries no balance sheet.'),
    ('r', 'The department', 'From the Cost Centres column on each transaction, which carries the department name outright. If it is ever blank the pack falls back to the bracket tag on the Job Numbers column. Three quarters of the P&L sits in accounts with no department in the name, so only the transaction knows which department earned or spent it.'),
    ('h', 'THINGS THAT WILL CATCH YOU OUT', ''),
    ('r', 'Sign convention', 'Amount is Credit less Debit. Income is positive, costs are negative, so gross profit is Income plus Cost of Sales rather than minus.'),
    ('r', 'Forecast months', 'On the P&L, months up to the reporting month come from the GL. Months after it come from Budget FY27, so you always see a full twelve and the shortfall ahead.'),
    ('r', 'EBITDA', 'Depreciation and amortisation is added back from the accounts flagged in the Addback column on Lists. Instant Asset Write Off is a separate toggle on Setup B12, because it is capital spending put through the P&L and it changes EBITDA either way.'),
    ('r', 'One financial year only', 'The GL paste is read for the reporting financial year only. Lines from another year are counted on Controls and ignored, so a stale paste cannot quietly double up.'),
    ('r', 'Risk column', 'No Activity, Low, Medium or High, off the thresholds on Setup. Below the materiality floor it is always Low.'),
    ('r', 'Capacity', 'GL_Paste holds 20,000 lines, about a full year at your volume. Setup tells you how many you have used.'),
]
rr = 4
for kind, a, b in rows:
    if kind == 'h':
        ws.cell(rr, 2, a).font = f(10, True, '1F3864'); rr += 1
    else:
        ws.cell(rr, 2, a).font = f(10, True)
        c = ws.cell(rr, 3, b); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[rr].height = 13 * (1 + len(b) // 105)
        rr += 1
    if kind == 'h' and rr > 5: ws.row_dimensions[rr - 1].height = 20

# ============================================================ 2. SETUP =====
st = wb.create_sheet('Setup')
title(st, 'Setup and control panel', 'The yellow cells are the only ones you type into. Everything else follows from them.')
widths(st, {'A': 42, 'B': 16, 'C': 96})
def srow(r, label, val, note='', nf=None, is_inp=False, bold=False):
    st.cell(r, 1, label).font = f(10, bold)
    c = st.cell(r, 2, val)
    if is_inp: inp(c, nf or 'General')
    else:
        c.font = f(10); c.number_format = nf or 'General'
    if note:
        n = st.cell(r, 3, note); n.font = f(9, False, '595959'); n.alignment = Alignment(wrap_text=True, vertical='top')
st.cell(4, 1, '1.  What am I reporting on').font = f(10, True, '1F3864')
srow(5, 'Reporting month', dt.datetime(2026, 9, 1), 'The one control you change each month.', 'mmm yyyy', True)
srow(6, 'Financial year starts in month', 7, 'July. Do not change this unless the company changes its year end.', '0', True)
srow(7, 'Financial year (the year it ends)', '=IF(MONTH($B$5)>=$B$6,YEAR($B$5)+1,YEAR($B$5))', '', '0')
srow(8, 'Period number in the financial year', '=MOD(MONTH($B$5)-$B$6+12,12)+1', '', '0')
srow(9, 'Prior financial year', '=$B$7-1', '', '0')
srow(10, 'Comparative month last year', '=EDATE($B$5,-12)', '', 'mmm yyyy')
srow(11, 'Hours in a working day', 8, 'Used by the utilisation sheet.', '0.0', True)
srow(12, 'Count Instant Asset Write Off as depreciation', 'Yes', 'Yes or No. It is capital spending expensed through the P&L. Yes adds it back, so EBITDA rises and EBIT does not move. No leaves it in EBITDA.', 'General', True)
srow(13, 'Days from the start of the year to the reporting month end',
     '=EOMONTH($B$5,0)-DATE($B$7-1,$B$6,1)+1', 'Drives debtor and creditor days.', '0')

st.cell(15, 1, '2.  Risk flag thresholds').font = f(10, True, '1F3864')
srow(16, 'Materiality floor ($)', 5000, 'Below this, a variance is always Low.', '$#,##0', True)
srow(17, 'Medium risk threshold (% variance)', 0.10, '', '0%', True)
srow(18, 'High risk threshold (% variance)', 0.25, '', '0%', True)
srow(19, 'High risk threshold ($ variance)', 50000, '', '$#,##0', True)

st.cell(21, 1, '3.  Where the pasted reports keep their figures').font = f(10, True, '1F3864')
st.cell(22, 1, 'Set these once. They only change if you change the way the Xero report is laid out.').font = f(9, False, '595959')
srow(23, 'P&L paste: column holding the reporting month', 2, 'Column B is 2, C is 3, and so on. The header text found in that column is shown on Controls so you can see you picked right.', '0', True)
srow(24, 'BS paste: column holding this month', 2, '', '0', True)
srow(25, 'BS paste: column holding last month', 3, 'Leave at 0 if your export has no comparison columns.', '0', True)
srow(26, 'BS paste: column holding the start of the year', 4, 'Leave at 0 if your export has no comparison columns.', '0', True)
srow(27, 'GL paste: expected column order',
     'Index, Account Code, Account Name, Source, Date, Contact, Debit, Credit, Job Numbers, Invoice Number, Reference, Description, Cost Centres',
     'Reference only. This is the order the Xero export comes out in. If Xero ever changes it, the unknown-account count on Controls will jump.')
st.cell(27, 2).font = f(9, False, '595959'); st.cell(27, 2).alignment = Alignment(wrap_text=True, vertical='top')
st.row_dimensions[27].height = 40

st.cell(29, 1, '4.  Data status  (calculated)').font = f(10, True, '1F3864')
srow(30, 'GL lines pasted', f'=COUNT(GL_Paste!$E$1:$E${GL_ROWS})', 'Rows with a real date in the Date column.', '#,##0')
srow(31, 'GL capacity', GL_ROWS, '', '#,##0')
srow(32, 'GL lines used by this financial year', f'=COUNTIF(Cleanup!$H${CL_R0}:$H${CL_R1},">0")', '', '#,##0')
srow(33, 'GL lines dated outside this financial year', '=$B$30-$B$32', 'These are ignored. If it is not nil, check what you pasted.', '#,##0')
srow(34, 'Lines with no department tag', f'=COUNTIF(Cleanup!$D${CL_R0}:$D${CL_R1},"UNALLOCATED")', '', '#,##0')
srow(35, 'Lines with an account not on Lists', f'=COUNTIF(Cleanup!$E${CL_R0}:$E${CL_R1},"UNKNOWN")', 'Every one of these is left out of the P&L. Add the account to Lists.', '#,##0')
srow(36, 'Earliest GL date', f'=IFERROR(MIN(GL_Paste!$E$1:$E${GL_ROWS}),0)', '', 'dd/mm/yyyy')
srow(37, 'Latest GL date', f'=IFERROR(MAX(GL_Paste!$E$1:$E${GL_ROWS}),0)', '', 'dd/mm/yyyy')
srow(38, 'BS lines pasted', '=COUNTA(BS_Paste!$A$1:$A$400)', '', '#,##0')
srow(39, 'P&L lines pasted', '=COUNTA(PL_Paste!$A$1:$A$300)', '', '#,##0')

# ====================================================== 3. PASTE ZONES =====
# Bare grids starting at A1. Excel refuses a whole-sheet paste anywhere else,
# so there is deliberately no title and no instruction above the data.
GLCOLS = ['Index', 'Account Code', 'Account Name', 'Source', 'Date', 'Contact',
          'Debit', 'Credit', 'Job Numbers', 'Invoice Number', 'Reference',
          'Description', 'Cost Centres']
gp = wb.create_sheet('GL_Paste')
widths(gp, {'A': 8, 'B': 12, 'C': 30, 'D': 12, 'E': 11, 'F': 28, 'G': 12, 'H': 12,
            'I': 16, 'J': 14, 'K': 20, 'L': 34, 'M': 16})
for r in range(1, GL_ROWS + 1):
    for c in range(1, 14):
        cell = gp.cell(r, c)
        cell.font = f(9); cell.fill = YELL; cell.border = BOX
        if c == 5: cell.number_format = 'dd/mm/yyyy'
        if c in (7, 8): cell.number_format = MON2
gp.freeze_panes = 'A2'

pp = wb.create_sheet('PL_Paste')
widths(pp, {'A': 44, **{CL(c): 15 for c in range(2, 17)}})
for r in range(1, 301):
    for c in range(1, 17):
        cell = pp.cell(r, c); cell.font = f(9); cell.fill = YELL; cell.border = BOX
        if c > 1: cell.number_format = MON2
pp.freeze_panes = 'B2'

bp = wb.create_sheet('BS_Paste')
widths(bp, {'A': 44, **{CL(c): 15 for c in range(2, 17)}})
for r in range(1, 401):
    for c in range(1, 17):
        cell = bp.cell(r, c); cell.font = f(9); cell.fill = YELL; cell.border = BOX
        if c > 1: cell.number_format = MON2
bp.freeze_panes = 'B2'

# ============================================================ 4. LISTS =====
ls = wb.create_sheet('Lists')
title(ls, 'Reference lists',
      'The account master, the balance sheet master, the department tags and the month list. '
      'Add a new Xero account here or its GL lines are left out of the P&L.')
widths(ls, {'A': 10, 'B': 34, 'C': 15, 'D': 32, 'E': 26, 'F': 14, 'H': 1,
            'I': 16, 'J': 16, 'L': 1, 'M': 14, 'O': 1, 'P': 11, 'Q': 22, 'S': 1,
            'T': 10, 'U': 34, 'V': 24, 'W': 14})
ls.cell(4, 1, 'Account master  (profit and loss)').font = f(10, True, '1F3864')
hdr(ls, 5, ['Code', 'Account name', 'Category', 'Subcategory', 'Type', 'Addback'])
LS_A0 = 6
LS_A1 = LS_A0 + len(MASTER) - 1
for i, m in enumerate(MASTER):
    r = LS_A0 + i
    for c, v in enumerate([m['code'], m['name'], m['cat'], m['sub'], m['type'] or '', m['addback']], start=1):
        cell = ls.cell(r, c, v); cell.font = f(9); cell.border = BOX
        if m['new']: cell.fill = NEWF
ls.cell(LS_A1 + 2, 1, 'Rows shaded cream are accounts live in Xero that the old pack did not carry. '
                      'Their category and subcategory are my best read of the account name and its Xero type. '
                      'Check them.').font = f(9, False, 'C00000')

ls.cell(4, 9, 'Department tags').font = f(10, True, '1F3864')
hdr(ls, 5, ['Tag on GL line', 'Department'], start=9)
TAGS = [('ONSITE', 'ONSITE'), ('PRODUCTION', 'PRODUCTION'), ('VIDEO', 'VIDEO'),
        ('INTEGRATION', 'INTEGRATION'), ('CONSULTING', 'CONSULTING'), ('CTS', 'ADMIN'),
        ('ONS', 'ONSITE'), ('PRD', 'PRODUCTION'), ('VID', 'VIDEO'),
        ('INT', 'INTEGRATION'), ('CONS', 'CONSULTING'), ('ADMIN', 'ADMIN')]
LS_T0 = 6; LS_T1 = LS_T0 + len(TAGS) - 1
for i, (t, d) in enumerate(TAGS):
    for c, v in enumerate([t, d], start=9):
        cell = ls.cell(LS_T0 + i, c, v); cell.font = f(9); cell.border = BOX

ls.cell(4, 13, 'Months in this financial year').font = f(10, True, '1F3864')
hdr(ls, 5, ['Month'], start=13)
LS_M0 = 6
for i in range(12):
    c = ls.cell(LS_M0 + i, 13, f'=EDATE(DATE(Setup!$B$7-1,Setup!$B$6,1),{i})')
    c.font = f(9); c.border = BOX; c.number_format = 'mmm yyyy'
LS_M1 = LS_M0 + 11

ls.cell(4, 16, 'Public holidays').font = f(10, True, '1F3864')
ls.cell(4, 17, 'Type the public holidays here. The utilisation block uses them.').font = f(9, False, '595959')
hdr(ls, 5, ['Date', 'Name'], start=16)
for i in range(30):
    inp(ls.cell(6 + i, 16), 'dd/mm/yyyy'); inp(ls.cell(6 + i, 17), 'General')

ls.cell(4, 20, 'Balance sheet master').font = f(10, True, '1F3864')
hdr(ls, 5, ['Code', 'Account name', 'Group', 'Flag'], start=20)
LS_B0 = 6; LS_B1 = LS_B0 + len(BSACC) - 1
for i, a in enumerate(BSACC):
    flag = ('Cash' if a['cash'] else 'Borrowings' if a['debt'] else 'Trade debtors' if a['debtors']
            else 'Trade creditors' if a['creditors'] else 'WIP' if a['wip'] else '')
    for c, v in enumerate([a['code'], a['name'], a['group'], flag], start=20):
        cell = ls.cell(LS_B0 + i, c, v); cell.font = f(9); cell.border = BOX

# department and subcategory index lists, so the engine can use numeric keys
ls.cell(4, 6, '').font = f(9)
IDX_COL = 25   # Y: department order,  Z: subcategory order
ls.cell(4, IDX_COL, 'Index lists  (do not edit)').font = f(10, True, '1F3864')
hdr(ls, 5, ['Department', 'Subcategory'], start=IDX_COL)
for i, d in enumerate(DEPTS):
    c = ls.cell(6 + i, IDX_COL, d); c.font = f(9); c.border = BOX
for i, (cat, sub) in enumerate(SUBS):
    c = ls.cell(6 + i, IDX_COL + 1, sub); c.font = f(9); c.border = BOX
LS_D0, LS_D1 = 6, 6 + len(DEPTS) - 1
LS_S0, LS_S1 = 6, 6 + len(SUBS) - 1
widths(ls, {CL(IDX_COL): 16, CL(IDX_COL + 1): 32})

for nm, ref in [
    ('AcctCode', f'Lists!$A${LS_A0}:$A${LS_A1}'), ('AcctName', f'Lists!$B${LS_A0}:$B${LS_A1}'),
    ('AcctCat', f'Lists!$C${LS_A0}:$C${LS_A1}'), ('AcctSub', f'Lists!$D${LS_A0}:$D${LS_A1}'),
    ('AcctAdd', f'Lists!$F${LS_A0}:$F${LS_A1}'),
    ('DeptTag', f'Lists!$I${LS_T0}:$I${LS_T1}'), ('DeptName', f'Lists!$J${LS_T0}:$J${LS_T1}'),
    ('Months', f'Lists!$M${LS_M0}:$M${LS_M1}'), ('Holidays', 'Lists!$P$6:$P$35'),
    ('BSCode', f'Lists!$T${LS_B0}:$T${LS_B1}'), ('BSName', f'Lists!$U${LS_B0}:$U${LS_B1}'),
    ('BSGroup', f'Lists!$V${LS_B0}:$V${LS_B1}'), ('BSFlag', f'Lists!$W${LS_B0}:$W${LS_B1}'),
    ('DeptList', f'Lists!${CL(IDX_COL)}${LS_D0}:${CL(IDX_COL)}${LS_D1}'),
    ('SubList', f'Lists!${CL(IDX_COL+1)}${LS_S0}:${CL(IDX_COL+1)}${LS_S1}'),
    ('AddList', f'Lists!$F${LS_A0}:$F${LS_A1}'),
]:
    wb.defined_names.add(DefinedName(nm, attr_text=ref))

# ========================================================== 5. CLEANUP =====
cu = wb.create_sheet('Cleanup')
title(cu, 'Cleanup  -  calculated, never type here',
      'One row per pasted GL line. Amount is Credit less Debit, so income is positive and costs are '
      'negative. Department comes from the Cost Centres column, falling back to the bracket tag on the '
      'job number. Key is a single number so the engine can add up fast. '
      'The filter buttons on row 4 are there so you can look up a line that will not classify.')
hdr(cu, 4, ['Amount', 'Month', 'Tag', 'Department', 'Subcategory', 'Dept #',
            'Sub #', 'Period', 'Addback #', 'Key', 'Addback key'])
widths(cu, {'A': 13, 'B': 11, 'C': 16, 'D': 15, 'E': 32, 'F': 8, 'G': 8, 'H': 8, 'I': 10,
            'J': 12, 'K': 13})
cu.freeze_panes = 'A5'
for i in range(GL_ROWS):
    r = CL_R0 + i; g = 1 + i          # Cleanup row 5 is GL_Paste row 1
    gate = f'N(GL_Paste!$E{g})=0'     # a title row has no real date, so N() is zero
    cu.cell(r, 1, f'=IF({gate},"",N(GL_Paste!$H{g})-N(GL_Paste!$G{g}))').number_format = MON2
    cu.cell(r, 2, f'=IF({gate},"",DATE(YEAR(GL_Paste!$E{g}),MONTH(GL_Paste!$E{g}),1))').number_format = 'mmm yyyy'
    cu.cell(r, 3, f'=IF({gate},"",IF(TRIM(GL_Paste!$M{g})<>"",UPPER(TRIM(GL_Paste!$M{g})),'
                  f'IFERROR(UPPER(MID(GL_Paste!$I{g},FIND("[",GL_Paste!$I{g})+1,'
                  f'FIND("]",GL_Paste!$I{g})-FIND("[",GL_Paste!$I{g})-1)),"")))')
    cu.cell(r, 4, f'=IF($C{r}="","",IFERROR(INDEX(DeptName,MATCH($C{r},DeptTag,0)),"UNALLOCATED"))')
    cu.cell(r, 5, f'=IF({gate},"",IFERROR(INDEX(AcctSub,MATCH(GL_Paste!$B{g}&"",AcctCode,0)),"UNKNOWN"))')
    cu.cell(r, 6, f'=IF($D{r}="",0,IFERROR(MATCH($D{r},DeptList,0),0))').number_format = '0'
    cu.cell(r, 7, f'=IF($E{r}="",0,IFERROR(MATCH($E{r},SubList,0),0))').number_format = '0'
    cu.cell(r, 8, f'=IF($B{r}="",0,IF(IF(MONTH($B{r})>=Setup!$B$6,YEAR($B{r})+1,YEAR($B{r}))'
                  f'<>Setup!$B$7,0,MOD(MONTH($B{r})-Setup!$B$6+12,12)+1))').number_format = '0'
    cu.cell(r, 9, f'=IF({gate},0,IFERROR(MATCH(INDEX(AcctAdd,MATCH(GL_Paste!$B{g}&"",AcctCode,0)),'
                  f'{{"D&A";"IAWO";"INTEREST_EXP";"INTEREST_INC";"TAX"}},0),0))').number_format = '0'
    cu.cell(r, 10, f'=IF($H{r}=0,0,IF($F{r}=0,0,IF($G{r}=0,0,$F{r}*10000+$G{r}*100+$H{r})))').number_format = '0'
    cu.cell(r, 11, f'=IF($H{r}=0,0,IF($F{r}=0,0,IF($I{r}=0,0,$I{r}*10000+$F{r}*100+$H{r})))').number_format = '0'
    for c in range(1, 12):
        cu.cell(r, c).font = f(9)
cu.auto_filter.ref = f'A4:K{CL_R1}'      # so the classified GL can be filtered on
CL_AMT = f'Cleanup!$A${CL_R0}:$A${CL_R1}'
CL_KEY = f'Cleanup!$J${CL_R0}:$J${CL_R1}'
CL_ADD = f'Cleanup!$K${CL_R0}:$K${CL_R1}'

# =========================================================== 6. ENGINE =====
en = wb.create_sheet('Engine')
title(en, 'Engine  -  calculated, never type here',
      'One row per department and subcategory, then one row per department and add-back, twelve '
      'months across. Read by the P&L, the Summary, the Dashboard and the charts.')
hdr(en, 4, ['Key', 'Department', 'Category', 'Line'])
for i in range(12):
    c = en.cell(4, 5 + i, f'=INDEX(Months,{i+1})')
    c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
    c.number_format = 'mmm yy'; c.alignment = Alignment(horizontal='center')
hdr(en, 4, ['FY total'], start=17)
widths(en, {'A': 40, 'B': 14, 'C': 15, 'D': 32, **{CL(c): 13 for c in range(5, 18)}})
en.freeze_panes = 'E5'
for d in DEPTS:
    di = DEPTS.index(d) + 1
    for cat, sub in SUBS:
        r = SUB_ROW[(d, cat, sub)]; si = [s for _, s in SUBS].index(sub) + 1
        en.cell(r, 1, f'{d}|{sub}').font = f(9, False, '808080')
        for c, v in zip((2, 3, 4), (d, cat, sub)): en.cell(r, c, v).font = f(9)
        for m in range(12):
            key = di * 10000 + si * 100 + (m + 1)
            cell = en.cell(r, 5 + m, f'=SUMIF({CL_KEY},{key},{CL_AMT})')
            cell.number_format = MON2; cell.font = f(9)
        en.cell(r, 17, f'=SUM(E{r}:P{r})').number_format = MON2
        en.cell(r, 17).font = f(9)
    for a in ADDBACKS:
        r = ADD_ROW[(d, a)]; ai = ADDBACKS.index(a) + 1
        en.cell(r, 1, f'{d}|{a}').font = f(9, False, '808080')
        for c, v in zip((2, 3, 4), (d, 'Addback', ADDLABEL[a])): en.cell(r, c, v).font = f(9)
        for m in range(12):
            key = ai * 10000 + di * 100 + (m + 1)
            cell = en.cell(r, 5 + m, f'=SUMIF({CL_ADD},{key},{CL_AMT})')
            cell.number_format = MON2; cell.font = f(9)
        en.cell(r, 17, f'=SUM(E{r}:P{r})').number_format = MON2
        en.cell(r, 17).font = f(9)

# ========================= 7. the shared 48-row department block builder ===
MCOLS = [CL(c) for c in range(2, 14)]        # B..M  the twelve months
C_FY, C_YTD, C_BUD, C_VAR, C_VARP = 14, 16, 17, 18, 19
C_PY, C_YOY, C_YOYP, C_FYB, C_FYV, C_RISK = 21, 22, 23, 25, 26, 27
PCT_PAIR = {'gp_pct': 'gp', 'ebitda_pct': 'ebitda', 'ebit_pct': 'ebit', 'net_pct': 'npat'}
DERIVED_KINDS = ({'sum_inc', 'sum_cos', 'sum_exp', 'gp', 'ebitda', 'ebit',
                 'npbt', 'npat', 'dept'} | set(PCT_PAIR))

def eng_ref(dept, cat, sub, col):
    return f'Engine!{col}{SUB_ROW[(dept, cat, sub)]}'

def add_ref(dept, a, col):
    return f'Engine!{col}{ADD_ROW[(dept, a)]}'

def engine_expr(dept, kind, ecol):
    """What this row is worth for one month, read straight off the Engine."""
    if isinstance(kind, tuple):
        return eng_ref(dept, kind[1], kind[2], ecol)
    if kind == 'oi_ex':
        return (f'({eng_ref(dept,"Other Income","Other Income",ecol)}'
                f'-{add_ref(dept,"INTEREST_INC",ecol)})')
    if kind == 'oe_ex':
        return (f'({eng_ref(dept,"Other Expenses","Other Expenses",ecol)}'
                f'-{add_ref(dept,"INTEREST_EXP",ecol)}-{add_ref(dept,"TAX",ecol)})')
    if kind == 'dep':
        return (f'({add_ref(dept,"D&A",ecol)}'
                f'+IF(Setup!$B$12="Yes",{add_ref(dept,"IAWO",ecol)},0))')
    if kind == 'int_inc': return add_ref(dept, 'INTEREST_INC', ecol)
    if kind == 'int_exp': return add_ref(dept, 'INTEREST_EXP', ecol)
    if kind == 'tax':     return add_ref(dept, 'TAX', ecol)
    return None

def derived(kind, mc, s):
    """Rows built from other rows of the same block. s = block start row."""
    R = lambda k: f'{mc}{s + off_of(k)}'
    if kind == 'sum_inc': return f'=SUM({mc}{s+2}:{mc}{s+8})'
    if kind == 'sum_cos': return f'=SUM({mc}{s+10}:{mc}{s+25})'
    if kind == 'sum_exp': return f'=SUM({mc}{s+29}:{mc}{s+34})'
    if kind == 'gp':      return f'={R("sum_inc")}+{R("sum_cos")}'
    if kind == 'ebitda':  return (f'={R("gp")}+{R("sum_exp")}+{R("oi_ex")}'
                                  f'+{R("oe_ex")}-{R("dep")}')
    if kind == 'ebit':    return f'={R("ebitda")}+{R("dep")}'
    if kind == 'npbt':    return f'={R("ebit")}+{R("int_inc")}+{R("int_exp")}'
    if kind == 'npat':    return f'={R("npbt")}+{R("tax")}'
    if kind == 'dept':    return f'={R("npat")}'
    if kind in PCT_PAIR:
        num, den = f'{mc}{s + off_of(PCT_PAIR[kind])}', f'{mc}{s + off_of("sum_inc")}'
        return f'=IF({den}=0,"",{num}/{den})'
    return None

def style_row(ws, r, lvl, kind, cols, pct):
    bold = lvl == 0 or kind in ('gp', 'ebitda', 'ebit', 'npbt', 'npat', 'sum_inc',
                                'sum_cos', 'sum_exp') or kind in PCT_PAIR
    lab = ws.cell(r, 1)
    lab.font = f(11 if lvl == 0 else 10, bold, '1F3864' if lvl == 0 else None)
    if lvl == 0: lab.fill = BAND
    elif kind in ('gp', 'ebitda', 'ebit', 'npat') or kind in PCT_PAIR: lab.fill = GREY
    for c in cols:
        cell = ws.cell(r, c)
        cell.font = f(10, bold); cell.number_format = PCT if pct else MONEY
        if lvl == 0: cell.fill = BAND
        elif kind in ('gp', 'ebitda', 'ebit', 'npat') or kind in PCT_PAIR: cell.fill = GREY
    if kind in ('gp', 'ebitda', 'ebit', 'npat'):
        for c in [1] + list(cols): ws.cell(r, c).border = TOPB

def build_pl(ws, dept_labels=True):
    """P&L FY27  -  actuals to the reporting month, budget after it."""
    title(ws, 'Profit & Loss  -  FY27 by department')
    ws['A3'] = ('=TEXT(Setup!$B$5,"mmm yyyy")&"   |   months to and including the reporting month '
                'come from the GL. Later months come from Budget FY27. Click the + and - buttons on '
                'the left to open a department."')
    ws['A3'].font = f(9, False, '595959')
    hdr(ws, 5, ['Line'])
    for i, mc in enumerate(MCOLS):
        c = ws.cell(5, 2 + i, f'=INDEX(Months,{i+1})')
        c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
        c.number_format = 'mmm yy'; c.alignment = Alignment(horizontal='center')
    for col, lab in [(C_FY, 'FY total'), (C_YTD, 'YTD actual'), (C_BUD, 'YTD budget'),
                     (C_VAR, 'Var $'), (C_VARP, 'Var %'), (C_PY, 'PY YTD'),
                     (C_YOY, 'YOY $'), (C_YOYP, 'YOY %'), (C_FYB, 'FY budget'),
                     (C_FYV, 'FY var $'), (C_RISK, 'Risk')]:
        hdr(ws, 5, [lab], start=col)
    widths(ws, {'A': 46, **{mc: 12 for mc in MCOLS}, 'N': 13, 'O': 1.5, 'P': 13, 'Q': 13,
                'R': 13, 'S': 9, 'T': 1.5, 'U': 13, 'V': 13, 'W': 9, 'X': 1.5, 'Y': 13,
                'Z': 13, 'AA': 12})
    ws.freeze_panes = 'B6'
    allcols = list(range(2, 14)) + [C_FY, C_YTD, C_BUD, C_VAR, C_VARP, C_PY, C_YOY, C_YOYP,
                                    C_FYB, C_FYV]
    for bi, s in enumerate(STARTS + [TOTAL_START]):
        is_total = bi == 6
        dept = None if is_total else DEPTS[bi]
        for off, lab, lvl, kind in BLOCK:
            r = s + off
            pct = kind in PCT_PAIR
            label = ('TOTAL COMPANY' if is_total else f'{dept}') + '  -  net profit after tax' \
                if kind == 'dept' else lab
            ws.cell(r, 1, label)
            for i, mc in enumerate(MCOLS):
                ecol = CL(5 + i)
                if is_total:
                    if pct:
                        v = derived(kind, mc, s)
                    else:
                        v = '=' + '+'.join(f'{mc}{st_ + off}' for st_ in STARTS)
                elif kind in DERIVED_KINDS:
                    v = derived(kind, mc, s)
                else:
                    ex = engine_expr(dept, kind, ecol)
                    v = f"=IF({mc}$5<=Setup!$B$5,{ex},'Budget FY27'!{mc}{r})"
                ws.cell(r, 2 + i, v)
            # FY total
            if pct:
                ws.cell(r, C_FY, f'=IF(N{s+off_of("sum_inc")}=0,"",'
                                 f'N{s+off_of(PCT_PAIR[kind])}/N{s+off_of("sum_inc")})')
            else:
                ws.cell(r, C_FY, f'=SUM(B{r}:M{r})')
            # YTD, budget, variance
            if pct:
                ws.cell(r, C_YTD, f'=IF(P{s+off_of("sum_inc")}=0,"",'
                                  f'P{s+off_of(PCT_PAIR[kind])}/P{s+off_of("sum_inc")})')
                ws.cell(r, C_BUD, f'=IF(Q{s+off_of("sum_inc")}=0,"",'
                                  f'Q{s+off_of(PCT_PAIR[kind])}/Q{s+off_of("sum_inc")})')
                ws.cell(r, C_VAR, f'=IF(OR(P{r}="",Q{r}=""),"",P{r}-Q{r})')
                ws.cell(r, C_PY, f'=IF(U{s+off_of("sum_inc")}=0,"",'
                                 f'U{s+off_of(PCT_PAIR[kind])}/U{s+off_of("sum_inc")})')
                ws.cell(r, C_FYB, f'=IF(Y{s+off_of("sum_inc")}=0,"",'
                                  f'Y{s+off_of(PCT_PAIR[kind])}/Y{s+off_of("sum_inc")})')
            else:
                ws.cell(r, C_YTD, f'=SUMIF($B$5:$M$5,"<="&Setup!$B$5,$B{r}:$M{r})')
                ws.cell(r, C_BUD, f"=SUMIF('Budget FY27'!$B$5:$M$5,\"<=\"&Setup!$B$5,"
                                  f"'Budget FY27'!$B{r}:$M{r})")
                ws.cell(r, C_VAR, f'=P{r}-Q{r}')
                ws.cell(r, C_VARP, f'=IF(Q{r}=0,"",R{r}/ABS(Q{r}))')
                ws.cell(r, C_PY, f"=SUMIF('Actual FY26'!$B$5:$M$5,\"<=\"&Setup!$B$10,"
                                 f"'Actual FY26'!$B{r}:$M{r})")
                ws.cell(r, C_YOY, f'=P{r}-U{r}')
                ws.cell(r, C_YOYP, f'=IF(U{r}=0,"",V{r}/ABS(U{r}))')
                ws.cell(r, C_FYB, f"=SUM('Budget FY27'!$B{r}:$M{r})")
                ws.cell(r, C_FYV, f'=N{r}-Y{r}')
                ws.cell(r, C_RISK,
                        f'=IF(AND($P{r}=0,$Q{r}=0),"No Activity",'
                        f'IF(ABS($R{r})<Setup!$B$16,"Low",'
                        f'IF(ABS($R{r})>=Setup!$B$19,"High",'
                        f'IF(N($S{r})>=Setup!$B$18,"High",'
                        f'IF(N($S{r})<=-Setup!$B$18,"High",'
                        f'IF(ABS(N($S{r}))>=Setup!$B$17,"Medium","Low"))))))')
                ws.cell(r, C_RISK).font = f(10, False)
                ws.cell(r, C_RISK).alignment = Alignment(horizontal='center')
            style_row(ws, r, lvl, kind, allcols, pct)
            for c in (C_VARP, C_YOYP):
                if not pct: ws.cell(r, c).number_format = PCT
            if pct:
                ws.cell(r, C_VAR).number_format = PCT
                for c in (C_VARP, C_YOY, C_YOYP, C_FYV): ws.cell(r, c).value = None
            ws.row_dimensions[r].outlineLevel = lvl
            ws.row_dimensions[r].hidden = lvl == 2
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False, applyStyles=False)

def build_input_sheet(ws, name, fill_from=None):
    """Budget FY27, Actual FY26, Budget FY26  -  same rows, typed numbers."""
    title(ws, name, 'Yellow cells are yours to type. Everything else adds itself up.')
    hdr(ws, 5, ['Line'])
    for i, mc in enumerate(MCOLS):
        c = ws.cell(5, 2 + i, f'=INDEX(Months,{i+1})' if 'FY27' in name
                    else f'=EDATE(INDEX(Months,{i+1}),-12)')
        c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
        c.number_format = 'mmm yy'; c.alignment = Alignment(horizontal='center')
    hdr(ws, 5, ['FY total'], start=C_FY)
    widths(ws, {'A': 46, **{mc: 12 for mc in MCOLS}, 'N': 13})
    ws.freeze_panes = 'B6'
    cols = list(range(2, 14)) + [C_FY]
    for bi, s in enumerate(STARTS + [TOTAL_START]):
        is_total = bi == 6
        dept = None if is_total else DEPTS[bi]
        for off, lab, lvl, kind in BLOCK:
            r = s + off
            pct = kind in PCT_PAIR
            label = ('TOTAL COMPANY' if is_total else f'{dept}') + '  -  net profit after tax' \
                if kind == 'dept' else lab
            ws.cell(r, 1, label)
            typed = (not is_total) and (kind in INPUT_KINDS or isinstance(kind, tuple))
            for i, mc in enumerate(MCOLS):
                if is_total:
                    v = derived(kind, mc, s) if pct else \
                        '=' + '+'.join(f'{mc}{st_ + off}' for st_ in STARTS)
                    ws.cell(r, 2 + i, v)
                elif typed:
                    cell = ws.cell(r, 2 + i)
                    if fill_from:
                        k = f'{dept}|{kind[1]}|{kind[2]}|{i}' if isinstance(kind, tuple) else None
                        if k and k in fill_from: cell.value = fill_from[k]
                    inp(cell, MONEY)
                else:
                    ws.cell(r, 2 + i, derived(kind, mc, s))
            if pct:
                ws.cell(r, C_FY, f'=IF(N{s+off_of("sum_inc")}=0,"",'
                                 f'N{s+off_of(PCT_PAIR[kind])}/N{s+off_of("sum_inc")})')
            else:
                ws.cell(r, C_FY, f'=SUM(B{r}:M{r})')
            style_row(ws, r, lvl, kind, [C_FY] if typed else cols, pct)
            if typed:
                for i in range(12): inp(ws.cell(r, 2 + i), MONEY)
            ws.row_dimensions[r].outlineLevel = lvl
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False, applyStyles=False)

# ========================================================== 8. SUMMARY =====
SUM_LINES = [
    ('{D}', 0, 'ebitda'), ('    Revenue', 1, 'sum_inc'), ('    Cost of sales', 1, 'sum_cos'),
    ('    Gross profit', 1, 'gp'), ('    Gross margin %', 1, 'gp_pct'),
    ('    Overheads', 1, 'sum_exp'), ('    EBITDA', 1, 'ebitda'), ('    EBITDA %', 1, 'ebitda_pct'),
    ('    Depreciation & amortisation', 1, 'dep'), ('    EBIT', 1, 'ebit'),
    ('    Net profit before tax', 1, 'npbt'),
]
SLEN = len(SUM_LINES)
S_STARTS = [6 + i * (SLEN + 1) for i in range(6)]
S_TOTAL = S_STARTS[-1] + SLEN + 1
S_LAST = S_TOTAL + SLEN - 1
SOFF = {k: i for i, (_, _, k) in enumerate(SUM_LINES)}
SPCT = {'gp_pct': 'gp', 'ebitda_pct': 'ebitda'}

sm = wb.create_sheet('Summary')
title(sm, 'Summary  -  department performance')
sm['A3'] = ('=TEXT(Setup!$B$5,"mmm yyyy")&"   |   the reporting month, the year to date and the '
            'full year, each against budget. Click the + and - buttons on the left to open a '
            'department."')
sm['A3'].font = f(9, False, '595959')
SCOLS = [('Month actual', 2, MONEY), ('Month budget', 3, MONEY), ('Var $', 4, MONEY),
         ('Var %', 5, PCT), ('YTD actual', 7, MONEY), ('YTD budget', 8, MONEY),
         ('Var $', 9, MONEY), ('Var %', 10, PCT), ('PY YTD', 12, MONEY), ('YOY %', 13, PCT),
         ('FY forecast', 15, MONEY), ('FY budget', 16, MONEY), ('Var $', 17, MONEY),
         ('Risk', 18, 'General')]
hdr(sm, 5, ['Line'])
for lab, col, _ in SCOLS: hdr(sm, 5, [lab], start=col)
widths(sm, {'A': 46, 'B': 13, 'C': 13, 'D': 13, 'E': 9, 'F': 1.5, 'G': 14, 'H': 14, 'I': 14,
            'J': 9, 'K': 1.5, 'L': 14, 'M': 9, 'N': 1.5, 'O': 14, 'P': 14, 'Q': 14, 'R': 12})
sm.freeze_panes = 'B6'
for bi, s in enumerate(S_STARTS + [S_TOTAL]):
    is_total = bi == 6
    pls = TOTAL_START if is_total else STARTS[bi]
    for i, (lab, lvl, kind) in enumerate(SUM_LINES):
        r = s + i
        pr = pls + off_of(kind)                       # matching row on P&L FY27
        pct = kind in SPCT
        sm.cell(r, 1, 'TOTAL COMPANY  -  EBITDA' if (is_total and lvl == 0) else
                      (DEPTS[bi] + '  -  EBITDA' if lvl == 0 else lab))
        if pct:
            rev, num = f'{s + SOFF["sum_inc"]}', f'{s + SOFF[SPCT[kind]]}'
            for col in (2, 3, 7, 8, 12, 15, 16):
                cc = CL(col)
                sm.cell(r, col, f'=IF({cc}{rev}=0,"",{cc}{num}/{cc}{rev})')
            sm.cell(r, 4, f'=IF(OR($B{r}="",$C{r}=""),"",$B{r}-$C{r})')
            sm.cell(r, 9, f'=IF(OR($G{r}="",$H{r}=""),"",$G{r}-$H{r})')
            sm.cell(r, 17, f'=IF(OR($O{r}="",$P{r}=""),"",$O{r}-$P{r})')
            sm.cell(r, 4).number_format = PCT; sm.cell(r, 9).number_format = PCT
            sm.cell(r, 17).number_format = PCT
        else:
            sm.cell(r, 2, f"=SUMIF('P&L FY27'!$B$5:$M$5,Setup!$B$5,'P&L FY27'!$B{pr}:$M{pr})")
            sm.cell(r, 3, f"=SUMIF('Budget FY27'!$B$5:$M$5,Setup!$B$5,'Budget FY27'!$B{pr}:$M{pr})")
            sm.cell(r, 4, f'=$B{r}-$C{r}')
            sm.cell(r, 5, f'=IF($C{r}=0,"",$D{r}/ABS($C{r}))')
            sm.cell(r, 7, f"='P&L FY27'!$P{pr}")
            sm.cell(r, 8, f"='P&L FY27'!$Q{pr}")
            sm.cell(r, 9, f'=$G{r}-$H{r}')
            sm.cell(r, 10, f'=IF($H{r}=0,"",$I{r}/ABS($H{r}))')
            sm.cell(r, 12, f"='P&L FY27'!$U{pr}")
            sm.cell(r, 13, f'=IF($L{r}=0,"",($G{r}-$L{r})/ABS($L{r}))')
            sm.cell(r, 15, f"='P&L FY27'!$N{pr}")
            sm.cell(r, 16, f"='P&L FY27'!$Y{pr}")
            sm.cell(r, 17, f'=$O{r}-$P{r}')
            sm.cell(r, 18, f"='P&L FY27'!$AA{pr}")
            sm.cell(r, 18).alignment = Alignment(horizontal='center')
        bold = lvl == 0 or kind in ('gp', 'ebitda', 'ebit', 'npbt') or pct
        lc = sm.cell(r, 1); lc.font = f(11 if lvl == 0 else 10, bold, '1F3864' if lvl == 0 else None)
        if lvl == 0: lc.fill = BAND
        for lab2, col, nf in SCOLS:
            cell = sm.cell(r, col)
            cell.font = f(10, bold)
            if nf != 'General': cell.number_format = PCT if pct and nf == MONEY else nf
            if lvl == 0: cell.fill = BAND
        if kind in ('gp', 'ebitda', 'ebit', 'npbt'):
            for c in [1] + [c for _, c, _ in SCOLS]: sm.cell(r, c).border = TOPB
        sm.row_dimensions[r].outlineLevel = lvl
sm.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False, applyStyles=False)

# ==================================================== 9. BALANCE SHEET =====
bs = wb.create_sheet('Balance_Sheet')
title(bs, 'Balance sheet',
      'Pulled from BS_Paste by matching the account name in column A of your export. '
      'Set which columns to read on Setup, rows 24 to 26.')
bs['A3'] = '=TEXT(Setup!$B$5,"mmm yyyy")'
bs['A3'].font = f(9, False, '595959')
hdr(bs, 5, ['Line', 'This month', 'Last month', 'Start of year',
            'Movement in month', 'Movement in year', '', 'Flag'])
widths(bs, {'A': 46, 'B': 15, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 1.5, 'H': 15})
bs.freeze_panes = 'B6'
GROUPS = ['Current Assets', 'Non-Current Assets', 'Current Liabilities',
          'Non-Current Liabilities', 'Equity']
BSPASTE = 'BS_Paste!$A$1:$P$400'
BSNAMES = 'BS_Paste!$A$1:$A$400'
grp_row = {}
r = 6
for g in GROUPS:
    members = [a for a in BSACC if a['group'] == g]
    grp_row[g] = r
    bs.cell(r, 1, g)
    first = r + 1; last = r + len(members)
    for c in (2, 3, 4): bs.cell(r, c, f'=SUM({CL(c)}{first}:{CL(c)}{last})')
    r += 1
    for a in members:
        bs.cell(r, 1, '    ' + a['name']).font = f(9)
        for c, setc in ((2, 'B$24'), (3, 'B$25'), (4, 'B$26')):
            bs.cell(r, c, f'=IF(Setup!${setc}=0,0,IFERROR(INDEX({BSPASTE},'
                          f'MATCH($A{r},{BSNAMES},0),Setup!${setc}),0))')
        bs.cell(r, 8, a['code'] + '  ' + (
            'Cash' if a['cash'] else 'Borrowings' if a['debt'] else 'Trade debtors' if a['debtors']
            else 'Trade creditors' if a['creditors'] else 'WIP' if a['wip'] else ''))
        bs.cell(r, 8).font = f(8, False, '808080')
        bs.row_dimensions[r].outlineLevel = 1
        bs.row_dimensions[r].hidden = True
        r += 1
    r += 1
# the lookup matches on the indented label, so strip the indent back out
for rr in range(6, r):
    v = bs.cell(rr, 1).value
    if isinstance(v, str) and v.startswith('    '):
        for c in (2, 3, 4):
            bs.cell(rr, c).value = bs.cell(rr, c).value.replace(f'$A{rr}', f'TRIM($A{rr})')
TA, TL, NA, TE, CHK = r, r + 1, r + 2, r + 3, r + 5
bs.cell(TA, 1, 'Total assets')
bs.cell(TL, 1, 'Total liabilities')
bs.cell(NA, 1, 'Net assets')
bs.cell(TE, 1, 'Total equity')
bs.cell(CHK, 1, 'Check  -  net assets less total equity')
for c in (2, 3, 4):
    cc = CL(c)
    bs.cell(TA, c, f'={cc}{grp_row["Current Assets"]}+{cc}{grp_row["Non-Current Assets"]}')
    bs.cell(TL, c, f'={cc}{grp_row["Current Liabilities"]}+{cc}{grp_row["Non-Current Liabilities"]}')
    bs.cell(NA, c, f'={cc}{TA}-{cc}{TL}')
    bs.cell(TE, c, f'={cc}{grp_row["Equity"]}')
    bs.cell(CHK, c, f'=ROUND({cc}{NA}-{cc}{TE},2)')
bs.cell(CHK, 5, '=IF(AND(ROUND(B%d,2)=0,ROUND(C%d,2)=0,ROUND(D%d,2)=0),"OK","CHECK")' % (CHK, CHK, CHK))
BS_LAST = r - 2
for rr in range(6, CHK + 1):
    lab = bs.cell(rr, 1)
    if lab.value is None: continue
    top = not str(lab.value).startswith('    ')
    lab.font = f(10, top)
    if top and rr < TA: lab.fill = BAND
    for c in (2, 3, 4, 5, 6):
        cell = bs.cell(rr, c); cell.font = f(10, top); cell.number_format = MONEY
        if top and rr < TA: cell.fill = BAND
    if rr <= BS_LAST or rr in (TA, TL, NA, TE):
        bs.cell(rr, 5, f'=$B{rr}-$C{rr}')
        bs.cell(rr, 6, f'=$B{rr}-$D{rr}')
    if rr in (TA, TL, NA, TE):
        for c in range(1, 7): bs.cell(rr, c).border = TOPB
bs.cell(CHK, 5).number_format = 'General'
bs.cell(CHK, 6).value = None
bs.cell(CHK, 5).font = f(10, True)
# key figures, by flag, for the dashboard
KF = CHK + 2
bs.cell(KF, 1, 'Key figures').font = f(10, True, '1F3864')
FLAGS = [('Cash and bank', 'Cash'), ('Trade debtors', 'Trade debtors'),
         ('Trade creditors', 'Trade creditors'), ('Work in progress', 'WIP'),
         ('Borrowings and lease liabilities', 'Borrowings')]
kf_row = {}
for i, (lab, flag) in enumerate(FLAGS):
    rr = KF + 1 + i; kf_row[flag] = rr
    bs.cell(rr, 1, lab).font = f(10)
    for c in (2, 3, 4):
        cc = CL(c)
        bs.cell(rr, c, f'=SUMIF($H$6:$H${BS_LAST},"*"&"{flag}",{cc}$6:{cc}${BS_LAST})')
        bs.cell(rr, c).number_format = MONEY; bs.cell(rr, c).font = f(10)
bs.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False, applyStyles=False)

# --------- two extra engine columns, so the controls need no array maths ---
hdr(en, 4, ['Reporting month', 'YTD'], start=18)
widths(en, {'R': 15, 'S': 15})
for r in range(ENG_R0, ENG_R1 + 1):
    en.cell(r, 18, f'=INDEX($E{r}:$P{r},Setup!$B$8)').number_format = MON2
    en.cell(r, 19, f'=SUMIF($E$4:$P$4,"<="&Setup!$B$5,$E{r}:$P{r})').number_format = MON2
    en.cell(r, 18).font = f(9); en.cell(r, 19).font = f(9)
EN_CAT = f'Engine!$C${ENG_R0}:$C${ENG_R1}'
EN_MTH = f'Engine!$R${ENG_R0}:$R${ENG_R1}'
EN_YTD = f'Engine!$S${ENG_R0}:$S${ENG_R1}'

# ========================================================= 10. CONTROLS ====
co = wb.create_sheet('Controls')
title(co, 'Controls  -  every line must read OK before you report',
      'The yellow cells hold the row labels as they appear in your Xero exports. '
      'If a label reads "not found", correct the text to match your export exactly.')
hdr(co, 5, ['#', 'What is being checked', 'Pack', 'Source', 'Difference', 'Status',
            '', 'Label in your Xero export', 'Found?'])
widths(co, {'A': 5, 'B': 62, 'C': 15, 'D': 15, 'E': 13, 'F': 11, 'G': 1.5, 'H': 30, 'I': 14})
co.freeze_panes = 'B6'
PLP = 'PL_Paste!$A$1:$P$300'; PLN = 'PL_Paste!$A$1:$A$300'
TOT = TOTAL_START
CAT_LBL = [('Income', 'Total Income'), ('Cost of Sales', 'Total Cost of Sales'),
           ('Expenses', 'Total Operating Expenses'), ('Other Income', 'Total Other Income'),
           ('Other Expenses', 'Total Other Expenses')]
r = 6; cno = 1
def ctl(label, pack, src, mode='match', note=None, gate='Setup!$B$30>0'):
    """mode: match (pack vs src must agree) | zero (pack must be nil) | info
    gate: while this is false the check reads NO DATA instead of a misleading OK."""
    global r, cno
    co.cell(r, 1, f'C{cno}').font = f(10, False, '808080')
    c = co.cell(r, 2, label); c.font = f(10); c.alignment = Alignment(wrap_text=True, vertical='top')
    if pack is not None:
        co.cell(r, 3, pack).number_format = MON2; co.cell(r, 3).font = f(10)
    if src is not None:
        co.cell(r, 4, src).number_format = MON2; co.cell(r, 4).font = f(10)
    if mode == 'match':
        co.cell(r, 5, f'=ROUND(N($C{r})-N($D{r}),2)').number_format = MON2
        co.cell(r, 6, f'=IF(NOT({gate}),"NO DATA",IF(ABS($E{r})<1,"OK","CHECK"))')
    elif mode == 'zero':
        co.cell(r, 5, f'=ROUND(N($C{r}),2)').number_format = MON2
        co.cell(r, 6, f'=IF(NOT({gate}),"NO DATA",IF(ABS($E{r})<0.5,"OK","CHECK"))')
    else:
        co.cell(r, 6, 'INFO')
    co.cell(r, 5).font = f(10)
    co.cell(r, 6).font = f(10, True); co.cell(r, 6).alignment = Alignment(horizontal='center')
    if note:
        co.cell(r, 8, note).font = f(9, False, '595959')
    out = r; r += 1; cno += 1
    return out

ctl('The cleanup layer adds up to the same figure as the engine. Proves no GL line was '
    'classified but then lost on the way into the engine.',
    f'=ROUND(SUMIF(Cleanup!$J${CL_R0}:$J${CL_R1},">0",{CL_AMT}),2)',
    f'=ROUND(SUM(Engine!$E${ENG_R0}:$P${ENG_R0 + len(DEPTS)*len(SUBS) - 1}),2)')
ctl('The engine adds up to the same figure as the P&L, year to date. Proves no line was left '
    'off the face of the report. The full-year columns cannot be compared this way, because '
    'months after the reporting month come from the budget, not the GL.',
    f'=ROUND(SUM(Engine!$S${ENG_R0}:$S${ENG_R0 + len(DEPTS)*len(SUBS) - 1}),2)',
    f"=ROUND('P&L FY27'!$P${TOT + off_of('npat')},2)")
lbl_rows = {}
for cat, deflbl in CAT_LBL:
    rr = ctl(f'{cat} for the reporting month agrees with the Xero P&L you pasted.',
             f'=SUMIF({EN_CAT},"{cat}",{EN_MTH})',
             None, gate='AND(Setup!$B$30>0,Setup!$B$39>0)')
    inp(co.cell(rr, 8), 'General'); co.cell(rr, 8).value = deflbl
    co.cell(rr, 4, f'=IF($H{rr}="",0,IFERROR(INDEX({PLP},MATCH($H{rr},{PLN},0),'
                   f'Setup!$B$23),0))')
    co.cell(rr, 4).number_format = MON2; co.cell(rr, 4).font = f(10)
    co.cell(rr, 9, f'=IF($H{rr}="","no label",IF(ISNA(MATCH($H{rr},{PLN},0)),'
                   f'"not found","found"))')
    co.cell(rr, 9).font = f(10); co.cell(rr, 9).alignment = Alignment(horizontal='center')
    lbl_rows[cat] = rr
ctl('Net profit for the reporting month agrees with the Xero P&L you pasted.',
    f'=SUM($C{lbl_rows["Income"]}:$C{lbl_rows["Other Expenses"]})',
    f'=SUM($D{lbl_rows["Income"]}:$D{lbl_rows["Other Expenses"]})',
    gate='AND(Setup!$B$30>0,Setup!$B$39>0)')
R_BSCHK = ctl('The balance sheet balances. Net assets equal total equity.',
              f'=ROUND(Balance_Sheet!$B${CHK},2)', None, 'zero', gate='Setup!$B$38>0')
ctl('Every GL line carries a department. Anything here is sitting in no department at all.',
    '=Setup!$B$34', None, 'zero')
ctl('Every GL line sits on an account the pack knows. Anything here is left out of the P&L '
    'altogether, so add the account to Lists.',
    '=Setup!$B$35', None, 'zero')
ctl('Nothing in the GL paste is dated outside this financial year, and nothing is dated after '
    'the reporting month. Either way those lines are left out.',
    f'=Setup!$B$33+COUNTIF(Cleanup!$H${CL_R0}:$H${CL_R1},">"&Setup!$B$8)', None, 'zero')
ctl('The GL paste is inside capacity.', '=MAX(0,Setup!$B$30-Setup!$B$31)', None, 'zero')
R_PY = ctl("Last year's actuals are loaded, so the PY and YOY columns mean something.",
           f"=ROUND(SUM('Actual FY26'!$B${TOT + off_of('sum_inc')}:$M${TOT + off_of('sum_inc')}),2)",
           None, 'info')
co.cell(R_PY, 6, f'=IF(ABS(N($C{R_PY}))>0,"OK","EMPTY")')
co.cell(R_PY, 6).font = f(10, True); co.cell(R_PY, 6).alignment = Alignment(horizontal='center')
R_PLH = ctl('The P&L paste column you chose on Setup B23 has this heading.', None, None, 'info')
co.cell(R_PLH, 3, '=IFERROR(INDEX(PL_Paste!$A$1:$P$10,1,Setup!$B$23)&"","")')
co.cell(R_PLH, 3).number_format = 'General'; co.cell(R_PLH, 3).font = f(10)
R_BSH = ctl('The BS paste columns you chose on Setup B24 to B26 have these headings.',
            None, None, 'info')
co.cell(R_BSH, 3, '=IFERROR(INDEX(BS_Paste!$A$1:$P$10,1,Setup!$B$24)&"","")')
co.cell(R_BSH, 4, '=IFERROR(INDEX(BS_Paste!$A$1:$P$10,1,Setup!$B$25)&" / "&'
                  'INDEX(BS_Paste!$A$1:$P$10,1,Setup!$B$26),"")')
for c in (3, 4):
    co.cell(R_BSH, c).number_format = 'General'; co.cell(R_BSH, c).font = f(10)
CTL_FIRST, CTL_LAST = 6, r - 1
co.cell(r + 1, 2, f'=IF(COUNTIF($F${CTL_FIRST}:$F${CTL_LAST},"NO DATA")>0,'
                  f'COUNTIF($F${CTL_FIRST}:$F${CTL_LAST},"NO DATA")&" check(s) have nothing to '
                  f'work with yet. Paste the reports first.",'
                  f'IF(COUNTIF($F${CTL_FIRST}:$F${CTL_LAST},"CHECK")=0,'
                  f'"All {CTL_LAST - CTL_FIRST + 1} checks are clean. Safe to report.",'
                  f'COUNTIF($F${CTL_FIRST}:$F${CTL_LAST},"CHECK")&" of '
                  f'{CTL_LAST - CTL_FIRST + 1} checks need attention"))')
co.cell(r + 1, 2).font = f(11, True, 'C00000')

# ======================================================== 11. DASHBOARD ====
db = wb.create_sheet('Dashboard')
title(db, 'Financial controller dashboard')
db['A3'] = '=TEXT(Setup!$B$5,"mmmm yyyy")&"   |   FY"&Setup!$B$7&"   |   period "&Setup!$B$8&" of 12"'
db['A3'].font = f(10, True, '1F3864')
db['A4'] = ('=IF(Setup!$B$30=0,"Nothing pasted yet. Paste the GL, the Xero P&L and the balance '
            'sheet, then read the Controls sheet.",'
            'IF(COUNTIF(Controls!$F$%d:$F$%d,"CHECK")=0,"All controls clean. Safe to report.",'
            'COUNTIF(Controls!$F$%d:$F$%d,"CHECK")&" control(s) need attention. '
            'Do not report these figures until the Controls sheet is clean."))'
            % (CTL_FIRST, CTL_LAST, CTL_FIRST, CTL_LAST))
db['A4'].font = f(10, True, 'C00000')
widths(db, {'A': 42, 'B': 15, 'C': 15, 'D': 14, 'E': 9, 'F': 15, 'G': 9, 'H': 15, 'I': 15,
            'J': 11, 'K': 2})
PLR = lambda kind: TOT + off_of(kind)
PLE = lambda cat, sub: TOT + off_eng(cat, sub)

def sect(row, text):
    c = db.cell(row, 1, text); c.font = f(11, True, 'FFFFFF'); c.fill = NAVY
    for cc in range(2, 11):
        db.cell(row, cc).fill = NAVY
    return row + 1

r = 6
r = sect(r, 'TRADING  -  year to date')
hdr(db, r, ['', 'YTD actual', 'YTD budget', 'Var $', 'Var %', 'PY YTD', 'YOY %',
            'FY forecast', 'FY budget', 'Risk'])
r += 1
HEAD = [('Revenue', 'sum_inc', False), ('Cost of sales', 'sum_cos', False),
        ('Gross profit', 'gp', False), ('Gross margin %', 'gp_pct', True),
        ('Overheads', 'sum_exp', False), ('EBITDA', 'ebitda', False),
        ('EBITDA %', 'ebitda_pct', True),
        ('Depreciation & amortisation', 'dep', False), ('EBIT', 'ebit', False),
        ('Net profit before tax', 'npbt', False), ('Net margin %', 'net_pct', True)]
HROW = {}
for lab, kind, pct in HEAD:
    pr = PLR(kind); HROW[kind] = r
    db.cell(r, 1, lab).font = f(10, kind in ('gp', 'ebitda', 'ebit', 'npbt') or pct)
    for c, src in ((2, 'P'), (3, 'Q'), (4, 'R'), (5, 'S'), (6, 'U'), (8, 'N'), (9, 'Y')):
        db.cell(r, c, f"='P&L FY27'!${src}${pr}")
    db.cell(r, 7, f'=IF(N($F{r})=0,"",($B{r}-$F{r})/ABS($F{r}))')
    db.cell(r, 10, f"='P&L FY27'!$AA${pr}")
    for c in range(2, 10):
        cell = db.cell(r, c)
        cell.font = f(10, kind in ('gp', 'ebitda', 'ebit', 'npbt') or pct)
        cell.number_format = PCT if (pct and c != 4) else (PCT if c in (5, 7) else MONEY)
    if pct:
        db.cell(r, 4).number_format = PCT
        for c in (5, 7, 10): db.cell(r, c).value = None
    db.cell(r, 10).alignment = Alignment(horizontal='center'); db.cell(r, 10).font = f(10)
    if kind in ('gp', 'ebitda', 'ebit', 'npbt'):
        for c in range(1, 11): db.cell(r, c).border = TOPB
    r += 1

r += 1
r = sect(r, 'BY DEPARTMENT  -  year to date')
hdr(db, r, ['', 'Revenue', 'Gross profit', 'Gross margin %', 'Overheads', 'EBITDA',
            'EBITDA %', 'EBITDA budget', 'Var $', 'Risk'])
r += 1
DEPT_FIRST = r
for i, d in enumerate(DEPTS):
    s = STARTS[i]
    db.cell(r, 1, d).font = f(10, True)
    db.cell(r, 2, f"='P&L FY27'!$P${s + off_of('sum_inc')}")
    db.cell(r, 3, f"='P&L FY27'!$P${s + off_of('gp')}")
    db.cell(r, 4, f'=IF($B{r}=0,"",$C{r}/$B{r})')
    db.cell(r, 5, f"='P&L FY27'!$P${s + off_of('sum_exp')}")
    db.cell(r, 6, f"='P&L FY27'!$P${s + off_of('ebitda')}")
    db.cell(r, 7, f'=IF($B{r}=0,"",$F{r}/$B{r})')
    db.cell(r, 8, f"='P&L FY27'!$Q${s + off_of('ebitda')}")
    db.cell(r, 9, f'=$F{r}-$H{r}')
    db.cell(r, 10, f"='P&L FY27'!$AA${s + off_of('ebitda')}")
    for c in range(2, 10):
        db.cell(r, c).number_format = PCT if c in (4, 7) else MONEY
        db.cell(r, c).font = f(10)
    db.cell(r, 10).alignment = Alignment(horizontal='center'); db.cell(r, 10).font = f(10)
    r += 1
DEPT_LAST = r - 1
db.cell(r, 1, 'TOTAL').font = f(10, True)
for c in (2, 3, 5, 6, 8, 9):
    db.cell(r, c, f'=SUM({CL(c)}{DEPT_FIRST}:{CL(c)}{DEPT_LAST})')
    db.cell(r, c).number_format = MONEY; db.cell(r, c).font = f(10, True)
db.cell(r, 4, f'=IF($B{r}=0,"",$C{r}/$B{r})').number_format = PCT
db.cell(r, 7, f'=IF($B{r}=0,"",$F{r}/$B{r})').number_format = PCT
db.cell(r, 4).font = f(10, True); db.cell(r, 7).font = f(10, True)
for c in range(1, 11): db.cell(r, c).border = TOPB
r += 2

r = sect(r, 'COST STRUCTURE  -  year to date, as a share of revenue')
hdr(db, r, ['', '$ YTD', '% of revenue', 'Budget %', 'Var pp'])
r += 1
REVR = HROW['sum_inc']
COSTS = [('Direct labour', [('Cost of Sales', s) for s in LABOUR]),
         ('Sub-contract labour', [('Cost of Sales', 'Sub-Contract Labour')]),
         ('Equipment purchases', [('Cost of Sales', 'Equipment - Purchase')]),
         ('Equipment hires', [('Cost of Sales', 'Equipment - Hires')]),
         ('Subscriptions & licences', [('Cost of Sales', 'Subscriptions & Licences')]),
         ('All other cost of sales', None),
         ('Overheads', None)]
for lab, members in COSTS:
    db.cell(r, 1, lab).font = f(10)
    if lab == 'Overheads':
        db.cell(r, 2, f"='P&L FY27'!$P${PLR('sum_exp')}")
        db.cell(r, 4, f"=IF('P&L FY27'!$Q${PLR('sum_inc')}=0,\"\","
                      f"'P&L FY27'!$Q${PLR('sum_exp')}/'P&L FY27'!$Q${PLR('sum_inc')})")
    elif lab == 'All other cost of sales':
        prev = f'=ROUND(\'P&L FY27\'!$P${PLR("sum_cos")}-SUM($B{r-5}:$B{r-1}),2)'
        db.cell(r, 2, prev)
        db.cell(r, 4, '')
    else:
        db.cell(r, 2, '=' + '+'.join(f"'P&L FY27'!$P${PLE(c, s)}" for c, s in members))
        db.cell(r, 4, '=IF(\'P&L FY27\'!$Q$%d=0,"",(%s)/\'P&L FY27\'!$Q$%d)'
                % (PLR('sum_inc'),
                   '+'.join(f"'P&L FY27'!$Q${PLE(c, s)}" for c, s in members),
                   PLR('sum_inc')))
    db.cell(r, 3, f'=IF($B{REVR}=0,"",$B{r}/$B{REVR})')
    db.cell(r, 5, f'=IF(OR($C{r}="",$D{r}=""),"",$C{r}-$D{r})')
    db.cell(r, 2).number_format = MONEY
    for c in (3, 4, 5): db.cell(r, c).number_format = PCT
    for c in range(2, 6): db.cell(r, c).font = f(10)
    r += 1
r += 1

r = sect(r, 'BALANCE SHEET, CASH AND WORKING CAPITAL')
hdr(db, r, ['', 'This month', 'Last month', 'Start of year', 'Movement in month'])
r += 1
B = lambda row: f'Balance_Sheet!$B${row}'
C_ = lambda row: f'Balance_Sheet!$C${row}'
D_ = lambda row: f'Balance_Sheet!$D${row}'
CA, CLb = grp_row['Current Assets'], grp_row['Current Liabilities']
BSROWS = [
    ('Cash and bank', kf_row['Cash'], MONEY), ('Trade debtors', kf_row['Trade debtors'], MONEY),
    ('Work in progress', kf_row['WIP'], MONEY),
    ('Trade creditors', kf_row['Trade creditors'], MONEY),
    ('Borrowings and lease liabilities', kf_row['Borrowings'], MONEY),
    ('Total current assets', CA, MONEY), ('Total current liabilities', CLb, MONEY),
    ('Total assets', TA, MONEY), ('Total liabilities', TL, MONEY),
    ('Net assets', NA, MONEY), ('Total equity', TE, MONEY),
]
DBR = {}
for lab, row, nf in BSROWS:
    DBR[lab] = r
    db.cell(r, 1, lab).font = f(10, lab.startswith(('Total', 'Net')))
    for c, g in ((2, B), (3, C_), (4, D_)):
        db.cell(r, c, '=' + g(row)); db.cell(r, c).number_format = nf
        db.cell(r, c).font = f(10, lab.startswith(('Total', 'Net')))
    db.cell(r, 5, f'=$B{r}-$C{r}').number_format = nf; db.cell(r, 5).font = f(10)
    r += 1
r += 1
hdr(db, r, ['Ratios and days', 'This month', '', '', 'How it is worked out'])
r += 1
EBYTD = f'$B${HROW["ebitda"]}'
RATIOS = [
    ('Working capital', f'=$B${DBR["Total current assets"]}-$B${DBR["Total current liabilities"]}',
     MONEY, 'Current assets less current liabilities.'),
    ('Current ratio', f'=IF($B${DBR["Total current liabilities"]}=0,"",'
     f'$B${DBR["Total current assets"]}/$B${DBR["Total current liabilities"]})', RATIO,
     'Current assets divided by current liabilities.'),
    ('Quick ratio', f'=IF($B${DBR["Total current liabilities"]}=0,"",'
     f'($B${DBR["Total current assets"]}-$B${DBR["Work in progress"]})'
     f'/$B${DBR["Total current liabilities"]})', RATIO,
     'Current assets less work in progress, divided by current liabilities.'),
    ('Debtor days', f'=IF($B${REVR}=0,"",$B${DBR["Trade debtors"]}/($B${REVR}/Setup!$B$13))',
     NUM1, 'Trade debtors divided by average daily revenue since 1 July. Debtors include GST and '
            'revenue does not, so this runs about 10 per cent high.'),
    ('Creditor days', f'=IF(($B${HROW["sum_cos"]}+$B${HROW["sum_exp"]})=0,"",'
     f'$B${DBR["Trade creditors"]}/(ABS($B${HROW["sum_cos"]}+$B${HROW["sum_exp"]})/Setup!$B$13))',
     NUM1, 'Trade creditors divided by average daily cost of sales plus overheads since 1 July. '
            'Creditors include GST.'),
    ('Net debt', f'=$B${DBR["Borrowings and lease liabilities"]}-$B${DBR["Cash and bank"]}',
     MONEY, 'Borrowings and lease liabilities less cash.'),
    ('Net debt to EBITDA', f'=IF(Setup!$B$8=0,"",IF({EBYTD}<=0,"",'
     f'$B{r+5}/({EBYTD}/Setup!$B$8*12)))', RATIO,
     'Net debt divided by EBITDA annualised from the year to date. Blank while EBITDA is negative.'),
    ('Gearing', f'=IF($B${DBR["Total equity"]}=0,"",$B{r+4}/$B${DBR["Total equity"]})',
     PCT, 'Net debt divided by total equity.'),
    ('Return on revenue', f'=IF($B${REVR}=0,"",$B${HROW["npbt"]}/$B${REVR})', PCT,
     'Net profit before tax divided by revenue, year to date.'),
]
RAT_FIRST = r
for lab, fx, nf, how in RATIOS:
    db.cell(r, 1, lab).font = f(10)
    db.cell(r, 2, fx); db.cell(r, 2).number_format = nf; db.cell(r, 2).font = f(10, True)
    c = db.cell(r, 5, how); c.font = f(9, False, '595959')
    c.alignment = Alignment(wrap_text=True, vertical='top')
    db.row_dimensions[r].height = 13 * (1 + len(how) // 95)
    r += 1
# net debt and net-debt-to-EBITDA reference rows fixed up now the rows are known
ND = RAT_FIRST + 5
db.cell(RAT_FIRST + 6, 2, f'=IF(Setup!$B$8=0,"",IF({EBYTD}<=0,"",$B{ND}/({EBYTD}/Setup!$B$8*12)))')
db.cell(RAT_FIRST + 7, 2, f'=IF($B${DBR["Total equity"]}=0,"",$B{ND}/$B${DBR["Total equity"]})')
DB_LAST = r

# ====================================================== 12. UTILISATION ====
ut = wb.create_sheet('Utilisation')
title(ut, 'Utilisation',
      'Working days and working hours calculate themselves from the calendar and the public '
      'holidays on Lists. Type the hours in the yellow cells.')
widths(ut, {'A': 18, 'B': 16, 'C': 18, **{CL(c): 12 for c in range(4, 16)}})
ut.cell(4, 1, 'Calendar').font = f(10, True, '1F3864')
for i in range(12):
    c = ut.cell(4, 4 + i, f'=INDEX(Months,{i+1})')
    c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
    c.number_format = 'mmm yy'; c.alignment = Alignment(horizontal='center')
ut.cell(5, 1, 'Working days').font = f(10)
ut.cell(6, 1, 'Working hours').font = f(10)
for i in range(12):
    cc = CL(4 + i)
    ut.cell(5, 4 + i, f'=NETWORKDAYS({cc}$4,EOMONTH({cc}$4,0),Holidays)').number_format = '0'
    ut.cell(6, 4 + i, f'={cc}5*Setup!$B$11').number_format = '#,##0'
    ut.cell(5, 4 + i).font = f(10); ut.cell(6, 4 + i).font = f(10)
hdr(ut, 8, ['FY (year it ends)', 'Department', 'Hour Type'])
for i in range(12):
    c = ut.cell(8, 4 + i, f'=INDEX(Months,{i+1})')
    c.font = f(10, True, 'FFFFFF'); c.fill = NAVY; c.border = BOX
    c.number_format = 'mmm yy'; c.alignment = Alignment(horizontal='center')
hdr(ut, 8, ['FY total'], start=16)
ut.freeze_panes = 'D9'
UT_ROW = {}
r = 9
for d in ['Onsite', 'Production', 'Video', 'Integration', 'Consulting', 'CTS']:
    base = r
    for kind in ['Chargeable', 'Non-chargeable', 'Leave']:
        ut.cell(r, 1, '=Setup!$B$7').font = f(10); ut.cell(r, 1).number_format = '0'
        ut.cell(r, 2, d).font = f(10)
        ut.cell(r, 3, kind).font = f(10)
        for i in range(12): inp(ut.cell(r, 4 + i), '#,##0.0')
        ut.cell(r, 16, f'=SUM(D{r}:O{r})').number_format = '#,##0.0'
        ut.cell(r, 16).font = f(10)
        r += 1
    ut.cell(r, 3, 'Utilisation %').font = f(10, True)
    ut.cell(r + 1, 3, 'FTE').font = f(10, True)
    for i in range(12):
        cc = CL(4 + i)
        ut.cell(r, 4 + i, f'=IF(({cc}{base}+{cc}{base+1})=0,"",{cc}{base}/({cc}{base}+{cc}{base+1}))')
        ut.cell(r, 4 + i).number_format = PCT; ut.cell(r, 4 + i).font = f(10, True)
        ut.cell(r + 1, 4 + i, f'=IF({cc}$6=0,"",({cc}{base}+{cc}{base+1})/{cc}$6)')
        ut.cell(r + 1, 4 + i).number_format = '0.0'; ut.cell(r + 1, 4 + i).font = f(10, True)
    ut.cell(r, 16, f'=IF((P{base}+P{base+1})=0,"",P{base}/(P{base}+P{base+1}))')
    ut.cell(r, 16).number_format = PCT; ut.cell(r, 16).font = f(10, True)
    UT_ROW[d] = r
    r += 3

# ============================================= 13. the three input sheets ==
pl = wb.create_sheet('P&L FY27')
bf27 = wb.create_sheet('Budget FY27')
af26 = wb.create_sheet('Actual FY26')
bf26 = wb.create_sheet('Budget FY26')
build_input_sheet(bf27, 'Budget FY27', fill_from=BUDGET)
build_input_sheet(af26, 'Actual FY26')
build_input_sheet(bf26, 'Budget FY26')
build_pl(pl)

# ============================================= 14. conditional formatting ==
from openpyxl.formatting.rule import FormulaRule
RED = PatternFill('solid', fgColor='FFC7CE')
AMB = PatternFill('solid', fgColor='FFEB9C')
GRN = PatternFill('solid', fgColor='C6EFCE')
rng = f'$A$6:$AA${PL_LAST}'
pl.conditional_formatting.add(rng, FormulaRule(formula=[f'$AA6="High"'], fill=RED, stopIfTrue=False))
pl.conditional_formatting.add(rng, FormulaRule(formula=[f'$AA6="Medium"'], fill=AMB, stopIfTrue=False))
srng = f'$A$6:$R${S_LAST}'
sm.conditional_formatting.add(srng, FormulaRule(formula=['$R6="High"'], fill=RED))
sm.conditional_formatting.add(srng, FormulaRule(formula=['$R6="Medium"'], fill=AMB))
for sheet, col, first, last in [(db, 'J', 7, DEPT_LAST), (co, 'F', CTL_FIRST, CTL_LAST)]:
    pass
co.conditional_formatting.add(f'$F${CTL_FIRST}:$F${CTL_LAST}',
                              FormulaRule(formula=[f'$F{CTL_FIRST}="CHECK"'], fill=RED))
co.conditional_formatting.add(f'$F${CTL_FIRST}:$F${CTL_LAST}',
                              FormulaRule(formula=[f'$F{CTL_FIRST}="OK"'], fill=GRN))
co.conditional_formatting.add(f'$F${CTL_FIRST}:$F${CTL_LAST}',
                              FormulaRule(formula=[f'$F{CTL_FIRST}="NO DATA"'], fill=AMB))
db.conditional_formatting.add(f'$J$7:$J${DEPT_LAST}',
                              FormulaRule(formula=['$J7="High"'], fill=RED))
db.conditional_formatting.add(f'$J$7:$J${DEPT_LAST}',
                              FormulaRule(formula=['$J7="Medium"'], fill=AMB))

# ============================================================= 15. CHARTS ==
ch1 = LineChart(); ch1.title = 'Revenue by month  -  actual and forecast against budget'
ch1.height, ch1.width = 7.5, 20; ch1.y_axis.numFmt = '$#,##0'
rev = PLR('sum_inc')
ch1.add_data(Reference(pl, min_col=2, max_col=13, min_row=rev, max_row=rev), from_rows=True)
ch1.add_data(Reference(bf27, min_col=2, max_col=13, min_row=rev, max_row=rev), from_rows=True)
ch1.set_categories(Reference(pl, min_col=2, max_col=13, min_row=5, max_row=5))
ch1.series[0].tx = None; ch1.series[1].tx = None
db.add_chart(ch1, f'A{DB_LAST + 2}')

ch2 = LineChart(); ch2.title = 'EBITDA by month  -  actual and forecast against budget'
ch2.height, ch2.width = 7.5, 20; ch2.y_axis.numFmt = '$#,##0'
eb = PLR('ebitda')
ch2.add_data(Reference(pl, min_col=2, max_col=13, min_row=eb, max_row=eb), from_rows=True)
ch2.add_data(Reference(bf27, min_col=2, max_col=13, min_row=eb, max_row=eb), from_rows=True)
ch2.set_categories(Reference(pl, min_col=2, max_col=13, min_row=5, max_row=5))
db.add_chart(ch2, f'A{DB_LAST + 18}')

ch3 = BarChart(); ch3.title = 'Gross margin % by department  -  year to date'
ch3.height, ch3.width = 7.5, 20; ch3.y_axis.numFmt = '0%'
ch3.add_data(Reference(db, min_col=4, min_row=DEPT_FIRST, max_row=DEPT_LAST), titles_from_data=False)
ch3.set_categories(Reference(db, min_col=1, min_row=DEPT_FIRST, max_row=DEPT_LAST))
db.add_chart(ch3, f'A{DB_LAST + 34}')

# ============================================================== 16. FINISH =
ORDER = ['README', 'Setup', 'GL_Paste', 'PL_Paste', 'BS_Paste', 'Controls', 'Dashboard',
         'Summary', 'P&L FY27', 'Balance_Sheet', 'Budget FY27', 'Actual FY26', 'Budget FY26',
         'Utilisation', 'Engine', 'Lists', 'Cleanup']
wb._sheets = [wb[n] for n in ORDER]
for ws_ in wb.worksheets:
    ws_.sheet_view.showGridLines = False
wb['Dashboard'].sheet_view.tabSelected = True
wb.active = ORDER.index('Dashboard')
# assigned last, because our summary rows sit ABOVE their detail
for n in ('P&L FY27', 'Summary', 'Balance_Sheet', 'Budget FY27', 'Actual FY26', 'Budget FY26'):
    wb[n].sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False,
                                               applyStyles=False)
wb.save(OUT)
print('wrote', OUT)
print('P&L rows 6..%d, total block at %d' % (PL_LAST, TOTAL_START))
print('Summary rows 6..%d, total block at %d' % (S_LAST, S_TOTAL))
print('Engine rows %d..%d   Cleanup rows %d..%d' % (ENG_R0, ENG_R1, CL_R0, CL_R1))
