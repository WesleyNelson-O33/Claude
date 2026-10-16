"""The charts for the two new sheets."""
import os, json
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
ACC = json.load(open(f'{HERE}/data/overhead_accounts.json'))
INK, WHITE, GREY = B['charcoal'], B['white'], B['grey']
NAVY, BLUE, PURP, AQUA = B['deep_blue'], B['primary_blue'], B['purple'], B['bright_teal']

wb = openpyxl.load_workbook(OUT)

# -------------------------------------------------- Top 10 YoY chart ------
ws = wb['Top 10 YoY']
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=4, max_row=14),
            titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=14))
ch.title = 'This year against the same three months last year'
ch.width, ch.height = 22, 12
ch.style = None
ch.y_axis.majorGridlines = None
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
ch.dataLabels.numFmt = '#,##0'
ch.legend = openpyxl.chart.legend.Legend(); ch.legend.position = 'b'
for s, col in zip(ch.series, [NAVY, GREY]):
    s.graphicalProperties = GraphicalProperties(solidFill=col)
    s.graphicalProperties.line = LineProperties(noFill=True)
ws.add_chart(ch, 'J4')

# ----------------------------------------------- Overheads MoM chart ------
ws = wb['Overheads MoM']
HR, TOT = 4, 5
SUB = [r for r in range(1, ws.max_row + 1)
       if ws.cell(r, 1).value in ACC]
P = ws.max_row + 2
ws.cell(P, 1, 'CHART DATA, the same figures as spend rather than as a negative')
for c in range(1, 14):
    ws.cell(P, c).fill = PatternFill('solid', fgColor=INK)
    ws.cell(P, c).font = Font(name='Calibri', size=10, bold=True, color=WHITE)
for i in range(12):
    ws.cell(P + 1, 2 + i, f'={gl(2+i)}${HR}').number_format = 'mmm\\-yy'
ws.cell(P + 1, 1, 'Line')
for j, sr in enumerate(SUB):
    r = P + 2 + j
    ws.cell(r, 1, f'=$A${sr}')
    for i in range(12):
        ws.cell(r, 2 + i, f'=-{gl(2+i)}${sr}').number_format = '#,##0'
LAST = P + 1 + len(SUB)

ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100
ch.add_data(Reference(ws, min_col=1, min_row=P + 2, max_col=13, max_row=LAST),
            from_rows=True, titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=P + 1, max_col=13))
ch.title = 'Overhead spend, month by month'
ch.width, ch.height = 30, 12
ch.style = None
ch.y_axis.majorGridlines = None
ch.legend = openpyxl.chart.legend.Legend(); ch.legend.position = 'b'
for s, col in zip(ch.series, [NAVY, BLUE, PURP, INK, GREY, AQUA]):
    s.graphicalProperties = GraphicalProperties(solidFill=col)
    s.graphicalProperties.line = LineProperties(noFill=True)
ws.add_chart(ch, f'A{LAST + 3}')

wb.save(OUT)
print(f'charts added. Overheads chart data rows {P+2} to {LAST}, {len(SUB)} series')
