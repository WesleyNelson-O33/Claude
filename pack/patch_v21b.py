"""The year on year chart for the top ten."""
import os, json
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
TEAL, GREY = B['dark_teal'], B['grey']

wb = openpyxl.load_workbook(OUT)
ws = wb['Top 10 Clients']
R0 = 106

ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=R0 + 2, max_row=R0 + 12),
            titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=R0 + 3, max_row=R0 + 12))
ch.title = 'Top ten this year against the same three months last year'
ch.width, ch.height = 22, 12
ch.style = None
ch.y_axis.majorGridlines = None
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
ch.dataLabels.numFmt = '#,##0'
ch.legend = openpyxl.chart.legend.Legend(); ch.legend.position = 'b'
for s, col in zip(ch.series, [TEAL, GREY]):
    s.graphicalProperties = GraphicalProperties(solidFill=col)
    s.graphicalProperties.line = LineProperties(noFill=True)
ws.add_chart(ch, f'J{R0}')

wb.save(OUT)
print('year on year chart added at J%d' % R0)
