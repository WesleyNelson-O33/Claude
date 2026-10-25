"""Put the headcount back to what the P&L analysis actually uses.

patch_v72 set the headcount to 49/18/3/1/1 over 72, on my reading that the
analysis's Staff key was 0.680556 / 0.25 / 0.041667 / 0.013889 / 0.013889.
That reading was wrong.  The analysis's own percentage columns on
'4. PL Per Department' hold 0.582090 / 0.358209 / 0.044776 / 0.014925 / 0,
which is 39 / 24 / 3 / 1 / 0 over 67.  That is the headcount the pack
already carried, so this restores it.
"""
import openpyxl

WB = 'pack/CTS Financial Controller Pack - v20.xlsx'
HEADCOUNT = {'Support': 39, 'Events': 24, 'Video': 3, 'Integration': 1,
             'Consulting': 0, 'Admin': 7}

wb = openpyxl.load_workbook(WB)
ws = wb['Overhead Split']
assert ws['A5'].value == 'Department' and ws['B5'].value == 'Headcount'

for r in range(6, 12):
    name = ws.cell(r, 1).value
    before = ws.cell(r, 2).value
    ws.cell(r, 2, HEADCOUNT[name])
    print(f'  {name:<12} {before} -> {HEADCOUNT[name]}')

wb.save(WB)
print('saved')
