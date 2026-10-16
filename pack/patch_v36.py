"""An office staff split matrix, so nothing has to be typed.

Gi fills that cell by hand off the labour hours report. The basis she names in
the video is labour hours: "this one here changes every single month based on
this here. That's labour hours. So update based on utilisation."

So the matrix allocates the admin employment pot by each department's share of
worked hours, chargeable plus non-chargeable, off the Utilisation tab. Leave is
out, because the report Gi uses sums worked hours.

It is laid out like her other three keys on this tab, the percentages are on
the face of it, and it recalculates every month on its own.

One thing to be straight about: this will not reproduce her historical figures
to the dollar. For August it gives Support 62,957 against her 57,727. Her
figure also carries a fixed split component that only shows up in the Other CC
tabs of the labour hours report, which is not in this pack.
"""
import os, json, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = f'{HERE}/CTS Financial Controller Pack - v20.xlsx'
B = json.load(open(f'{HERE}/data/brand.json'))
STAFF = json.load(open(f'{HERE}/data/staffsplit.json'))
POT = STAFF['pot']
INK, WHITE, GREY, GREYL, AQUA = (B['charcoal'], B['white'], B['grey'],
                                 B['grey_lt'], B['bright_teal'])
COST = '[Red]"($"#,##0\\);\\$#,##0;\\-'
PCT = '0.0%;[Red](0.0%);\\-'
HRS = '#,##0.0;[Red](#,##0.0);\\-'
# department -> its chargeable and non-chargeable rows on the Utilisation tab
UTIL = {'Support': (9, 10), 'Events': (15, 16), 'Video': (21, 22),
        'Integration': (27, 28), 'Consulting': (33, 34)}
OMM = "'Overheads MoM'"
UT = "Utilisation"
MON = range(3, 15)                      # C..N, the month columns on this tab
UCOL = lambda c: gl(c + 1)              # C here is D on the Utilisation tab

wb = openpyxl.load_workbook(OUT)
ws = wb['Overhead Split']
thin = Side(style='thin', color=GREY)
hair = Side(style='medium', color=AQUA)
S = 110

def banner(r, text):
    ws.cell(r, 1, text)
    for c in range(1, 15):
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).font = Font(name='Calibri', size=12, bold=True, color=WHITE)

def head(r, first):
    ws.cell(r, 1, first)
    for c in MON:
        ws.cell(r, c, f'={gl(c)}$22').number_format = 'mmm\\-yy'
    for c in range(1, 15):
        ws.cell(r, c).font = Font(name='Calibri', size=10, bold=True, color=WHITE)
        ws.cell(r, c).fill = PatternFill('solid', fgColor=INK)
        ws.cell(r, c).alignment = Alignment(horizontal='right' if c > 1 else 'left')
        ws.cell(r, c).border = Border(bottom=hair)

def row(r, label, body, fmt, bold=False, fill=None):
    ws.cell(r, 1, label)
    for c in MON:
        ws.cell(r, c, body(gl(c), c)).number_format = fmt
    for c in range(1, 15):
        ws.cell(r, c).font = Font(name='Calibri', size=11, bold=bold, color=INK)
        ws.cell(r, c).border = Border(bottom=thin)
        if fill:
            ws.cell(r, c).fill = PatternFill('solid', fgColor=fill)
        if c > 1:
            ws.cell(r, c).alignment = Alignment(horizontal='right')

banner(S, '6.  OFFICE STAFF SPLIT, THE ADMIN EMPLOYMENT EXPENSE PUSHED INTO EACH DEPARTMENT')
ws.cell(S + 1, 1, 'Allocated on each department’s share of worked hours, chargeable plus '
                  'non-chargeable, off the Utilisation tab. Leave is out of it. '
                  'Nothing here is typed.')
ws.cell(S + 1, 1).font = Font(name='Calibri', size=9, italic=True, color=INK)

H0 = S + 2
head(H0, 'Worked hours')
HR = {}
for i, (d, (ch, nc)) in enumerate(UTIL.items()):
    r = H0 + 1 + i
    row(r, '    ' + d, lambda c, cc, a=ch, b=nc:
        f'={UT}!{UCOL(cc)}${a}+{UT}!{UCOL(cc)}${b}', HRS)
    HR[d] = r
HTOT = H0 + 6
row(HTOT, '    Total worked hours',
    lambda c, cc: '=' + '+'.join(f'{c}{HR[d]}' for d in UTIL), HRS,
    bold=True, fill=GREYL)

P0 = HTOT + 2
head(P0, 'Share of worked hours')
SH = {}
for i, d in enumerate(UTIL):
    r = P0 + 1 + i
    row(r, '    ' + d, lambda c, cc, x=d: f'=IFERROR({c}{HR[x]}/{c}{HTOT},0)', PCT)
    SH[d] = r
STOT = P0 + 6
row(STOT, '    Total', lambda c, cc: '=' + '+'.join(f'{c}{SH[d]}' for d in UTIL),
    PCT, bold=True, fill=GREYL)

A0 = STOT + 2
head(A0, 'Office staff charged across')
AM = {}
POTREF = lambda c: f'SUMIF({OMM}!$B$4:$M$4,{c}$22,{OMM}!$B${POT}:$M${POT})'
for i, d in enumerate(UTIL):
    r = A0 + 1 + i
    row(r, '    ' + d, lambda c, cc, x=d: f'={POTREF(c)}*{c}{SH[x]}', COST)
    AM[d] = r
ATOT = A0 + 6
row(ATOT, '    Total charged across',
    lambda c, cc: '=' + '+'.join(f'{c}{AM[d]}' for d in UTIL), COST,
    bold=True, fill=GREYL)
row(ATOT + 1, '    The admin pot', lambda c, cc: f'={POTREF(c)}', COST, bold=True)
row(ATOT + 2, '    Check, should be nil',
    lambda c, cc: f'=ROUND({c}{ATOT}-{c}{ATOT+1},2)', COST, bold=True)

for c in MON:
    ws.column_dimensions[gl(c)].width = 12
print(f'matrix written on Overhead Split, rows {S} to {ATOT+2}')
print('  hours', HR)
print('  share', SH)
print('  amount', AM)

# ---- the Overheads MoM office staff rows now read the matrix, not a keyboard
omm = wb['Overheads MoM']
OS = "'Overhead Split'"
n = 0
for d, r in STAFF['office'].items():
    for i in range(12):
        c = 2 + i
        omm.cell(r, c, f'=SUMIF({OS}!$C$22:$N$22,{gl(c)}$4,'
                       f'{OS}!$C${AM[d]}:$N${AM[d]})')
        omm.cell(r, c).number_format = COST
        omm.cell(r, c).fill = PatternFill()          # no longer an input
        n += 1
print(f'Overheads MoM: {n} office staff cells now calculated')
omm.cell(131, 1, 'Both sides are calculated. Office staff comes off the matrix on '
                 'the Overhead Split tab, section 6.')

json.dump({**STAFF, 'matrix_amount': AM, 'matrix_share': SH},
          open(f'{HERE}/data/staffsplit.json', 'w'), indent=1)
wb.save(OUT)
print('saved')
