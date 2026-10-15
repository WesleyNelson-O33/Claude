"""The overhead split matrix, worked into the P&L month by month.

Where this comes from. The reporting video, at 50:24, explains that whatever
sits in admin is split across the departments with an XLOOKUP against a set
of preset percentages, and that the percentages are fixed for the financial
year and live only in the spreadsheet, never in Xero. The matrix itself is
on the "4. PL Per Department" tab of her PL_Analysis file. There are three
keys, and every overhead account is assigned one of them:

  3 Way        Support .33  Events .26  Video .08  Integration .05  Consulting .28
  Staff        headcount, admin left out of the share
  Office Dept  Support .0179  Events .875  Video .0713  Integration .0179
               Consulting .0179

The Staff key is headcount driven. Her figures are Support 39, Events 24,
Video 3, Integration 1, Consulting 0 and Admin 7. Admin is the cost being
spread, so it is left out of the denominator, which is why the shares are
over 67 and not 74. That is also why Consulting is nil: nobody works in it.
At 51:20 the video says to change it when someone is hired into that role.

The indirect wages and salaries all sit on the Staff key, which is the piece
she has been asking for.

Sixty three accounts come across with their key. Fifty one of them have a
balance in the current GL. The rest are either nil this period or, like
Legal Fees, split by hand upstream, which the video covers at 54:08.
"""
import os, json, shutil, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as gl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = f'{HERE}/CTS Financial Controller Pack - v16.xlsx'
OUT = f'{HERE}/CTS Financial Controller Pack - v17.xlsx'
SPLIT = json.load(open(f'{HERE}/data/overhead_split.json'))

TEAL, AQUA, GREY, GREY_LT, MID, INK, WHITE = (
    '005358', '3FD0C9', 'E3E6EA', 'F4F4F4', '5F6670', '1A1A1A', 'FFFFFF')
MONEY = '\\$#,##0;[Red]"($"#,##0\\);\\-'
PCT = '0.0%;[Red](0.0%);\\-'
DEPTS = ['Support', 'Events', 'Video', 'Integration', 'Consulting']
HEAD = {'Support': 39, 'Events': 24, 'Video': 3, 'Integration': 1, 'Consulting': 0}
ADMIN_HEAD = 7
KEY3 = [0.33, 0.26, 0.08, 0.05, 0.28]
KEYOFF = [0.0179, 0.875, 0.0713, 0.0179, 0.0179]
CLN = 'Cleanup!$A$5:$A$20004'
CLI = 'Cleanup!$I$5:$I$20004'
CLB = 'Cleanup!$B$5:$B$20004'
CLD = 'Cleanup!$D$5:$D$20004'
CLE = 'Cleanup!$E$5:$E$20004'
OVH = ('Expenses', 'Other Expenses', 'Other Income')

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

# ------------------------------------- the account name, on to Cleanup -----
cu = wb['Cleanup']
cu['I4'] = 'Account name'
cu['I4']._style = copy.copy(cu['H4']._style)
for r in range(5, 20005):
    g = r - 4
    cu.cell(r, 9, f'=IF(N(\'P&L Analysis\'!$D{g})=0,"",TRIM(\'P&L Analysis\'!$B{g}))')
    cu.cell(r, 9)._style = copy.copy(cu.cell(r, 8)._style)
cu.column_dimensions['I'].width = 34
print('Cleanup: account name column written')

# ------------------------------------------------- the Overhead Split ------
if 'Overhead Split' in wb.sheetnames:
    del wb['Overhead Split']
ws = wb.create_sheet('Overhead Split', wb.sheetnames.index('Utilisation') + 1)
ws.sheet_properties.tabColor = TEAL

def title(c, t, sz=14):
    c.value = t
    c.font = Font(name='Calibri', sz=sz, bold=True, color=TEAL)

def head(r, c0, labels, width=None):
    for i, t in enumerate(labels):
        c = ws.cell(r, c0 + i, t)
        c.font = Font(name='Calibri', sz=10, bold=True, color=WHITE)
        c.fill = PatternFill('solid', fgColor=TEAL)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 28

title(ws['A1'], 'Overhead Split')
ws['A2'] = ('Admin overhead is spread across the departments on preset percentages, '
            'the same way the PL_Analysis file does it. Set the headcount and the two '
            'typed keys once a year. Everything else follows.')
ws['A2'].font = Font(name='Calibri', sz=9, italic=True, color=MID)

title(ws['A4'], '1.  Headcount, which drives the Staff key', 11)
head(5, 1, ['Department', 'Headcount', 'Share of staff'])
for i, d in enumerate(DEPTS):
    r = 6 + i
    ws.cell(r, 1, d).font = Font(name='Calibri', sz=10, color=INK)
    c = ws.cell(r, 2, HEAD[d])
    c.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
    c.fill = PatternFill('solid', fgColor=GREY)
    c.alignment = Alignment(horizontal='center')
    s = ws.cell(r, 3, f'=IF($B$12=0,0,$B{r}/$B$12)')
    s.number_format = PCT
    s.alignment = Alignment(horizontal='center')
ws.cell(11, 1, 'Admin').font = Font(name='Calibri', sz=10, italic=True, color=MID)
ws.cell(11, 2, ADMIN_HEAD).fill = PatternFill('solid', fgColor=GREY)
ws.cell(11, 2).font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
ws.cell(11, 2).alignment = Alignment(horizontal='center')
ws.cell(11, 3, 'the cost being spread, so left out of the share').font = Font(
    name='Calibri', sz=9, italic=True, color=MID)
ws.cell(12, 1, 'Total excluding Admin').font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
ws.cell(12, 2, '=SUM(B6:B10)').font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
ws.cell(12, 2).alignment = Alignment(horizontal='center')
ws.cell(12, 3, '=SUM(C6:C10)').number_format = PCT
ws.cell(12, 3).alignment = Alignment(horizontal='center')

title(ws['A14'], '2.  The three keys', 11)
head(15, 1, ['Key'] + DEPTS + ['Total'])
for i, (name, vals) in enumerate([('3 Way', KEY3), ('Staff', None), ('Office Dept', KEYOFF)]):
    r = 16 + i
    ws.cell(r, 1, name).font = Font(name='Calibri', sz=10, bold=True, color=INK)
    for j in range(5):
        c = ws.cell(r, 2 + j)
        if vals is None:
            c.value = f'=$C${6 + j}'
        else:
            c.value = vals[j]
            c.fill = PatternFill('solid', fgColor=GREY)
            c.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
        c.number_format = PCT
        c.alignment = Alignment(horizontal='center')
    t = ws.cell(r, 7, f'=SUM(B{r}:F{r})')
    t.number_format = PCT
    t.alignment = Alignment(horizontal='center')
    t.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
ws.cell(17, 8, 'calculated from the headcount above').font = Font(
    name='Calibri', sz=9, italic=True, color=MID)

# --------------------------------------- the accounts and their months ----
A0 = 23
accounts = sorted(SPLIT)
A1_ = A0 + len(accounts)        # one extra row for the catch-all
title(ws['A20'], '3.  Each overhead account coded to admin, its key, and what it cost each month', 11)
head(22, 1, ['Account coded to admin', 'Key'] + [''] * 12 + [''] + DEPTS)
for j in range(12):
    c = ws.cell(22, 3 + j, f"='P&L FY27'!{gl(2 + j)}$5")
    c.number_format = 'mmm-yy'
    c.font = Font(name='Calibri', sz=10, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=TEAL)
    c.alignment = Alignment(horizontal='center')
dv = DataValidation(type='list', formula1='"3 Way,Staff,Office Dept"', allow_blank=False)
ws.add_data_validation(dv)
for i, a in enumerate(accounts):
    r = A0 + i
    ws.cell(r, 1, a).font = Font(name='Calibri', sz=9, color=INK)
    k = ws.cell(r, 2, SPLIT[a])
    k.font = Font(name='Calibri', sz=9, bold=True, color=TEAL)
    k.fill = PatternFill('solid', fgColor=GREY)
    k.alignment = Alignment(horizontal='center')
    dv.add(k)
    for j in range(12):
        c = ws.cell(r, 3 + j,
                    f'=SUMIFS({CLN},{CLI},$A{r},{CLB},{gl(3 + j)}$22,{CLD},"ADMIN")')
        c.number_format = MONEY
        c.font = Font(name='Calibri', sz=9, color=INK)
    for j in range(5):
        c = ws.cell(r, 16 + j,
                    f'=IFERROR(INDEX(${gl(2 + j)}$16:${gl(2 + j)}$18,'
                    f'MATCH($B{r},$A$16:$A$18,0)),0)')
        c.number_format = PCT
        c.font = Font(name='Calibri', sz=9, color=MID)
        c.alignment = Alignment(horizontal='center')

# ------------------------------------------- what each department gets ----
R = A1_ + 2
title(ws.cell(R, 1), '4.  Admin overhead spread across the departments', 11)
head(R + 1, 1, ['Department'] + [''] * 12)
for j in range(12):
    c = ws.cell(R + 1, 3 + j, f'={gl(3 + j)}$22')
    c.number_format = 'mmm-yy'
    c.font = Font(name='Calibri', sz=10, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=TEAL)
    c.alignment = Alignment(horizontal='center')
ALLOC0 = R + 2
for i, d in enumerate(DEPTS):
    r = ALLOC0 + i
    ws.cell(r, 1, d).font = Font(name='Calibri', sz=10, color=INK)
    for j in range(12):
        c = ws.cell(r, 3 + j,
                    f'=SUMPRODUCT({gl(3 + j)}${A0}:{gl(3 + j)}${A1_},'
                    f'${gl(16 + i)}${A0}:${gl(16 + i)}${A1_})')
        c.number_format = MONEY
        c.font = Font(name='Calibri', sz=10, color=INK)
rt = ALLOC0 + 5
ws.cell(rt, 1, 'Total spread').font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
for j in range(12):
    c = ws.cell(rt, 3 + j, f'=SUM({gl(3 + j)}{ALLOC0}:{gl(3 + j)}{rt - 1})')
    c.number_format = MONEY
    c.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
    c.border = Border(top=Side('thin', color=TEAL))
ws.cell(rt + 1, 1, 'Everything coded to admin').font = Font(name='Calibri', sz=10, color=MID)
for j in range(12):
    c = ws.cell(rt + 1, 3 + j, f'=SUMIFS({CLN},{CLD},"ADMIN",{CLB},{gl(3 + j)}$22)')
    c.number_format = MONEY
    c.font = Font(name='Calibri', sz=10, color=MID)
ws.cell(rt + 2, 1, 'Check, should be nil').font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
for j in range(12):
    c = ws.cell(rt + 2, 3 + j, f'=ROUND({gl(3 + j)}{rt}-{gl(3 + j)}{rt + 1},2)')
    c.number_format = MONEY
    c.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)


# ------------------------------- what each department already carries -----
OWN = rt + 4
title(ws.cell(OWN, 1), "5.  Each department's own overhead, which stays where it is", 11)
head(OWN + 1, 1, ['Department'] + [''] * 12)
for j in range(12):
    c = ws.cell(OWN + 1, 3 + j, f'={gl(3 + j)}$22')
    c.number_format = 'mmm-yy'
    c.font = Font(name='Calibri', sz=10, bold=True, color=WHITE)
    c.fill = PatternFill('solid', fgColor=TEAL)
    c.alignment = Alignment(horizontal='center')
OWN0 = OWN + 2
TAG = {'Support': 'ONSITE', 'Events': 'PRODUCTION', 'Video': 'VIDEO',
       'Integration': 'INTEGRATION', 'Consulting': 'CONSULTING'}
for i, d in enumerate(DEPTS):
    r = OWN0 + i
    ws.cell(r, 1, d).font = Font(name='Calibri', sz=10, color=INK)
    for j in range(12):
        parts = '+'.join(
            f'SUMIFS({CLN},{CLD},"{TAG[d]}",{CLB},{gl(3 + j)}$22,{CLE},"{k}")' for k in OVH)
        c = ws.cell(r, 3 + j, '=' + parts)
        c.number_format = MONEY
        c.font = Font(name='Calibri', sz=10, color=INK)
rt2 = OWN0 + 5
ws.cell(rt2, 1, 'Total').font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
for j in range(12):
    c = ws.cell(rt2, 3 + j, f'=SUM({gl(3 + j)}{OWN0}:{gl(3 + j)}{rt2 - 1})')
    c.number_format = MONEY
    c.font = Font(name='Calibri', sz=10, bold=True, color=TEAL)
    c.border = Border(top=Side('thin', color=TEAL))

CATCH = A1_
ws.cell(CATCH, 1, 'Everything else coded to admin').font = Font(
    name='Calibri', sz=9, bold=True, color=TEAL)
k = ws.cell(CATCH, 2, '3 Way')
k.font = Font(name='Calibri', sz=9, bold=True, color=TEAL)
k.fill = PatternFill('solid', fgColor=GREY)
k.alignment = Alignment(horizontal='center')
dv.add(k)
for j in range(12):
    c = ws.cell(CATCH, 3 + j,
                f'=SUMIFS({CLN},{CLD},"ADMIN",{CLB},{gl(3 + j)}$22)'
                f'-SUM({gl(3 + j)}${A0}:{gl(3 + j)}${CATCH - 1})')
    c.number_format = MONEY
    c.font = Font(name='Calibri', sz=9, bold=True, color=TEAL)
for j in range(5):
    c = ws.cell(CATCH, 16 + j,
                f'=IFERROR(INDEX(${gl(2 + j)}$16:${gl(2 + j)}$18,'
                f'MATCH($B{CATCH},$A$16:$A$18,0)),0)')
    c.number_format = PCT
    c.font = Font(name='Calibri', sz=9, color=MID)
    c.alignment = Alignment(horizontal='center')
ws.cell(CATCH, 21, 'anything admin that is not named above lands here, so nothing is lost. '
                   'If it is more than small change, add the account to the list with its key.'
        ).font = Font(name='Calibri', sz=9, italic=True, color=MID)

ws.column_dimensions['A'].width = 36
ws.column_dimensions['B'].width = 13
for j in range(12):
    ws.column_dimensions[gl(3 + j)].width = 12
ws.column_dimensions['O'].width = 3
for j in range(5):
    ws.column_dimensions[gl(16 + j)].width = 11
ws.freeze_panes = 'C23'
ws.sheet_view.showGridLines = False
print(f'Overhead Split: {len(accounts)} accounts, allocation rows {ALLOC0}..{ALLOC0+4}')

json.dump({'alloc_first_row': ALLOC0, 'own_first_row': OWN0, 'depts': DEPTS},
          open(f'{HERE}/data/overhead_rows.json', 'w'))
wb.save(OUT)
print('saved', OUT)
