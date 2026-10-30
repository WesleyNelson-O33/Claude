"""Fifth pass on Danica's CBA AGM 2026 Costs and Hours 29102026 file. Rebuilds Staff Hours by Date from her
Employment Hero Timesheet tabs (Jan 2025 to 30 Oct 2026), every non-rejected line at location CBA AGM 2026 [PRD] or
[VID], de-duplicated on Timesheet Id. Repoints the Summary hours lines. Leaves her other edits and the WIP lines alone."""
import sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC, OUT = sys.argv[1:3]
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
AMBER = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HRS = '0.00;(0.00);"-"'
TAB = "Staff Hours by Date"

wb = openpyxl.load_workbook(SRC)


def put(ws, ref, v, bold=False, fmt=None, fill=None, wrap=False, head=False, size=10):
    c = ws[ref]
    c.value = v
    c.font = Font(name=F, bold=bold or head, size=size, color="FFFFFF" if head else "000000")
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    c.fill = HEAD if head else (fill or PatternFill())
    c.border = BOX if (head or fmt) else Border()
    if fmt:
        c.number_format = fmt


def hhmm(td):
    if not isinstance(td, datetime.timedelta):
        return None
    m = int(td.total_seconds() // 60)
    return "%d:%02d" % (m // 60, m % 60)


# ---------------------------------------------------------------- read the timesheet tabs
lines, seen, tabs = [], set(), []
for ws in wb.worksheets:
    if not ws.title.startswith("Timesheet "):
        continue
    tabs.append(ws.title)
    rows = ws.iter_rows(values_only=True)
    h = next(rows)
    ix = {k: i for i, k in enumerate(h) if k}
    for r in rows:
        g = lambda k: r[ix[k]] if k in ix else None
        loc = str(g("Location") or "")
        if "CBA AGM 2026" not in loc or g("Status") == "Rejected" or g("Timesheet Id") in seen:
            continue
        seen.add(g("Timesheet Id"))
        assert g("Unit Type") == "hours"
        name = "%s %s" % (g("First Name"), str(g("Surname")).title())
        lines.append([g("Start Date"), name, loc.split(" / ")[-1], hhmm(g("Start Time")), hhmm(g("End Time")),
                      hhmm(g("Unpaid Break Duration")), round(g("Duration").total_seconds() / 3600, 2), g("Status"),
                      g("Work Type") or "", g("Timesheet Id"), ws.title])
lines.sort(key=lambda x: (x[0], x[1], x[3] or ""))
assert len(tabs) == 8

# ---------------------------------------------------------------- rebuild Staff Hours by Date
idx = wb.sheetnames.index(TAB)
del wb[TAB]
sd = wb.create_sheet(TAB, idx)
sd.sheet_view.showGridLines = False
put(sd, "A1", "CBA AGM 2026, job 2506208. Staff hours by date", bold=True, size=13)
put(sd, "A2", "From the Employment Hero Timesheet tabs, 1 January 2025 to 30 October 2026. Every timesheet at location "
              "CBA AGM 2026 [PRD] or CBA AGM 2026 [VID]. Rejected timesheets are left out. One line that appears in two "
              "tabs is counted once.", wrap=True, fill=AMBER)
sd.merge_cells("A2:K2"); sd.row_dimensions[2].height = 30
H = ["Date", "Employee", "Location", "Start", "Finish", "Unpaid break", "Hours worked", "Status", "Work type",
     "Timesheet Id", "Source tab"]
FM = ["DD/MM/YYYY", "@", "@", "@", "@", "@", HRS, "@", "@", "0", "@"]
for i, h in enumerate(H):
    put(sd, "%s3" % chr(65 + i), h, head=True, wrap=True)
for r, vals in enumerate(lines, 4):
    for i, v in enumerate(vals):
        put(sd, "%s%d" % (chr(65 + i), r), v, fmt=FM[i])
first, last = 4, 3 + len(lines)
put(sd, "F%d" % (last + 1), "Total", bold=True)
put(sd, "G%d" % (last + 1), "=SUM(G%d:G%d)" % (first, last), bold=True, fmt=HRS)
TOTAL = "'%s'!G%d" % (TAB, last + 1)
G, A, B, C = ("$%s$%d:$%s$%d" % (c, first, c, last) for c in "GABC")

r = last + 4
put(sd, "A%d" % r, "Hours by employee", bold=True, size=12)
for col, h in zip("ABC", ["Employee", "Hours worked", "Days"]):
    put(sd, "%s%d" % (col, r + 1), h, head=True)
tot = {}
for x in lines:
    tot[x[1]] = tot.get(x[1], 0) + x[6]
names = sorted(tot, key=lambda n: -tot[n])
emp_first = r + 2
for i, n in enumerate(names, emp_first):
    put(sd, "A%d" % i, n, fmt="@")
    put(sd, "B%d" % i, "=SUMIFS(%s,%s,A%d)" % (G, B, i), fmt=HRS)
    put(sd, "C%d" % i, "=SUMPRODUCT((%s=A%d)/COUNTIFS(%s,%s,%s,%s))" % (B, i, B, B, A, A), fmt="0")
e = emp_first + len(names)
put(sd, "A%d" % e, "Total", bold=True)
put(sd, "B%d" % e, "=SUM(B%d:B%d)" % (emp_first, e - 1), bold=True, fmt=HRS)

r = e + 3
put(sd, "A%d" % r, "By month", bold=True, size=12)
for col, h in zip("ABCD", ["Month", "Production [PRD]", "Video [VID]", "Total"]):
    put(sd, "%s%d" % (col, r + 1), h, head=True)
months = sorted({datetime.date(x[0].year, x[0].month, 1) for x in lines})
m_first = r + 2
for i, m in enumerate(months, m_first):
    nxt = datetime.date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    rng = '%s,">="&A%d,%s,"<"&DATE(%d,%d,1)' % (A, i, A, nxt.year, nxt.month)
    put(sd, "A%d" % i, m, fmt="MMM YYYY")
    put(sd, "B%d" % i, '=SUMIFS(%s,%s,%s,"*[PRD]*")' % (G, rng, C), fmt=HRS)
    put(sd, "C%d" % i, '=SUMIFS(%s,%s,%s,"*[VID]*")' % (G, rng, C), fmt=HRS)
    put(sd, "D%d" % i, "=B%d+C%d" % (i, i), fmt=HRS)
me = m_first + len(months)
put(sd, "A%d" % me, "Total", bold=True)
for col in "BCD":
    put(sd, "%s%d" % (col, me), "=SUM(%s%d:%s%d)" % (col, m_first, col, me - 1), bold=True, fmt=HRS)
for col, w in zip("ABCDEFGHIJK", [12, 30, 22, 8, 8, 9, 10, 11, 22, 12, 26]):
    sd.column_dimensions[col].width = w
sd.freeze_panes = "A4"

# ---------------------------------------------------------------- Summary
sm = wb["Summary"]
assert sm["A7"].value == "Staff hours, earlier monthly reports" and sm["A25"].value == "Staff hours by employee"
assert sm["B21"].value == "='Staff Hours'!D22"
sm["A7"].value = "Staff hours, EH timesheets"
sm["B7"].value = "=" + TOTAL
sm["C7"].value = "Hours. Every EH timesheet at CBA AGM 2026, PRD and VID, January 2025 to 30 October 2026."
assert sm["A13"].value.startswith("3. Staff hours by person and date are still to come")
sm["A13"].value = ("3. Staff hours by person and date are from the Employment Hero timesheet exports, 1 January 2025 "
                   "to 30 October 2026. See the Staff Hours by Date tab.")

for rng in list(sm.merged_cells.ranges):
    if rng.min_row >= 25:
        sm.unmerge_cells(str(rng))
for row in sm.iter_rows(min_row=25, max_row=max(sm.max_row, 60)):
    for c in row:
        c.value = None
        c.fill, c.border = PatternFill(), Border()
D = "'%s'!" % TAB
r = 25
put(sm, "A%d" % r, "Staff hours by employee", bold=True, size=12)
for col, h in zip("ABCDE", ["Employee", "Hours worked", "Days", "First date", "Last date"]):
    put(sm, "%s%d" % (col, r + 1), h, head=True, wrap=True)
for k, i in enumerate(range(r + 2, r + 2 + len(names))):
    put(sm, "A%d" % i, "=%sA%d" % (D, emp_first + k), fmt="@")
    put(sm, "B%d" % i, "=%sB%d" % (D, emp_first + k), fmt=HRS)
    put(sm, "C%d" % i, "=%sC%d" % (D, emp_first + k), fmt="0")
    put(sm, "D%d" % i, "=_xlfn.MINIFS(%s%s,%s%s,A%d)" % (D, A, D, B, i), fmt="DD/MM/YYYY")
    put(sm, "E%d" % i, "=_xlfn.MAXIFS(%s%s,%s%s,A%d)" % (D, A, D, B, i), fmt="DD/MM/YYYY")
t = r + 2 + len(names)
put(sm, "A%d" % t, "Total", bold=True, fmt="@")
put(sm, "B%d" % t, "=SUM(B%d:B%d)" % (r + 2, t - 1), bold=True, fmt=HRS)
put(sm, "A%d" % (t + 1), "From the Employment Hero timesheet exports. Hours by date are on the Staff Hours by Date tab.", wrap=True)
sm.merge_cells("A%d:E%d" % (t + 1, t + 1))

wb.save(OUT)
print(len(lines), "lines", round(sum(x[6] for x in lines), 2), "hours;", len(names), "staff;", tabs)
