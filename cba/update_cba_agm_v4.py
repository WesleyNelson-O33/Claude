"""Fourth pass: add the Deaf Services INV409016 bill (Xero, dated 26 Oct 2026, CBA AGM 2026) to Supplier Costs,
above the ExhibitionCo line, and repoint every total and Summary link that the extra row moves."""
import sys, copy, datetime
import openpyxl

SRC, OUT = sys.argv[1:3]
wb = openpyxl.load_workbook(SRC)
sc, sm = wb["Supplier Costs"], wb["Summary"]

# Layout as it stands: bills rows 4 to 16, totals 17, paid/unpaid 19-20, Encore PO row 24, bills plus PO 26
assert sc["C16"].value == "INV8728" and sc["F17"].value == "=SUM(F4:F16)" and sc["F24"].value == 18318.43
assert sc["F26"].value == "=F17+F24" and sm["B3"].value == "='Supplier Costs'!F17"

sc.unmerge_cells("A28:J28")
sc.insert_rows(16)
for col in range(1, 13):
    src, dst = sc.cell(15, col), sc.cell(16, col)
    dst._style = copy.copy(src._style)
sc.row_dimensions[16].height = sc.row_dimensions[15].height
vals = [datetime.datetime(2026, 10, 26), "Deaf Services Limited", "INV409016",
        "Onsite interpreting 14/10/2026, 7:30am to 12:30pm (Regina Sapko)", "51400", 282.50, 28.25, "=F16+G16", "Unpaid",
        "Bill says DL/TW yet to approve and confirm CC."]
for col, v in enumerate(vals, 1):
    sc.cell(16, col).value = v

# Rows from 17 down moved by one. Rewrite every formula in the moved block and the Summary links.
sc["H17"] = "=F17+G17"
for col in "FGH":
    sc["%s18" % col] = "=SUM(%s4:%s17)" % (col, col)
    sc["%s27" % col] = "=%s18+%s25" % (col, col)
sc["F20"] = '=SUMIFS(F4:F17,I4:I17,"Paid")'
sc["F21"] = '=SUMIFS(F4:F17,I4:I17,"Unpaid")'
sc["H25"] = "=F25+G25"
assert sc["C17"].value == "INV8728" and sc["E18"].value == "Total bills" and sc["F25"].value == 18318.43
assert sc["E27"].value == "Bills plus Encore PO amount"
sc.merge_cells("A29:J29")

sm["B3"] = "='Supplier Costs'!F18"
sm["B4"] = "='Supplier Costs'!F25"
sm["B5"] = "='Supplier Costs'!F27"
for col, s in zip("BCD", "FGH"):
    assert sm["%s18" % col].value == "='Supplier Costs'!%s17" % s
    sm["%s18" % col] = "='Supplier Costs'!%s18" % s
    sm["%s19" % col] = "='Supplier Costs'!%s25" % s

wb.save(OUT)
