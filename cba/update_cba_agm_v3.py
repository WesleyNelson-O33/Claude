"""Third pass on Danica's CBA AGM 2026 Costs and Hours v2. Keeps her edits (ExhibitionCo INV8728 row, renumbered notes)
and adds staff hours by employee to the Summary tab, pulled by formula from the Staff Hours by Date tab."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC, OUT = sys.argv[1:3]
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
AMBER = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HRS = '0.00;(0.00);"-"'

wb = openpyxl.load_workbook(SRC)
sm, sd = wb["Summary"], wb["Staff Hours by Date"]


def put(ref, v, bold=False, fmt=None, fill=None, wrap=False, head=False, size=10):
    c = sm[ref]
    c.value = v
    c.font = Font(name=F, bold=bold or head, size=size, color="FFFFFF" if head else "000000")
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    if head:
        c.fill = HEAD
    elif fill:
        c.fill = fill
    if head or fmt:
        c.border = BOX
    if fmt:
        c.number_format = fmt


# Labels: the bills line now includes ExhibitionCo INV8728, which is not in Xero yet
assert sm["A3"].value == "Supplier bills in Xero" and sm["A18"].value == "Supplier bills in Xero"
assert sm["B3"].value == "='Supplier Costs'!F17"
sm["A3"].value = sm["A18"].value = "Supplier bills"
sm["C3"].value = "Ex GST. Includes ExhibitionCo INV8728, received 29/10/2026 and not yet in Xero."
assert sm["C5"].value == "Ex GST. Xero bills plus the Encore PO amount."
sm["C5"].value = "Ex GST. Supplier bills plus the Encore PO amount."
assert sm["A6"].value == "Staff hours seen so far"
sm["A6"].value = "Staff hours, earlier monthly reports"

# Staff Hours by Date layout as Danica left it: data rows 4 to 28, employee list A34:A40
assert sd["A3"].value == "Date" and sd["F29"].value == "Total" and sd["A33"].value == "Employee" and sd["A41"].value == "Total"
D = "'Staff Hours by Date'!"
names = [sd["A%d" % r].value for r in range(34, 41)]

r = 23
put("A%d" % r, "Staff hours by employee", bold=True, size=12)
for col, h in zip("ABCDE", ["Employee", "Hours worked", "Days", "First date", "Last date"]):
    put("%s%d" % (col, r + 1), h, head=True, wrap=True)
for i, n in enumerate(names, r + 2):
    put("A%d" % i, "=%sA%d" % (D, 34 + i - (r + 2)), fmt="@")
    put("B%d" % i, "=SUMIFS(%s$G$4:$G$28,%s$B$4:$B$28,A%d)" % (D, D, i), fmt=HRS)
    put("C%d" % i, "=COUNTIFS(%s$B$4:$B$28,A%d)" % (D, i), fmt="0")
    put("D%d" % i, "=_xlfn.MINIFS(%s$A$4:$A$28,%s$B$4:$B$28,A%d)" % (D, D, i), fmt="DD/MM/YYYY")
    put("E%d" % i, "=_xlfn.MAXIFS(%s$A$4:$A$28,%s$B$4:$B$28,A%d)" % (D, D, i), fmt="DD/MM/YYYY")
t = r + 2 + len(names)
put("A%d" % t, "Total", bold=True, fmt="@")
put("B%d" % t, "=SUM(B%d:B%d)" % (r + 2, t - 1), bold=True, fmt=HRS)
put("C%d" % t, "=SUM(C%d:C%d)" % (r + 2, t - 1), bold=True, fmt="0")
put("A%d" % (t + 1), "From the Roster vs Timesheet report, timesheets at CBA AGM 2026 only. It leaves out hours worked "
                     "without a roster, including Duncan Lugstein and Daniel Sobkowski. The full timesheet export is still to come.",
    wrap=True, fill=AMBER)
sm.merge_cells("A%d:E%d" % (t + 1, t + 1)); sm.row_dimensions[t + 1].height = 45
sm.column_dimensions["E"].width = 12

wb.save(OUT)
print(names)
