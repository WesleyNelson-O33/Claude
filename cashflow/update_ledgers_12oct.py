import openpyxl, datetime, copy, re, sys
from openpyxl.styles import PatternFill
SRC, OUT = sys.argv[1], sys.argv[2]
D=datetime.datetime
wb=openpyxl.load_workbook(SRC)
rec=wb['Receivables']; pb=wb['Payables']; cf=wb['Cash Flow']
GREY=PatternFill('solid', fgColor='FFEDEDED')

# ================= RECEIVABLES =================
# 1. mark what has been paid since 28 Sep (expected date -> the day it landed)
paid={5:(D(2026,10,2),'Paid 2 Oct'),6:(D(2026,10,7),'Paid 7 Oct'),7:(D(2026,10,8),'Paid 8 Oct'),10:(D(2026,9,30),'Paid 30 Sep'),
      20:(D(2026,9,29),'Paid 29 Sep'),21:(D(2026,9,29),'Paid 29 Sep'),28:(D(2026,9,29),'Paid 29 Sep'),29:(D(2026,9,29),'Paid 29 Sep'),
      30:(D(2026,10,6),'Paid 6 Oct'),31:(D(2026,9,29),'Paid 29 Sep'),32:(D(2026,9,29),'Paid 29 Sep'),33:(D(2026,10,1),'Paid 1 Oct'),
      34:(D(2026,10,8),'Paid 8 Oct'),35:(D(2026,10,9),'Paid 9 Oct'),37:(D(2026,9,29),'Paid 29 Sep')}
for r,(dt,txt) in paid.items():
    rec[f'G{r}']=txt; rec[f'H{r}']=dt; rec[f'K{r}']=f'{txt} - in the bank, in the actuals.'
    for c in range(1,12): rec.cell(r,c).fill=GREY
# 2. still open at 12 Oct - move the date where nothing has come in
upd={11:(D(2026,10,16),'Open','Still unpaid at 12 Oct (due 30 Sep). Aware usually pay on the due date - call it Fri 16 Oct and chase with the big one.'),
     16:(D(2026,10,16),'Overdue','Still not in at 12 Oct - chase. Fri 16 Oct.'),
     17:(D(2026,10,16),'Overdue','Still not in at 12 Oct - chase with the other one. Fri 16 Oct.'),
     18:(D(2026,10,21),'Overdue','Six weeks overdue at 12 Oct. Xero shows Gen-e owing $39,106 across the three, so one has a credit on it. Chase. Week 3.'),
     22:(D(2026,10,30),'Open','Staff recovery - goes through payroll, no cash. Pushed to end of month.'),
     24:(D(2026,10,21),'Overdue','Six weeks overdue at 12 Oct - usually a PO / portal issue. Week 3.'),
     25:(D(2026,10,21),'Overdue','Six weeks overdue at 12 Oct - chase. Week 3.'),
     26:(D(2026,10,21),'Overdue','Assume it comes with the Aug one. Week 3.'),
     41:(D(2026,10,16),'Overdue','Still not in at 12 Oct - chase. Fri 16 Oct.')}
for r,(dt,st,txt) in upd.items():
    rec[f'H{r}']=dt; rec[f'G{r}']=st; rec[f'K{r}']=txt
# 3. the Sep contract invoices were raised 30 Sep - replace the 'To be raised' rows with the real ones
def setrow(r,vals,note):
    for c,v in zip('ABCDEFGH',vals): rec[f'{c}{r}']=v
    rec[f'I{r}']=f'=H{r}-E{r}'
    rec[f'J{r}']=f'=IF(H{r}<Assumptions!$B$5,"Before",IF(H{r}>Assumptions!$C$48,"After",INT((H{r}-Assumptions!$B$5)/7)+1))'
    rec[f'K{r}']=note
# move totals 53-56 down to 66-69 and copy row style for the new rows
for r in range(53,57):
    for c in range(1,12):
        s=rec.cell(r,c); d=rec.cell(r+13,c); d.value=s.value; d._style=copy.copy(s._style); s.value=None
for r in range(53,66):
    for c in range(1,12): rec.cell(r,c)._style=copy.copy(rec.cell(52,c)._style)
rows={
 43:(('Deloitte Services Trust','INV-10638','DTTL Sep contract (PO) - raised 30 Sep',D(2026,9,30),D(2026,10,30),112683.22,'Paid 8 Oct',D(2026,10,8)),'Paid 8 Oct - Deloitte pay about a week after we invoice.'),
 44:(('Deloitte Services Trust','INV-10665','DTTL Sep additionals - raised 30 Sep',D(2026,9,30),D(2026,10,30),18891.66,'Paid 8 Oct',D(2026,10,8)),'Paid 8 Oct with the contract one.'),
 45:(('Commonwealth Bank (CBA Place South)','INV-10643','85720001CBPSCSJUN27 - Sep contract',D(2026,9,30),D(2026,10,30),24619.84,'Current',D(2026,10,16)),'Aug one was paid 8 days after invoice. Not in yet at 12 Oct - call it Fri 16 Oct.'),
 46:(('Commonwealth Bank (CBA Place South)','INV-10657 to 10662','Sep dedicated tech + after hours (6 small invoices)',D(2026,9,30),D(2026,10,30),3432.0,'Current',D(2026,10,16)),'$1,188 + 4 x $467.50 + $374. With the contract one.'),
 47:(('Commonwealth Bank (CBA Place South)','INV-10675','Live event support Sep (PO3317501)',D(2026,9,30),D(2026,10,30),5398.25,'Current',D(2026,10,16)),'With the contract one.'),
 48:(('Commonwealth Bank (CBA Square)','INV-10637','85720001CBSCSJUN27 - Sep contract',D(2026,9,30),D(2026,10,30),11792.0,'Current',D(2026,10,16)),'Same as Place South.'),
 49:(('Commonwealth Bank (CBA Melbourne)','INV-10644','85720001MEL435CSJUN27 - Sep contract',D(2026,9,30),D(2026,10,30),12995.72,'Current',D(2026,10,16)),'Same. (Their $93.50 one was paid 6 Oct.)'),
 50:(('Bankwest','INV-10640','85720001BWPCSJUN27 - Sep additionals',D(2026,9,30),D(2026,10,30),4752.0,'Current',D(2026,10,16)),'Bankwest paid the Sep contract one ($13,156) on 2 Oct. This is the additionals - give it to Fri 16 Oct.'),
 51:(('PwC Services','INV-10655','PO2601702104 - Sep contract',D(2026,9,30),D(2026,10,30),13097.7,'Current',D(2026,10,30)),'PwC paid the Aug ones on 29 Sep (the due date). On the due date, week 5.'),
 52:(('PwC Services','INV-10654','PO2601703509 - Sep',D(2026,9,30),D(2026,10,30),8415.0,'Current',D(2026,10,30)),'On the due date, week 5.'),
 53:(('PwC Services','INV-10641','PO2601702371 - Sep contract',D(2026,9,30),D(2026,10,30),7865.0,'Current',D(2026,10,30)),'On the due date, week 5.'),
 54:(('Aware Super','INV-10652','Annual National Maintenance - September',D(2026,9,30),D(2026,10,30),12100.0,'Current',D(2026,10,30)),'Aware pay on the due date - but the Aug one is still sitting there, so watch it.'),
 55:(('Generation-e Productivity Solutions','INV-10642','PAU26-407 - Sep',D(2026,9,30),D(2026,10,30),13627.68,'Current',D(2026,11,20)),'Slow payers - 3 weeks after due, so after this forecast.'),
 56:(('Automic Group','INV-10663','September All Hands 25.9.26',D(2026,9,29),D(2026,10,29),8775.8,'Current',D(2026,10,29)),'New customer - on the due date, week 5.'),
 57:(('ASX Operations','INV-10668','External listing test days 29.9.26',D(2026,9,30),D(2026,10,30),5747.5,'Current',D(2026,11,2)),'Due Fri 30 Oct, ASX run a day or two late - Mon 2 Nov.'),
 58:(('Global Philanthropic','INV-10669','Talking Philanthropy event B-roll',D(2026,9,30),D(2026,10,30),1373.9,'Current',D(2026,10,30)),'On the due date, week 5.'),
 59:(('Global Philanthropic','INV-10670','Talking Philanthropy Aus 18.09.2026',D(2026,9,30),D(2026,10,30),4343.9,'Current',D(2026,10,30)),'On the due date, week 5.'),
 60:(('Australian Payments Network','INV-10671','AV support 21-22 Sep',D(2026,9,30),D(2026,10,30),748.0,'Current',D(2026,10,30)),'Paid the last one 11 days late - may slip past the forecast. Left on the due date.'),
 61:(('ICC Sydney','INV-10672','Speaker presentation technicians Sep (TTS)',D(2026,9,30),D(2026,10,30),4100.26,'Current',D(2026,10,30)),'ICC pay on the due date.'),
 62:(('ICC Sydney','INV-10674','Speaker presentation technicians Sep (CommBio)',D(2026,9,30),D(2026,10,30),2448.88,'Current',D(2026,10,30)),'ICC pay on the due date.'),
 63:(("St Luke's Liverpool",'INV-10676','Digital mixer equipment',D(2026,9,30),D(2026,10,30),4988.5,'Current',D(2026,10,30)),'New - on the due date, week 5.'),
}
for r,(vals,note) in rows.items(): setrow(r,vals,note)
for r in (43,44):
    for c in range(1,12): rec.cell(r,c).fill=GREY
# the Aug Bankwest / CBA Brisbane / CBA Melbourne small ones that were paid in week 1 are already in the week 1 actuals
rec['A64']=None
rec['A2']=rec['A2'].value+' Updated 12 Oct: paid invoices are greyed out with the date they landed, the Sep contract invoices (raised 30 Sep) replace the "to be raised" rows, and the expected date has been moved on anything that has not come in.'
rec['A67']='of which still open in Xero at 12 Oct (should be $583,424.49)'
rec['F67']='=SUMIFS(F5:F65,G5:G65,"<>Paid*")'
# extend every range that pointed at rows 5-52
def fix(f):
    if not isinstance(f,str) or not f.startswith('='): return f
    f=re.sub(r'(\$?[FGH]\$?5:\$?[FGH]\$?)52\b', r'\g<1>65', f)
    f=re.sub(r'\bF53\b','F66',f).replace('F55','F68')
    return f
for ws in (cf,rec):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('='): c.value=fix(c.value)
rec['F66']='=SUM(F5:F65)'; rec['F68']='=SUMIFS(F5:F65,H5:H65,">="&Assumptions!$B$5,H5:H65,"<="&Assumptions!$C$48)'; rec['F69']='=F66-F68'

# ================= PAYABLES =================
paidb={20:'PAID 29 Sep (Xero shows the whole bill as paid 8 Oct - see note on the hold).',24:'PAID 29 Sep - PL card',26:'PAID 8 Oct',27:'PAID 22 Sep - PL card',28:'PAID 1 Oct',30:'PAID 18 Sep - card',31:'PAID 23 Sep - card',32:'PAID 8 Oct - BPAY',34:'PAID 8 Oct',35:'PAID 30 Sep - DD',36:'PAID 1 Oct',37:'PAID 18 Sep - card',38:'PAID 1 Oct',39:'PAID 22 Sep - card',40:'Refunded 24 Sep - card',41:'PAID - card',42:'PAID 17 Sep - card',43:'PAID 24 Sep - card',49:'PAID 8 Oct (INV-19664)',50:'PAID 6 Oct',52:'PAID 8 Oct - BPAY (TTI20227626284, $136)',53:'Open - PL card, in the card line',54:'Open - PL card, in the card line'}
dates={26:D(2026,10,8),28:D(2026,10,1),32:D(2026,10,8),34:D(2026,10,8),36:D(2026,10,1),38:D(2026,10,1),49:D(2026,10,8),50:D(2026,10,6),52:D(2026,10,8)}
for r,txt in paidb.items():
    pb[f'J{r}']=txt
    if txt.startswith('PAID') or txt.startswith('Refunded'):
        for c in range(1,11): pb.cell(r,c).fill=GREY
for r,dt in dates.items(): pb[f'G{r}']=dt
pb['B49']='INV-19664 - SOLO standard data pack October'; pb['I49']='Xero bill'
pb['B52']='TTI20227626284 - Nighthawk charges October'; pb['D52']=136; pb['C52']=D(2026,10,12); pb['I52']='Xero bill'
pb['J21']='HELD - agreed 28 Sep, released on Aware sign off and handover docs. Xero shows INV41565 fully paid on 8 Oct ($178,119.25) - if the $44,873 did go out, it needs to go into week 2 actuals.'
pb['J23']='Still open in Xero at 12 Oct (the Oct one, 42 of 60, was paid 6 Oct). Pay this week.'; pb['G23']=D(2026,10,14)
pb['J47']='Template $3,850 a month. Nothing paid in Xero to 12 Oct - check if the Oct invoice has come in.'; pb['G47']=D(2026,10,16)
pb['J51']='Quarterly, due 8 Oct. No bill in Xero at 12 Oct - check with DL.'; pb['G51']=D(2026,10,16)
# payments this week block -> week 2 (8 - 14 Oct)
pb['A4']='PAYMENTS THIS WEEK (8 - 14 Oct)'
block=[
 ('Brookvale Insurance Brokers','Six annual renewals - INV1918462, 1918336, 1918436, 1920058, 1920113, 1914759',D(2026,10,14),'=SUM(D61:D66)','EFT','AGREED - PAY',D(2026,10,14),'GC approved all six. One payment Wed 14 Oct - $66,238.98. Not in the 29 Sep forecast at all.'),
 ('=A45','=B45','=C45','=D45','=E45','Due this week','=G45','First Focus Azure - DD on the due date'),
 ('=A51','=B51','=C51','=D51','=E51','Due this week','=G51','CEO Institute quarterly - no bill in Xero yet, check with DL'),
 ('=A23','=B23','=C23','=D23','=E23','Due this week','=G23','Employsure Sep (41 of 60) - still open, overdue'),
 ('=A26','=B26','=C26','=D26','=E26','PAID 8 Oct','=G26','Eptura - paid'),
 ('=A49','=B49','=C49','=D49','=E49','PAID 8 Oct','=G49','LiveU data pack - paid'),
 ('=A32','=B32','=C32','=D32','=E32','PAID 8 Oct','=G32','Optus mobiles - BPAY, paid'),
 ('=A52','=B52','=C52','=D52','=E52','PAID 8 Oct','=G52','Telstra Nighthawk - BPAY, paid'),
 ('=A34','=B34','=C34','=D34','=E34','PAID 8 Oct','=G34','Posh plants - paid'),
]
plain=[copy.copy(pb.cell(7,c)._style) for c in range(1,16)]
green=[copy.copy(pb.cell(6,c)._style) for c in range(1,16)]
for i,row in enumerate(block):
    r=6+i
    for c in range(1,16): pb.cell(r,c)._style=copy.copy(green[c-1] if i==0 else plain[c-1])
    for c,v in zip('ABCDEFGJ',row): pb[f'{c}{r}']=v
    if row[5].startswith('PAID'):
        for c in range(1,11): pb.cell(r,c).fill=GREY
pb['A16']='of which still to go out this week (agreed + due, excludes what is already paid)'
# Cash Flow risks - same four, facts brought up to 12 Oct
cf['A69']="1. Aware Super $225,887 (INV-10596) is the big one - still not paid at 12 Oct and now expected 9 Nov, after this forecast. Complete AV materials $133,246 went out 29 Sep. If Aware slip again the held labour stays held and we watch the cheque account. Build-up on the Complete AV tab."
cf['A71']="3. Ventia's $58,025 landed 29 Sep. Still chasing their other three ($36,677), Downer ($39,664 - both overdue), Gen-e, Microsoft and ON24 - all a month or more old. Chase this week."
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('saved',OUT)
