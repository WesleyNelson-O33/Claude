"""Adds the Zoho auto-paste feed to a live filled tracker without touching anything else
(its own Read Me is kept; a Zoho section is appended). Optionally pastes a Zoho export in.

Usage: python add_zoho_feed_live.py <tracker.xlsx> <out.xlsx> [zoho_export.xlsx]
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import zoho_feed

SRC, OUT = sys.argv[1:3]
ZOHO = sys.argv[3] if len(sys.argv) > 3 else None
wb = openpyxl.load_workbook(SRC)
assert "Zoho Paste" not in wb.sheetnames, "already has the Zoho feed"
zoho_feed.apply(wb)
if ZOHO:
    src = openpyxl.load_workbook(ZOHO).active
    dst = wb["Zoho Paste"]
    for row in src.iter_rows(min_row=1, max_row=src.max_row, max_col=8):
        for c in row:
            dst.cell(c.row, c.column).value = c.value

rm = wb["Read Me"]
r0 = rm.max_row + 2
lines = [
    ("WORK WON  -  MONTHLY ZOHO DUMP", True),
    ("Work Won fills itself from Zoho. Nobody types on it. Once a month (or whenever you want it current):", False),
    ("1. In Zoho run 'All Deals by Stage' (year to date plus all open deals) and export it to Excel.", False),
    ("2. On the Zoho Paste sheet select A1:H3001 and press Delete. In the export click A1, press Ctrl+Shift+End, Ctrl+C. Back on Zoho Paste click A1 and Ctrl+V.", False),
    ("3. The panel on the right of Zoho Paste must read OK: it compares the deals it read with Zoho's own Record Count.", False),
    ("4. Work Won then shows every deal won, lost or closing this financial year plus every open deal, with what Xero has invoiced on the job (and its V job) and what is still to invoice. Month-End section 6 and the FY Summary pipeline and forecast update too.", False),
    ("5. Won Overrides is for what Zoho cannot tell us - the cost centre of a Multi-Service deal, a different expected invoice month, or a note. It is keyed on the Opportunity Number so it survives every new dump. Zoho departments map to cost centres on Lists AG:AH.", False),
    ("Expected Invoice Month = the event date in the deal name if there is one (for example 'ASX - AGM - 22.10.26'), otherwise the Zoho closing date, unless Won Overrides says otherwise.", False),
]
for i, (t, h) in enumerate(lines):
    c = rm.cell(r0 + i, 2, t)
    if h:
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    else:
        c.font = Font(name="Arial", size=11, color="262626")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rm.row_dimensions[r0 + i].height = 15 * (len(t) // 105 + 1) + 2
wb.save(OUT)
print("saved", OUT)
