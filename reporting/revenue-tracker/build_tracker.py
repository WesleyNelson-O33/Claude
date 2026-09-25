"""Builds FY27_Revenue_Tracker_v4.xlsx.

Departments own the job (Onsite / Production / Consulting). Every department row
has a fixed Row ID and a matching row on the Finance sheet, so the job appears on
Finance by itself. Finance types three things per invoice: Invoice No, Invoice
Date and Ex GST (one set per month of FY27). The invoice details flow back to the
department sheets.

Usage: python build_tracker.py <old_v3_workbook.xlsx> <output.xlsx>
The old workbook is only read (for the Lists sheet, the department headings,
column widths and number formats). It is never saved.
"""
import sys
import datetime as dt

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.utils import column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC, OUT = sys.argv[1], sys.argv[2]
PASSWORD = "CTS1234"          # same password as v3
N = int(sys.argv[3]) if len(sys.argv) > 3 else 1500  # job rows per department
DEF_N = int(sys.argv[4]) if len(sys.argv) > 4 else 500  # deferral rows
DEPTS = [("Onsite", "ONS", "tbl_Onsite"),
         ("Production", "PRD", "tbl_Production"),
         ("Consulting", "CON", "tbl_Consulting")]
MONTHS = ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]

# ---------------------------------------------------------------- styles
FONT = "Arial"
NAVY, SLATE, GREY_TXT = "3E5066", "5B708A", "595959"
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, color=GREY_TXT)
F_TOT = Font(name=FONT, size=11, bold=True, color=NAVY)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_IN = Font(name=FONT, size=10, color="000000")
F_CALC = Font(name=FONT, size=10, color=GREY_TXT)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_SEC = Font(name=FONT, size=11, bold=True, color="FFFFFF")
FILL_IN_HDR = PatternFill("solid", fgColor=NAVY)      # header over a typed column
FILL_CALC_HDR = PatternFill("solid", fgColor=SLATE)   # header over a formula column
FILL_FIN_HDR = PatternFill("solid", fgColor="B7791F")  # header over a Finance-typed column
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")
FILL_TOT = PatternFill("solid", fgColor="EDF1F6")
FILL_FIN = PatternFill("solid", fgColor="FFF9E6")     # Finance / Xero typing cells
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '$#,##0.00;[Red]-$#,##0.00'
DATE = "dd-mmm-yy"
WRAP = Alignment(wrap_text=True, vertical="top")
HDR_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")

src = openpyxl.load_workbook(SRC)        # read only - never saved
wb = openpyxl.Workbook()
wb.remove(wb.active)


def title_block(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws.row_dimensions[2].height = 30
    ws["A2"].alignment = Alignment(wrap_text=False, vertical="top")


def hdr(cell, text, kind="calc"):
    cell.value = text
    cell.font = F_HDR
    cell.fill = {"in": FILL_IN_HDR, "calc": FILL_CALC_HDR, "fin": FILL_FIN_HDR}[kind]
    cell.alignment = HDR_ALIGN
    cell.border = BORDER


def style_body(cell, kind, fmt="General"):
    cell.border = BORDER
    cell.number_format = fmt
    if kind == "in":
        cell.font, cell.fill = F_IN, FILL_WHITE
        cell.protection = Protection(locked=False)
    elif kind == "fin":
        cell.font, cell.fill = F_IN, FILL_FIN
        cell.protection = Protection(locked=False)
    else:
        cell.font, cell.fill = F_CALC, FILL_CALC
        cell.protection = Protection(locked=True)


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.password = PASSWORD
    p.autoFilter = False        # filtering allowed
    p.formatColumns = False     # column widths allowed
    p.formatRows = False
    p.sort = True               # sorting blocked (keeps rows lined up)
    p.insertRows = True
    p.deleteRows = True


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)


def add_name(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def red_if(ws, rng, formula, fill="F8D7DA", color="9C0006", bold=False):
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[formula], fill=PatternFill("solid", fgColor=fill),
        font=Font(name=FONT, color=color, bold=bold)))


# ====================================================================== Lists
def build_lists():
    old = src["Lists"]
    ws = wb.create_sheet("Lists")
    for row in old.iter_rows(min_row=1, max_row=old.max_row, max_col=23):
        for c in row:
            if c.value is None:
                continue
            n = ws.cell(c.row, c.column, c.value)
            n.font = Font(name=FONT, size=10, bold=bool(c.font.b),
                          color=(c.font.color.rgb if c.font.color and isinstance(c.font.color.rgb, str) else None))
            n.number_format = c.number_format
            if c.fill and c.fill.fgColor and isinstance(c.fill.fgColor.rgb, str) and c.fill.fgColor.rgb != "00000000":
                n.fill = PatternFill("solid", fgColor=c.fill.fgColor.rgb)
    for k, v in old.column_dimensions.items():
        ws.column_dimensions[k].width = v.width
    # job numbers must stay text so they match the department sheets
    for r in range(5, 1501):
        for col in ("P", "Q", "R"):
            c = ws[f"{col}{r}"]
            c.number_format = "@"
            if c.value is not None:
                c.value = str(c.value)
            c.protection = Protection(locked=False)
    # FY27 month list and first month setting
    ws["Y4"] = "FY27 Months"
    ws["Y4"].font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    ws["Y4"].fill = FILL_IN_HDR
    ws["T12"] = "FY first month (month end)"
    ws["U12"] = dt.datetime(2026, 7, 31)
    ws["U12"].number_format = DATE
    ws["U12"].fill = FILL_FIN
    ws["U12"].protection = Protection(locked=False)
    ws["T13"] = ("Change U12 to 31-Jul of a new year to roll the whole tracker to the next financial "
                 "year. Every month heading on Finance, FY Summary and Month-End follows.")
    ws["T15"], ws["U15"] = "Deferred revenue account (balance sheet)", "11300 Work in Progress"
    ws["T16"], ws["U16"] = "Deferred / prepaid cost account (balance sheet)", "11300 Work in Progress"
    ws["T17"] = ("The balance sheet accounts the Deferral Journal uses. v3 used 11300 for both. If you hold income in "
                 "advance or prepayments in their own accounts, type those accounts here and every journal line follows.")
    for c in ("U15", "U16"):
        ws[c].fill = FILL_FIN
        ws[c].protection = Protection(locked=False)
    for c in ("T12", "T13", "T15", "T16", "T17", "U15", "U16"):
        ws[c].font = Font(name=FONT, size=10)
    for i in range(12):
        c = ws.cell(5 + i, 25, f"=EOMONTH($U$12,{i})")
        c.number_format = "mmm-yy"
        c.font = F_CALC
    ws.column_dimensions["Y"].width = 12
    # unlock the list areas so new values can be added under protection
    for col in "ABCDEFGHIJKLMNOV W".replace(" ", ""):
        for r in range(5, 120):
            ws[f"{col}{r}"].protection = Protection(locked=False)
    for r in range(5, 12):
        ws[f"U{r}"].protection = Protection(locked=False)
    names = {
        "lst_CostCentre": "$A$5:$A$10", "lst_InvoiceType": "$B$5:$B$14", "lst_TaxCode": "$C$5:$C$8",
        "lst_Status": "$D$5:$D$13", "lst_YN": "$E$5:$E$7", "lst_Ariba": "$F$5:$F$10",
        "lst_WIPType": "$G$5:$G$13", "lst_WIPGL": "$H$5:$H$7", "lst_Team": "$I$5:$I$8",
        "lst_RevCost": "$J$5:$J$6", "lst_CCforGL": "$V$5:$V$10", "lst_CCtoGL": "$W$5:$W$10",
        "lst_WonSource": "$M$5:$M$8", "lst_WonStatus": "$O$5:$O$8", "set_GSTRate": "$U$8",
        "set_QwilrBase": "$U$5", "set_RMSBase": "$U$6", "set_WIPAsset": "$U$9", "set_WIPIncome": "$U$10",
        "lst_RevGL": "$K$5:$K$39", "lst_Months": "$N$5:$N$100", "lst_Jobs": "$P$5:$P$1500",
        "lst_JobName": "$Q$5:$Q$1500", "lst_JobCC": "$R$5:$R$1500",
        "lst_FY27Months": "$Y$5:$Y$16", "set_FYFirstMonth": "$U$12", "lst_GLName": "$L$5:$L$39",
        "set_DefRevAcct": "$U$15", "set_DefCostAcct": "$U$16",
    }
    for n, ref in names.items():
        add_name(n, f"Lists!{ref}")
    ws.freeze_panes = "A5"
    protect(ws)


# ====================================================== department sheets
# Columns whose formula changes in v4. Everything else keeps the v3 role:
# typed (unlocked) or formula (locked) exactly as the old file had it.
def dept_formulas(dept, cols):
    """Return {header: formula_template} with {r} as row number."""
    c = lambda h: cols[h]              # column letter for a header
    rid = f"${c('Row ID')}{{r}}"
    fin = lambda fld: f"INDEX(tbl_Finance[{fld}],${c('Finance Row')}{{r}})"
    job = f"${c('Job Number')}{{r}}"
    guard = lambda body: f'=IF({job}="","",IFERROR({body},""))'
    f = {
        "Finance Row": f'=IFERROR(MATCH({rid},tbl_Finance[Row ID],0),"")',
        "Invoice Date": guard(f'IF({fin("Latest Invoice Date")}="","",{fin("Latest Invoice Date")})'),
        "Invoice No": guard(f'{fin("Invoice Numbers")}&""'),
        "Job in Xero?": f'=IF({job}="","",IF(COUNTIF(lst_Jobs,{job}&"")>0,"OK","CHECK"))',
        "Job Name (Xero)": f'=IF({job}="","",IFERROR(INDEX(lst_JobName,MATCH({job}&"",lst_Jobs,0)),""))',
        "Lines on Job": guard(f'N({fin("Invoice Count")})'),
        "Revenue Ex GST": guard(f'N({fin("Xero Invoiced Ex GST")})'),
        "Invoiced": guard(f'{fin("Compare")}&""'),
        "To Invoice": f'=IF({job}="","",MAX(0,N({c("Expected Revenue Ex GST")}{{r}})-N({c("Revenue Ex GST")}{{r}})))',
        "Cost Centres": (f'=IF({job}="","",IF({c("Cost Centre")}{{r}}="","(no cost centre)",{c("Cost Centre")}{{r}}'
                         f'&" - GL "&IFERROR(INDEX(lst_CCtoGL,MATCH({c("Cost Centre")}{{r}},lst_CCforGL,0)),"?")))'),
    }
    # completeness check - the fields Finance needs before it can invoice
    need = [("Client", "Client"), ("Job Description", "Description"), ("Cost Centre", "Cost Centre"),
            ("Invoice Type", "Invoice Type"), ("Tax Code", "Tax Code"),
            ("Expected Revenue Ex GST", "Expected Revenue")]
    miss = "&".join(f'IF({c(h)}{{r}}="",", {lab}","")' for h, lab in need)
    dup = f'IF(COUNTIF(tbl_Finance[Job Number],{job}&"")>1,", Job number used twice","")'
    f["Not Yet on Finance"] = (f'=IF({job}="","",IF(({miss}&{dup})="","Complete",'
                               f'"Fix: "&MID({miss}&{dup},3,300)))')
    if dept == "Onsite":
        f["Open Qwilr"] = '=IF(S{r}="","",HYPERLINK(IF(LEFT(S{r},4)="http",S{r},set_QwilrBase&S{r}),"Open quote"))'
    if dept == "Production":
        f["Total Inc GST"] = guard(f'N({fin("Inc GST")})')
        f["Open in Current RMS"] = ('=IF(W{r}="","",IF(AND(LEFT(W{r},4)<>"http",set_RMSBase=""),"",'
                                    'HYPERLINK(IF(LEFT(W{r},4)="http",W{r},set_RMSBase&W{r}),"Open in RMS")))')
        f["Open Qwilr"] = '=IF(Y{r}="","",HYPERLINK(IF(LEFT(Y{r},4)="http",Y{r},set_QwilrBase&Y{r}),"Open quote"))'
        f["Zoho Number"] = '=IF(C{r}="","",C{r})'
        f["Value Before Discount"] = '=IF(C{r}="","",N(L{r})+N(AC{r}))'
        f["Discount %"] = '=IFERROR(IF(N(AC{r})=0,"",AC{r}/AD{r}),"")'
        f["Video Revenue"] = '=IF(C{r}="","",IF(H{r}="VIDEO",N(L{r}),0))'
        f["Video % of Revenue"] = '=IFERROR(IF(N(L{r})=0,"",AJ{r}/L{r}),"")'
        f["Total Expense"] = '=IF(C{r}="","",N(AF{r})+N(AG{r})+N(AH{r})+N(AI{r}))'
        f["Margin"] = '=IF(C{r}="","",MAX(N(L{r}),N(O{r}))-N(AL{r}))'
        f["Margin %"] = '=IFERROR(IF(MAX(N(L{r}),N(O{r}))=0,"",AM{r}/MAX(N(L{r}),N(O{r}))),"")'
    if dept == "Consulting":
        f["Open Qwilr"] = '=IF(R{r}="","",HYPERLINK(IF(LEFT(R{r},4)="http",R{r},set_QwilrBase&R{r}),"Open quote"))'
        f["Revenue Split Check"] = ('=IF(C{r}="","",IF(N(T{r})+N(U{r})+N(V{r})=0,"Not split",'
                                    'IF(ROUND(N(T{r})+N(U{r})+N(V{r}),2)=ROUND(MAX(N(L{r}),N(O{r})),2),"OK","MISMATCH")))')
        f["Total Expense"] = '=IF(C{r}="","",N(X{r})+N(Y{r})+N(Z{r}))'
        f["Margin"] = '=IF(C{r}="","",MAX(N(L{r}),N(O{r}))-N(AA{r}))'
        f["Margin %"] = '=IFERROR(IF(MAX(N(L{r}),N(O{r}))=0,"",AB{r}/MAX(N(L{r}),N(O{r}))),"")'
    return f


DEPT_SUB = {
    "Onsite": "Onsite / Support jobs. One row per job number. Fill in every white cell. Invoice Date, Invoice No, "
              "Revenue Ex GST and Invoiced come back from the Finance sheet once Finance enters the Xero invoice.",
    "Production": "Production and Video jobs. One row per job number (a video part invoiced on its own V job number "
                  "gets its own row, cost centre VIDEO). Fill in every white cell. Invoice details come back from Finance.",
    "Consulting": "Consulting and Integration jobs. One row per job number. Fill in every white cell and split the "
                  "revenue across Labour / Equipment / Subscription. Invoice details come back from Finance.",
}
HDR_NOTES = {
    "Invoice Date": "From Finance: the latest Xero invoice date on this job.",
    "Invoice No": "From Finance: every Xero invoice number raised on this job.",
    "Lines on Job": "From Finance: how many invoices have been entered against this job.",
    "Revenue Ex GST": "From Finance: total ex GST invoiced in Xero for this job, including prior years.",
    "Invoiced": "From Finance: Not invoiced / Dept higher - to invoice / Agrees / Xero higher - check dept.",
    "To Invoice": "Expected Revenue Ex GST less what Xero has invoiced. This is what is still to be billed.",
    "Expected Revenue Ex GST": "YOU TYPE: what the whole job is worth ex GST. Finance compares Xero to this number.",
    "Not Yet on Finance": "Checks the information Finance needs. Reads Complete, or Fix: and what is missing or wrong.",
    "Cost Centres": "Your cost centre and the revenue GL it posts to.",
    "Row ID": "Fixed link to this job's row on the Finance sheet. Never type over it, never delete the row.",
}


def build_dept(dept, prefix, tname):
    old = src[dept]
    ws = wb.create_sheet(dept)
    headers = [c.value for c in old[4] if c.value is not None]
    headers += ["Row ID", "Finance Row"]
    cols = {h: CL(i + 1) for i, h in enumerate(headers)}
    last = CL(len(headers))
    title_block(ws, old["A1"].value, DEPT_SUB[dept])
    ws["A3"] = f'="Expected Ex GST  "&TEXT(SUBTOTAL(109,{tname}[Expected Revenue Ex GST]),"$#,##0.00")&"      Invoiced in Xero  "&TEXT(SUBTOTAL(109,{tname}[Revenue Ex GST]),"$#,##0.00")&"      Still to invoice  "&TEXT(SUBTOTAL(109,{tname}[To Invoice]),"$#,##0.00")'
    ws.merge_cells("A3:P3")
    ws["A3"].font, ws["A3"].fill = F_TOT, FILL_TOT
    forms = dept_formulas(dept, cols)
    for i, h in enumerate(headers):
        col = i + 1
        L = CL(col)
        oldc = old.cell(5, col) if h not in ("Row ID", "Finance Row") else None
        typed = h not in forms and h != "Row ID" and (oldc is not None and not oldc.protection.locked)
        kind = "in" if typed else "calc"
        hdr(ws.cell(4, col), h, kind)
        if h in HDR_NOTES:
            ws.cell(4, col).comment = Comment(HDR_NOTES[h], "Tracker")
        fmt = oldc.number_format if oldc is not None else ("0" if h == "Finance Row" else "@")
        if h in ("Invoiced", "Not Yet on Finance", "Cost Centres", "Invoice No"):
            fmt = "@" if h != "Not Yet on Finance" else "General"
        if h == "Invoice Date":
            fmt = DATE
        w = old.column_dimensions[L].width if L in old.column_dimensions else 14
        ws.column_dimensions[L].width = w or 14
        for r in range(5, 5 + N):
            cell = ws.cell(r, col)
            if h == "Row ID":
                cell.value = f"{prefix}-{r - 4:04d}"
            elif h in forms:
                cell.value = forms[h].format(r=r)
            style_body(cell, kind, fmt if h not in ("Not Yet on Finance",) else "General")
    ws.column_dimensions["M"].width = 24
    ws.column_dimensions["P"].width = 30
    ws.column_dimensions["Q"].width = 20
    ws.column_dimensions[cols["Row ID"]].width = 10
    ws.column_dimensions[cols["Finance Row"]].hidden = True
    ws.row_dimensions[4].height = 34
    add_table(ws, tname, f"A4:{last}{4 + N}")
    ws.freeze_panes = "D5"
    # dropdowns (same lists as v3)
    dvs = [("Cost Centre", "lst_CostCentre"), ("Invoice Type", "lst_InvoiceType"), ("Tax Code", "lst_TaxCode")]
    if dept == "Production":
        dvs.append(("Job Closed", "lst_YN"))
    for h, lst in dvs:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True)
        dv.add(f"{cols[h]}5:{cols[h]}{4 + N}")
        ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Expected Revenue Ex GST must be a number.", showErrorMessage=True)
    dv.add(f"{cols['Expected Revenue Ex GST']}5:{cols['Expected Revenue Ex GST']}{4 + N}")
    ws.add_data_validation(dv)
    rng = lambda h: f"{cols[h]}5:{cols[h]}{4 + N}"
    red_if(ws, rng("Job in Xero?"), f'{cols["Job in Xero?"]}5="CHECK"')
    red_if(ws, rng("Not Yet on Finance"), f'LEFT({cols["Not Yet on Finance"]}5,4)="Fix:"')
    inv = cols["Invoiced"]
    red_if(ws, rng("Invoiced"), f'{inv}5="Agrees"', fill="D4EDDA", color="155724")
    red_if(ws, rng("Invoiced"), f'LEFT({inv}5,4)="Xero"')
    red_if(ws, rng("Invoiced"), f'OR(LEFT({inv}5,4)="Dept",LEFT({inv}5,3)="Not")', fill="FFF3CD", color="856404")
    if dept == "Consulting":
        red_if(ws, rng("Revenue Split Check"), f'{cols["Revenue Split Check"]}5="MISMATCH"')
    protect(ws)
    return cols


# ================================================================ Finance
FIN_FIRST, FIN_HDR = 7, 6
FIN_LAST = FIN_FIRST + 3 * N - 1
FIN_COLS = [  # header, kind, number format, width
    ("Row ID", "calc", "@", 10), ("Department", "calc", "@", 12), ("Job Number", "calc", "@", 13),
    ("Client", "calc", "@", 24), ("Job in Xero?", "calc", "@", 9), ("Job Name (Xero)", "calc", "@", 26),
    ("Job Description", "calc", "@", 28), ("Cost Centre", "calc", "@", 13), ("Invoice Type", "calc", "@", 13),
    ("Tax Code", "calc", "@", 11), ("Revenue GL", "calc", "@", 10), ("PO or Quote Ref", "calc", "@", 16),
    ("Dept Notes", "calc", "@", 26), ("Dept Expected Ex GST", "calc", MONEY, 15),
    ("Xero Invoiced Ex GST", "calc", MONEY, 15), ("Variance Dept vs Xero", "calc", MONEY, 15),
    ("Compare", "calc", "@", 22), ("Issue", "calc", "@", 34),
    ("Prior Years Invoice Nos", "fin", "@", 16), ("Prior Years Ex GST", "fin", MONEY, 14),
]
for m in MONTHS:
    FIN_COLS += [(f"{m} Invoice No", "fin", "@", 12), (f"{m} Invoice Date", "fin", DATE, 11),
                 (f"{m} Ex GST", "fin", MONEY, 13)]
FIN_COLS += [("GST", "calc", MONEY, 12), ("Inc GST", "calc", MONEY, 13), ("Invoice Count", "calc", "0", 9),
             ("Latest Invoice Date", "calc", DATE, 11), ("Invoice Numbers", "calc", "@", 26),
             ("Finance Notes", "fin", "@", 30), ("Dept Row", "calc", "0", 8)]
FC = {h: CL(i + 1) for i, (h, *_rest) in enumerate(FIN_COLS)}
FIN_NOTES = {
    "Row ID": "Fixed key to the department row. Never change it.",
    "Department": "Which department sheet this row belongs to. Fixed.",
    "Job Number": "From the department sheet. Finance never types it.",
    "Dept Expected Ex GST": "What the department says the whole job is worth (Expected Revenue Ex GST).",
    "Xero Invoiced Ex GST": "Prior Years Ex GST plus the twelve FY27 monthly Ex GST cells Finance entered from Xero.",
    "Prior Years Invoice Nos": "Jobs brought in from 2024 / FY25 / FY26: every invoice number raised before 1 July 2026, separated by commas.",
    "Prior Years Ex GST": "Total ex GST invoiced on this job before 1 July 2026, per Xero. Keeps the variance right for older jobs.",
    "Variance Dept vs Xero": "Dept Expected less Xero Invoiced. Positive = still to invoice. Negative = Xero is higher than the department expected.",
    "Compare": "Agrees / Dept higher - to invoice / Not invoiced - to invoice / Xero higher - check dept / No dept value.",
    "Issue": "Anything wrong on this row. Filter this column to non-blanks and fix each one.",
    "Dept Row": "Helper: the row on the department sheet this Finance row reads from.",
    "Finance Notes": "Finance's own notes. Shown only here.",
}


def fin_pull(field_by_dept, numeric=False):
    """Formula body pulling one field from the right department table."""
    parts = []
    for dept, _p, tname in DEPTS:
        fld = field_by_dept.get(dept)
        parts.append(f"INDEX({tname}[{fld}],${FC['Dept Row']}{{r}})" if fld else '""')
    body = f'IF(${FC["Department"]}{{r}}="Onsite",{parts[0]},IF(${FC["Department"]}{{r}}="Production",{parts[1]},{parts[2]}))'
    if numeric:
        return f'=IF(${FC["Dept Row"]}{{r}}="","",IF({body}="","",{body}))'
    return f'=IF(${FC["Dept Row"]}{{r}}="","",{body}&"")'


def build_finance():
    ws = wb.create_sheet("Finance", 1)
    title_block(ws, "Finance - Job & Invoice Register",
                "Every job the departments set up appears here by itself. Finance types three things per invoice, "
                "straight off Xero, in the month the invoice is dated: Invoice No, Invoice Date and Ex GST (yellow). "
                "Everything else reads itself.")
    ws["A3"] = ('="Dept Expected  "&TEXT(SUBTOTAL(109,tbl_Finance[Dept Expected Ex GST]),"$#,##0.00")'
                '&"      Xero Invoiced  "&TEXT(SUBTOTAL(109,tbl_Finance[Xero Invoiced Ex GST]),"$#,##0.00")'
                '&"      Variance  "&TEXT(SUBTOTAL(109,tbl_Finance[Variance Dept vs Xero]),"$#,##0.00")'
                '&"      Rows with an issue  "&COUNTIF(tbl_Finance[Issue],"?*")')
    ws.merge_cells("A3:R3")
    ws["A3"].font, ws["A3"].fill = F_TOT, FILL_TOT
    ws["A4"] = ("YELLOW = Finance types (Invoice No, Invoice Date, Ex GST).  GREY = pulled from the department "
                "sheets or calculated.  Tip: filter Job Number and untick (Blanks) to see only live jobs.  "
                "Two invoices on one job in the same month: type both numbers in one cell (INV-1001, INV-1002), "
                "the later date and the combined Ex GST.")
    ws["A4"].font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)
    ws.merge_cells(f"{FC['Prior Years Invoice Nos']}5:{FC['Prior Years Ex GST']}5")
    pc = ws[f"{FC['Prior Years Invoice Nos']}5"]
    pc.value, pc.font, pc.fill, pc.alignment = "Before FY27", F_HDR, FILL_FIN_HDR, HDR_ALIGN
    # month banner row 5 (merged over each 3-column block)
    for i, m in enumerate(MONTHS):
        a = FC[f"{m} Invoice No"]
        c = FC[f"{m} Ex GST"]
        ws.merge_cells(f"{a}5:{c}5")
        cell = ws[f"{a}5"]
        cell.value = f"=INDEX(lst_FY27Months,{i + 1})"
        cell.number_format = "mmmm yyyy"
        cell.font = F_HDR
        cell.fill = FILL_FIN_HDR
        cell.alignment = HDR_ALIGN
    for i, (h, kind, fmt, w) in enumerate(FIN_COLS):
        col = i + 1
        hdr(ws.cell(FIN_HDR, col), h, kind)
        if h in FIN_NOTES:
            ws.cell(FIN_HDR, col).comment = Comment(FIN_NOTES[h], "Tracker")
        ws.column_dimensions[CL(col)].width = w
    ws.row_dimensions[FIN_HDR].height = 34

    inv_no = [FC[f"{m} Invoice No"] for m in MONTHS]
    inv_dt = [FC[f"{m} Invoice Date"] for m in MONTHS]
    inv_am = [FC[f"{m} Ex GST"] for m in MONTHS]
    pno, pam = FC["Prior Years Invoice Nos"], FC["Prior Years Ex GST"]
    first_in, last_in = pno, inv_am[-1]
    r_ = "{r}"
    anyin = f"COUNTA(${first_in}{r_}:${last_in}{r_})>0"
    job = f"${FC['Job Number']}{r_}"
    F = {
        "Dept Row": (f'=IFERROR(MATCH($A{r_},IF($B{r_}="Onsite",tbl_Onsite[Row ID],IF($B{r_}="Production",'
                     f'tbl_Production[Row ID],tbl_Consulting[Row ID])),0),"")'),
        "Job Number": fin_pull({d: "Job Number" for d, *_ in DEPTS}),
        "Client": fin_pull({d: "Client" for d, *_ in DEPTS}),
        "Job in Xero?": f'=IF({job}="","",IF(COUNTIF(lst_Jobs,{job})>0,"OK","CHECK"))',
        "Job Name (Xero)": f'=IF({job}="","",IFERROR(INDEX(lst_JobName,MATCH({job},lst_Jobs,0)),""))',
        "Job Description": fin_pull({d: "Job Description" for d, *_ in DEPTS}),
        "Cost Centre": fin_pull({d: "Cost Centre" for d, *_ in DEPTS}),
        "Invoice Type": fin_pull({d: "Invoice Type" for d, *_ in DEPTS}),
        "Tax Code": fin_pull({d: "Tax Code" for d, *_ in DEPTS}),
        "Revenue GL": f'=IF({FC["Cost Centre"]}{r_}="","",IFERROR(INDEX(lst_CCtoGL,MATCH({FC["Cost Centre"]}{r_},lst_CCforGL,0)),""))',
        "PO or Quote Ref": fin_pull({"Onsite": "PO / Reference", "Production": "Current RMS No", "Consulting": "Qwilr Quote"}),
        "Dept Notes": fin_pull({d: "Notes" for d, *_ in DEPTS}),
        "Dept Expected Ex GST": fin_pull({d: "Expected Revenue Ex GST" for d, *_ in DEPTS}, numeric=True),
        "Xero Invoiced Ex GST": f'=IF(AND({job}="",NOT({anyin})),"",SUM({",".join(c + r_ for c in [pam] + inv_am)}))',
    }
    E, X, V = (f'{FC["Dept Expected Ex GST"]}{r_}', f'{FC["Xero Invoiced Ex GST"]}{r_}',
               f'{FC["Variance Dept vs Xero"]}{r_}')
    F["Variance Dept vs Xero"] = f'=IF({X}="","",N({E})-{X})'
    F["Compare"] = (f'=IF({X}="","",IF({job}="","Invoice on empty job row",IF({E}="",'
                    f'IF(ROUND({X},2)=0,"No dept value","Invoiced - no dept value"),'
                    f'IF(ROUND({X},2)=0,"Not invoiced - to invoice",IF(ROUND({V},2)=0,"Agrees",'
                    f'IF({V}>0,"Dept higher - to invoice","Xero higher - check dept"))))))')
    # Issue column: every problem on the row, joined
    incomplete = "+".join(
        f'IF(OR(COUNTA({a}{r_}:{c}{r_})=0,AND(COUNTA({a}{r_}:{c}{r_})=3,ISNUMBER({c}{r_}))),0,1)'
        for a, c in zip(inv_no, inv_am)) + (f'+IF(OR(AND({pno}{r_}<>"",NOT(ISNUMBER({pam}{r_}))),'
                                            f'AND({pno}{r_}="",{pam}{r_}<>"")),1,0)')
    outside = "+".join(
        f'IF({d}{r_}="",0,IFERROR(--(EOMONTH({d}{r_},0)<>{a}$5),1))' for a, d in zip(inv_no, inv_dt))
    checks = [
        (f'${FC["Dept Row"]}{r_}=""', "Row ID not found on the department sheet - a row was deleted"),
        (f'AND({job}="",{anyin})', "Invoice entered but the department row has no job number"),
        (f'{FC["Job in Xero?"]}{r_}="CHECK"', "Job number not in the Xero job list"),
        (f'AND({job}<>"",COUNTIF(${FC["Job Number"]}${FIN_FIRST}:${FC["Job Number"]}${FIN_LAST},{job})>1)',
         "Job number on more than one row"),
        (f'AND({job}<>"",{FC["Cost Centre"]}{r_}="")', "Dept has not set a cost centre"),
        (f'AND({job}<>"",{FC["Tax Code"]}{r_}="")', "Dept has not set a tax code"),
        (f'AND({job}<>"",{E}="")', "Dept has not entered an expected value"),
        (f'({incomplete})>0', "An invoice is incomplete (needs No, Date and a number in Ex GST)"),
        (f'({outside})>0', "An invoice date is not in the month it is typed under"),
    ]
    body = "&".join(f'IF({cond},"; {txt}","")' for cond, txt in checks)
    F["Issue"] = f'=IF(AND({job}="",NOT({anyin}),${FC["Dept Row"]}{r_}<>""),"",MID({body},3,500))'
    F["GST"] = f'=IF({X}="","",IF({FC["Tax Code"]}{r_}="GST 10%",ROUND({X}*set_GSTRate,2),0))'
    F["Inc GST"] = f'=IF({X}="","",{X}+{FC["GST"]}{r_})'
    F["Invoice Count"] = (f'=IF({X}="","",' + "+".join(
        f'IF({c}{r_}="",0,LEN({c}{r_})-LEN(SUBSTITUTE({c}{r_},",",""))+1)' for c in [pno] + inv_no) + ")")
    F["Latest Invoice Date"] = f'=IF(COUNT({",".join(c + r_ for c in inv_dt)})=0,"",MAX({",".join(c + r_ for c in inv_dt)}))'
    F["Invoice Numbers"] = "=MID(" + "&".join(f'IF({c}{r_}="","",", "&{c}{r_})' for c in [pno] + inv_no) + ",3,500)"

    for i in range(3 * N):
        r = FIN_FIRST + i
        d_name, prefix, _t = DEPTS[i // N]
        for j, (h, kind, fmt, _w) in enumerate(FIN_COLS):
            cell = ws.cell(r, j + 1)
            if h == "Row ID":
                cell.value = f"{prefix}-{i % N + 1:04d}"
            elif h == "Department":
                cell.value = d_name
            elif h in F:
                cell.value = F[h].format(r=r)
            style_body(cell, kind, fmt)
    add_table(ws, "tbl_Finance", f"A{FIN_HDR}:{CL(len(FIN_COLS))}{FIN_LAST}")
    ws.column_dimensions[FC["Dept Row"]].hidden = True
    ws.freeze_panes = f"E{FIN_FIRST}"
    # validation: date must sit inside its month; amount must be a number
    for a, d, am in zip(inv_no, inv_dt, inv_am):
        dv = DataValidation(type="date", operator="between", formula1=f"EOMONTH(${a}$5,-1)+1",
                            formula2=f"${a}$5", allow_blank=True, showErrorMessage=True,
                            errorTitle="Wrong month", error="This date is not in the month this column is for. "
                                                          "Type the invoice under the month it is dated.")
        dv.add(f"{d}{FIN_FIRST}:{d}{FIN_LAST}")
        ws.add_data_validation(dv)
        dv2 = DataValidation(type="decimal", operator="between", formula1="-99999999", formula2="99999999",
                             allow_blank=True, showErrorMessage=True, error="Ex GST must be a number.")
        dv2.add(f"{am}{FIN_FIRST}:{am}{FIN_LAST}")
        ws.add_data_validation(dv2)
    rng = lambda h: f"{FC[h]}{FIN_FIRST}:{FC[h]}{FIN_LAST}"
    cmp_ = FC["Compare"]
    red_if(ws, rng("Compare"), f'{cmp_}{FIN_FIRST}="Agrees"', fill="D4EDDA", color="155724")
    red_if(ws, rng("Compare"), f'OR(LEFT({cmp_}{FIN_FIRST},4)="Xero",LEFT({cmp_}{FIN_FIRST},7)="Invoice")')
    red_if(ws, rng("Compare"), f'OR(LEFT({cmp_}{FIN_FIRST},4)="Dept",LEFT({cmp_}{FIN_FIRST},3)="Not")',
           fill="FFF3CD", color="856404")
    red_if(ws, rng("Issue"), f'{FC["Issue"]}{FIN_FIRST}<>""')
    red_if(ws, rng("Job in Xero?"), f'{FC["Job in Xero?"]}{FIN_FIRST}="CHECK"')
    red_if(ws, rng("Variance Dept vs Xero"),
           f'AND(ISNUMBER({FC["Variance Dept vs Xero"]}{FIN_FIRST}),ROUND({FC["Variance Dept vs Xero"]}{FIN_FIRST},2)<>0)',
           fill="FFF3CD", color="856404", bold=True)
    protect(ws)


# ============================================================ FY Summary
def build_summary():
    ws = wb.create_sheet("FY Summary", 2)
    title_block(ws, "FY27 Revenue Summary",
                "Invoiced revenue (ex GST) by month, straight off the Finance sheet. Nothing is typed here.")
    ws.column_dimensions["A"].width = 30
    for i in range(2, 16):
        ws.column_dimensions[CL(i)].width = 14

    def month_header(row, first):
        hdr(ws.cell(row, 1), first)
        for i, m in enumerate(MONTHS):
            c = ws.cell(row, 2 + i, f"=INDEX(lst_FY27Months,{i + 1})")
            hdr(c, c.value)
            c.number_format = "mmm-yy"
        hdr(ws.cell(row, 14), "FY Total")

    def section(row, text):
        ws.cell(row, 1, text).font = F_SEC
        for c in range(1, 15):
            ws.cell(row, c).fill = FILL_IN_HDR

    # block 1 - by cost centre
    section(4, "1.  INVOICED EX GST BY COST CENTRE")
    month_header(5, "Cost Centre")
    for k in range(6):
        r = 6 + k
        ws.cell(r, 1, f"=INDEX(lst_CostCentre,{k + 1})")
        for i, m in enumerate(MONTHS):
            ws.cell(r, 2 + i, f'=SUMIFS(tbl_Finance[{m} Ex GST],tbl_Finance[Cost Centre],$A{r})')
    ws.cell(12, 1, "No cost centre set (fix on dept sheet)")
    for i, m in enumerate(MONTHS):
        L = CL(2 + i)
        ws.cell(12, 2 + i, f"=SUM(tbl_Finance[{m} Ex GST])-SUM({L}6:{L}11)")
    ws.cell(13, 1, "TOTAL")
    for c in range(2, 14):
        L = CL(c)
        ws.cell(13, c, f"=SUM({L}6:{L}12)")
    for r in range(6, 14):
        ws.cell(r, 14, f"=SUM(B{r}:M{r})")
    # block 2 - by department
    section(15, "2.  INVOICED EX GST BY DEPARTMENT")
    month_header(16, "Department")
    for k, (d, *_x) in enumerate(DEPTS):
        r = 17 + k
        ws.cell(r, 1, d)
        for i, m in enumerate(MONTHS):
            ws.cell(r, 2 + i, f'=SUMIFS(tbl_Finance[{m} Ex GST],tbl_Finance[Department],$A{r})')
    ws.cell(20, 1, "TOTAL")
    for c in range(2, 14):
        L = CL(c)
        ws.cell(20, c, f"=SUM({L}17:{L}19)")
    for r in range(17, 21):
        ws.cell(r, 14, f"=SUM(B{r}:M{r})")
    ws.cell(21, 1, "Check - department total = cost centre total")
    ws.cell(21, 14, '=IF(ROUND(N20-N13,2)=0,"Agrees","CHECK")')
    # block 3 - jobs: expected vs Xero
    section(23, "3.  JOBS  -  DEPARTMENT EXPECTED vs XERO INVOICED (whole of FY27)")
    heads = ["Department", "Live jobs", "Dept Expected Ex GST", "Xero Invoiced Ex GST", "Variance Dept vs Xero",
             "Not invoiced (count)", "Not invoiced ($ expected)", "Part invoiced (count)",
             "Part invoiced ($ to go)", "Xero higher (count)", "Xero higher ($)", "Rows with an issue"]
    for j, h in enumerate(heads):
        hdr(ws.cell(24, 1 + j), h)
    ws.row_dimensions[24].height = 40
    for k, (d, *_x) in enumerate(DEPTS):
        r = 25 + k
        ws.cell(r, 1, d)
        crit = f'tbl_Finance[Department],$A{r}'
        ws.cell(r, 2, f'=COUNTIFS({crit},tbl_Finance[Job Number],"?*")')
        ws.cell(r, 3, f'=SUMIFS(tbl_Finance[Dept Expected Ex GST],{crit})')
        ws.cell(r, 4, f'=SUMIFS(tbl_Finance[Xero Invoiced Ex GST],{crit})')
        ws.cell(r, 5, f'=C{r}-D{r}')
        ws.cell(r, 6, f'=COUNTIFS({crit},tbl_Finance[Compare],"Not invoiced*")')
        ws.cell(r, 7, f'=SUMIFS(tbl_Finance[Dept Expected Ex GST],{crit},tbl_Finance[Compare],"Not invoiced*")')
        ws.cell(r, 8, f'=COUNTIFS({crit},tbl_Finance[Compare],"Dept higher*")')
        ws.cell(r, 9, f'=SUMIFS(tbl_Finance[Variance Dept vs Xero],{crit},tbl_Finance[Compare],"Dept higher*")')
        ws.cell(r, 10, f'=COUNTIFS({crit},tbl_Finance[Compare],"Xero higher*")')
        ws.cell(r, 11, f'=SUMIFS(tbl_Finance[Variance Dept vs Xero],{crit},tbl_Finance[Compare],"Xero higher*")')
        ws.cell(r, 12, f'=COUNTIFS({crit},tbl_Finance[Issue],"?*")')
    ws.cell(28, 1, "TOTAL")
    for c in range(2, 13):
        L = CL(c)
        ws.cell(28, c, f"=SUM({L}25:{L}27)")
    ws.cell(29, 1, "Revenue still to invoice (Not invoiced $ + Part invoiced $)")
    ws.cell(29, 5, "=G28+I28")
    # styling
    for row in ws.iter_rows(min_row=6, max_row=29, max_col=14):
        for c in row:
            if c.value is None or c.row in (15, 16, 23, 24):
                continue
            c.font = Font(name=FONT, size=10, bold=c.row in (13, 20, 28, 29) or c.column == 1 and c.row in (21,))
            c.border = BORDER
            if c.column > 1 and not (24 <= c.row <= 28 and c.column in (2, 6, 8, 10, 12)):
                c.number_format = MONEY
            if c.row in (13, 20, 28):
                c.fill = FILL_TOT
    for r in range(25, 29):
        for c in (2, 6, 8, 10, 12):
            ws.cell(r, c).number_format = "0"
    ws.freeze_panes = "B4"
    protect(ws)


# ============================================================ WIP Movements
WIP_N = 300
WIP_COLS = [("Month", "in", "mmm-yy", 10), ("Job Number", "in", "@", 13), ("Job in Xero?", "calc", "@", 9),
            ("Client", "calc", "@", 24), ("Cost Centre", "in", "@", 13), ("Xero Invoice No", "in", "@", 13),
            ("Description", "in", "@", 30), ("Type", "in", "@", 22), ("Revenue or Cost", "in", "@", 10),
            ("Amount", "in", MONEY, 14), ("P&L Account", "in", "@", 22), ("Debit Account", "calc", "@", 24),
            ("Credit Account", "calc", "@", 24), ("Journal Amount", "calc", MONEY, 14),
            ("Journal Narration", "calc", "@", 40), ("Journal Ref", "in", "@", 12), ("Posted to Xero", "in", "@", 9),
            ("Notes", "in", "@", 30)]


def build_wip():
    ws = wb.create_sheet("WIP Movements")
    title_block(ws, "WIP Movements",
                "Every journal that moves revenue between the P&L and GL 11300 Work in Progress. SIGN RULE: + = revenue "
                "put INTO this month (work done, not yet invoiced, or a deferral released). - = revenue pushed OUT "
                "of this month (invoiced in advance). One row per month per movement.")
    ws["A3"] = '="Amount  "&TEXT(SUBTOTAL(109,tbl_WIP[Amount]),"$#,##0.00")'
    ws.merge_cells("A3:L3")
    ws["A3"].font, ws["A3"].fill = F_TOT, FILL_TOT
    W = {h: CL(i + 1) for i, (h, *_x) in enumerate(WIP_COLS)}
    f = {
        "Job in Xero?": '=IF(B{r}="","",IF(COUNTIF(lst_Jobs,B{r}&"")>0,"OK","CHECK"))',
        "Client": '=IF(B{r}="","",IFERROR(INDEX(tbl_Finance[Client],MATCH(B{r}&"",tbl_Finance[Job Number],0)),""))',
        "Debit Account": ('=IF(N(J{r})=0,"",IF(J{r}>0,set_WIPAsset,IF(I{r}="Cost",IF(K{r}="","(set the P&L Account)",K{r}),'
                          'IF(K{r}="",set_WIPIncome,K{r}))))'),
        "Credit Account": ('=IF(N(J{r})=0,"",IF(J{r}>0,IF(I{r}="Cost",IF(K{r}="","(set the P&L Account)",K{r}),'
                           'IF(K{r}="",set_WIPIncome,K{r})),set_WIPAsset))'),
        "Journal Amount": '=IF(N(J{r})=0,"",ABS(J{r}))',
        "Journal Narration": ('=IF(N(J{r})=0,"","WIP "&H{r}&IF(B{r}="",""," - "&B{r})&IF(D{r}="",""," - "&D{r})'
                              '&IF(A{r}="",""," - "&TEXT(A{r},"mmm yyyy")))'),
    }
    for i, (h, kind, fmt, w) in enumerate(WIP_COLS):
        hdr(ws.cell(4, i + 1), h, kind)
        ws.column_dimensions[CL(i + 1)].width = w
        for r in range(5, 5 + WIP_N):
            c = ws.cell(r, i + 1)
            if h in f:
                c.value = f[h].format(r=r)
            style_body(c, kind, fmt)
    ws.row_dimensions[4].height = 34
    add_table(ws, "tbl_WIP", f"A4:{CL(len(WIP_COLS))}{4 + WIP_N}")
    for h, lst in [("Month", "lst_Months"), ("Cost Centre", "lst_CostCentre"), ("Type", "lst_WIPType"),
                   ("Revenue or Cost", "lst_RevCost"), ("Posted to Xero", "lst_YN")]:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True)
        dv.add(f"{W[h]}5:{W[h]}{4 + WIP_N}")
        ws.add_data_validation(dv)
    red_if(ws, f"C5:C{4 + WIP_N}", 'C5="CHECK"')
    red_if(ws, f"Q5:Q{4 + WIP_N}", 'Q5="N"')
    ws.freeze_panes = "D5"
    protect(ws)


# ============================================================== WIP Summary
def build_wip_summary():
    ws = wb.create_sheet("WIP Summary")
    title_block(ws, "WIP Balance by Job",
                "Job-level WIP for the month selected on Month-End. Paste the matching balance from the WIP Schedule "
                "file into column G and the variance shows which job is out.")
    ws["A4"], ws["C4"], ws["D4"] = "Month", "='Month-End'!C4", "<- set this on the Month-End sheet"
    ws["C4"].number_format = "mmm-yy"
    rows = [("Opening total", '=SUMIFS(tbl_WIP[Amount],tbl_WIP[Month],"<"&$C$4)'),
            ("Movement this month", '=SUMIFS(tbl_WIP[Amount],tbl_WIP[Month],$C$4)'),
            ("Closing total", "=B6+B7"),
            ("Of which is on a job Xero does not have - fix these",
             '=SUMIFS(tbl_WIP[Amount],tbl_WIP[Job in Xero?],"CHECK",tbl_WIP[Month],"<="&$C$4)'),
            ("Of which has no job number", '=SUMIFS(tbl_WIP[Amount],tbl_WIP[Job Number],"",tbl_WIP[Month],"<="&$C$4)'),
            ("Closing listed below, job by job", "=B8-B9-B10")]
    for k, (lab, fm) in enumerate(rows):
        ws.cell(6 + k, 1, lab)
        ws.cell(6 + k, 2, fm).number_format = MONEY
    heads = ["Job Number", "Job Name (Xero)", "Opening", "Movement", "Closing", "Any balance?",
             "Per WIP Schedule (paste in)", "Variance"]
    for j, h in enumerate(heads):
        hdr(ws.cell(13, 1 + j), h, "in" if h.startswith("Per WIP") else "calc")
    n_jobs = 1496
    for i in range(n_jobs):
        r = 14 + i
        vals = [f'=IF(INDEX(lst_Jobs,{i + 1})="","",INDEX(lst_Jobs,{i + 1})&"")',
                f'=IF(A{r}="","",INDEX(lst_JobName,{i + 1})&"")',
                f'=IF(A{r}="","",SUMIFS(tbl_WIP[Amount],tbl_WIP[Job Number],A{r},tbl_WIP[Month],"<"&$C$4))',
                f'=IF(A{r}="","",SUMIFS(tbl_WIP[Amount],tbl_WIP[Job Number],A{r},tbl_WIP[Month],$C$4))',
                f'=IF(A{r}="","",C{r}+D{r})',
                f'=IF(A{r}="","",IF(AND(ROUND(C{r},2)=0,ROUND(D{r},2)=0),"","yes"))',
                None,
                f'=IF(OR(A{r}="",G{r}=""),"",E{r}-G{r})']
        for j, v in enumerate(vals):
            c = ws.cell(r, 1 + j, v)
            style_body(c, "in" if j == 6 else "calc", MONEY if j in (2, 3, 4, 6, 7) else "@")
    for col, w in zip("ABCDEFGH", (13, 36, 14, 14, 14, 10, 16, 14)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].width = 44
    ws.auto_filter.ref = f"A13:H{13 + n_jobs}"
    red_if(ws, f"H14:H{13 + n_jobs}", 'AND(H14<>"",ROUND(H14,2)<>0)')
    ws.freeze_panes = "A14"
    protect(ws)


# ================================================================ Work Won
def build_won():
    old = src["Work Won"]
    ws = wb.create_sheet("Work Won")
    title_block(ws, old["A1"].value,
                "Paste the month's won opportunities here, one row each. Month-End compares them against what was "
                "invoiced. Put the job number on the row and it shows what Xero has invoiced on that job so far.")
    ws["A3"] = ('="Value Ex GST  "&TEXT(SUBTOTAL(109,tbl_Won[Value Ex GST]),"$#,##0.00")&"      On the Finance Sheet  "'
                '&TEXT(SUBTOTAL(109,tbl_Won[On the Finance Sheet]),"$#,##0.00")')
    ws.merge_cells("A3:L3")
    ws["A3"].font, ws["A3"].fill = F_TOT, FILL_TOT
    heads = [c.value for c in old[4] if c.value is not None]
    f = {"Job in Xero?": '=IF(E{r}="","",IF(COUNTIF(lst_Jobs,E{r}&"")>0,"OK","CHECK"))',
         "On the Finance Sheet": '=IF(E{r}="","",SUMIFS(tbl_Finance[Xero Invoiced Ex GST],tbl_Finance[Job Number],E{r}&""))',
         "Still to Invoice": '=IF(J{r}<>"Won","",IF(E{r}="",N(I{r}),N(I{r})-N(M{r})))'}
    for i, h in enumerate(heads):
        L = CL(i + 1)
        oldc = old.cell(5, i + 1)
        kind = "calc" if h in f else "in"
        hdr(ws.cell(4, i + 1), h, kind)
        ws.column_dimensions[L].width = old.column_dimensions[L].width or 14
        for r in range(5, 605):
            c = ws.cell(r, i + 1)
            if h in f:
                c.value = f[h].format(r=r)
            style_body(c, kind, oldc.number_format)
    ws.row_dimensions[4].height = 34
    add_table(ws, "tbl_Won", f"A4:{CL(len(heads))}604")
    for col, lst in [("A", "lst_Months"), ("B", "lst_WonSource"), ("G", "lst_CostCentre"), ("J", "lst_WonStatus"),
                     ("L", "lst_Months")]:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True)
        dv.add(f"{col}5:{col}604")
        ws.add_data_validation(dv)
    red_if(ws, "F5:F604", 'F5="CHECK"')
    ws.freeze_panes = "C5"
    protect(ws)


# ================================================================ Month-End
def build_month_end():
    ws = wb.create_sheet("Month-End", 1)
    title_block(ws, "Month-End Revenue Close",
                "Pick the month, type the Xero figures into the yellow cells, and work down. Everything else calculates.")
    widths = [44, 16, 14, 15, 16, 16, 18, 18, 15, 22]
    for i, w in enumerate(widths):
        ws.column_dimensions[CL(i + 1)].width = w
    ws["A4"], ws["D4"] = "Month being closed", "<- pick a FY27 month end"
    ws["C4"] = dt.datetime(2026, 8, 31)
    ws["C4"].number_format = "mmm-yy"
    ws["C4"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
    ws["C4"].fill = FILL_FIN
    ws["C4"].protection = Protection(locked=False)
    ws["A4"].font = F_BOLD
    dv = DataValidation(type="list", formula1="lst_FY27Months", allow_blank=False)
    dv.add("C4")
    ws.add_data_validation(dv)
    ws["E4"] = '=IF(ISNUMBER(MATCH($C$4,lst_FY27Months,0)),"","CHECK - pick a month inside FY27")'
    ws["E4"].font = Font(name=FONT, size=10, bold=True, color="9C0006")
    mi = "MATCH($C$4,lst_FY27Months,0)"
    inputs = []

    def section(row, text):
        ws.merge_cells(f"A{row}:J{row}")
        ws.cell(row, 1, text).font = F_SEC
        ws.cell(row, 1).fill = FILL_IN_HDR

    def heads(row, items):
        for j, h in enumerate(items):
            if h:
                hdr(ws.cell(row, 1 + j), h, "fin" if "type in" in h.lower() or "per " in h.lower() else "calc")

    # ---- 1 revenue by cost centre
    section(6, "1.  REVENUE BY COST CENTRE  -  tracker vs Xero P&L")
    heads(7, ["Cost Centre", "Invoiced Ex GST", "GST", "Inc GST", "WIP Movement", "Deferral Movement",
              "Revenue Recognised", "Xero Revenue (type in)", "Variance", "Check"])
    dgrid = f"Deferrals!${DEF_G0}${DEF_HDR - 1}:${DEF_G1}${DEF_HDR - 1}"
    dvals = f"Deferrals!${DEF_G0}${DEF_FIRST}:${DEF_G1}${DEF_LAST}"
    dtype = f"Deferrals!$A${DEF_FIRST}:$A${DEF_LAST}"
    dcc = f"Deferrals!${DC['Cost Centre']}${DEF_FIRST}:${DC['Cost Centre']}${DEF_LAST}"
    for k in range(7):
        r = 8 + k
        if k < 6:
            ws.cell(r, 1, f"=INDEX(lst_CostCentre,{k + 1})")
            ws.cell(r, 2, f"=IFERROR(INDEX('FY Summary'!$B${6 + k}:$M${6 + k},{mi}),0)")
            gst = "+".join(f'SUMIFS(tbl_Finance[{m} Ex GST],tbl_Finance[Cost Centre],$A{r},tbl_Finance[Tax Code],"GST 10%")*({mi}={i + 1})'
                           for i, m in enumerate(MONTHS))
            ws.cell(r, 3, f"=IFERROR(ROUND(({gst})*set_GSTRate,2),0)")
            ws.cell(r, 5, f'=SUMIFS(tbl_WIP[Amount],tbl_WIP[Month],$C$4,tbl_WIP[Cost Centre],$A{r},tbl_WIP[Revenue or Cost],"<>Cost")')
            ws.cell(r, 6, f'=SUMPRODUCT(({dtype}="Revenue")*({dcc}=$A{r})*({dgrid}=$C$4)*{dvals})')
        else:
            ws.cell(r, 1, "No cost centre set")
            ws.cell(r, 2, f"=IFERROR(INDEX('FY Summary'!$B$12:$M$12,{mi}),0)")
            ws.cell(r, 3, 0)
            ws.cell(r, 5, f'=SUMIFS(tbl_WIP[Amount],tbl_WIP[Month],$C$4,tbl_WIP[Revenue or Cost],"<>Cost")-SUM(E8:E13)')
            ws.cell(r, 6, f'=SUMPRODUCT(({dtype}="Revenue")*({dgrid}=$C$4)*{dvals})-SUM(F8:F13)')
        ws.cell(r, 4, f"=B{r}+C{r}")
        ws.cell(r, 7, f"=B{r}+E{r}+F{r}")
        ws.cell(r, 9, f'=IF($H{r}="","",$G{r}-$H{r})')
        ws.cell(r, 10, f'=IF($H{r}="",IF(ROUND(G{r},2)=0,"","Enter Xero figure"),IF(ROUND($I{r},2)=0,"Reconciled","CHECK - "&TEXT($I{r},"$#,##0.00")))')
        inputs.append(f"H{r}")
    ws["A15"] = "TOTAL"
    for c in "BCDEFGHI":
        ws[f"{c}15"] = f"=SUM({c}8:{c}14)"
    ws["J15"] = '=IF(COUNT($H8:$H14)=0,"Enter Xero figures",IF(ROUND($I15,2)=0,"Reconciled","CHECK - "&TEXT($I15,"$#,##0.00")))'
    ws["A16"] = ("Invoiced Ex GST is the month's column on Finance. WIP Movement is WIP Movements rows. Deferral Movement is the "
                 "revenue deferred (-) or released (+) this month from the Deferrals sheet. Revenue Recognised = all three, and "
                 "ties to the Xero P&L once both journals are posted. GST is worked on the month total (rounding of a few cents).")
    ws["A16"].font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)

    # ---- 2 WIP
    section(18, "2.  WORK IN PROGRESS AND DEFERRALS  -  GL 11300 (debit balance)")
    heads(19, ["", "Total", "Revenue (WIP + deferrals)", "Cost (WIP + deferrals)", "", "", "", "", "Check"])
    wip = lambda crit, op: [f"=C{{r}}+D{{r}}",
                            f'=SUMIFS(tbl_WIP[Amount],{crit},tbl_WIP[Revenue or Cost],"<>Cost")'
                            f'+SUMPRODUCT(({dtype}="Revenue")*({dgrid}{op}$C$4)*{dvals})',
                            f'=SUMIFS(tbl_WIP[Amount],{crit},tbl_WIP[Revenue or Cost],"Cost")'
                            f'-SUMPRODUCT(({dtype}="Cost")*({dgrid}{op}$C$4)*{dvals})']
    for r, lab, crit, op in [(20, "Opening balance (all months before this one)", 'tbl_WIP[Month],"<"&$C$4', "<"),
                             (21, "Movement this month", "tbl_WIP[Month],$C$4", "=")]:
        ws.cell(r, 1, lab)
        for j, fm in enumerate(wip(crit, op)):
            ws.cell(r, 2 + j, fm.format(r=r))
    ws["A22"] = "Closing balance per this tracker"
    ws["A27"] = ("Only compare this to 11300 if the deferral accounts on Lists (U15, U16) are also 11300. "
                 "Negative = credit balance (revenue deferred). Positive = debit (accrued revenue or prepaid cost).")
    ws["A27"].font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)
    for c in "BCD":
        ws[f"{c}22"] = f"={c}20+{c}21"
    ws["A23"], ws["A24"] = "Closing balance per the WIP Schedule file (type in)", "Variance - tracker vs WIP Schedule"
    ws["B24"] = '=IF(B23="","",B22-B23)'
    ws["I24"] = '=IF(B23="","Enter WIP Schedule figure",IF(ROUND(B24,2)=0,"Reconciled","CHECK"))'
    ws["A25"], ws["A26"] = "Closing balance per Xero GL 11300 (type in)", "Variance - tracker vs Xero"
    ws["B26"] = '=IF(B25="","",B22-B25)'
    ws["I26"] = '=IF(B25="","Enter Xero 11300 balance",IF(ROUND(B26,2)=0,"Reconciled","CHECK"))'
    inputs += ["B23", "B25"]

    # ---- 3 missed revenue
    section(28, "3.  MISSED REVENUE CHECK  -  department expected vs Xero (whole FY to date)")
    heads(29, ["Item", "Jobs", "Ex GST", "", "", "", "", "", "Status"])
    miss = [("Jobs set up but not invoiced at all", 'COUNTIF(tbl_Finance[Compare],"Not invoiced*")',
             'SUMIFS(tbl_Finance[Dept Expected Ex GST],tbl_Finance[Compare],"Not invoiced*")'),
            ("Jobs part invoiced - department expects more", 'COUNTIF(tbl_Finance[Compare],"Dept higher*")',
             'SUMIFS(tbl_Finance[Variance Dept vs Xero],tbl_Finance[Compare],"Dept higher*")'),
            ("Jobs where Xero is higher than the department expected", 'COUNTIF(tbl_Finance[Compare],"Xero higher*")',
             'SUMIFS(tbl_Finance[Variance Dept vs Xero],tbl_Finance[Compare],"Xero higher*")'),
            ("Invoiced but the department has no expected value", 'COUNTIF(tbl_Finance[Compare],"Invoiced - no dept value")',
             'SUMIFS(tbl_Finance[Xero Invoiced Ex GST],tbl_Finance[Compare],"Invoiced - no dept value")'),
            ("Invoice typed on a row with no job", 'COUNTIF(tbl_Finance[Compare],"Invoice on empty job row")',
             'SUMIFS(tbl_Finance[Xero Invoiced Ex GST],tbl_Finance[Compare],"Invoice on empty job row")')]
    for k, (lab, cnt, amt) in enumerate(miss):
        r = 30 + k
        ws.cell(r, 1, lab)
        ws.cell(r, 2, f"={cnt}")
        ws.cell(r, 3, f"={amt}")
        ws.cell(r, 9, f'=IF(B{r}=0,"Clear","REVIEW")')
    ws["A35"] = "Revenue still to invoice (first two lines)"
    ws["C35"] = "=C30+C31"
    ws["A36"] = ("Filter Compare on the Finance sheet to see each job. Not invoiced and Dept higher are revenue at risk "
                 "of being missed. Chase the department or raise the invoice.")
    ws["A36"].font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)

    # ---- 4 data checks
    section(38, "4.  DATA CHECKS  -  every count should be zero before you close")
    heads(39, ["Check", "Count", "", "", "", "", "", "", "Status"])
    fin_range = f"$C${FIN_FIRST}:$C${FIN_LAST}"
    checks = [
        ("Finance rows with anything in the Issue column", 'COUNTIF(tbl_Finance[Issue],"?*")'),
        ("Job numbers not found in the Xero job list", 'COUNTIF(tbl_Finance[Job in Xero?],"CHECK")'),
        ("Job numbers used on more than one row (any department)",
         f"SUMPRODUCT((Finance!{fin_range}<>\"\")*(COUNTIF(Finance!{fin_range},Finance!{fin_range})>1))"),
        ("Department rows missing information Finance needs",
         'COUNTIF(tbl_Onsite[Not Yet on Finance],"Fix:*")+COUNTIF(tbl_Production[Not Yet on Finance],"Fix:*")'
         '+COUNTIF(tbl_Consulting[Not Yet on Finance],"Fix:*")'),
        ("Incomplete invoices (No, Date or Ex GST missing)", 'COUNTIF(tbl_Finance[Issue],"*incomplete*")'),
        ("Invoice dates typed under the wrong month", 'COUNTIF(tbl_Finance[Issue],"*not in the month*")'),
        ("Department rows deleted (Row ID missing)", 'COUNTIF(tbl_Finance[Issue],"*Row ID not found*")'),
        ("Consulting revenue split does not equal the job value", 'COUNTIF(tbl_Consulting[Revenue Split Check],"MISMATCH")'),
        ("WIP movements with an amount but no month or no job",
         'COUNTIFS(tbl_WIP[Month],"",tbl_WIP[Amount],"<>")+COUNTIFS(tbl_WIP[Job Number],"",tbl_WIP[Amount],"<>")'),
        ("WIP job numbers not found in the Xero job list", 'COUNTIF(tbl_WIP[Job in Xero?],"CHECK")'),
        ("WIP rows with an amount but Revenue or Cost not set", 'COUNTIFS(tbl_WIP[Amount],"<>",tbl_WIP[Revenue or Cost],"")'),
        ("WIP cost rows with no P&L Account", 'COUNTIFS(tbl_WIP[Revenue or Cost],"Cost",tbl_WIP[P&L Account],"",tbl_WIP[Amount],"<>")'),
        ("Deferral rows with anything in the Issue column", 'COUNTIF(tbl_Def[Issue],"?*")'),
        ("WIP rows this month not posted to Xero", 'COUNTIFS(tbl_WIP[Month],$C$4,tbl_WIP[Amount],"<>",tbl_WIP[Posted to Xero],"<>Y")'),
    ]
    for k, (lab, fm) in enumerate(checks):
        r = 40 + k
        ws.cell(r, 1, lab)
        ws.cell(r, 2, f"={fm}")
        ws.cell(r, 9, f'=IF(B{r}=0,"Clear","REVIEW")')
    last_chk = 40 + len(checks) - 1

    # ---- 5 WIP journal
    j0 = last_chk + 2
    section(j0, "5.  WIP JOURNAL FOR THE MONTH  -  post this one journal in Xero")
    heads(j0 + 1, ["Account", "Cost Centre", "Debit", "Credit"])
    r = j0 + 2
    ws.cell(r, 1, "=set_WIPAsset")
    ws.cell(r, 2, '="(no tracking)"')
    ws.cell(r, 3, '=IF(ROUND(SUM(E8:E14),2)>0,SUM(E8:E14),"")')
    ws.cell(r, 4, '=IF(ROUND(SUM(E8:E14),2)<0,-SUM(E8:E14),"")')
    for k in range(7):
        rr = r + 1 + k
        ws.cell(rr, 1, "=set_WIPIncome")
        ws.cell(rr, 2, f"=$A${8 + k}")
        ws.cell(rr, 3, f'=IF(ROUND($E${8 + k},2)<0,-$E${8 + k},"")')
        ws.cell(rr, 4, f'=IF(ROUND($E${8 + k},2)>0,$E${8 + k},"")')
    tot = r + 8
    ws.cell(tot, 1, "TOTAL")
    ws.cell(tot, 3, f"=SUM(C{r}:C{tot - 1})")
    ws.cell(tot, 4, f"=SUM(D{r}:D{tot - 1})")
    ws.cell(tot, 5, f'=IF(ROUND(C{tot},2)=ROUND(D{tot},2),"Balanced","CHECK - does not balance")')
    ws.cell(tot + 1, 1, "Narration to use")
    ws.cell(tot + 1, 3, '="WIP movement "&TEXT($C$4,"mmmm yyyy")')
    ws.cell(tot + 2, 1, "WIP Movements revenue rows only. Deferrals have their own journal on the Deferral Journal sheet. WIP cost rows post line by line from WIP Movements.")
    ws.cell(tot + 2, 1).font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)

    # ---- 6 work won
    w0 = tot + 4
    section(w0, "6.  WORK WON  -  ZOHO / Current RMS / Qwilr vs what was invoiced")
    heads(w0 + 1, ["Source", "Opportunities", "Value Ex GST", "", "", "", "", "", "Check"])
    for k in range(4):
        rr = w0 + 2 + k
        ws.cell(rr, 1, f"=INDEX(lst_WonSource,{k + 1})")
        ws.cell(rr, 2, f'=COUNTIFS(tbl_Won[Month],$C$4,tbl_Won[Source],$A{rr},tbl_Won[Status],"Won")')
        ws.cell(rr, 3, f'=SUMIFS(tbl_Won[Value Ex GST],tbl_Won[Month],$C$4,tbl_Won[Source],$A{rr},tbl_Won[Status],"Won")')
    rr = w0 + 6
    ws.cell(rr, 1, "TOTAL WON THIS MONTH")
    ws.cell(rr, 2, f"=SUM(B{w0 + 2}:B{w0 + 5})")
    ws.cell(rr, 3, f"=SUM(C{w0 + 2}:C{w0 + 5})")
    ws.cell(rr + 1, 1, "Invoiced this month (Finance, Ex GST)")
    ws.cell(rr + 1, 3, "=$B$15")
    ws.cell(rr + 2, 1, "Won less invoiced")
    ws.cell(rr + 2, 3, f"=C{rr}-C{rr + 1}")
    ws.cell(rr + 3, 1, "Won this month, still not invoiced against the job")
    ws.cell(rr + 3, 3, '=SUMIFS(tbl_Won[Still to Invoice],tbl_Won[Month],$C$4,tbl_Won[Status],"Won")')
    ws.cell(rr + 4, 1, "Won with a job number Xero does not have (count)")
    ws.cell(rr + 4, 3, '=COUNTIFS(tbl_Won[Month],$C$4,tbl_Won[Status],"Won",tbl_Won[Job in Xero?],"CHECK")')

    # ---- 7 sign-off
    s0 = rr + 6
    section(s0, "7.  SIGN-OFF")
    for k, (lab, role) in enumerate([("Prepared by", "Accounts Assistant"), ("Reviewed by", "Finance Operations Manager"),
                                     ("Date closed", ""), ("Notes / carried forward items", "")]):
        ws.cell(s0 + 1 + k, 1, lab)
        ws.merge_cells(f"B{s0 + 1 + k}:D{s0 + 1 + k}")
        ws.cell(s0 + 1 + k, 5, role)
        inputs.append(f"B{s0 + 1 + k}")

    # styling pass
    for row in ws.iter_rows(min_row=7, max_row=s0 + 4, max_col=10):
        for c in row:
            if c.font.b and c.font.color is not None and c.font.color.rgb in ("00FFFFFF", "FFFFFFFF"):
                continue
            if isinstance(c.value, str) and c.font.i:
                continue
            c.font = Font(name=FONT, size=10, bold=(ws.cell(c.row, 1).value in ("TOTAL", "TOTAL WON THIS MONTH")))
            if c.column in (2, 3, 4, 5, 6, 7, 8, 9) and c.value is not None and not (
                    39 <= c.row <= last_chk or 29 <= c.row <= 34 and c.column == 2):
                c.number_format = MONEY
    for a in inputs:
        ws[a].fill = FILL_FIN
        ws[a].protection = Protection(locked=False)
        ws[a].border = BORDER
    ws["B25"].number_format = ws["B23"].number_format = MONEY
    for r in range(8, 15):
        ws[f"H{r}"].number_format = MONEY
    ws[f"B{s0 + 3}"].number_format = DATE
    for rng_, f_ in [("J8:J15", 'LEFT(J8,5)="CHECK"'), ("I24:I26", 'I24="CHECK"'),
                     (f"I30:I{last_chk}", 'I30="REVIEW"')]:
        red_if(ws, rng_, f_)
    for rng_, f_ in [("J8:J15", 'J8="Reconciled"'), ("I24:I26", 'I24="Reconciled"'), (f"I30:I{last_chk}", 'I30="Clear"')]:
        red_if(ws, rng_, f_, fill="D4EDDA", color="155724")
    red_if(ws, f"E{tot}", f'LEFT(E{tot},5)="CHECK"')
    ws.freeze_panes = "A5"
    protect(ws)


# ================================================================ Deferrals
DEF_HDR, DEF_FIRST = 6, 7
DEF_LAST = DEF_FIRST + DEF_N - 1
GRID_MONTHS = [dt.date(2024 + (6 + i) // 12, (6 + i) % 12 + 1, 1) for i in range(72)]   # Jul-24 .. Jun-30
DEF_COLS = [  # header, kind, format, width
    ("Type", "fin", "@", 9), ("Invoice or Bill No", "fin", "@", 13), ("Invoice or Bill Date", "fin", DATE, 11),
    ("Defer Start", "fin", "mmm-yy", 9), ("Defer End", "fin", "mmm-yy", 9),
    ("Job Number Override", "fin", "@", 12), ("Amount Override", "fin", MONEY, 13), ("P&L GL Override", "fin", "@", 10),
    ("Notes", "fin", "@", 24),
    ("Finance Row", "calc", "0", 7), ("Project Number", "calc", "@", 13), ("Project Name", "calc", "@", 30),
    ("Department", "calc", "@", 12), ("Client", "calc", "@", 20), ("Cost Centre", "calc", "@", 12),
    ("Amount Ex GST", "calc", MONEY, 13), ("Invoice Month", "calc", "mmm-yy", 9), ("Start Month", "calc", "mmm-yy", 9),
    ("End Month", "calc", "mmm-yy", 9), ("Months", "calc", "0", 7), ("Per Month", "calc", MONEY, 12),
    ("P&L Account", "calc", "@", 30), ("Balance Sheet Account", "calc", "@", 24),
    ("Opening Deferred", "calc", MONEY, 13), ("Movement This Month", "calc", MONEY, 13),
    ("Closing Deferred", "calc", MONEY, 13), ("Issue", "calc", "@", 36), ("Journal Seq", "calc", "0", 7),
]
DEF_FIXED = len(DEF_COLS)
DEF_COLS += [(m.strftime("%b-%y"), "calc", '#,##0.00;[Red]-#,##0.00;"-"', 10) for m in GRID_MONTHS]
DC = {h: CL(i + 1) for i, (h, *_x) in enumerate(DEF_COLS)}
DEF_G0, DEF_G1 = CL(DEF_FIXED + 1), CL(len(DEF_COLS))
DEF_NOTES = {
    "Type": "Revenue (a sales invoice) or Cost (a supplier bill).",
    "Invoice or Bill No": "Exactly as in Xero. For revenue, the tracker finds the job and the amount on the Finance sheet.",
    "Invoice or Bill Date": "The date on the Xero invoice or bill. The full amount hits the P&L in this month.",
    "Defer Start": "Optional. First month the revenue/cost belongs to. Leave blank to start in the invoice month.",
    "Defer End": "The last month the revenue/cost belongs to. Type any date in that month.",
    "Job Number Override": "COST: type the job number. REVENUE: leave blank - it is found from the invoice number.",
    "Amount Override": "COST: type the bill amount ex GST (positive). REVENUE: leave blank unless the invoice shares a cell on Finance with another invoice, or only part of it is deferred.",
    "P&L GL Override": "COST: pick the expense GL. REVENUE: leave blank to use the job's revenue GL.",
    "Opening Deferred": "Amount still deferred at the start of the journal month (Deferral Journal C4).",
    "Movement This Month": "Journal month movement. Negative = deferred out of the month. Positive = released into the month.",
    "Closing Deferred": "Amount still deferred at the end of the journal month.",
    "Journal Seq": "Line number of this row on the Deferral Journal for the journal month.",
}


def build_deferrals():
    ws = wb.create_sheet("Deferrals")
    title_block(ws, "Deferrals - revenue and cost spread by month",
                "Finance types the yellow cells: Revenue or Cost, the Xero invoice or bill number, its date, and the month "
                "the deferral ends. The schedule on the right spreads it evenly, month by month, until it ends.")
    ws["A3"] = ('="Journal month  "&TEXT(\'Deferral Journal\'!$C$4,"mmm-yy")&"      Revenue deferred at month end  "'
                f'&TEXT(SUMIFS(tbl_Def[Closing Deferred],tbl_Def[Type],"Revenue"),"$#,##0.00")&"      Cost deferred at month end  "'
                f'&TEXT(SUMIFS(tbl_Def[Closing Deferred],tbl_Def[Type],"Cost"),"$#,##0.00")&"      Rows with an issue  "'
                '&COUNTIF(tbl_Def[Issue],"?*")')
    ws.merge_cells("A3:Q3")
    ws["A3"].font, ws["A3"].fill = F_TOT, FILL_TOT
    ws["A4"] = ("REVENUE: type Type, Invoice No, Date and Defer End - the job, amount and GL come from Finance. "
                "COST: also type Job Number, Amount and the expense GL.  Schedule sign: - = deferred out of that month, "
                "+ = released into it.")
    ws["A4"].font = Font(name=FONT, size=9, italic=True, color=GREY_TXT)
    # grid month dates (row 5) - the schedule reads these
    ws.merge_cells(f"A5:{CL(DEF_FIXED)}5")
    ws["A5"] = "Schedule by month  ->"
    ws["A5"].font, ws["A5"].fill, ws["A5"].alignment = F_HDR, FILL_CALC_HDR, Alignment(horizontal="right")
    for i, m in enumerate(GRID_MONTHS):
        c = ws.cell(DEF_HDR - 1, DEF_FIXED + 1 + i, f"=EOMONTH(DATE({m.year},{m.month},1),0)")
        c.number_format, c.font, c.fill, c.alignment = "mmm-yy", F_HDR, FILL_CALC_HDR, HDR_ALIGN
    for i, (h, kind, fmt, w) in enumerate(DEF_COLS):
        hdr(ws.cell(DEF_HDR, i + 1), h, kind)
        if h in DEF_NOTES:
            ws.cell(DEF_HDR, i + 1).comment = Comment(DEF_NOTES[h], "Tracker")
        ws.column_dimensions[CL(i + 1)].width = w
    ws.row_dimensions[DEF_HDR].height = 34

    r_ = "{r}"
    A, B, Cc, D, E, Fo, Go, Ho = (f"${DC[h]}{r_}" for h in (
        "Type", "Invoice or Bill No", "Invoice or Bill Date", "Defer Start", "Defer End",
        "Job Number Override", "Amount Override", "P&L GL Override"))
    J, K, P, Q, R, S, T, U = (f"${DC[h]}{r_}" for h in (
        "Finance Row", "Project Number", "Amount Ex GST", "Invoice Month", "Start Month", "End Month",
        "Months", "Per Month"))
    X, Y = f"${DC['Opening Deferred']}{r_}", f"${DC['Movement This Month']}{r_}"
    JM = "'Deferral Journal'!$C$4"
    inv_cols = ["Prior Years Invoice Nos"] + [f"{m} Invoice No" for m in MONTHS]
    amt_cols = ["Prior Years Ex GST"] + [f"{m} Ex GST" for m in MONTHS]
    key = f'SUBSTITUTE({B}," ","")'
    found = "+".join(f'ISNUMBER(SEARCH(","&{key}&",",","&SUBSTITUTE(tbl_Finance[{c}]," ","")&","))' for c in inv_cols)
    inv_row = f"SUMPRODUCT(MAX((({found})>0)*(ROW(tbl_Finance[Row ID])-{FIN_HDR})))"
    exact_amt = "+".join(f'SUMPRODUCT(--(SUBSTITUTE(tbl_Finance[{n}]," ","")={key}),tbl_Finance[{a}])'
                         for n, a in zip(inv_cols, amt_cols))
    exact_cnt = "+".join(f'COUNTIF(tbl_Finance[{n}],{B})' for n in inv_cols)
    fpull = lambda fld: f'IF({J}="","",INDEX(tbl_Finance[{fld}],{J})&"")'
    anydata = f"COUNTA(${DC['Type']}{r_}:${DC['P&L GL Override']}{r_})>0"
    gl_name = lambda code: f'IFERROR(" "&INDEX(lst_GLName,MATCH({code},lst_RevGL,0)),"")'
    F = {
        "Finance Row": (f'=IF({Fo}<>"",IFERROR(MATCH({Fo}&"",tbl_Finance[Job Number],0),""),'
                        f'IF(OR({A}<>"Revenue",{B}=""),"",IFERROR(1/(1/{inv_row}),"")))'),
        "Project Number": f'=IF({Fo}<>"",{Fo}&"",{fpull("Job Number")})',
        "Project Name": (f'=IF({K}="","",IFERROR(INDEX(lst_JobName,MATCH({K},lst_Jobs,0))&"",'
                         f'{fpull("Job Description")}))'),
        "Department": "=" + fpull("Department"),
        "Client": "=" + fpull("Client"),
        "Cost Centre": "=" + fpull("Cost Centre"),
        "Amount Ex GST": (f'=IF({Go}<>"",{Go},IF(OR({A}<>"Revenue",{B}=""),"",'
                          f'IF(({exact_cnt})=1,{exact_amt},"")))'),
        "Invoice Month": f'=IF({Cc}="","",EOMONTH({Cc},0))',
        "Start Month": f'=IF({Q}="","",IF({D}="",{Q},EOMONTH({D},0)))',
        "End Month": f'=IF({E}="","",EOMONTH({E},0))',
        "Months": f'=IF(OR({R}="",{S}=""),"",IF({S}<{R},"",(YEAR({S})-YEAR({R}))*12+MONTH({S})-MONTH({R})+1))',
        "Per Month": f'=IF(OR({T}="",{P}=""),"",ROUND({P}/{T},2))',
        "P&L Account": (f'=IF({A}="","",IF({Ho}<>"",{Ho}&{gl_name(Ho + "&" + chr(34) * 2)},IF({A}="Cost","(pick the cost GL)",'
                        f'IF({J}="","(job not found)",INDEX(tbl_Finance[Revenue GL],{J})&'
                        f'{gl_name("INDEX(tbl_Finance[Revenue GL]," + J + ")")}))))'),
        "Balance Sheet Account": f'=IF({A}="","",IF({A}="Cost",set_DefCostAcct,set_DefRevAcct))',
        "Opening Deferred": f'=IF({U}="","",-SUMIF(${DEF_G0}${DEF_HDR - 1}:${DEF_G1}${DEF_HDR - 1},"<"&{JM},{DEF_G0}{r_}:{DEF_G1}{r_}))',
        "Movement This Month": f'=IF({U}="","",SUMIF(${DEF_G0}${DEF_HDR - 1}:${DEF_G1}${DEF_HDR - 1},{JM},{DEF_G0}{r_}:{DEF_G1}{r_}))',
        "Closing Deferred": f'=IF({U}="","",{X}-{Y})',
        "Journal Seq": f'=IF(ROUND(N({Y}),2)=0,"",COUNTIF({DC["Journal Seq"]}${DEF_HDR}:{DC["Journal Seq"]}{{rm1}},">0")+1)',
    }
    first_m, last_m = f"${DEF_G0}${DEF_HDR - 1}", f"${DEF_G1}${DEF_HDR - 1}"
    checks = [
        (f'{A}=""', "Pick Revenue or Cost"),
        (f'{B}=""', "No invoice or bill number"),
        (f'{Cc}=""', "No invoice or bill date"),
        (f'{E}=""', "No defer end"),
        (f'AND({R}<>"",{S}<>"",{S}<{R})', "Defer end is before the start"),
        (f'AND({A}="Revenue",{Fo}="",{J}="",{B}<>"")', "Invoice number not found on Finance - check it, or type the job number"),
        (f'AND({Fo}<>"",{J}="")', "Job number not on any department sheet"),
        (f'AND({A}="Revenue",{B}<>"",{P}="")', "Amount not found (invoice shares a cell on Finance) - type it in Amount Override"),
        (f'AND({A}="Cost",{Fo}="")', "Cost needs a job number"),
        (f'AND({A}="Cost",{Go}="")', "Cost needs an amount"),
        (f'AND({A}="Cost",{Ho}="")', "Cost needs a P&L GL"),
        (f'AND({J}<>"",{fpull("Revenue GL")}="",{A}="Revenue",{Ho}="")', "Job has no cost centre, so no revenue GL"),
        (f'OR(AND({Q}<>"",OR({Q}<{first_m},{Q}>{last_m})),AND({S}<>"",{S}>{last_m}),AND({R}<>"",{R}<{first_m}))',
         "Dates fall outside the Jul-24 to Jun-30 schedule"),
        (f'AND({B}<>"",COUNTIF(tbl_WIP[Xero Invoice No],{B})>0)', "Also on WIP Movements - it would count twice"),
        (f'AND({B}<>"",COUNTIF(${DC["Invoice or Bill No"]}${DEF_FIRST}:${DC["Invoice or Bill No"]}${DEF_LAST},{B})>1)',
         "Same invoice or bill number on two deferral rows"),
    ]
    F["Issue"] = (f'=IF(NOT({anydata}),"",MID(' + "&".join(f'IF({c},"; {t}","")' for c, t in checks) + ',3,500))')
    for i in range(DEF_N):
        r = DEF_FIRST + i
        for j, (h, kind, fmt, _w) in enumerate(DEF_COLS):
            c = ws.cell(r, j + 1)
            if j >= DEF_FIXED:
                L = CL(j + 1)
                m = f"{L}${DEF_HDR - 1}"
                c.value = (f'=IF(OR({U.format(r=r)}="",{Q.format(r=r)}=""),0,'
                           f'IF(AND({m}>={R.format(r=r)},{m}<={S.format(r=r)}),'
                           f'IF({m}={S.format(r=r)},{P.format(r=r)}-{U.format(r=r)}*({T.format(r=r)}-1),{U.format(r=r)}),0)'
                           f'-IF({Q.format(r=r)}={m},{P.format(r=r)},0))')
            elif h in F:
                c.value = F[h].format(r=r, rm1=r - 1)
            style_body(c, kind, fmt)
    add_table(ws, "tbl_Def", f"A{DEF_HDR}:{DEF_G1}{DEF_LAST}")
    ws.column_dimensions[DC["Finance Row"]].hidden = True
    ws.freeze_panes = f"{DC['Project Number']}{DEF_FIRST}"
    for h, lst in [("Type", "lst_RevCost"), ("P&L GL Override", "lst_RevGL")]:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True)
        dv.add(f"{DC[h]}{DEF_FIRST}:{DC[h]}{DEF_LAST}")
        ws.add_data_validation(dv)
    for h in ("Invoice or Bill Date", "Defer Start", "Defer End"):
        dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)",
                            allow_blank=True, showErrorMessage=True, error="Type a date.")
        dv.add(f"{DC[h]}{DEF_FIRST}:{DC[h]}{DEF_LAST}")
        ws.add_data_validation(dv)
    red_if(ws, f"{DC['Issue']}{DEF_FIRST}:{DC['Issue']}{DEF_LAST}", f'{DC["Issue"]}{DEF_FIRST}<>""')
    # highlight the journal month column in the schedule
    red_if(ws, f"{DEF_G0}{DEF_FIRST}:{DEF_G1}{DEF_LAST}", f"{DEF_G0}${DEF_HDR - 1}={JM}",
           fill="FFF3CD", color="000000")
    protect(ws)


def build_def_journal():
    ws = wb.create_sheet("Deferral Journal")
    title_block(ws, "Deferral Journal",
                "Pick the month. Every deferral that moves in that month is listed with its project, department, the "
                "account to debit, the account to credit and the amount. The Xero lines on the right key straight in.")
    ws["A4"], ws["A4"].font = "Journal month", F_BOLD
    ws["C4"] = dt.datetime(2026, 8, 31)
    ws["C4"].number_format = "mmm-yy"
    ws["C4"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
    ws["C4"].fill = FILL_FIN
    ws["C4"].protection = Protection(locked=False)
    dv = DataValidation(type="list", formula1="lst_Months", allow_blank=False)
    dv.add("C4")
    ws.add_data_validation(dv)
    ws["D4"] = "<- normally the same month as Month-End"
    ws["D4"].font = F_SUB
    ws["F4"] = '=IF($C$4=\'Month-End\'!$C$4,"","Note: Month-End is on "&TEXT(\'Month-End\'!$C$4,"mmm-yy"))'
    ws["F4"].font = Font(name=FONT, size=10, bold=True, color="9C0006")
    mv = "tbl_Def[Movement This Month]"
    summ = [("Revenue deferred (pushed out of this month)", f'-SUMIFS({mv},tbl_Def[Type],"Revenue",{mv},"<0")'),
            ("Revenue released (brought into this month)", f'SUMIFS({mv},tbl_Def[Type],"Revenue",{mv},">0")'),
            ("Cost deferred (pushed out of this month)", f'-SUMIFS({mv},tbl_Def[Type],"Cost",{mv},"<0")'),
            ("Cost released (brought into this month)", f'SUMIFS({mv},tbl_Def[Type],"Cost",{mv},">0")'),
            ("Net effect on revenue this month (+ up / - down)", f'SUMIFS({mv},tbl_Def[Type],"Revenue")'),
            ("Net effect on cost this month (+ up / - down)", f'SUMIFS({mv},tbl_Def[Type],"Cost")')]
    for k, (lab, fm) in enumerate(summ):
        ws.cell(5 + k, 1, lab).font = Font(name=FONT, size=10)
        c = ws.cell(5 + k, 4, "=" + fm)
        c.number_format, c.font = MONEY, Font(name=FONT, size=10, bold=True)
    top = 12
    ws.cell(top, 1, "JOURNAL LINES").font = F_SEC
    for c in range(1, 16):
        ws.cell(top, c).fill = FILL_IN_HDR
    cols = [("Line", 6), ("Month", 9), ("Type", 9), ("Invoice or Bill No", 13), ("Project Number", 13),
            ("Project Name", 30), ("Department", 12), ("Client", 20), ("Cost Centre", 12), ("Debit Account", 30),
            ("Credit Account", 30), ("Amount", 13), ("Deferred or Released", 11), ("Narration", 50),
            ("Still Deferred After", 13), ("Deferral Row", 7)]
    hr = top + 1
    for j, (h, w) in enumerate(cols):
        hdr(ws.cell(hr, j + 1), h)
        ws.column_dimensions[CL(j + 1)].width = w
    ws.row_dimensions[hr].height = 34
    ws.column_dimensions["P"].hidden = True
    g = lambda fld, r: f"INDEX(tbl_Def[{fld}],$P{r})"
    first = hr + 1
    for i in range(DEF_N):
        r = first + i
        m = g("Movement This Month", r)
        pl, bs = g("P&L Account", r), g("Balance Sheet Account", r)
        vals = [
            f'=IF({i + 1}<=MAX(tbl_Def[Journal Seq]),{i + 1},"")',
            f'=IF($A{r}="","",$C$4)',
            f'=IF($A{r}="","",{g("Type", r)}&"")',
            f'=IF($A{r}="","",{g("Invoice or Bill No", r)}&"")',
            f'=IF($A{r}="","",{g("Project Number", r)}&"")',
            f'=IF($A{r}="","",{g("Project Name", r)}&"")',
            f'=IF($A{r}="","",{g("Department", r)}&"")',
            f'=IF($A{r}="","",{g("Client", r)}&"")',
            f'=IF($A{r}="","",{g("Cost Centre", r)}&"")',
            f'=IF($A{r}="","",IF($C{r}="Revenue",IF({m}<0,{pl},{bs}),IF({m}<0,{bs},{pl})))',
            f'=IF($A{r}="","",IF($C{r}="Revenue",IF({m}<0,{bs},{pl}),IF({m}<0,{pl},{bs})))',
            f'=IF($A{r}="","",ABS({m}))',
            f'=IF($A{r}="","",IF({m}<0,"Deferred","Released"))',
            (f'=IF($A{r}="","",$C{r}&" "&LOWER($M{r})&" - "&$D{r}&" - "&$E{r}&" "&$F{r}&" - "&$G{r}'
             f'&" - "&TEXT($B{r},"mmm yyyy"))'),
            f'=IF($A{r}="","",{g("Closing Deferred", r)})',
            f'=IF($A{r}="","",MATCH($A{r},tbl_Def[Journal Seq],0))',
        ]
        fmts = ["0", "mmm-yy", "@", "@", "@", "@", "@", "@", "@", "@", "@", MONEY, "@", "@", MONEY, "0"]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 1, v)
            style_body(c, "calc", fmts[j])
    last = first + DEF_N - 1
    ws.cell(top + 0, 12, f'="Total  "&TEXT(SUM(L{first}:L{last}),"$#,##0.00")').font = F_SEC
    ws.auto_filter.ref = f"A{hr}:O{last}"
    # Xero manual journal lines: two per entry
    xc = 18   # column R
    ws.cell(top, xc, "XERO MANUAL JOURNAL LINES  -  debits = credits").font = F_SEC
    for c in range(xc, xc + 6):
        ws.cell(top, c).fill = FILL_IN_HDR
    xh = [("Line", 6), ("Account", 30), ("Description", 50), ("Tracking - Cost Centre", 14), ("Debit", 13), ("Credit", 13)]
    for j, (h, w) in enumerate(xh):
        hdr(ws.cell(hr, xc + j), h)
        ws.column_dimensions[CL(xc + j)].width = w
    for i in range(2 * DEF_N):
        r = first + i
        src_r = f"INDEX($A${first}:$O${last},INT(($R{r}+1)/2),{{col}})"
        odd = f"ISODD($R{r})"
        vals = [f'=IF(INT(({i + 1}+1)/2)<=MAX(tbl_Def[Journal Seq]),{i + 1},"")',
                f'=IF($R{r}="","",IF({odd},{src_r.format(col=10)},{src_r.format(col=11)}))',
                f'=IF($R{r}="","",{src_r.format(col=14)})',
                f'=IF($R{r}="","",{src_r.format(col=9)})',
                f'=IF($R{r}="","",IF({odd},{src_r.format(col=12)},""))',
                f'=IF($R{r}="","",IF({odd},"",{src_r.format(col=12)}))']
        fmts = ["0", "@", "@", "@", MONEY, MONEY]
        for j, v in enumerate(vals):
            c = ws.cell(r, xc + j, v)
            style_body(c, "calc", fmts[j])
    xl = first + 2 * DEF_N - 1
    ws.cell(first - 3, xc, "Debits").font = F_BOLD
    ws.cell(first - 3, xc + 1, f"=SUM(V{first}:V{xl})").number_format = MONEY
    ws.cell(first - 3, xc + 2, "Credits").font = F_BOLD
    ws.cell(first - 3, xc + 3, f"=SUM(W{first}:W{xl})").number_format = MONEY
    ws.cell(first - 3, xc + 4, f'=IF(ROUND(S{first - 3}-U{first - 3},2)=0,"Balanced","CHECK - does not balance")').font = F_BOLD
    red_if(ws, f"V{first - 3}", f'LEFT(V{first - 3},5)="CHECK"')
    ws.freeze_panes = f"A{first}"
    protect(ws)


# ================================================================== Read Me
README = [
    ("FY27 Revenue Tracker - v4", None),
    ("Corporate Technology Services Pty Ltd. Replaces FY27_Revenue_Tracker_v3.", None),
    ("", None),
    ("THE RULE", "h"),
    ("The job number and the department are the source of truth.", None),
    ("The department sets the job up. Finance never types a job number.", None),
    ("Finance types three things per invoice, straight off Xero: the invoice number, the invoice date and the amount excluding GST.", None),
    ("Everything else flows by itself, both ways.", None),
    ("", None),
    ("WHO TYPES WHERE", "h"),
    ("Department managers (Onsite, Production, Consulting): fill in every white cell on your own sheet. One row per job number. Start at the top and use the next empty row.", None),
    ("The cells Finance needs before it can invoice are: Job Number, Client, Job Description, Cost Centre, Invoice Type, Tax Code and Expected Revenue Ex GST.", None),
    ("The Not Yet on Finance column checks those for you. It reads Complete, or Fix: followed by what is missing or wrong.", None),
    ("Finance: on the Finance sheet, type only in the yellow cells. Find the job, go to the month the invoice is dated, and type the Invoice No, Invoice Date and Ex GST.", None),
    ("Grey cells are formulas. They are locked so nobody can wipe them by accident.", None),
    ("", None),
    ("HOW THE INFORMATION FLOWS", "h"),
    ("Department to Finance: every department row has a fixed Row ID (last column). The Finance sheet has a matching row for every Row ID, so as soon as a department types a job, the job number, client, description, cost centre, invoice type, tax code, revenue GL, PO or quote reference, notes and the expected value appear on Finance.", None),
    ("Finance to department: once Finance types an invoice, the department sheet shows the Invoice Date (latest), every Invoice No, the number of invoices, Revenue Ex GST (what Xero has invoiced), Invoiced (the status) and To Invoice (what is left).", None),
    ("So department managers can see what has been invoiced without asking Finance.", None),
    ("", None),
    ("THE VARIANCE", "h"),
    ("On Finance, Dept Expected Ex GST sits next to Xero Invoiced Ex GST, then Variance Dept vs Xero, then Compare.", None),
    ("Variance = Dept Expected less Xero Invoiced.", None),
    ("Positive means the department expects more than Xero has invoiced. That is revenue still to invoice, and the row reads Dept higher - to invoice.", None),
    ("Zero reads Agrees. Negative reads Xero higher - check dept, which means either the department value is out of date or the job was over-billed.", None),
    ("Nothing invoiced yet reads Not invoiced - to invoice.", None),
    ("Worked example: Production sets up job 26073110 at $12,000. In August Finance types INV-5001, 20-Aug-26, $5,000 under August. Variance is $7,000 and Compare reads Dept higher - to invoice. In September Finance types INV-5090, 15-Sep-26, $7,000 under September. Variance is nil and Compare reads Agrees. The Production sheet shows Invoice No INV-5001, INV-5090 and Invoice Date 15-Sep-26.", None),
    ("", None),
    ("MAKING SURE NO REVENUE IS MISSED", "h"),
    ("Month-End section 3 counts and values every job that is Not invoiced, Dept higher, or Xero higher.", None),
    ("FY Summary section 3 shows the same thing by department, and the total still to invoice.", None),
    ("On Finance, filter the Compare column to Not invoiced or Dept higher to get the list to chase.", None),
    ("The Issue column on Finance lists everything wrong on a row: job not in Xero, job number used twice, department information missing, an incomplete invoice, a date typed under the wrong month, or a deleted department row.", None),
    ("", None),
    ("RULES THAT KEEP IT ACCURATE", "h"),
    ("Never delete a job row on a department sheet. If a job is cancelled, set Expected Revenue Ex GST to 0 and say so in Notes.", None),
    ("Never type over a Row ID.", None),
    ("Sorting is switched off on the department sheets and Finance so rows cannot be shuffled. Filtering works as normal.", None),
    ("If someone unprotects a sheet and sorts it anyway, nothing breaks - the Row ID travels with the row and Finance follows it. If a row is deleted, the Issue column on Finance says so.", None),
    ("Two invoices for one job in the same month: type both numbers in the one Invoice No cell (INV-1001, INV-1002), the later date, and the combined Ex GST.", None),
    ("A credit note: type it in the month it is dated, as a negative Ex GST. If that month already has an invoice for the job, net them in the one cell and list both numbers.", None),
    ("A video part of a production job that Xero holds on its own V job number (for example 26073110V) goes on its own Production row with cost centre VIDEO.", None),
    ("", None),
    ("JOBS BROUGHT IN FROM 2024, FY25 AND FY26", "h"),
    ("Departments set older jobs up exactly like new ones, with the whole job value in Expected Revenue Ex GST.", None),
    ("On Finance, type everything Xero invoiced on that job before 1 July 2026 into the two Before FY27 cells: Prior Years Invoice Nos (separated by commas) and Prior Years Ex GST (the total).", None),
    ("The variance then compares the department value with everything Xero has ever invoiced on the job, not just FY27. The monthly revenue on FY Summary and Month-End stays FY27 only.", None),
    ("", None),
    ("DEFERRALS  -  REVENUE AND COST", "h"),
    ("Use the Deferrals sheet when an invoice or a supplier bill covers more than one month.", None),
    ("Revenue: type Revenue, the Xero invoice number, the invoice date and the Defer End month. That is all. The project number, project name, department, client, cost centre, amount and revenue GL are found on the Finance sheet from the invoice number.", None),
    ("Cost: type Cost, the bill number, the bill date, the Defer End month, the job number, the bill amount ex GST and the expense GL. Bills are not on the Finance sheet, so these have to be typed.", None),
    ("Defer Start is optional. Leave it blank and the deferral starts in the invoice month.", None),
    ("The schedule to the right of each row works out every month from the start to the end. The amount is spread evenly and the last month takes the rounding. A minus figure is revenue or cost pushed out of that month. A plus figure is revenue or cost released into it.", None),
    ("Worked example: invoice INV-8001, $12,000 ex GST, dated 15 Sep 2026, Defer End Aug 2027. Xero puts the full $12,000 into September. The schedule shows September -11,000 (keep $1,000, defer $11,000), then +1,000 every month from October 2026 to August 2027. By August 2027 nothing is left deferred.", None),
    ("The Deferral Journal sheet: pick the month and every deferral that moves that month is listed with Project Number, Project Name, Department, Client, Cost Centre, the account to debit, the account to credit, the amount and a narration. The Xero lines on the right are the same journal as two lines per entry, ready to key in.", None),
    ("Revenue deferred: debit the job's revenue GL, credit the deferred revenue account. Revenue released: the other way round.", None),
    ("Cost deferred: debit the prepaid cost account, credit the expense GL. Cost released: the other way round.", None),
    ("The two balance sheet accounts are on Lists U15 and U16. They are set to 11300 Work in Progress, the same as v3. Change them there if you use separate income-in-advance or prepayment accounts.", None),
    ("Month-End section 1 adds the revenue deferral movement to revenue recognised, so it ties to the Xero P&L once the journal is posted. Section 2 includes the deferral balances.", None),
    ("Do not also put a deferred invoice on WIP Movements. It would count twice, and the Issue column will say so.", None),
    ("The schedule runs from July 2024 to June 2030.", None),
    ("", None),
    ("THE SHEETS", "h"),
    ("Finance: one row per job, all departments. Finance types the yellow invoice cells only.", None),
    ("Deferrals: the deferral register and the month-by-month schedule. Deferral Journal: the journal for any month.", None),
    ("Month-End: pick the month. Revenue by cost centre against the Xero P&L, WIP, the missed revenue check, data checks, the WIP journal, work won and sign-off.", None),
    ("FY Summary: invoiced revenue by month, by cost centre and by department, and expected vs invoiced by department.", None),
    ("Onsite, Production, Consulting: the department sheets. Headings are unchanged from v3.", None),
    ("WIP Movements: manual journals between the P&L and GL 11300.", None),
    ("WIP Summary: WIP balance job by job, to tie back to the WIP Schedule file.", None),
    ("Work Won: the month's won opportunities, compared against what was invoiced.", None),
    ("Lists: every dropdown, the Xero job list, GST rate, WIP accounts and the first month of the financial year.", None),
    ("", None),
    ("LOCKS, SIZE AND NEXT YEAR", "h"),
    ("Every sheet is protected with the password CTS1234, the same as v3. The workbook structure is protected with the same password.", None),
    ("Each department sheet has 1,500 job rows, and the Finance sheet has a matching row for every one (4,500). The Deferrals sheet has 500 rows. The file is built by build_tracker.py, so if you ever need more rows it is rebuilt bigger rather than extended by hand.", None),
    ("To roll to FY28: save a copy, clear the white and yellow cells, and change Lists cell U12 to 31-Jul-27. Every month heading follows.", None),
    ("", None),
    ("WHAT CHANGED FROM v3", "h"),
    ("Finance used to type the job number and cost centre on every invoice line. Now the job and cost centre come from the department, and Finance types only Invoice No, Invoice Date and Ex GST.", None),
    ("The Finance sheet is now one row per job with twelve monthly invoice slots, instead of one row per invoice line.", None),
    ("Department headings are unchanged. A Row ID column was added at the end of each department sheet to link the row to Finance.", None),
    ("On the department sheets: Invoiced now shows the status text from Finance. To Invoice is Expected Revenue less Revenue Ex GST. Not Yet on Finance is now the completeness check. Cost Centres shows your cost centre and its revenue GL.", None),
    ("Deferrals now have their own register (Deferrals) and their own journal (Deferral Journal), for revenue and cost, with the project number, project name and department on every journal line.", None),
    ("Finance has two Before FY27 cells per job for invoicing done before 1 July 2026.", None),
]


def build_readme():
    ws = wb.create_sheet("Read Me", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 120
    r = 2
    for text, kind in README:
        c = ws.cell(r, 2, text)
        if r == 2:
            c.font = F_TITLE
        elif kind == "h":
            c.font = F_SEC
            c.fill = FILL_IN_HDR
        else:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    protect(ws)


# ==================================================================== build
build_lists()
dept_cols = {}
for d, p, t in DEPTS:
    dept_cols[d] = build_dept(d, p, t)
build_finance()
build_summary()
build_month_end()
build_wip()
build_deferrals()
build_def_journal()
build_wip_summary()
build_won()
build_readme()
order = ["Read Me", "Finance", "Month-End", "Deferrals", "Deferral Journal", "FY Summary", "Onsite", "Production", "Consulting",
         "WIP Movements", "WIP Summary", "Work Won", "Lists"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.security = WorkbookProtection(workbookPassword=PASSWORD, lockStructure=True)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
