"""Fills a copy of the v4 tracker with a worked example for review.

Job numbers, job names, cost centres and GL codes are real (Xero job list on the
v3 Lists sheet). Clients are read off the job names. Every amount, invoice number,
bill number and date is MADE UP for the example.

Usage: python build_example.py <blank_v4.xlsx> <out.xlsx> [xero_figures.json]
"""
import datetime as dt
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

SRC, OUT = sys.argv[1], sys.argv[2]
XERO = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
D = lambda y, m, d: dt.datetime(y, m, d)
wb = openpyxl.load_workbook(SRC)

# ------------------------------------------------------------- departments
ONSITE = [  # C Job, F Client, G Desc, H CC, I Type, J Tax, O Expected, R PO, W Notes
    ("1144", "Bankwest", "Half-time tech, Perth - Jul to Dec 2026", "ONSITE", "Contract", "GST 10%", 24000, "BW-PO-4471",
     "Invoice $4,000 monthly in arrears"),
    ("1350", "CBA", "Brisbane onsite support - Jul/Aug 2026", "ONSITE", "Contract", "GST 10%", 7500, "CBA-88213",
     "Two monthly invoices"),
    ("20250731", "PWC", "National pool of hours - started 2025", "ONSITE", "Ad-hoc", "GST 10%", 24000, "PWC-NPH-25",
     "Job from FY26 - $18,000 already invoiced before 1 Jul 2026"),
    ("2026021101", "AusPayNet", "Annual AV support & maintenance Jul-26 to Jun-27", "ONSITE", "Contract", "GST 10%",
     12000, "APN-2026-Q", "Invoiced annually in advance - defer over 12 months"),
    ("1530", "Deloitte (DTTL)", "Brisbane onsite tech - August cover", "ONSITE", "Ad-hoc", "GST 10%", 5500, "",
     "Work done in August, invoice not raised yet"),
    ("2192", "Bankwest", "Additional tech - Aug event + 3 months remote support", "ONSITE", "Ad-hoc", "GST 10%", 3000,
     "BW-PO-4502", "Two invoices raised same day: $1,800 event + $1,200 support Aug-Oct"),
]
PRODUCTION = [  # C, F, G, H, I, J, O, AF cross hire, AI labour, V event date, AT notes
    ("26073110", "Garvan Institute", "Weizmann dinner event 20.8.26 - AV production", "PRODUCTION", "Project", "GST 10%",
     15000, 2500, 3200, D(2026, 8, 20), "Deposit Aug, balance after event"),
    ("26073110V", "Garvan Institute", "Weizmann dinner event 20.8.26 - videography", "VIDEO", "Project", "GST 10%",
     3000, 0, 900, D(2026, 8, 20), "Quoted $3,000 - extra edit hours billed"),
    ("26073111", "Downer", "FYR & employee webcast 25.8.26", "PRODUCTION", "Project", "GST 10%", 9800, 1200, 2100,
     D(2026, 8, 25), ""),
    ("26091601", "Automic", "September all hands 25.9.26", "PRODUCTION", "Project", "GST 10%", 7200, 800, 1500,
     D(2026, 9, 25), "Event today - invoice after event"),
    ("2501738", "APA", "HYR25 & FYR25 results webcasts", "PRODUCTION", "Project", "GST 10%", 18500, 0, 0,
     D(2025, 8, 20), "FY26 job - fully invoiced before 1 Jul 2026"),
]
CONSULTING = [  # C, F, G, H, I, J, O, T lab, U equip, V subs, X ext lab, Y equip exp, Z subs exp, AD notes
    ("24121001", "Bankwest", "Perth studio equipment supply & install", "INTEGRATION", "Progress Claim", "GST 10%", 45000,
     12000, 33000, 0, 4000, 24000, 0, "Progress claim 1 in FY26, final claim Jul-26"),
    ("25041502", "NSW Property (PDNSW)", "Desk booking licence extension Aug-26 to Jul-27", "INTEGRATION", "Subscription",
     "GST 10%", 9600, 0, 0, 9600, 0, 0, 7200, "12-month licence invoiced upfront - defer"),
    ("2722", "PWC", "Monthly consulting retainer FY27", "CONSULTING", "Contract", "GST 10%", 30000, 30000, 0, 0, 0, 0, 0,
     "$2,500 a month"),
    ("3040", "Corrs Chambers Westgarth", "Melbourne office relocation - AV design", "CONSULTING", "Milestone", "GST 10%",
     12000, 12000, 0, 0, 3000, 0, 0, "Milestone 1 invoiced"),
    ("3099", "Test Client", "EXAMPLE ERROR - job not in Xero and no tax code", "CONSULTING", "Project", None, 2000,
     2000, 0, 0, 0, 0, 0, "Deliberate mistake to show the checks"),
]


def put(ws, row, **cells):
    for col, v in cells.items():
        if v not in (None, ""):
            ws[f"{col}{row}"] = v


on, pr, co = wb["Onsite"], wb["Production"], wb["Consulting"]
for i, (j, cl, ds, cc, it, tx, ex, po, nt) in enumerate(ONSITE):
    put(on, 5 + i, C=j, F=cl, G=ds, H=cc, I=it, J=tx, O=ex, R=po, W=nt, V="Onsite Manager")
for i, (j, cl, ds, cc, it, tx, ex, xh, lab, ed, nt) in enumerate(PRODUCTION):
    put(pr, 5 + i, C=j, F=cl, G=ds, H=cc, I=it, J=tx, O=ex, AF=xh, AI=lab, V=ed, AT=nt, S=cl, AB="N")
for i, (j, cl, ds, cc, it, tx, ex, t, u, v, x, y, z, nt) in enumerate(CONSULTING):
    put(co, 5 + i, C=j, F=cl, G=ds, H=cc, I=it, J=tx, O=ex, T=t, U=u, V=v, X=x, Y=y, Z=z, AD=nt)

# ----------------------------------------------------------------- Finance
fi = wb["Finance"]
H = {c.value: c.column_letter for c in fi[6]}
rowof = {}
for r in range(7, fi.max_row + 1):
    rowof[fi[f"A{r}"].value] = r
ROW = {j: rowof[f"ONS-{i + 1:04d}"] for i, (j, *_x) in enumerate(ONSITE)}
ROW.update({j: rowof[f"PRD-{i + 1:04d}"] for i, (j, *_x) in enumerate(PRODUCTION)})
ROW.update({j: rowof[f"CON-{i + 1:04d}"] for i, (j, *_x) in enumerate(CONSULTING)})


def inv(job, month, no, date, amt):
    r = ROW[job]
    fi[f"{H[month + ' Invoice No']}{r}"] = no
    fi[f"{H[month + ' Invoice Date']}{r}"] = date
    fi[f"{H[month + ' Ex GST']}{r}"] = amt


def prior(job, nos, amt):
    r = ROW[job]
    fi[f"{H['Prior Years Invoice Nos']}{r}"] = nos
    fi[f"{H['Prior Years Ex GST']}{r}"] = amt


inv("1144", "Jul", "INV-10101", D(2026, 7, 31), 4000)
inv("1144", "Aug", "INV-10245", D(2026, 8, 31), 4000)
inv("1144", "Sep", "INV-10390", D(2026, 9, 24), 4000)
inv("1350", "Jul", "INV-10102", D(2026, 7, 31), 3750)
inv("1350", "Aug", "INV-10246", D(2026, 8, 31), 3750)
prior("20250731", "INV-9120, INV-9187", 18000)
inv("20250731", "Aug", "INV-10230", D(2026, 8, 14), 6000)
inv("2026021101", "Jul", "INV-10088", D(2026, 7, 1), 12000)
inv("2192", "Aug", "INV-10231, INV-10232", D(2026, 8, 15), 3000)
inv("26073110", "Aug", "INV-10205", D(2026, 8, 5), 6000)
inv("26073110", "Sep", "INV-10301", D(2026, 9, 2), 9000)
inv("26073110V", "Sep", "INV-10302", D(2026, 9, 2), 3500)
inv("26073111", "Aug", "INV-10270", D(2026, 8, 27), 9800)
prior("2501738", "INV-8801, INV-9004", 18500)
prior("24121001", "INV-9350", 30000)
inv("24121001", "Jul", "INV-10120", D(2026, 7, 22), 15000)
inv("25041502", "Aug", "INV-10210", D(2026, 8, 3), 9600)
inv("2722", "Jul", "INV-10103", D(2026, 7, 31), 2500)
inv("2722", "Aug", "INV-10247", D(2026, 8, 31), 2500)
inv("2722", "Sep", "INV-10391", D(2026, 9, 24), 2500)
inv("3040", "Aug", "INV-10312", D(2026, 9, 3), 6000)   # deliberate: September date typed under August
fi[f"{H['Finance Notes']}{ROW['2192']}"] = "INV-10231 $1,800 event, INV-10232 $1,200 support Aug-Oct (deferred)"

# --------------------------------------------------------------- Deferrals
df = wb["Deferrals"]
DH = {c.value: c.column_letter for c in df[6]}


def defer(row, **k):
    for key, v in k.items():
        df[f"{DH[key]}{row}"] = v


defer(7, **{"Type": "Revenue", "Invoice or Bill No": "INV-10088", "Invoice or Bill Date": D(2026, 7, 1),
            "Defer End": D(2027, 6, 30), "Notes": "AusPayNet annual support in advance"})
defer(8, **{"Type": "Revenue", "Invoice or Bill No": "INV-10210", "Invoice or Bill Date": D(2026, 8, 3),
            "Defer End": D(2027, 7, 31), "P&L GL Override": "41175",
            "Notes": "PDNSW 12-month licence - GL override: INTEGRATION maps to 42300, licences belong in 41175"})
defer(9, **{"Type": "Revenue", "Invoice or Bill No": "INV-10232", "Invoice or Bill Date": D(2026, 8, 15),
            "Defer End": D(2026, 10, 31), "Amount Override": 1200,
            "Notes": "Shares a cell with INV-10231 on Finance, so the amount is typed"})
defer(10, **{"Type": "Cost", "Invoice or Bill No": "BILL-5521", "Invoice or Bill Date": D(2026, 8, 12),
             "Defer End": D(2026, 12, 31), "Job Number Override": "1144", "Amount Override": 1800,
             "P&L GL Override": "45010", "Notes": "Qantas flights & accommodation prepaid for the Perth contract"})

# ----------------------------------------------------------- WIP Movements
wp = wb["WIP Movements"]
put(wp, 5, A=D(2026, 8, 31), B="1530", E="ONSITE", G="August onsite cover worked, invoice not raised",
    H="Accrual - unbilled work", I="Revenue", J=2750, Q="N")
put(wp, 6, A=D(2026, 9, 30), B="1530", E="ONSITE", G="Reverse August accrual", H="Reversal of prior accrual",
    I="Revenue", J=-2750, Q="N")

# --------------------------------------------------------------- Work Won
ww = wb["Work Won"]
put(ww, 5, A=D(2026, 8, 31), B="Current RMS", C="RMS-7781", D="Garvan Institute", E="26073110", G="PRODUCTION",
    H="Weizmann dinner", I=15000, J="Won", K=D(2026, 7, 28), L=D(2026, 8, 31))
put(ww, 6, A=D(2026, 8, 31), B="Qwilr", C="Q-2291", D="Automic", E="26091601", G="PRODUCTION",
    H="September all hands", I=7200, J="Won", K=D(2026, 8, 18), L=D(2026, 9, 30))

# ------------------------------------------------------- Month-End (Aug-26)
me = wb["Month-End"]
me["C4"] = D(2026, 8, 31)
wb["Deferral Journal"]["C4"] = D(2026, 8, 31)
for cell, v in XERO.items():
    me[cell] = v
if XERO:
    me["B106"] = "Example - Accounts Assistant"

# ------------------------------------------------------------ Example Guide
from openpyxl.styles import Border, Side
PRIMARY, ACCENT, TINT, LINE = "1F3864", "BF8F00", "D9E2F3", "C9D3E3"
g = wb.create_sheet("Example Guide", 0)
for col, w in zip("ABCDE", (5, 30, 58, 72, 2)):
    g.column_dimensions[col].width = w
for c in range(1, 6):
    for r in (1, 2, 3):
        g.cell(r, c).fill = PatternFill("solid", fgColor=PRIMARY)
g["B2"] = "WORKED EXAMPLE  -  for review only, not live data"
g["B2"].font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
g.row_dimensions[2].height = 34
g["B3"] = "Real job numbers from your Xero job list. Every amount, invoice number, bill number and date is made up."
g["B3"].font = Font(name="Arial", size=11, italic=True, color=TINT)
g.row_dimensions[3].height = 22
g["B5"] = ("Month-End and the Deferral Journal are set to August 2026. Work down the list: open the sheet named in "
           "column B, find the job, and check it shows what column D says.")
g["B5"].font = Font(name="Arial", size=11, color="262626")
g["B5"].alignment = Alignment(wrap_text=True, vertical="center")
g.merge_cells("B5:D5")
g.row_dimensions[5].height = 32
heads = ["#", "Job / where to look", "What was entered", "What you should see"]
for j, h in enumerate(heads):
    c = g.cell(7, 1 + j, h)
    c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ACCENT if j == 3 else PRIMARY)
    c.alignment = Alignment(vertical="center", horizontal="center" if j == 0 else None, indent=0 if j == 0 else 1)
g.row_dimensions[7].height = 26
SCEN = [
    ("1144 Bankwest - Onsite + Finance",
     "Dept: expected $24,000 (Jul-Dec contract). Finance: $4,000 in Jul, Aug and Sep.",
     "Finance: Xero $12,000, Variance $12,000, 'Dept higher - to invoice'. Onsite sheet shows all three invoice numbers, latest date 24-Sep-26, To Invoice $12,000. Normal - the contract is only half way."),
    ("1350 CBA - Onsite + Finance", "Dept $7,500. Finance $3,750 Jul + $3,750 Aug.",
     "'Agrees', variance nil, on both sheets."),
    ("20250731 PWC - job from FY26",
     "Dept $24,000. Finance: Before FY27 = INV-9120, INV-9187, $18,000; plus $6,000 in Aug.",
     "Xero $24,000 = Agrees. Without the Before FY27 cells this job would wrongly show $18,000 still to invoice."),
    ("2026021101 AusPayNet - revenue deferral",
     "Finance: $12,000 in Jul (annual support in advance). Deferrals row 1: INV-10088, 01-Jul-26, end Jun-27.",
     "Deferrals finds the job, name, dept, cost centre and $12,000 by itself. Schedule: Jul -11,000, then +1,000 every month to Jun-27. Deferral Journal (Aug): Dr 11300 / Cr 41100 $1,000 released."),
    ("1530 DTTL - not invoiced + WIP accrual",
     "Dept $5,500, no invoice. WIP Movements: +$2,750 Aug accrual, -$2,750 Sep reversal.",
     "Finance: 'Not invoiced - to invoice'. Month-End section 3 counts it. Month-End section 1 ONSITE WIP Movement = +$2,750."),
    ("2192 Bankwest - ONE INVOICE SHARING A CELL",
     "Two invoices dated the same day in August: INV-10231 $1,800 + INV-10232 $1,200. Finance typed both numbers in the one August cell and $3,000.",
     "Finance is right ($3,000, Agrees). Only INV-10232 is deferred, and the tracker cannot tell how much of the $3,000 is INV-10232 - so Deferrals row 3 has $1,200 typed in Amount Override. Remove that override and the row flags 'Amount not found'."),
    ("BILL-5521 - cost deferral (job 1144)",
     "Deferrals row 4: Cost, bill $1,800 dated 12-Aug-26, end Dec-26, job 1144, GL 45010.",
     "Schedule: Aug -1,440, then +360 Sep to Dec. Deferral Journal (Aug): Dr 11300 / Cr 45010 Freight & Travel - ONS $1,440, with project 1144 Bankwest PER Half Tech, Onsite."),
    ("26073110 Garvan - staged production", "Dept $15,000. Finance $6,000 Aug deposit + $9,000 Sep.",
     "Agrees. Production sheet: margin $15,000 - $5,700 costs = $9,300 (62%)."),
    ("26073110V Garvan video - over-invoiced", "Dept $3,000 (VIDEO). Finance $3,500 Sep.",
     "'Xero higher - check dept', variance -$500. Either the dept updates its value or the invoice is wrong."),
    ("26073111 Downer", "Dept $9,800. Finance $9,800 Aug.", "Agrees."),
    ("26091601 Automic", "Dept $7,200. Event is today, no invoice yet.", "'Not invoiced - to invoice'. Work Won shows $7,200 won, $0 invoiced."),
    ("2501738 APA - FY26 job, all invoiced",
     "Dept $18,500. Finance: Before FY27 INV-8801, INV-9004 $18,500. Nothing in FY27.",
     "Agrees. No FY27 revenue - FY Summary monthly figures ignore it."),
    ("24121001 BW studio - progress claims",
     "Dept $45,000 split $12,000 labour / $33,000 equipment. Before FY27 $30,000 + $15,000 in Jul.",
     "Agrees. Consulting sheet: Revenue Split Check OK, margin $45,000 - $28,000 = $17,000."),
    ("25041502 PDNSW licence - revenue deferral",
     "Finance $9,600 Aug. Deferrals row 2: INV-10210, end Jul-27, P&L GL Override 41175.",
     "Schedule: Aug -8,800, then +800 a month Sep-26 to Jul-27. Month-End section 1 INTEGRATION: invoiced $9,600, deferral -$8,800, recognised $800. Journal uses 41175 Subscriptions & Licences - Income because of the override; without it the tracker would use 42300 Installation Labour (the INTEGRATION GL on Lists)."),
    ("2722 PWC retainer", "Dept $30,000. Finance $2,500 Jul, Aug, Sep.", "'Dept higher - to invoice' $22,500 - expected for a retainer."),
    ("3040 Corrs - DELIBERATE MISTAKE",
     "INV-10312 dated 03-Sep-26 typed under AUGUST.",
     "Finance Issue: 'An invoice date is not in the month it is typed under'. Month-End section 4 counts it. (In real use the date box would also refuse it - the example was typed in behind the validation.)"),
    ("3099 - DELIBERATE MISTAKE", "Job not in the Xero list, no tax code.",
     "Consulting 'Not Yet on Finance' = Fix: Tax Code. Finance Issue: job not in Xero list, no tax code. Month-End section 4 counts both."),
    ("Month-End (Aug-26)",
     "Xero revenue by cost centre typed in to match, except CONSULTING which is typed $250 low.",
     "Every cost centre reads Reconciled except CONSULTING: 'CHECK - $250.00'. That is what a break looks like. Section 3 lists $55,200 still to invoice. Section 4 flags the two deliberate mistakes and the WIP accrual not yet posted to Xero."),
    ("Deferral Journal (Aug-26)", "Nothing typed - month set to Aug-26.",
     "Four journal lines (two revenue deferred, one revenue released, one cost deferred), each with project number, name, department, debit account, credit account. Xero lines on the right balance."),
]
for k, (a, b, c) in enumerate(SCEN):
    r = 8 + k
    mistake = "MISTAKE" in a
    band = PatternFill("solid", fgColor="FCE4D6" if mistake else ("F4F7FB" if k % 2 else "FFFFFF"))
    for j, v in enumerate([k + 1, a, b, c]):
        cell = g.cell(r, 1 + j, v)
        cell.font = Font(name="Arial", size=10, bold=(j in (0, 1)), color=PRIMARY if j in (0, 1) else "262626")
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j == 0 else None,
                                   indent=0 if j == 0 else 1)
        cell.fill = band
        cell.border = Border(bottom=Side(style="thin", color=LINE))
g.freeze_panes = "A8"
g.sheet_view.showGridLines = False
g.sheet_properties.tabColor = ACCENT
wb.active = 0
wb.save(OUT)
print("saved", OUT)
