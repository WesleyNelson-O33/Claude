"""Loads a Zoho 'All Deals by Stage' export into the Work Won sheet.

Loads every deal closing in the tracker's financial year (won, lost and open) plus every
open deal regardless of date, so the pipeline is complete.

Usage: python populate_won.py <tracker.xlsx> <zoho_export.xlsx> <out.xlsx> [log.txt]
"""
import calendar
import datetime as dt
import re
import sys

import openpyxl

TRACKER, ZOHO, OUT = sys.argv[1:4]
LOG = sys.argv[4] if len(sys.argv) > 4 else None

CC = {"Event Production": "PRODUCTION", "Video Production": "VIDEO", "Support": "ONSITE",
      "Integration": "INTEGRATION", "Consulting": "CONSULTING"}
STATUS = {"Closed Won": "Won", "Closed Lost": "Lost"}

wb = openpyxl.load_workbook(TRACKER)
fy_start = wb["Lists"]["U12"].value.replace(day=1)          # 1 Jul
fy_end = dt.datetime(fy_start.year + 1, 6, 30)


def eom(d):
    return dt.datetime(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def event_date(name, close):
    """Event date written in the deal name, e.g. 'ASX - AGM - 22.10.26'. Zoho's closing date is when
    the deal was won, so the event date is the better guide to the invoice month. Ignored if it is
    implausible (before the deal closed or more than 18 months after)."""
    for dd, mm, yy in re.findall(r"(?<![\d.])(\d{1,2})\.(\d{1,2})\.(\d{4}|\d{2})(?![\d.])", str(name)):
        y = int(yy) + (2000 if len(yy) == 2 else 0)
        try:
            d = dt.datetime(y, int(mm), int(dd))
        except ValueError:
            continue
        if eom(close) <= eom(d) <= close + dt.timedelta(days=548):
            return d
    return None


zws = openpyxl.load_workbook(ZOHO).active
hdr_row = next(r for r in range(1, 20) if zws.cell(r, 1).value == "Stage")
heads = [c.value for c in zws[hdr_row]]
col = {h: i for i, h in enumerate(heads)}
deals, stage = [], None
for row in zws.iter_rows(min_row=hdr_row + 1, values_only=True):
    if row[0]:
        stage = re.sub(r"\s*\(\s*\d+\s*\)\s*$", "", str(row[0])).strip()
    if not row[col["Deal Name"]]:
        continue
    amt = row[col["Amount (Record Currency)"]]
    deals.append(dict(stage=stage, name=row[col["Deal Name"]], client=row[col["Account Name"]],
                      amt=None if amt in ("", None) else float(amt), dept=row[col["CTS DEPARTMENT"]],
                      close=row[col["Closing Date"]], opp=str(row[col["Opportunity Number"]] or "").strip()))

jobs = {str(c.value) for c in wb["Lists"]["P"][4:] if c.value}
opp_count = {}
for d in deals:
    opp_count[d["opp"]] = opp_count.get(d["opp"], 0) + 1

load = [d for d in deals if STATUS.get(d["stage"], "Open") == "Open" or fy_start <= d["close"] <= fy_end]
load.sort(key=lambda d: (d["close"], d["opp"]))

ws = wb["Work Won"]
assert ws["P4"].value == "Pipeline Stage"
today = dt.datetime.now()
log = []
for i, d in enumerate(load):
    r = 5 + i
    status = STATUS.get(d["stage"], "Open")
    notes = [f"Zoho: {d['dept']}"]
    cc = CC.get(d["dept"])
    if cc is None:
        notes.append("Multi-Service - pick the cost centre")
    if d["amt"] is None:
        notes.append("No amount in Zoho")
    if opp_count[d["opp"]] > 1:
        notes.append("Opportunity number used on more than one Zoho deal")
    if status == "Open" and eom(d["close"]) < eom(today):
        notes.append("Closing date has passed - update Zoho")
    m = eom(d["close"])
    ev = event_date(d["name"], d["close"])
    inv_m = eom(ev) if ev else m
    if ev and inv_m != m:
        notes.append(f"Invoice month from event date in deal name ({ev:%d/%m/%Y}); Zoho closing {d['close']:%d/%m/%Y}")
    vals = {1: m, 2: "ZOHO", 3: d["opp"], 4: d["client"], 5: d["opp"] or None, 7: cc, 8: d["name"],
            9: d["amt"], 10: status, 11: d["close"] if status == "Won" else None, 12: inv_m,
            15: "; ".join(notes), 16: d["stage"]}
    for c, v in vals.items():
        ws.cell(r, c).value = v
    if status == "Won" and d["opp"] not in jobs:
        log.append(f"Row {r}: won deal {d['opp']} {d['name']} - job not on the Lists job list yet")
wb.save(OUT)

by = {}
for d in load:
    k = STATUS.get(d["stage"], "Open")
    n, v = by.get(k, (0, 0.0))
    by[k] = (n + 1, v + (d["amt"] or 0))
out = [f"Loaded {len(load)} deals into Work Won rows 5-{4 + len(load)}"]
out += [f"  {k}: {n} deals, ${v:,.2f}" for k, (n, v) in sorted(by.items())]
out += log
print("\n".join(out[:8]))
if LOG:
    open(LOG, "w").write("\n".join(out) + "\n")
