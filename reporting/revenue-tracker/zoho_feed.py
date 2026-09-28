"""Zoho dump -> Work Won, all by formula.

apply(wb) adds to a tracker workbook:
  * 'Zoho Paste'    - paste the Zoho 'All Deals by Stage' export into A1 exactly as it comes out of Zoho.
                      Helper columns fill the stage down, read the dates and flag which deals belong on Work Won.
  * 'Won Overrides' - the few things Zoho cannot tell us, keyed by Opportunity Number so they survive every
                      new dump: cost centre (Multi-Service deals), expected invoice month, notes.
  * Work Won        - every column is a formula off the two sheets above. Nothing is typed there any more.
  * Lists AD:AE     - Zoho department -> cost centre.

Used by build_tracker.py (new builds) and patch_zoho_feed.py (existing filled trackers).
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

FONT = "Arial"
NAVY, SLATE, ACCENT = "1F3864", "1F4E79", "BF8F00"
FILL_TITLE = PatternFill("solid", fgColor=NAVY)
FILL_CALC_HDR = PatternFill("solid", fgColor=SLATE)
FILL_FIN_HDR = PatternFill("solid", fgColor=ACCENT)
FILL_CALC = PatternFill("solid", fgColor="F4F7FB")
FILL_FIN = PatternFill("solid", fgColor="FFF2CC")
FILL_TOT = PatternFill("solid", fgColor="D9E2F3")
HAIR = Side(style="hair", color="C9D3E3")
THIN = Side(style="thin", color="C9D3E3")
ROW_LINE = Border(bottom=HAIR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '$#,##0.00;[Red]-$#,##0.00;"-"'
PASSWORD = "CTS1234"
ZP_ROWS = 3000          # rows the paste area and helpers cover
WON_LAST = 604          # Work Won rows 5..604
OVR_ROWS = 300

ZOHO_DEPT_CC = [("Event Production", "PRODUCTION"), ("Video Production", "VIDEO"), ("Support", "ONSITE"),
                ("Integration", "INTEGRATION"), ("Consulting", "CONSULTING"), ("Multi-Service", "")]


def _name(wb, name, ref):
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _protect(ws):
    p = ws.protection
    p.sheet = True
    p.password = PASSWORD
    p.autoFilter = False
    p.formatCells = False
    p.formatColumns = False
    p.formatRows = False
    p.sort = True
    p.insertRows = True
    p.deleteRows = True


def _hdr(c, text, fill):
    c.value = text
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER


def _title(ws, title, sub, width_cols):
    for col in range(1, width_cols + 1):
        for r in (1, 2):
            ws.cell(r, col).fill = FILL_TITLE
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color="FFFFFF")
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="D9E2F3")
    ws.row_dimensions[1].height = 30


# ------------------------------------------------------------------ Lists
def _lists(wb):
    ws = wb["Lists"]
    ws["AD4"], ws["AE4"] = "Zoho CTS Department", "Cost Centre"
    for c in ("AD4", "AE4"):
        ws[c].font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        ws[c].fill = FILL_TITLE
    for i in range(8):
        dept, cc = ZOHO_DEPT_CC[i] if i < len(ZOHO_DEPT_CC) else (None, None)
        for col, v in ((30, dept), (31, cc)):
            c = ws.cell(5 + i, col, v)
            c.font = Font(name=FONT, size=10)
            c.fill = FILL_FIN
            c.protection = Protection(locked=False)
    ws["AD14"] = ("Maps Zoho's CTS DEPARTMENT to a tracker cost centre for Work Won. Multi-Service has no single "
                  "cost centre - set it per deal on Won Overrides. Add a row if Zoho gets a new department.")
    ws["AD14"].font = Font(name=FONT, size=9, italic=True, color="595959")
    ws.column_dimensions["AD"].width = 22
    ws.column_dimensions["AE"].width = 14
    _name(wb, "lst_ZohoDept", "Lists!$AD$5:$AD$12")
    _name(wb, "lst_ZohoCC", "Lists!$AE$5:$AE$12")


# ------------------------------------------------------------ Zoho Paste
def _zoho_paste(wb, pos):
    if "Zoho Paste" in wb.sheetnames:
        del wb["Zoho Paste"]
    ws = wb.create_sheet("Zoho Paste", pos)
    ws.sheet_properties.tabColor = ACCENT
    last = ZP_ROWS + 1
    # paste area A:H - cream, unlocked
    for r in range(1, last + 1):
        for col in range(1, 9):
            c = ws.cell(r, col)
            c.fill = FILL_FIN
            c.protection = Protection(locked=False)
    for L, w in zip("ABCDEFGH", (22, 44, 30, 14, 18, 13, 16, 4)):
        ws.column_dimensions[L].width = w
        ws.column_dimensions[L].protection = Protection(locked=False)   # a whole-column paste still works
    # helper columns J:T (formulas)
    heads = {"J": "Stage (filled down)", "K": "Status", "L": "Closing Date", "M": "On Work Won? (1 = yes)",
             "N": "Running count", "O": "Work Won row", "P": "Last word of deal name", "Q": "Event day",
             "R": "Event month", "S": "Event year", "T": "Event date in deal name"}
    for col, t in heads.items():
        _hdr(ws[f"{col}1"], t, FILL_CALC_HDR)
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["J"].width = 24
    ws.column_dimensions["P"].width = 14
    ws.row_dimensions[1].height = 42
    fy0 = "(EOMONTH(set_FYFirstMonth,-1)+1)"
    fy1 = "EOMONTH(set_FYFirstMonth,11)"
    for r in range(2, last + 1):
        p = r - 1
        f = {
            "J": f'=IF(OR(zp_Hdr=0,ROW()<=zp_Hdr),"",IF(TRIM($A{r}&"")<>"",TRIM(IFERROR(LEFT($A{r},FIND("(",$A{r})-1),$A{r})),'
                 f'IF(ROW()-1<=zp_Hdr,"",J{p})))',
            "K": f'=IF(OR(J{r}="",TRIM($B{r}&"")=""),"",IF(J{r}="Closed Won","Won",IF(J{r}="Closed Lost","Lost","Open")))',
            "L": f'=IF(K{r}="","",IF(ISNUMBER($F{r}),$F{r},IFERROR(DATEVALUE($F{r}&""),"")))',
            "M": f'=IF(K{r}="",0,IF(K{r}="Open",1,IF(AND(ISNUMBER(L{r}),L{r}>={fy0},L{r}<={fy1}),1,0)))',
            "N": f'=N(N{p})+M{r}',
            "O": f'=IF(M{r}=1,N{r},"")',
            "P": f'=IF(K{r}="","",TRIM(RIGHT(SUBSTITUTE(TRIM($B{r}&"")," ",REPT(" ",99)),99)))',
            "Q": f'=IFERROR(VALUE(LEFT(P{r},FIND(".",P{r})-1)),"")',
            "R": f'=IFERROR(VALUE(MID(P{r},FIND(".",P{r})+1,FIND(".",P{r},FIND(".",P{r})+1)-FIND(".",P{r})-1)),"")',
            "S": f'=IFERROR(VALUE(MID(P{r},FIND(".",P{r},FIND(".",P{r})+1)+1,9)),"")',
            "T": (f'=IFERROR(IF(OR(K{r}="",Q{r}="",R{r}="",S{r}="",L{r}=""),"",IF(AND(R{r}>=1,R{r}<=12,Q{r}>=1,Q{r}<=31),'
                  f'IF(AND(DATE(IF(S{r}<100,2000+S{r},S{r}),R{r},Q{r})>=EOMONTH(L{r},-1)+1,'
                  f'DATE(IF(S{r}<100,2000+S{r},S{r}),R{r},Q{r})<=L{r}+548),DATE(IF(S{r}<100,2000+S{r},S{r}),R{r},Q{r}),""),"")),"")'),
        }
        for col, fx in f.items():
            c = ws[f"{col}{r}"]
            c.value = fx
            c.font = Font(name=FONT, size=9, color="595959")
            c.fill = FILL_CALC
            if col in ("L", "T"):
                c.number_format = "dd-mmm-yy"
    # control panel V:W
    panel = [
        ("ZOHO DUMP - HOW TO", None),
        ("1. In Zoho run the report 'All Deals by Stage' for the year to date plus all open deals, and export it to Excel.", None),
        ("2. Here: select A1:H3001 and press Delete to clear last month's dump.", None),
        ("3. In the export click cell A1, press Ctrl+Shift+End, then Ctrl+C. Come back here, click A1 and press Ctrl+V.", None),
        ("4. Check the three lines below read OK. Work Won, Month-End and the FY Summary forecast update by themselves.", None),
        ("5. Cost centres for Multi-Service deals and any invoice month you want to change go on the Won Overrides sheet.", None),
        ("", None),
        ("Header row found at row", "=zp_Hdr"),
        ("Deals in the paste", f'=COUNTIF($K$2:$K${last},"?*")'),
        ("Record Count Zoho printed", f'=IFERROR(VALUE(TRIM(MID(INDEX($A$1:$A$12,MATCH("*Record Count*",$A$1:$A$12,0)),FIND(":",INDEX($A$1:$A$12,MATCH("*Record Count*",$A$1:$A$12,0)))+1,20))),"")'),
        ("Check", f'=IF(zp_Hdr=0,"PASTE THE ZOHO EXPORT INTO A1",IF(W10="","OK - no record count to compare",IF(W9=W10,"OK - every deal read","CHECK - clear A1:H3001 and paste again")))'),
        ("Deals on Work Won (year to date + all open)", f'=MAX($N$2:$N${last})'),
        ("Work Won has room for", WON_LAST - 4),
        ("Room check", '=IF(W12<=W13,"OK","CHECK - more deals than Work Won rows")'),
    ]
    for i, (lab, fx) in enumerate(panel):
        r = 1 + i
        a = ws.cell(r, 22, lab)
        if i == 0:
            a.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
            a.fill = FILL_TITLE
            ws.cell(r, 23).fill = FILL_TITLE
        elif fx is None:
            a.font = Font(name=FONT, size=10, color="262626")
            a.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=22, end_row=r, end_column=23)
            ws.row_dimensions[r].height = 30 if lab else 8
        else:
            a.font = Font(name=FONT, size=10, bold=True, color="262626")
            w = ws.cell(r, 23, fx)
            w.font = Font(name=FONT, size=10, bold=True, color=NAVY)
            w.fill = FILL_TOT
            w.border = BORDER
    ws["Y1"] = f'=IFERROR(MATCH("Stage",$A$1:$A${last},0),0)'
    ws["Y1"].font = Font(name=FONT, size=8, color="FFFFFF")
    _name(wb, "zp_Hdr", "'Zoho Paste'!$Y$1")
    ws.column_dimensions["U"].width = 2
    ws.column_dimensions["V"].width = 46
    ws.column_dimensions["W"].width = 34
    ws.conditional_formatting.add("W11", FormulaRule(formula=['LEFT(W11,5)="CHECK"'], fill=PatternFill("solid", fgColor="F8D7DA"),
                                                     font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add("W14", FormulaRule(formula=['LEFT(W14,5)="CHECK"'], fill=PatternFill("solid", fgColor="F8D7DA"),
                                                     font=Font(color="9C0006", bold=True)))
    ws.freeze_panes = "A2"
    _protect(ws)
    return ws


# --------------------------------------------------------- Won Overrides
def _overrides(wb, pos, keep=None):
    if "Won Overrides" in wb.sheetnames:
        del wb["Won Overrides"]
    ws = wb.create_sheet("Won Overrides", pos)
    ws.sheet_properties.tabColor = ACCENT
    _title(ws, "Won Overrides",
           "Only for what Zoho cannot tell us. Keyed on the Opportunity Number, so it survives every new Zoho dump.", 6)
    heads = [("Opportunity Number", FILL_FIN_HDR, 18), ("Deal (from Zoho)", FILL_CALC_HDR, 50),
             ("Cost Centre", FILL_FIN_HDR, 16), ("Expected Invoice Month", FILL_FIN_HDR, 16),
             ("Notes", FILL_FIN_HDR, 50), ("Found in Zoho?", FILL_CALC_HDR, 12)]
    for i, (t, fill, w) in enumerate(heads):
        _hdr(ws.cell(4, 1 + i), t, fill)
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[4].height = 32
    last = 4 + OVR_ROWS
    for r in range(5, last + 1):
        for col in (1, 3, 4, 5):
            c = ws.cell(r, col)
            c.fill = FILL_FIN
            c.protection = Protection(locked=False)
            c.border = ROW_LINE
            c.font = Font(name=FONT, size=10)
        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 4).number_format = "mmm-yy"
        b = ws.cell(r, 2, f'=IF($A{r}="","",IFERROR(INDEX(\'Zoho Paste\'!$B:$B,MATCH($A{r}&"",\'Zoho Paste\'!$G:$G,0))&"",""))')
        f = ws.cell(r, 6, f'=IF($A{r}="","",IF($B{r}="","CHECK","OK"))')
        for c in (b, f):
            c.font = Font(name=FONT, size=10, color="262626")
            c.fill = FILL_CALC
            c.border = ROW_LINE
    for (col, lst) in (("C", "lst_CostCentre"), ("D", "lst_Months")):
        dv = DataValidation(type="list", formula1=lst, allow_blank=True)
        dv.add(f"{col}5:{col}{last}")
        ws.add_data_validation(dv)
    t = Table(displayName="tbl_WonOvr", ref=f"A4:F{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['F5="CHECK"'], fill=PatternFill("solid", fgColor="F8D7DA"),
                                                              font=Font(color="9C0006")))
    for i, row in enumerate(keep or []):
        for j, v in enumerate(row):
            ws.cell(5 + i, (1, 3, 4, 5)[j]).value = v
    ws.freeze_panes = "B5"
    _protect(ws)
    return ws


# ------------------------------------------------------------- Work Won
def _work_won(wb):
    ws = wb["Work Won"]
    ws["A2"] = ("Fills itself from the Zoho dump on the Zoho Paste sheet - nothing is typed here. Year-to-date deals "
                "(won, lost, open) plus every open deal. Change a cost centre or invoice month on Won Overrides.")
    head = {ws.cell(4, c).value: c for c in range(1, 19)}
    zp = "'Zoho Paste'!"
    ix = lambda col, r: f"INDEX({zp}${col}:${col},$S{r})"
    ovr = lambda fld, r: f"INDEX(tbl_WonOvr[{fld}],MATCH($E{r},tbl_WonOvr[Opportunity Number],0))"
    f = {
        "Month": lambda r: f'=IF($S{r}="","",IFERROR(EOMONTH({ix("L", r)},0),""))',
        "Source": lambda r: f'=IF($S{r}="","","ZOHO")',
        "Reference": lambda r: f'=IF($S{r}="","",{ix("G", r)}&"")',
        "Client": lambda r: f'=IF($S{r}="","",{ix("C", r)}&"")',
        "Job Number": lambda r: f'=IF($S{r}="","",{ix("G", r)}&"")',
        "Cost Centre": lambda r: (f'=IF($S{r}="","",IF(IFERROR({ovr("Cost Centre", r)}&"","")<>"",{ovr("Cost Centre", r)}&"",'
                                  f'IFERROR(INDEX(lst_ZohoCC,MATCH({ix("E", r)}&"",lst_ZohoDept,0))&"","")))'),
        "Description": lambda r: f'=IF($S{r}="","",{ix("B", r)}&"")',
        "Value Ex GST": lambda r: (f'=IF($S{r}="","",IF(TRIM({ix("D", r)}&"")="","",'
                                   f'IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE({ix("D", r)}&"","$",""),",","")),"")))'),
        "Status": lambda r: f'=IF($S{r}="","",{ix("K", r)})',
        "Date Won": lambda r: f'=IF($J{r}="Won",{ix("L", r)},"")',
        "Expected Invoice Month": lambda r: (f'=IF($S{r}="","",IFERROR(IF(N(IFERROR({ovr("Expected Invoice Month", r)},0))>0,'
                                             f'EOMONTH({ovr("Expected Invoice Month", r)},0),'
                                             f'EOMONTH(IF(ISNUMBER({ix("T", r)}),{ix("T", r)},{ix("L", r)}),0)),""))'),
        "Notes": lambda r: (f'=IF($S{r}="","",MID(IF(IFERROR({ovr("Notes", r)}&"","")<>"","; "&{ovr("Notes", r)},"")'
                            f'&IF(AND($G{r}="",{ix("E", r)}&""<>""),"; "&{ix("E", r)}&" - set the cost centre on Won Overrides","")'
                            f'&IF($I{r}="","; No amount in Zoho","")'
                            f'&IF(COUNTIF({zp}$G$2:$G${ZP_ROWS + 1},$E{r})>1,"; Opportunity number is on more than one Zoho deal","")'
                            f'&IF(AND($J{r}="Open",{ix("L", r)}<EOMONTH(TODAY(),-1)+1),"; Closing date has passed - update Zoho","")'
                            f'&IF(ISNUMBER({ix("T", r)}),IF(EOMONTH({ix("T", r)},0)<>EOMONTH({ix("L", r)},0),"; Invoice month taken from the event date in the deal name",""),""),3,400))'),
        "Pipeline Stage": lambda r: f'=IF($S{r}="","",{ix("J", r)})',
    }
    for h, fn in f.items():
        col = head[h]
        hc = ws.cell(4, col)
        hc.fill = FILL_CALC_HDR
        for r in range(5, WON_LAST + 1):
            c = ws.cell(r, col)
            c.value = fn(r)
            c.fill = FILL_CALC
            c.font = Font(name=FONT, size=10, color="262626")
            c.protection = Protection(locked=True)
            if h in ("Month", "Expected Invoice Month"):
                c.number_format = "mmm-yy"
            elif h == "Date Won":
                c.number_format = "dd-mmm-yy"
            elif h == "Value Ex GST":
                c.number_format = MONEY
    # hidden helper: which Zoho Paste row feeds this Work Won row
    _hdr(ws.cell(4, 19), "Zoho Row", FILL_CALC_HDR)
    for r in range(5, WON_LAST + 1):
        c = ws.cell(r, 19, f'=IFERROR(MATCH(ROW()-4,{zp}$O$1:$O${ZP_ROWS + 1},0),"")')
        c.font = Font(name=FONT, size=9, color="595959")
    ws.column_dimensions["S"].hidden = True
    ws.column_dimensions[chr(64 + head["Notes"])].width = 60
    ws.data_validations.dataValidation = []      # nothing to pick any more


def apply(wb, keep_overrides=None):
    _lists(wb)
    pos = wb.sheetnames.index("Work Won") + 1
    _zoho_paste(wb, pos)
    _overrides(wb, pos + 1, keep_overrides)
    _work_won(wb)
