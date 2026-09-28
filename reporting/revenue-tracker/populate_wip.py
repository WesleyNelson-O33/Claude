"""Loads the WIP Schedule (GL 11300) into the FY27 tracker.

* Straight-line licence / subscription deferrals (revenue and cost) go on the Deferrals sheet,
  so the tracker works out every month's release itself.
* Everything else (accruals, accrued costs, package deferrals, pre-production work) goes on
  WIP Movements: an opening balance at 30 Jun 2026, then the Jul / Aug / Sep movements.
* September movements that repeat an August reversal on a line already cleared in August are
  NOT loaded - they are copy errors in the schedule (listed in the log).

Usage: python populate_wip.py <tracker.xlsx> <wip_schedule.xlsx> <out.xlsx>
"""
import datetime as dt
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL

TRK, SCH, OUT = sys.argv[1:4]
log = []
wb = openpyxl.load_workbook(TRK)
sv = openpyxl.load_workbook(SCH, data_only=True)["Schedule Updated"]
num = lambda v: v if isinstance(v, (int, float)) else 0.0


def eom(d):
    n = dt.datetime(d.year + d.month // 12, d.month % 12 + 1, 1)
    return n - dt.timedelta(days=1)


def add_months(d, k):
    y, m = divmod(d.year * 12 + d.month - 1 + k, 12)
    return eom(dt.datetime(y, m + 1, 1))


# movement columns: every dated header in row 5 (I, K, M ... BY)
MV = [(c, eom(sv.cell(5, c).value)) for c in range(CI("I"), CI("BZ") + 1) if isinstance(sv.cell(5, c).value, dt.datetime)]
MCOL = {m: c for c, m in MV}
JUN, JUL, AUG, SEP = (dt.datetime(2026, 6, 30), dt.datetime(2026, 7, 31), dt.datetime(2026, 8, 31), dt.datetime(2026, 9, 30))
BAL = {JUN: "BT", JUL: "BV", AUG: "BX", SEP: "BZ"}
MOVE = {JUL: "BU", AUG: "BW", SEP: "BY"}

DEPT = {"PRODUCTION": "PRODUCTION", "PRD": "PRODUCTION", "VIDEO": "VIDEO", "VID": "VIDEO", "INTEGRATION": "INTEGRATION",
        "INT": "INTEGRATION", "WORKTECH": "INTEGRATION", "CONSULTING": "CONSULTING", "ONSITE": "ONSITE", "ONS": "ONSITE",
        "CTS": "CTS"}


def gl(v):
    if v is None:
        return None
    return str(v).replace("-", "").strip()


# explicit terms where the schedule has no release history yet (or the wrong one) - from the line descriptions
EXPLICIT = {  # schedule row: (invoice month, start, end)
    404: (JUL, dt.datetime(2027, 3, 31), dt.datetime(2028, 2, 29)),
    405: (JUL, dt.datetime(2027, 3, 31), dt.datetime(2028, 2, 29)),
    471: (SEP, SEP, dt.datetime(2027, 8, 31)),
    472: (SEP, SEP, dt.datetime(2027, 8, 31)),
    # Appspace: licence runs Aug-26 to Jul-27 but was invoiced in Sep. Released Sep-26 to Jul-27 so the
    # closed month of August still ties to the schedule; nothing is lost - it is fully released by Jul-27.
    473: (SEP, SEP, dt.datetime(2027, 7, 31)),
    474: (SEP, SEP, dt.datetime(2027, 7, 31)),
}

lines = []
for r in range(6, sv.max_row + 1):
    job = sv[f"A{r}"].value
    if job in (None, "") and not sv[f"D{r}"].value:
        continue
    vals = {m: num(sv[f"{c}{r}"].value) for m, c in BAL.items()}
    mv = {m: num(sv[f"{c}{r}"].value) for m, c in MOVE.items()}
    if all(abs(v) < 0.005 for v in list(vals.values()) + list(mv.values())):
        continue
    desc = str(sv[f"D{r}"].value or "").strip()
    cmt = str(sv[f"CB{r}"].value or "").strip()
    g = gl(sv[f"CC{r}"].value)
    text = (desc + " " + cmt).lower()
    if g and g[0] == "4":
        kind = "Revenue"
    elif g and g[0] in "56":
        kind = "Cost"
    else:
        kind = "Cost" if "expense" in text else "Revenue"
    job = re.sub(r"\s+", " ", str(job or "")).strip().split(" ")[0]
    lines.append(dict(row=r, job=job, client=str(sv[f"B{r}"].value or "").strip(), dept=str(sv[f"C{r}"].value or "").strip(),
                      desc=desc, cmt=cmt, gl=g, kind=kind, E=num(sv[f"E{r}"].value), bal=vals, mv=mv,
                      hist={m: num(sv.cell(r, c).value) for c, m in MV}))


def is_deferral(ln):
    t = (ln["desc"] + " " + ln["cmt"]).lower()
    if ln["row"] in EXPLICIT:
        return True
    return ("recognise" in t and "labour hours" not in t) and (ln["gl"] or "")[:3] in ("411", "515")


# ------------------------------------------------------------ tracker sheets
fi = wb["Finance"]
FH = {c.value: c.column_letter for c in fi[6]}
fin_inv = {}
for r in range(7, fi.max_row + 1):
    no = fi[f"{FH['Xero Invoice No']}{r}"].value
    if no:
        fin_inv[str(no).strip()] = fi[f"{FH['Xero Invoiced Ex GST']}{r}"].value
df = wb["Deferrals"]
DH = {c.value: c.column_letter for c in df[6]}
wp = wb["WIP Movements"]
WH = {c.value: c.column_letter for c in wp[4]}
# clear what was loaded from v2 before
for r in range(5, wp.max_row + 1):
    for c in wp[r]:
        if not c.protection.locked and c.value is not None and not (isinstance(c.value, str) and c.value.startswith("=")):
            c.value = None


def inv_no(ln):
    m = re.search(r"INV[\s-]?(\d{4,6})", ln["desc"], re.I)
    if m:
        return f"INV-{m.group(1)}"
    m = re.search(r"Eptura[^0-9]*(?:INV-)?(\d{4,6})", ln["desc"], re.I) or re.search(r"(?:Bill|INV)[\s#-]*([0-9]{4,9})", ln["desc"] + " " + ln["cmt"], re.I)
    if m:
        return f"BILL-{m.group(1)}"
    return f"WIP-{ln['row']}"


def deferral_terms(ln):
    if ln["row"] in EXPLICIT:
        return EXPLICIT[ln["row"]]
    rel_sign = 1 if ln["kind"] == "Revenue" else -1
    hist = [(m, v) for m, v in sorted(ln["hist"].items()) if abs(v) > 0.004]
    rel = [round(abs(v), 6) for m, v in hist if v * rel_sign > 0]
    ini = [(m, v) for m, v in hist if v * rel_sign < 0]
    if not rel or not ini:
        return None
    r_amt = Counter(rel).most_common(1)[0][0]
    m0, v0 = ini[0]
    E = abs(ln["E"]) or abs(v0)
    n = round(E / r_amt)
    first_rel = next(m for m, v in hist if v * rel_sign > 0)
    start = first_rel if abs(abs(v0) - E) < 0.05 else m0
    ln["per"] = r_amt
    if abs(abs(v0) - E) < 0.05:
        ln["deferred"] = round(abs(v0), 6)      # the amount actually put on the schedule (unrounded)
    return (m0, start, add_months(start, n - 1))


def sim_deferral(amount, inv_m, start, end, kind, per=None):
    """Debit balance of GL 11300 for this deferral at each month end (the tracker's own maths)."""
    months = (end.year - start.year) * 12 + end.month - start.month + 1
    per = per or round(amount / months, 2)
    out, cum = {}, 0.0
    m = min(inv_m, start)
    while m <= max(end, SEP):
        rec = 0.0
        if start <= m <= end:
            rec = amount - per * (months - 1) if m == end else per
        cum += rec - (amount if m == inv_m else 0)
        out[m] = cum if kind == "Revenue" else -cum
        m = add_months(m, 1)
    return lambda d: round(out.get(d, 0.0 if d < min(inv_m, start) else cum if kind == "Revenue" else -cum), 2)


drow, wrow = 7, 5
seen = {}
pending_adj = []
mismatch, dups, def_count, wip_count = [], [], 0, 0
for ln in lines:
    cc = DEPT.get(ln["dept"].upper())
    if is_deferral(ln):
        terms = deferral_terms(ln)
        if not terms:
            log.append(f"Row {ln['row']} {ln['job']}: deferral terms not found - loaded to WIP Movements instead")
        else:
            inv_m, start, end = terms
            amount = ln.get("deferred") or round(abs(ln["E"]), 6)
            no = inv_no(ln)
            key = (no, ln["job"], ln["kind"])
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 1:
                no = f"{no} ({seen[key]})"
            on_fin = ln["kind"] == "Revenue" and no in fin_inv and abs(num(fin_inv[no]) - amount) < 0.01
            put = {"Type": ln["kind"], "Invoice or Bill No": no, "Invoice or Bill Date": inv_m,
                   "Defer Start": start, "Defer End": end,
                   "Notes": f"From WIP Schedule row {ln['row']}: {ln['desc'][:120]}"}
            if not on_fin:
                put.update({"Job Number Override": ln["job"], "Amount Override": amount,
                            "P&L GL Override": ln["gl"] or ("41175" if ln["kind"] == "Revenue" else "51500"),
                            "Cost Centre Override": cc})
            for h, v in put.items():
                if v not in (None, ""):
                    df[f"{DH[h]}{drow}"] = v
            # does the tracker's maths reproduce the schedule?
            if ln.get("per"):      # the schedule's own monthly amount - Excel's rounding can differ by a cent
                df[f"{DH['Per Month Override']}{drow}"] = ln["per"]
            f = sim_deferral(amount, inv_m, start, end, ln["kind"], df[f"{DH['Per Month Override']}{drow}"].value)
            gap = round(ln["bal"][JUN] - f(JUN), 6)
            if 0.005 <= abs(gap) <= 0.05 and all(abs(ln["bal"][m] - f(m) - gap) < 0.000001 for m in (JUL, AUG)):
                pending_adj.append((ln, gap, cc))
                f0 = f
                f = lambda d, f0=f0, g=gap: round(f0(d) + g, 2)
            for m in (JUN, JUL, AUG, SEP):
                if abs(f(m) - ln["bal"][m]) > 0.004:
                    mismatch.append((ln["row"], ln["job"], ln["kind"], m.strftime("%b-%y"), f(m), round(ln["bal"][m], 2)))
            drow += 1
            def_count += 1
            continue
    # ---------------- WIP Movements: opening balance + FY27 movements
    entries = []
    if abs(ln["bal"][JUN]) > 0.004:
        entries.append((JUN, round(ln["bal"][JUN], 6), "Migrated opening balance"))
    run = ln["bal"][JUN]
    for m in (JUL, AUG, SEP):
        v = round(ln["mv"][m], 6)
        if abs(v) < 0.005:
            continue
        if m == SEP and abs(ln["bal"][AUG]) < 0.005 and abs(v - round(ln["mv"][AUG], 2)) < 0.005:
            dups.append((ln["row"], ln["job"], ln["desc"][:70], v))
            continue
        reversing = run * v < 0 and abs(v) <= abs(run) + 0.01
        if ln["kind"] == "Revenue":
            t = ("Release of deferral" if run < 0 else "Reversal of prior accrual") if reversing else \
                ("Accrual - unbilled work" if v > 0 else "Deferral - invoiced in advance")
        else:
            t = ("Reversal of prior accrual" if run < 0 else "Release of deferred cost") if reversing else \
                ("Deferred cost - paid in advance" if v > 0 else "Accrued cost - incurred not billed")
        entries.append((m, v, t))
        run += v
    for m, v, t in entries:
        put = {"Month": m, "Job Number": ln["job"], "Cost Centre": cc, "Description": ln["desc"][:250], "Type": t,
               "Revenue or Cost": ln["kind"], "Amount": v, "P&L Account": ln["gl"], "Notes": ln["cmt"][:250]}
        for h, val in put.items():
            if val not in (None, ""):
                wp[f"{WH[h]}{wrow}"] = val
        wrow += 1
        wip_count += 1

for ln, gap, cc in pending_adj:
    put = {"Month": JUN, "Job Number": ln["job"], "Cost Centre": cc, "Type": "Adjustment / correction",
           "Description": f"Rounding adjustment per WIP Schedule row {ln['row']} (one month released {abs(gap):.2f} differently)",
           "Revenue or Cost": ln["kind"], "Amount": gap, "P&L Account": ln["gl"],
           "Notes": "Keeps the tracker equal to the WIP Schedule to the cent"}
    for h, val in put.items():
        if val not in (None, ""):
            wp[f"{WH[h]}{wrow}"] = val
    wrow += 1
    wip_count += 1
    log.append(f"Rounding adjustment {gap:+.6f} added for schedule row {ln['row']} job {ln['job']}")
log.append(f"Deferrals sheet: {def_count} lines. WIP Movements: {wip_count} rows.")
log.append("")
log.append("SEPTEMBER LINES NOT LOADED - the schedule repeats an August reversal on a line already cleared in August:")
for d in dups:
    log.append(f"  row {d[0]}  job {d[1]}  {d[2]}  Sep {d[3]:,.2f}")
log.append(f"  Total not loaded: {sum(d[3] for d in dups):,.2f}")
log.append("")
log.append("DEFERRAL LINES WHERE THE SCHEDULE DIFFERS FROM A STRAIGHT-LINE RELEASE (tracker vs schedule balance):")
for m in mismatch:
    log.append(f"  row {m[0]}  job {m[1]}  {m[2]}  {m[3]}: tracker {m[4]:,.2f}  schedule {m[5]:,.2f}  diff {m[4] - m[5]:,.2f}")
wb.save(OUT)
open(OUT + ".log.txt", "w").write("\n".join(log))
print(f"saved {OUT} | deferrals {def_count} | wip rows {wip_count} | sep dups {len(dups)} | mismatches {len(mismatch)}")
