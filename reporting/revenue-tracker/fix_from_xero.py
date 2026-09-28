"""Applies the Xero check corrections to the Finance sheet (finance input cells only) and adds
any job numbers from Xero's 'Job Numbers' tracking category that are missing from Lists.

Usage: python fix_from_xero.py <tracker.xlsx> <xero_invoices.json> <xero_tracking.json> <out.xlsx> <log.txt>
"""
import datetime as dt
import json
import re
import sys

import openpyxl

TR, XJ, TJ, OUT, LOG = sys.argv[1:6]
wb = openpyxl.load_workbook(TR)
fin = wb["Finance"]
H = {c.value: c.column for c in fin[6] if c.value}
INV, DATE, AMT, NOTE = (H[k] for k in ("Xero Invoice No", "Xero Invoice Date", "Xero Invoiced Ex GST", "Finance Notes"))
row_of = {fin.cell(r, 1).value: r for r in range(7, fin.max_row + 1) if fin.cell(r, 1).value}
xero = {i["invoice_number"]: i for i in json.load(open(XJ))}
log = []


def d(s):
    return dt.datetime.strptime(s, "%Y-%m-%d")


def note(r, text):
    old = fin.cell(r, NOTE).value
    fin.cell(r, NOTE).value = f"{old}; {text}" if old else text


def setrow(rid, why, inv=..., date=..., amt=...):
    r = row_of[rid]
    for col, v, lab in ((INV, inv, "Invoice No"), (DATE, date, "Date"), (AMT, amt, "Ex GST")):
        if v is ...:
            continue
        old = fin.cell(r, col).value
        fin.cell(r, col).value = v
        log.append(f"{rid}  {lab}: {old!r} -> {v!r}")
    note(r, "Xero check 28-Sep-26: " + why)
    log.append(f"{rid}  why: {why}\n")


X = lambda n: xero[n]
# ---- wrong invoice numbers / double counts
setrow("ONS-0032", "INV-10582 is not in Xero. CBA Melbourne August $11,960 is INV-10615", inv="INV-10615")
setrow("ONS-0033", "INV-10615 is the CBA Melbourne invoice. Aware Super August $11,000 is INV-10583",
       inv="INV-10583", date=d(X("INV-10583")["invoice_date"]))
setrow("ONS-0023", "INV-10598 is job 2609401 in Xero and is already on ONS-0050 - it was counted twice ($8,500). "
       "No Xero invoice found for 26073102 Jul-Sep", inv=None, date=None, amt=None)
setrow("ONS-0063", "INV-10645 in Xero is Bankwest BCEC WGEA Report Launch (job 26092101) $950, not this PwC job. "
       "No approved PwC invoice for 26073101 yet", inv=None)
# ---- amounts to Xero
setrow("PRD-0016", "Xero INV-10559 is $8,882.50 ex GST; the $332.50 loading reversal is the separate credit note row PRD-0017, "
       "so $8,550 here counted the credit twice", amt=8882.50)
setrow("CON-0024", "Xero INV-10596 is $205,352.17 ex GST", amt=205352.17)
setrow("ONS-0011", "Xero INV-10536 is $11,860.56 ex GST", amt=11860.56)
setrow("PRD-0049", "Xero INV-10523 $39.15 ex GST", amt=39.15, date=d(X("INV-10523")["invoice_date"]))
setrow("PRD-0050", "Xero INV-10621 $60.58 ex GST dated 31-Aug-26", amt=60.58, date=d(X("INV-10621")["invoice_date"]))
for rid in ("PRD-0042", "PRD-0082"):
    setrow(rid, "invoice number typed as in Xero", inv="INV-10567")
# ---- September production invoices raised in Xero but not typed on Finance
sept = [("PRD-0051", "INV-10624", 2674), ("PRD-0084", "INV-10624", 1500), ("PRD-0052", "INV-10623", 6260),
        ("PRD-0085", "INV-10623", 1508), ("PRD-0053", "INV-10626", 495), ("PRD-0054", "INV-10625", 865),
        ("PRD-0055", "INV-10627", 2232), ("PRD-0086", "INV-10627", 1029), ("PRD-0056", "INV-10632", 7527.25),
        ("PRD-0087", "INV-10632", 1578.10), ("PRD-0057", "INV-10628", 570), ("PRD-0058", "INV-10629", 380),
        ("PRD-0059", "INV-10634", 380), ("PRD-0060", "INV-10633", 4665), ("PRD-0061", "INV-10630", 946.25)]
extra = {"PRD-0056": "INV-10632 total $9,105.35 is $910.54 under the department's $10,015.89 (video line kept at $1,578.10)",
         "PRD-0060": "Xero INV-10633 (ISE) $4,665 vs department $4,729",
         "PRD-0061": "Xero INV-10630 (ICAS) is on job 26082502, not 26082503 - department to check the job number"}
for rid, inv, amt in sept:
    setrow(rid, "invoice found in Xero" + (" - " + extra[rid] if rid in extra else ""),
           inv=inv, date=d(X(inv)["invoice_date"]), amt=amt)

# ---- every other matched invoice: date to the Xero invoice date
skip = {"PRD-0038", "PRD-0079", "PRD-0037"}          # nil package rows, dated on purpose
for rid, r in row_of.items():
    raw = fin.cell(r, INV).value
    if not raw or rid in skip or not str(raw).startswith("INV"):
        continue
    nums = re.findall(r"\d{4,6}", str(raw))
    cands = [x for k, x in xero.items() if re.sub(r"\D", "", k) in nums]
    if not cands:
        continue
    dates = {x["invoice_date"] for x in cands}
    if len(dates) != 1:
        continue
    xd = d(dates.pop())
    cur = fin.cell(r, DATE).value
    if cur is not None and cur != xd:
        fin.cell(r, DATE).value = xd
        log.append(f"{rid}  Date: {cur:%d-%b-%y} -> {xd:%d-%b-%y} (Xero {raw})")

# ---- job list: add Xero tracking job numbers missing from Lists
CC = {"ONS": "ONSITE", "PRD": "PRODUCTION", "VID": "VIDEO", "INT": "INTEGRATION", "CONS": "CONSULTING", "CON": "CONSULTING", "CTS": "CTS"}
lst = wb["Lists"]
have = {str(lst.cell(r, 16).value) for r in range(5, 1501) if lst.cell(r, 16).value}
nxt = max(r for r in range(5, 1501) if lst.cell(r, 16).value) + 1
opts = next(t for t in json.load(open(TJ))["tracking_categories"] if t["name"] == "Job Numbers")["options"]
added = []
for o in opts:
    m = re.match(r"\s*(\S+)\s*-\s*(.*)$", o["name"])
    if not m or m.group(1) in have:
        continue
    job, name = m.group(1), m.group(2).strip()
    tag = re.search(r"\[\s*(ONS|PRD|VID|INT|CONS|CON|CTS)\b", name) or re.search(r"\b(ONS|PRD|VID|INT|CONS|CTS)\b", name)
    cc = CC.get(tag.group(1)) if tag else None
    for c, v in ((16, job), (17, re.sub(r"\s*\[[^\]]*\]\s*$", "", name)), (18, cc)):
        lst.cell(nxt, c).value = v
        lst.cell(nxt, c).number_format = "@"
    added.append(f"{job}  {name}  -> {cc}")
    have.add(job)
    nxt += 1
log.append(f"\nLists: added {len(added)} job numbers from Xero tracking 'Job Numbers' (rows up to {nxt - 1})")
log += ["  " + a for a in added]
wb.save(OUT)
open(LOG, "w").write("\n".join(log) + "\n")
print(len(added), "jobs added;", sum(1 for l in log if "->" in l), "cell changes")
