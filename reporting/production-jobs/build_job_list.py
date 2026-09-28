"""Clean workbook from the production job list pasted into chat (FY27 Jul-Sep 2026).
Each data row is followed by two empty rows. Net Total, Discount %, Margin and Margin %
are live formulas; the pasted figures are checked against them.

Money columns: Video Total, Discounts Included, Cross Hire Expense, Labour Expense (Internal).
Blank cells in the paste were placed by their spacing; every row was then checked with
Net = Invoice - Discounts and Margin = Invoice - Cross Hire - Labour.
"""
import datetime as dt
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1]
D = lambda s: dt.datetime.strptime(s, "%d-%b-%y") if s else None
_ = None
# #, date, dept, client, job, project, event date, posted, inv no, inv value, email, zoho, closed, current no,
# hours (prod labour, PM, video filming, video editing, video PM), video total, discounts, cross hire, labour,
# pasted: net, disc%, margin, margin%
ROWS = [
    (2, "07-Jul-26", "Production", "RSNSW", "26032301", "Ordinary General Meeting | 1.7.26", "01-Jul-26", "Y", "INV-10519", 3071.00, "lindsay.botten@outlook.com", "26032301", "Y", 5924, (10, _, 4, _, _), (1029.00, 334.00, _, 977.23), (2737.00, .122, 2093.77, .682)),
    (3, "07-Jul-26", "Production", "PWC", "26062201", "Tax and Legal Webcast 2.7.26", "02-Jul-26", "Y", "INV-10515", 1925.00, "libby.kovac@au.pwc.com", "26062201", "Y", 6076, (8, 1, 8, _, _), (970.00, _, _, 821.56), (1925.00, 0, 1103.44, .573)),
    (6, "13-Jul-26", "Video", "Litecard", "2607701", "Litecard - Event Photographer - 16.7.26", "16-Jul-26", "Y", "INV-10522", 550.00, "brian@litecard.com.au", "2607701", "Y", 6102, (_, _, 4, _, _), (550.00, _, _, 312.18), (550.00, 0, 237.82, .432)),
    (7, "31-Jul-26", "Production", "Microsoft", "26052703", "Microsoft Town Hall 21.7.26", "21-Jul-26", "Y", "INV-10554", 12587.00, "helenkingham@microsoft.com", "26052703", "Y", 6034, (44, _, 4, _, _), (1702.00, 580.00, 631.50, 3148.70), (12007.00, .048, 8806.80, .700)),
    (9, "31-Jul-26", "Production", "Arup", "2605403", "Oves Awards Night 23.7.26", "23-Jul-26", "Y", "INV-10540", 5636.00, "Timothy.Hicks@arup.com", "2605403", "Y", 5956, (16, _, _, _, _), (176.00, 839.00, _, 1375.96), (4797.00, .175, 4260.04, .756)),
    (18, "13-Aug-26", "Production", "RSNSW", "2512804", "Ordinary General Meeting 5.8.26", "05-Aug-26", "Y", "INV-10561", 3261.00, "lindsay.botten@outlook.com", "2512804", "Y", 5826, (12, _, 4, _, _), (1029.00, 334.00, _, _), (2927.00, .114, 3261.00, 1.0)),
    (20, "13-Aug-26", "Production", "PwC", "26072902", "AI Townhall 6.8.26 (PO2601703328)", "06-Aug-26", "Y", "INV-10562", 1650.00, "libby.kovac@au.pwc.com", "26072902", "Y", 6128, (10, 1, 4, _, _), (485.00, _, 1092.36, _), (1650.00, 0, 557.64, .338)),
    (22, "21-Aug-26", "Production", "CEO Institute", "26063002", "CEO Connect 11.8.26", "11-Aug-26", "Y", "INV-10581", 8260.35, "taniaa@ceoinstitute.com", "26063002", "Y", 6084, (37, 2, 7, _, _), (1977.35, 2104.40, _, _), (6155.95, .342, 8260.35, 1.0)),
    (23, "25-Aug-26", "Production", "PwC", "26072903", "GB Webcast 11.8.26", "11-Aug-26", "Y", "INV-10577", 4300.00, "libby.kovac@au.pwc.com", "26072903", "Y", 6134, (9, _, 9, _, _), (1065.00, _, 1092.36, _), (4300.00, 0, 3207.64, .746)),
    (24, "25-Aug-26", "Production", "MUFG", "26061104", "MUFG - Suncorp FYR", "12-Aug-26", "Y", "INV-10584, INV-10585", 28652.25, "events@cm.mpms.mufg.com", "26061104", "Y", "Proj: 0016", (104, 12, 8, _, _), (3017.00, 5271.50, 980.16, _), (23380.75, .225, 27672.09, .966)),
    (25, "13-Aug-26", "Production", "CBA", "26061105", "CBA FYR 2026 (Order PO3308779)", "12-Aug-26", "Y", "INV-10563", 12170.00, "Rebecca.Rodger@cba.com.au", "26061105", "Y", 6093, (16, _, _, _, _), (850.00, 1680.00, 3111.20, _), (10490.00, .160, 9058.80, .744)),
    (26, "21-Aug-26", "Production", "AGL", "2605401", "AGL FYR - Investor Briefing", "12-Aug-26", "Y", "INV-10579", 14271.50, "KDraisey@agl.com.au", "2605401", "Y", "Proj: 0019", (18, _, _, 8, _), (1806.00, 2351.50, 3520.25, _), (11920.00, .197, 10751.25, .753)),
    (28, "21-Aug-26", "Production", "ASX", "2608604", "ASX Employee Townhall 13.8.26", "13-Aug-26", "Y", "INV-10599", 1510.00, "Aimee.Kuipers@asx.com.au", "2608604", "Y", 6138, (6, 2, 4, _, _), (485.00, 25.00, _, _), (1485.00, .017, 1510.00, 1.0)),
    (29, "21-Aug-26", "Production", "ASX", "26052702", "ASX FYR Briefing 13.8.26", "13-Aug-26", "Y", "INV-10576", 37579.00, "simon.starr@asx.com.au", "26052702", "Y", 6045, (78, 6, _, _, _), (3357.00, 2181.00, 5005.78, _), (35398.00, .062, 32573.22, .867)),
    (30, "27-Aug-26", "Video", "ASX", "26072906", "ASX Advisory Group on Corporate Governance", "18-Aug-26", "Y", "INV-10600", 4601.70, "vicky.duan@asx.com.au", "26072906", "Y", 6130, (8, _, 8, 8, _), (4221.70, 148.80, _, _), (4452.90, .033, 4601.70, 1.0)),
    (34, "25-Aug-26", "Production", "TLC - Melbourne", "26061106", "TLC Full Year Results", "19-Aug-26", "Y", "INV-10591", 13661.00, "sharne.flanagan@thelotterycorporation.com", "26061106", "Y", 5853, (25, 2, _, 4, _), (674.00, 611.50, 4272.26, _), (13049.50, .047, 9388.74, .687)),
    (35, "25-Aug-26", "Production", "Downer", "26061107", "Downer FYR26 20.8.26", "20-Aug-26", "Y", "INV-10587", 11644.25, "mitchell.dale@downergroup.com", "26061107", "Y", 6092, (18, 3, _, _, _), (263.50, 885.75, 1221.04, _), (10758.50, .082, 10423.21, .895)),
    (_, "", "Production", "APA", "26061108", "APA FYR Town Hall 20.8.26", "20-Aug-26", "Y", "Package", 12600.60, "steph.bennett@apa.com.au", "26061108", "Y", 6073, (45, 6, 5, _, _), (1261.60, 1651.40, _, _), (10949.20, .151, 12600.60, 1.0)),
    (39, "31-Aug-26", "Production", "Ventia", "26061109", "Ventia HYR Employee Town Hall 24.8.26 (PO 4701366307)", "24-Aug-26", "Y", "INV-10593", 13401.50, "amy.johnson@ventia.com", "26061109", "Y", 6095, (41, _, 4, _, _), (1771.00, 524.50, 1350.00, _), (12877.00, .041, 12051.50, .899)),
    (41, "31-Aug-26", "Production", "Downer - Melbourne", "26073111", "Employee Results Briefing", "25-Aug-26", "Y", "INV-10611", 24413.75, "mitchell.dale@downergroup.com", "26073111", "Y", 6144, (37, 3, 5, _, _), (1785.00, 2281.25, 14754.30, _), (22132.50, .103, 9659.45, .396)),
    (42, "19-Aug-26", "Production", "Symal - Melbourne", "2606904", "Symal FYR 2026", "25-Aug-26", "Y", "INV10567", 13347.00, "pia.witt@symal.com.au", "2606904", "Y", 6044, (22, 2, _, _, _), (86.50, 3563.00, 1874.00, _), (9784.00, .364, 11473.00, .860)),
    (48, "15-Jul-26", "Production", "APA", "26061609/26061608/26061607/26061108", "APA - All Hands Package FY27 (PO-00030766)", "N/A", "Y", "INV-10521", 54051.40, _, "26061609", "Y", "Proj: 0017", (_, _, _, _, _), (4308.00, _, _, _), (54051.40, 0, 54051.40, 1.0)),
    (_, "", "Production", "ASX", "26061602", "CEO Connect (Extended Results)", "01-Sep-26", "Y", _, 4174.00, "oliver.mariano@asx.com.au", "26061602", "Y", 5849, (16, _, _, 9, _), (1500.00, 261.00, _, _), (3913.00, .067, 4174.00, 1.0)),
    (_, "", "Production", "ASX", "26071701", "ASX Live 1.9.26", "01-Sep-26", "Y", _, 7768.00, "Aimee.Kuipers@asx.com.au", "26071701", "Y", 6143, (31, _, 4, _, _), (1508.00, 180.00, _, _), (7588.00, .024, 7768.00, 1.0)),
    (_, "", "Production", "RSNSW", "26032302", "Ordinary General Meeting 2.9.26", "02-Sep-26", "Y", _, 3261.00, "lindsay.botten@outlook.com", "26032302", "Y", 5925, (12, _, 4, _, _), (1029.00, 334.00, _, _), (2927.00, .114, 3261.00, 1.0)),
    (_, "", "Production", "Ventia", "26072102", "RUOK? Day 10.9.26", "10-Sep-26", "Y", _, 10015.89, "amy.johnson@ventia.com", "26072102", "Y", 6122, (28, _, 4, _, _), (1578.10, 2005.65, _, _), (8010.24, .250, 10015.89, 1.0)),
]
HEADERS = ["#", "Date", "Department", "Client", "Job Number", "Project Name", "Event Date", "Invoice Posted to Xero",
           "Invoice Number", "Invoice Value", "Client Email", "Zoho Number", "Closed", "Current Number",
           "Production Labour Hours", "Project Management", "Video Filming", "Video Editing", "Video Project Management",
           "Video Total", "Discounts Included", "Cross Hire Expense", "Labour Expense (Internal)", "Net Total",
           "Discounts as % of Net Total", "Margin", "Margin %"]
WIDTHS = [5, 11, 12, 20, 16, 44, 11, 10, 20, 14, 36, 14, 8, 12, 11, 11, 10, 10, 11, 13, 13, 13, 14, 14, 11, 14, 9]
MONEY = '$#,##0.00;[Red]-$#,##0.00;"-"'
NAVY, TINT, LINE = "1F3864", "D9E2F3", "C9D3E3"
F = lambda **k: Font(name="Arial", size=10, **k)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Production Jobs"
ws.sheet_view.showGridLines = False
ncol = len(HEADERS)
for c in range(1, ncol + 1):
    ws.cell(1, c).fill = PatternFill("solid", fgColor=NAVY)
ws["A1"] = "Production & Video Jobs  -  July to September 2026"
ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 32
ws["A2"] = ("Net Total = Invoice Value - Discounts.  Discounts % = Discounts / Net Total.  Margin = Invoice Value - "
            "Cross Hire - Labour.  Margin % = Margin / Invoice Value.  (These four are formulas.)")
ws["A2"].font = F(italic=True, color="595959")
HR = 4
for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
    c = ws.cell(HR, i, h)
    c.font = F(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.column_dimensions[CL(i)].width = w
ws.row_dimensions[HR].height = 42
L = {h: CL(i) for i, h in enumerate(HEADERS, 1)}
mismatches = []
r = HR + 1
first = r
for row in ROWS:
    (n, date, dept, client, job, proj, ev, posted, inv, val, email, zoho, closed, cur, hrs, money, pasted) = row
    vals = [n, D(date), dept, client, job, proj, D(ev) if ev not in ("N/A", "") else ev, posted, inv, val, email, zoho,
            closed, cur, *hrs, *money]
    for i, v in enumerate(vals, 1):
        if v is not None:
            ws.cell(r, i, v)
    iv, dc, ch, lb = L["Invoice Value"], L["Discounts Included"], L["Cross Hire Expense"], L["Labour Expense (Internal)"]
    ws[f"{L['Net Total']}{r}"] = f"={iv}{r}-N({dc}{r})"
    ws[f"{L['Discounts as % of Net Total']}{r}"] = f'=IFERROR(N({dc}{r})/{L["Net Total"]}{r},0)'
    ws[f"{L['Margin']}{r}"] = f"={iv}{r}-N({ch}{r})-N({lb}{r})"
    ws[f"{L['Margin %']}{r}"] = f'=IFERROR({L["Margin"]}{r}/{iv}{r},0)'
    # check against the pasted figures
    net = val - (money[1] or 0)
    mar = val - (money[2] or 0) - (money[3] or 0)
    for lab, calc, pv in [("Net Total", net, pasted[0]), ("Margin", mar, pasted[2]),
                          ("Disc %", round((money[1] or 0) / net, 3), pasted[1]), ("Margin %", round(mar / val, 3), pasted[3])]:
        if abs(calc - pv) > (0.0015 if "%" in lab else 0.005):
            mismatches.append((n, job, lab, calc, pv))
    for i in range(1, ncol + 1):
        c = ws.cell(r, i)
        h = HEADERS[i - 1]
        c.font = F(bold=(h in ("Job Number", "Invoice Value", "Margin")))
        c.border = Border(bottom=Side(style="thin", color=LINE), top=Side(style="thin", color=LINE))
        c.fill = PatternFill("solid", fgColor="FFFFFF" if (r - first) % 6 == 0 else "F4F7FB")
        c.alignment = Alignment(vertical="center", wrap_text=h in ("Project Name", "Job Number", "Invoice Number"))
        if h in ("Date", "Event Date"):
            c.number_format = "dd-mmm-yy"
        elif h in ("Invoice Value", "Video Total", "Discounts Included", "Cross Hire Expense",
                   "Labour Expense (Internal)", "Net Total", "Margin"):
            c.number_format = MONEY
        elif "%" in h:
            c.number_format = "0.0%"
        elif h in ("Job Number", "Zoho Number", "Invoice Number"):
            c.number_format = "@"
        if h in ("#", "Invoice Posted to Xero", "Closed", "Current Number") or i >= 15 and "%" not in h and h not in (
                "Video Total", "Discounts Included", "Cross Hire Expense", "Labour Expense (Internal)", "Net Total", "Margin"):
            c.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 30
    r += 3          # the row, then two empty rows
last = r - 3
# totals
tr = r
ws.cell(tr, 1, "TOTAL").font = F(bold=True, color=NAVY)
for h in ("Invoice Value", "Production Labour Hours", "Project Management", "Video Filming", "Video Editing",
          "Video Project Management", "Video Total", "Discounts Included", "Cross Hire Expense",
          "Labour Expense (Internal)", "Net Total", "Margin"):
    c = ws[f"{L[h]}{tr}"]
    c.value = f"=SUM({L[h]}{first}:{L[h]}{last})"
    c.number_format = MONEY if h not in ("Production Labour Hours", "Project Management", "Video Filming",
                                         "Video Editing", "Video Project Management") else "0"
ws[f"{L['Discounts as % of Net Total']}{tr}"] = f'=IFERROR({L["Discounts Included"]}{tr}/{L["Net Total"]}{tr},0)'
ws[f"{L['Margin %']}{tr}"] = f'=IFERROR({L["Margin"]}{tr}/{L["Invoice Value"]}{tr},0)'
for i in range(1, ncol + 1):
    c = ws.cell(tr, i)
    c.fill = PatternFill("solid", fgColor=TINT)
    c.font = F(bold=True, color=NAVY)
    c.border = Border(top=Side(style="medium", color=NAVY))
    if "%" in HEADERS[i - 1]:
        c.number_format = "0.0%"
ws.freeze_panes = ws.cell(HR + 1, 7)
ws.auto_filter.ref = f"A{HR}:{CL(ncol)}{last}"
wb.save(OUT)
print("rows", len(ROWS), "mismatches", mismatches)
