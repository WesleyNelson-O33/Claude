"""Load a synthetic GL into a copy of the pack, recalc, and tie every headline
number back to an independent Python calculation."""
import json, random, shutil, subprocess, sys, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

random.seed(7)
HERE = Path(__file__).parent
SRC = HERE / "GL Month-on-Month & Year-on-Year Analysis.xlsx"
TEST = HERE / "_verify.xlsx"
RECALC = "/root/.claude/skills/synced/e9aa5ee9-a99d-40c9-9bb1-cc5bc5010cf8_723c92f1-3883-471f-8787-0ed0739f050b/xlsx/scripts/recalc.py"
ACCOUNTS = json.loads((HERE.parent / "data/accounts.json").read_text())

CY, PER, PY = 2027, 3, 2026          # FY27, period 3 = September 2026
BS = [("Trade Debtors", "Current Assets", "Admin"),
      ("Cash at Bank", "Current Assets", "Admin"),
      ("Trade Creditors", "Current Liabilities", "Admin")]

def fy_period(d):
    fy = d.year + (1 if d.month >= 7 else 0)
    return fy, (d.month - 7) % 12 + 1

# months: all of FY26 (Jul25-Jun26) plus FY27 P1-P3 (Jul-Sep 26)
months = [dt.date(2025, m, 15) for m in range(7, 13)] + \
         [dt.date(2026, m, 15) for m in range(1, 10)]

pool = [a for a in ACCOUNTS if a["group"] in ("Income", "Cost of Sales", "Expenses", "Other Income")]
chosen = random.sample(pool, 55)
sign = {"Income": -1, "Other Income": -1, "Cost of Sales": 1, "Expenses": 1,
        "Current Assets": 1, "Current Liabilities": -1}

rows = []           # (date, account, desc, ref, contact, div, debit, credit)
for d in months:
    net = 0.0
    for k, a in enumerate(chosen):
        base = {"Income": 90000, "Other Income": 1200, "Cost of Sales": 52000, "Expenses": 9000}[a["group"]]
        amt = round(base * random.uniform(0.55, 1.45) / 6, 2)
        if a["account"] == chosen[0]["account"] and d >= dt.date(2026, 8, 1):
            amt *= 2.4                                    # a deliberate High-risk spike
        amt = round(amt, 2)
        dr, cr = (amt, 0.0) if a["group"] in ("Cost of Sales", "Expenses") else (0.0, amt)
        rows.append((d, a["account"], f"{a['group']} posting {k}", f"JNL{d:%y%m}{k:03d}",
                     "Various", a["division"], dr, cr))
        net += dr - cr
    # balancing entry keeps debits = credits
    rows.append((d, "Cash at Bank", "Net cash movement", f"JNL{d:%y%m}BAL", "Bank", "Admin",
                 round(-net, 2) if net < 0 else 0.0, round(net, 2) if net > 0 else 0.0))

shutil.copy(SRC, TEST)
wb = load_workbook(TEST)
gl, coa, st = wb["GL_Data"], wb["COA_Mapping"], wb["Setup"]
st["B7"], st["B8"], st["B12"] = CY, PER, "GL_Data"
for i, (name, grp, div) in enumerate(BS):                 # add BS accounts in blank COA rows
    r = 7 + len(ACCOUNTS) + i
    coa.cell(r, 1, name); coa.cell(r, 3, grp); coa.cell(r, 4, div)
for i, row in enumerate(rows):
    r = 8 + i
    for j, v in enumerate(row):
        gl.cell(r, 1 + j, v if not (j in (6, 7) and v == 0.0) else None)
wb.save(TEST)
print(f"wrote {len(rows)} GL lines")

res = subprocess.run([sys.executable, RECALC, str(TEST), "600"], capture_output=True, text=True)
print(res.stdout.strip()[:400])

# ---- independent expectation -------------------------------------------
grp_of = {a["account"]: a["group"] for a in ACCOUNTS}
div_of = {a["account"]: a["division"] for a in ACCOUNTS}
for n, g, d in BS:
    grp_of[n], div_of[n] = g, d

agg = defaultdict(float)          # (account, fy, per) -> reported amount
for d, acct, *_x, div, dr, cr in rows:
    fy, per = fy_period(d)
    agg[(acct, fy, per)] += (dr - cr) * sign[grp_of[acct]]

def grp_total(group, fy, periods):
    return sum(v for (a, f, p), v in agg.items()
               if grp_of[a] == group and f == fy and p in periods)

def div_total(group, division, fy, periods):
    return sum(v for (a, f, p), v in agg.items()
               if grp_of[a] == group and div_of[a] == division and f == fy and p in periods)

exp = {
    "B7":  grp_total("Income", CY, {PER}),
    "C7":  grp_total("Income", PY, {PER}),
    "F7":  grp_total("Income", CY, {PER - 1}),
    "I7":  grp_total("Income", CY, set(range(1, PER + 1))),
    "J7":  grp_total("Income", PY, set(range(1, PER + 1))),
    "M7":  grp_total("Income", PY, set(range(1, 13))),
    "B8":  grp_total("Cost of Sales", CY, {PER}),
    "I12": grp_total("Expenses", CY, set(range(1, PER + 1))),
    "B11": grp_total("Other Income", CY, {PER}),
    "I43": grp_total("Current Assets", CY, set(range(1, PER + 1))),
    "I46": grp_total("Current Liabilities", CY, set(range(1, PER + 1))),
}
exp["B9"] = exp["B7"] - exp["B8"]
exp["D7"] = exp["B7"] - exp["C7"]
exp["G7"] = exp["B7"] - exp["F7"]
exp["K7"] = exp["I7"] - exp["J7"]
for i, d in enumerate(["Onsite", "Production", "Video", "Consulting", "Integration", "Admin", "Unallocated"]):
    exp[f"B{23+i}"] = div_total("Income", d, CY, {PER})

vb = load_workbook(TEST, data_only=True)
sm, mm, yy, ct, se = vb["Summary"], vb["MoM_Analysis"], vb["YoY_Analysis"], vb["Controls"], vb["Setup"]
fails = 0
print("\n--- Summary ties ---")
for ref, want in exp.items():
    got = sm[ref].value or 0
    ok = abs(got - want) < 0.05
    fails += not ok
    print(f"{'OK ' if ok else 'BAD'} Summary!{ref:5} expected {want:>14,.2f}  got {got:>14,.2f}")

# account-level spot check on the account we deliberately spiked
spike = chosen[0]["account"]
row = next(r for r in range(6, 256) if mm.cell(r, 1).value == spike)
m_cur = agg.get((spike, CY, PER), 0.0)
m_pri = agg.get((spike, CY, PER - 1), 0.0)
y_cy = sum(agg.get((spike, CY, p), 0.0) for p in range(1, PER + 1))
y_py = sum(agg.get((spike, PY, p), 0.0) for p in range(1, PER + 1))
print(f"\n--- account detail: {spike} (row {row}) ---")
for label, got, want in [("MoM current month", mm.cell(row, 17).value, m_cur),
                         ("MoM prior month", mm.cell(row, 18).value, m_pri),
                         ("MoM movement $", mm.cell(row, 19).value, m_cur - m_pri),
                         ("YoY CY YTD", yy.cell(row, 8).value, y_cy),
                         ("YoY PY YTD", yy.cell(row, 9).value, y_py),
                         ("YoY YTD var $", yy.cell(row, 10).value, y_cy - y_py)]:
    ok = abs((got or 0) - want) < 0.05
    fails += not ok
    print(f"{'OK ' if ok else 'BAD'} {label:20} expected {want:>13,.2f}  got {(got or 0):>13,.2f}")
print("   MoM flag:", mm.cell(row, 24).value, "| YoY flag:", yy.cell(row, 16).value,
      "| movement type:", yy.cell(row, 15).value)

print("\n--- Controls ---")
for r in range(6, 25):
    ref, desc, val, stat = ct.cell(r, 1).value, ct.cell(r, 2).value, ct.cell(r, 3).value, ct.cell(r, 5).value
    if ref:
        bad = stat == "FAIL"
        fails += bad
        print(f"{'!! ' if bad else '   '}{ref:4} {stat:7} {str(val):>10}  {desc[:62]}")
print("\nGL lines loaded:", se["B21"].value, "| Dr less Cr:", se["B30"].value)
print("\nRESULT:", "ALL TIES PASS" if fails == 0 else f"{fails} MISMATCHES")
