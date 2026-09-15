"""Regression: the manual GL_Data path must still tie after the rebuild."""
import json, random, shutil, subprocess, sys, re, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
random.seed(3)
HERE = Path(__file__).parent
TEST = HERE / "_verify_gl.xlsx"
RECALC = str(next(Path("/root/.claude/skills/synced").glob("*/xlsx/scripts/recalc.py")))
ACCOUNTS = json.loads((HERE.parent / "data/accounts.json").read_text())
CY, PER, MONTH = 2027, 3, "Sep 2026"
sign = {"Income": -1, "Other Income": -1, "Cost of Sales": 1, "Expenses": 1}
net = {"Income": 1, "Other Income": 1, "Cost of Sales": -1, "Expenses": -1}
months = [dt.date(2025, m, 12) for m in range(7, 13)] + [dt.date(2026, m, 12) for m in range(1, 10)]
chosen = random.sample([a for a in ACCOUNTS if a["group"] in sign], 50)
rows = []
for a in chosen:
    for d in months:
        amt = round(random.uniform(2000, 40000), 2)
        dr, cr = (amt, 0.0) if a["group"] in ("Cost of Sales", "Expenses") else (0.0, amt)
        rows.append((d, a["account"], "line", "J1", "x", a["division"], dr, cr))
shutil.copy(HERE / "GL Month-on-Month & Year-on-Year Analysis.xlsx", TEST)
wb = load_workbook(TEST); gl, st = wb["GL_Data"], wb["Setup"]
st["B5"], st["B12"] = MONTH, "GL_Data"
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        if not (j in (6, 7) and v == 0.0):
            gl.cell(8 + i, 1 + j, v)
wb.save(TEST)
subprocess.run([sys.executable, RECALC, str(TEST), "900"], capture_output=True, text=True)

grp_of = {a["account"]: a["group"] for a in ACCOUNTS}
div_of = {a["account"]: a["division"] for a in ACCOUNTS}
def rdept(a): return "Overheads" if (grp_of[a] == "Expenses" and div_of[a] == "Unallocated") else div_of[a]
agg = defaultdict(float)
for d, acct, *_r, dr, cr in rows:
    fy, per = d.year + (1 if d.month >= 7 else 0), (d.month - 7) % 12 + 1
    agg[(acct, fy, per)] += (dr - cr) * sign[grp_of[acct]]
YTD = set(range(1, PER + 1))
def tot(dept=None, group=None):
    return sum(v for (a, f, p), v in agg.items() if f == CY and p in YTD
               and (dept is None or rdept(a) == dept) and (group is None or grp_of[a] == group))
def netc(dept):
    return sum(v * net[grp_of[a]] for (a, f, p), v in agg.items() if f == CY and p in YTD and rdept(a) == dept)

vb = load_workbook(TEST, data_only=True); sm, se = vb["Summary"], vb["Setup"]
dept_rows, comp, cur = {}, {}, None
for r in range(6, sm.max_row + 1):
    v = sm.cell(r, 1).value
    if not isinstance(v, str): continue
    m = re.match(r"^([A-Z][A-Za-z ]+?)\s+\(net contribution\)$", v.strip())
    if m:
        cur = m.group(1).strip().title(); dept_rows[cur] = {"__net__": r}; continue
    lab = v.strip()
    if cur and lab in ("Income", "Cost of Sales", "Direct Expenses"): dept_rows[cur][lab] = r
    elif cur is None and lab in ("Income", "Cost of Sales", "Operating Expenses"): comp[lab] = r
fails = 0
def chk(l, got, want):
    global fails
    ok = abs((got or 0) - want) < 0.05; fails += not ok
    print(f"{'OK ' if ok else 'BAD'} {l:44} exp {want:>13,.2f}  got {(got or 0):>13,.2f}")
print("GL lines loaded:", se["B21"].value, "| FY", se["B7"].value, "period", se["B8"].value)
chk("company Income YTD", sm.cell(comp["Income"], 9).value, tot(group="Income"))
chk("company Cost of Sales YTD", sm.cell(comp["Cost of Sales"], 9).value, tot(group="Cost of Sales"))
chk("company Operating Expenses YTD", sm.cell(comp["Operating Expenses"], 9).value, tot(group="Expenses"))
for d in ["Onsite", "Consulting", "Admin", "Overheads", "Unallocated"]:
    chk(f"{d}: Income YTD", sm.cell(dept_rows[d]["Income"], 9).value, tot(d, "Income"))
    chk(f"{d}: net contribution YTD", sm.cell(dept_rows[d]["__net__"], 9).value, netc(d))
print("\nRESULT:", "GL_DATA PATH OK" if fails == 0 else f"{fails} PROBLEMS")
