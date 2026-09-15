"""Paste a deliberately filthy Xero-style export into Raw_Paste and prove the
cleaning keeps exactly the real transactions and nothing else."""
import json, random, shutil, subprocess, sys, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

random.seed(11)
HERE = Path(__file__).parent
SRC = HERE / "GL Month-on-Month & Year-on-Year Analysis.xlsx"
TEST = HERE / "_verify_dirty.xlsx"
RECALC = "/root/.claude/skills/synced/e9aa5ee9-a99d-40c9-9bb1-cc5bc5010cf8_723c92f1-3883-471f-8787-0ed0730f050b/xlsx/scripts/recalc.py"
RECALC = str(next(Path("/root/.claude/skills/synced").glob("*/xlsx/scripts/recalc.py")))
ACCOUNTS = json.loads((HERE.parent / "data/accounts.json").read_text())
CY, PER, PY = 2027, 3, 2026
RAW_R0 = 1

months = [dt.date(2025, m, 12) for m in range(7, 13)] + [dt.date(2026, m, 12) for m in range(1, 10)]
pool = [a for a in ACCOUNTS if a["group"] in ("Income", "Cost of Sales", "Expenses", "Other Income")]
chosen = random.sample(pool, 40)
sign = {"Income": -1, "Other Income": -1, "Cost of Sales": 1, "Expenses": 1}

# real transactions, tracked independently
txns = []                                   # (date, account, debit, credit)
for a in chosen:
    for d in months:
        for _ in range(random.randint(1, 3)):
            base = {"Income": 30000, "Other Income": 600, "Cost of Sales": 18000, "Expenses": 3200}[a["group"]]
            amt = round(base * random.uniform(0.4, 1.6), 2)
            dr, cr = (amt, 0.0) if a["group"] in ("Cost of Sales", "Expenses") else (0.0, amt)
            txns.append((d, a["account"], dr, cr))

# ---- now write it out the way Xero actually does, junk and all -----------
dump = []                                   # list of 7-column rows
JUNK = {"title": 0, "blank": 0, "heading": 0, "subtotal": 0, "nil": 0}
def put(*cells):
    dump.append(list(cells) + [None] * (7 - len(cells)))

put("Corporate Technology Services Pty Ltd"); JUNK["title"] += 1
put("Account Transactions"); JUNK["title"] += 1
put("For the period 1 July 2025 to 30 September 2026"); JUNK["title"] += 1
put(); JUNK["blank"] += 1
put("Date", "Source", "Description", "Reference", "Debit", "Credit", "Running Balance"); JUNK["heading"] += 1
put(); JUNK["blank"] += 1

by_acct = defaultdict(list)
for t in txns:
    by_acct[t[1]].append(t)
for acct in sorted(by_acct):
    put(acct); JUNK["heading"] += 1
    put("Opening Balance", None, None, None, None, None, round(random.uniform(0, 9000), 2)); JUNK["subtotal"] += 1
    run = 0.0
    tdr = tcr = 0.0
    for d, _a, dr, cr in sorted(by_acct[acct]):
        run += dr - cr
        tdr += dr; tcr += cr
        put(d, random.choice(["Invoice", "Bill", "Manual Journal", "Payment"]),
            f"{acct} activity", f"INV-{random.randint(10000,99999)}",
            dr or None, cr or None, round(run, 2))
    put(sorted(by_acct[acct])[0][0], "Bill", "Cancelled - nil value", "VOID-001", None, None, round(run, 2))
    JUNK["nil"] += 1
    put(f"Total {acct}", None, None, None, round(tdr, 2), round(tcr, 2)); JUNK["subtotal"] += 1
    put("Closing Balance", None, None, None, None, None, round(run, 2)); JUNK["subtotal"] += 1
    put(); JUNK["blank"] += 1
net_by_month = defaultdict(float)
for d, _a, dr, cr in txns:
    net_by_month[d] += dr - cr
put("Cash at Bank"); JUNK["heading"] += 1
put("Opening Balance", None, None, None, None, None, 0.0); JUNK["subtotal"] += 1
run = 0.0
for d in sorted(net_by_month):
    v = round(-net_by_month[d], 2)
    run += v
    put(d, "Payment", "Net cash movement", "BANK-001",
        v if v > 0 else None, -v if v < 0 else None, round(run, 2))
    txns.append((d, "Cash at Bank", v if v > 0 else 0.0, -v if v < 0 else 0.0))
put("Total Cash at Bank", None, None, None, 0.0, 0.0); JUNK["subtotal"] += 1
put(); JUNK["blank"] += 1
for _ in range(25):
    put(); JUNK["blank"] += 1

print(f"dump rows: {len(dump)}   real transactions: {len(txns)}   junk: {JUNK}")

shutil.copy(SRC, TEST)
wb = load_workbook(TEST)
rp, st = wb["Raw_Paste"], wb["Setup"]
st["B7"], st["B8"], st["B12"] = CY, PER, "Raw_Paste"
coa = wb["COA_Mapping"]
for i, (name, grp, div) in enumerate([("Cash at Bank", "Current Assets", "Admin")]):
    r = 7 + len(ACCOUNTS) + i
    coa.cell(r, 1, name); coa.cell(r, 3, grp); coa.cell(r, 4, div)
st["B53"], st["B54"], st["B55"] = "Section headings", 1, 1    # account heading and date both in column A
st["B56"], st["B57"], st["B58"] = 3, 4, 2                      # description, reference, contact
st["B59"], st["B60"], st["B61"] = "Debit and Credit", 5, 6
st["B62"], st["B63"] = "Yes", "Yes"
for i, row in enumerate(dump):
    for j, v in enumerate(row):
        if v is not None:
            rp.cell(RAW_R0 + i, 1 + j, v)
wb.save(TEST)

res = subprocess.run([sys.executable, RECALC, str(TEST), "900"], capture_output=True, text=True)
print(res.stdout.strip()[:300])

# ---- independent expectation from the transactions only ------------------
grp_of = {a["account"]: a["group"] for a in ACCOUNTS}
grp_of["Cash at Bank"] = "Current Assets"
sign["Current Assets"] = 1
div_of = {a["account"]: a["division"] for a in ACCOUNTS}
def fy_period(d):
    return d.year + (1 if d.month >= 7 else 0), (d.month - 7) % 12 + 1
agg = defaultdict(float)
for d, acct, dr, cr in txns:
    fy, per = fy_period(d)
    agg[(acct, fy, per)] += (dr - cr) * sign[grp_of[acct]]
def grp_total(group, fy, periods):
    return sum(v for (a, f, p), v in agg.items() if grp_of[a] == group and f == fy and p in periods)

exp = {
    "B7": grp_total("Income", CY, {PER}), "C7": grp_total("Income", PY, {PER}),
    "F7": grp_total("Income", CY, {PER - 1}), "I7": grp_total("Income", CY, set(range(1, PER + 1))),
    "J7": grp_total("Income", PY, set(range(1, PER + 1))), "M7": grp_total("Income", PY, set(range(1, 13))),
    "B8": grp_total("Cost of Sales", CY, {PER}), "I8": grp_total("Cost of Sales", CY, set(range(1, PER + 1))),
    "I11": grp_total("Other Income", CY, set(range(1, PER + 1))),
    "I12": grp_total("Expenses", CY, set(range(1, PER + 1))),
}
exp["B9"] = exp["B7"] - exp["B8"]

vb = load_workbook(TEST, data_only=True)
sm, se, ct, cu = vb["Summary"], vb["Setup"], vb["Controls"], vb["Cleanup"]
fails = 0
print("\n--- cleaning results ---")
labels = ["rows found", "kept", "dropped headings", "dropped subtotals", "dropped nil value", "dropped text only",
          "kept debits", "kept credits", "kept Dr less Cr", "kept unmapped"]
got_kept = se["B34"].value
for i, lab in enumerate(labels):
    print(f"   {lab:22} {se.cell(33 + i, 2).value}")
ok = got_kept == len(txns)
fails += not ok
print(f"{'OK ' if ok else 'BAD'} kept lines {got_kept} vs real transactions {len(txns)}")
ok = abs((se['B41'].value or 0)) < 0.01
fails += not ok
print(f"{'OK ' if ok else 'BAD'} kept debits = kept credits (diff {se['B41'].value})")
ok = (se['B42'].value or 0) == 0
fails += not ok
print(f"{'OK ' if ok else 'BAD'} unmapped kept lines = {se['B42'].value}")

print("\n--- Summary ties (subtotal rows would roughly double these) ---")
for ref, want in exp.items():
    got = sm[ref].value or 0
    ok = abs(got - want) < 0.05
    fails += not ok
    print(f"{'OK ' if ok else 'BAD'} Summary!{ref:4} expected {want:>14,.2f}  got {got:>14,.2f}")

print("\n--- Cleanup preview (first 5 kept lines) ---")
for r in range(6, 11):
    print("   ", [cu.cell(r, c).value for c in range(20, 27)])

print("\n--- Controls ---")
for r in range(6, 30):
    if ct.cell(r, 1).value:
        stat = ct.cell(r, 5).value
        fails += stat == "FAIL"
        print(f"{'!! ' if stat == 'FAIL' else '   '}{ct.cell(r,1).value:5}{str(stat):8}{str(ct.cell(r,3).value):>10}  {ct.cell(r,2).value[:58]}")
print("\nRESULT:", "ALL CHECKS PASS" if fails == 0 else f"{fails} PROBLEMS")
