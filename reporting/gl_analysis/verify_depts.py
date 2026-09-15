"""Dirty Xero dump -> department-grouped reports. Ties every department block
back to an independent calculation, and checks the collapse structure survived."""
import json, random, shutil, subprocess, sys, re, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

random.seed(11)
HERE = Path(__file__).parent
SRC = HERE / "GL Month-on-Month & Year-on-Year Analysis.xlsx"
TEST = HERE / "_verify_depts.xlsx"
RECALC = str(next(Path("/root/.claude/skills/synced").glob("*/xlsx/scripts/recalc.py")))
ACCOUNTS = json.loads((HERE.parent / "data/accounts.json").read_text())
_OV = json.loads((HERE.parent / "data/account_overrides.json").read_text())
for _a in ACCOUNTS:
    if _a["account"] in _OV["division"]:
        _a["division"] = _OV["division"][_a["account"]]["to"]
    if _a["account"] in _OV["group"]:
        _a["group"] = _OV["group"][_a["account"]]["to"]
CY, PER, PY, MONTH_LABEL = 2027, 3, 2026, "Sep 2026"

months = [dt.date(2025, m, 12) for m in range(7, 13)] + [dt.date(2026, m, 12) for m in range(1, 10)]
pool = [a for a in ACCOUNTS if a["group"] in ("Income", "Cost of Sales", "Expenses", "Other Income")]
chosen = random.sample(pool, 60)
sign = {"Income": -1, "Other Income": -1, "Cost of Sales": 1, "Expenses": 1, "Current Assets": 1,
        "Depreciation & Amortisation": 1, "Finance Costs": 1, "Income Tax Expense": 1, "Equity": -1}
net = {"Income": 1, "Other Income": 1, "Cost of Sales": -1, "Expenses": -1, "Current Assets": 0,
       "Depreciation & Amortisation": -1, "Finance Costs": -1, "Income Tax Expense": -1, "Equity": 0}

txns = []
for a in chosen:
    for d in months:
        for _ in range(random.randint(1, 2)):
            base = {"Income": 30000, "Other Income": 600, "Cost of Sales": 18000, "Expenses": 3200}[a["group"]]
            amt = round(base * random.uniform(0.4, 1.6), 2)
            dr, cr = (amt, 0.0) if a["group"] in ("Cost of Sales", "Expenses") else (0.0, amt)
            txns.append((d, a["account"], dr, cr))

dump = []
def put(*cells): dump.append(list(cells) + [None] * (7 - len(cells)))
put("General Ledger Detail"); put("Corporate Technology Services Pty Ltd")
put("For the period 1 July 2025 to 30 September 2026"); put()
put("Date", "Source", "Description", "Reference", "Debit", "Credit", "Running Balance"); put()
by = defaultdict(list)
for t in txns: by[t[1]].append(t)
for acct in sorted(by):
    put(acct); put("Opening Balance", None, None, None, None, None, 1234.0)
    run = tdr = tcr = 0.0
    for d, _a, dr, cr in sorted(by[acct]):
        run += dr - cr; tdr += dr; tcr += cr
        put(d, "Invoice", f"{acct} activity", f"INV-{random.randint(10000,99999)}", dr or None, cr or None, round(run, 2))
    put(sorted(by[acct])[0][0], "Bill", "Cancelled - nil", "VOID", None, None, round(run, 2))
    put(f"Total {acct}", None, None, None, round(tdr, 2), round(tcr, 2))
    put("Closing Balance", None, None, None, None, None, round(run, 2)); put()
nbm = defaultdict(float)
for d, _a, dr, cr in txns: nbm[d] += dr - cr
put("Cash at Bank"); run = 0.0
for d in sorted(nbm):
    v = round(-nbm[d], 2); run += v
    put(d, "Payment", "Net cash movement", "BANK", v if v > 0 else None, -v if v < 0 else None, round(run, 2))
    txns.append((d, "Cash at Bank", v if v > 0 else 0.0, -v if v < 0 else 0.0))
put("Total Cash at Bank", None, None, None, 0.0, 0.0)
for _ in range(20): put()

shutil.copy(SRC, TEST)
wb = load_workbook(TEST)
rp, st, coa = wb["Raw_Paste"], wb["Setup"], wb["COA_Mapping"]
st["B5"], st["B12"] = MONTH_LABEL, "Raw_Paste"
st["B53"], st["B54"], st["B55"] = "Section headings", 1, 1
st["B56"], st["B57"], st["B58"] = 3, 4, 2
st["B59"], st["B60"], st["B61"] = "Debit and Credit", 5, 6
st["B62"], st["B63"] = "Yes", "Yes"
r = 7 + len(ACCOUNTS)
coa.cell(r, 1, "Cash at Bank"); coa.cell(r, 3, "Current Assets"); coa.cell(r, 4, "Admin"); coa.cell(r, 10, "No")
for i, row in enumerate(dump):
    for j, v in enumerate(row):
        if v is not None: rp.cell(1 + i, 1 + j, v)
wb.save(TEST)
print(f"dump rows {len(dump)}  real transactions {len(txns)}")
res = subprocess.run([sys.executable, RECALC, str(TEST), "900"], capture_output=True, text=True)
print(res.stdout.strip()[:220])

# ---- independent expectation --------------------------------------------
grp_of = {a["account"]: a["group"] for a in ACCOUNTS}; grp_of["Cash at Bank"] = "Current Assets"
div_of = {a["account"]: a["division"] for a in ACCOUNTS}; div_of["Cash at Bank"] = "Admin"
def rdept(acct):
    return "Admin" if (grp_of[acct] == "Expenses" and div_of[acct] == "Unallocated") else div_of[acct]
def fyp(d): return d.year + (1 if d.month >= 7 else 0), (d.month - 7) % 12 + 1
agg = defaultdict(float)
for d, acct, dr, cr in txns:
    fy, per = fyp(d)
    agg[(acct, fy, per)] += (dr - cr) * sign[grp_of[acct]]
def tot(fy, periods, dept=None, group=None):
    return sum(v for (a, f, p), v in agg.items() if f == fy and p in periods
               and (dept is None or rdept(a) == dept) and (group is None or grp_of[a] == group))
def netc(fy, periods, dept):
    return sum(v * net[grp_of[a]] for (a, f, p), v in agg.items()
               if f == fy and p in periods and rdept(a) == dept)
CURP, YTD = {PER}, set(range(1, PER + 1))

vb = load_workbook(TEST, data_only=True)
sm, mm, yy, ct, se, en = (vb["Summary"], vb["MoM_Analysis"], vb["YoY_Analysis"],
                          vb["Controls"], vb["Setup"], vb["Data_Engine"])
fails = 0
def chk(label, got, want, tol=0.05):
    global fails
    ok = abs((got or 0) - want) < tol
    fails += not ok
    print(f"{'OK ' if ok else 'BAD'} {label:52} exp {want:>14,.2f}  got {(got or 0):>14,.2f}")

print(f"\n--- Setup derived from the month dropdown '{MONTH_LABEL}' ---")
for lab, got, want in (("FY", se['B7'].value, CY), ("period", se['B8'].value, PER), ("prior FY", se['B9'].value, PY)):
    ok = got == want; fails += not ok
    print(f"{'OK ' if ok else 'BAD'} {lab:12} expected {want}  got {got}")
print(f"    comparative month label: {se['B11'].value}")

# locate Summary rows by section
dept_rows, comp = {}, {}
cur = None
for r in range(6, sm.max_row + 1):
    v = sm.cell(r, 1).value
    if not isinstance(v, str): continue
    m = re.match(r"^([A-Z][A-Za-z ]+?)\s+\(net contribution\)$", v.strip())
    if m:
        cur = m.group(1).strip().title(); dept_rows[cur] = {"__net__": r}; continue
    lab = v.strip()
    if cur and lab in ("Income", "Cost of Sales", "Other Income", "Direct Expenses", "Gross Profit"):
        dept_rows[cur][lab] = r
    elif cur is None and lab in ("Income", "Cost of Sales", "Operating Expenses", "Other Income"):
        comp[lab] = r

print("\n--- Summary: company ---")
chk("Income CY month", sm.cell(comp["Income"], 2).value, tot(CY, CURP, group="Income"))
chk("Income CY YTD", sm.cell(comp["Income"], 9).value, tot(CY, YTD, group="Income"))
chk("Cost of Sales CY YTD", sm.cell(comp["Cost of Sales"], 9).value, tot(CY, YTD, group="Cost of Sales"))
chk("Operating Expenses CY YTD", sm.cell(comp["Operating Expenses"], 9).value, tot(CY, YTD, group="Expenses"))

print("\n--- Summary: each department (CY YTD) ---")
for dept in ["Onsite", "Production", "Video", "Consulting", "Integration", "Admin", "Unallocated"]:
    d = dept_rows.get(dept)
    if not d:
        print(f"BAD department block missing: {dept}"); fails += 1; continue
    chk(f"{dept}: Income", sm.cell(d["Income"], 9).value, tot(CY, YTD, dept, "Income"))
    chk(f"{dept}: Cost of Sales", sm.cell(d["Cost of Sales"], 9).value, tot(CY, YTD, dept, "Cost of Sales"))
    chk(f"{dept}: Direct Expenses", sm.cell(d["Direct Expenses"], 9).value, tot(CY, YTD, dept, "Expenses"))
    chk(f"{dept}: NET CONTRIBUTION", sm.cell(d["__net__"], 9).value, netc(CY, YTD, dept))

print("\n--- overheads landed in Admin ---")
oh = [a for a in ACCOUNTS if a["group"] == "Expenses" and a["division"] == "Unallocated"]
print(f"    {len(oh)} accounts flagged Overhead, all reporting under Admin")
chk("Unallocated has no expenses left", sm.cell(dept_rows["Unallocated"]["Direct Expenses"], 9).value, 0.0)
chk("Admin carries the overheads", sm.cell(dept_rows["Admin"]["Direct Expenses"], 9).value,
    tot(CY, YTD, "Admin", "Expenses"))

print("\n--- MoM sheet: department and group rows ---")
cur = None
checked = 0
for r in range(6, mm.max_row + 1):
    v = mm.cell(r, 1).value
    if not isinstance(v, str): continue
    m = re.match(r"^([A-Z][A-Za-z ]+?)\s+\(net contribution\)$", v.strip())
    if m:
        cur = m.group(1).strip().title()
        chk(f"MoM {cur}: current month net", mm.cell(r, 17).value, netc(CY, CURP, cur)); checked += 1
    elif cur and v.strip() == "Income":
        chk(f"MoM {cur}: Income current month", mm.cell(r, 17).value, tot(CY, CURP, cur, "Income")); checked += 1
print(f"    ({checked} department/group rows checked)")

print("\n--- account rows resolved under the right department ---")
found = 0
for r in range(6, mm.max_row + 1):
    if mm.row_dimensions[r].outlineLevel == 2 and mm.cell(r, 1).value:
        nm = str(mm.cell(r, 1).value).strip()
        if nm in grp_of:
            want_dept = rdept(nm)
            got_dept = mm.cell(r, 3).value
            ok = got_dept == want_dept
            fails += not ok
            found += 1
            if found <= 5 or not ok:
                print(f"{'OK ' if ok else 'BAD'} {nm[:38]:40} -> {got_dept} (expected {want_dept})")
print(f"    {found} account rows resolved")

print("\n--- Controls ---")
for r in range(6, 34):
    if ct.cell(r, 1).value:
        stt = ct.cell(r, 5).value
        fails += stt == "FAIL"
        print(f"{'!! ' if stt == 'FAIL' else '   '}{ct.cell(r,1).value:5}{str(stt):8}{str(ct.cell(r,3).value):>10}  {str(ct.cell(r,2).value)[:56]}")
print("\nRESULT:", "ALL CHECKS PASS" if fails == 0 else f"{fails} PROBLEMS")
