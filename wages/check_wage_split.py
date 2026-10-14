"""Check the wage and salary split of each pay run reclass journal against Employment Hero.

Rule (from Danica): client job = Direct, internal code = Indirect; casual = Wages,
full or part time = Salaries; department from the job's tag.
Inputs are the EH Payroll Split File (EH data (2), Employee Details Report) and the
Xero General Ledger Detail export. Output is a workbook in wages/out (employee data,
not committed).
"""
import re, sys, datetime
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font

SPLIT, GLX, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
REFS = {1: ['#69527'], 2: ['#68743'], 3: ['#69396'], 4: ['#69669', '#69801', '#69802'],
        5: ['#70260'], 6: ['#70657']}
INTERNAL = ['non-chargeable', 'quotes', 'sales', 'marketing', 'training', 'mgmt', 'human resources',
            'office / admin', 'cts day', '(including prod backup)', 'video dept', 'internal']

def dept(loc):
    m = re.search(r'\[(ONS|PRD|VID|INT|CONS|CTS)\]', loc) or re.search(r'(ONS|PRD|VID|INT|CONS|CTS)\[', loc)
    if m:
        return {'ONS': 'ONS', 'PRD': 'PRD', 'VID': 'VID', 'INT': 'INTEGRATION', 'CONS': 'CONS', 'CTS': 'ADMIN'}[m.group(1)]
    u = loc.upper()
    for k, v in [('ONSITE', 'ONS'), ('ONS ', 'ONS'), ('PRODUCTION', 'PRD'), ('VIDEO', 'VID'),
                 ('INTEGR', 'INTEGRATION'), ('CONSULT', 'CONS'), ('ADMIN', 'ADMIN')]:
        if k in u:
            return v
    return '?'

def side(loc, lid):
    l = loc.lower()
    if str(lid).isdigit() and len(str(lid)) == 4 and str(lid).startswith('9'):
        return 'Indirect'
    if any(w in l for w in INTERNAL):
        return 'Indirect'
    if re.fullmatch(r'(PRODUCTION|INTEGRATION|CONSULTING|ONSITE|VIDEO)( \[[A-Z]+\])?', loc.strip(), re.I):
        return 'Indirect'
    return 'Direct'

wb = openpyxl.load_workbook(SPLIT, read_only=True, data_only=True)
et = {}
for r in list(wb['Employee Details Report'].iter_rows(values_only=True))[7:]:
    if r[1]:
        et[r[1]] = (r[7], f'{r[2]} {r[3]}')
rows = [r for r in list(wb['EH data (2)'].iter_rows(values_only=True))[2:] if r[1] in REFS]

gl = defaultdict(float)
acct = None
for r in openpyxl.load_workbook(GLX, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True):
    if r[0] and not isinstance(r[0], datetime.datetime) and r[1] is None and r[4] is None:
        acct = r[0]
        continue
    if isinstance(r[0], datetime.datetime) and acct and ('Wages' in acct or 'Salaries' in acct):
        for p, rs in REFS.items():
            if str(r[3]).strip() in rs:
                gl[(p, acct)] += (r[4] or 0) - (r[5] or 0)

detail, should = [], defaultdict(float)
for r in rows:
    loc = str(r[14]).strip()
    kind = 'Wages' if et[r[4]][0] == 'Casual' else 'Salaries'
    a = f'{side(loc, r[13])} {kind} - {dept(loc)}'
    should[(r[1], a)] += r[18] or 0
    detail.append((f'3700{r[1]}', et[r[4]][1], et[r[4]][0], loc, r[9], r[18] or 0, a))

out = openpyxl.Workbook()
ws = out.active
ws.title = 'Summary'
ws.append(['Account', 'Posted Jul-Sep', 'Should be', 'Change needed'])
accts = sorted({a for _, a in gl} | {a for _, a in should})
for a in accts:
    po = sum(gl[(p, a)] for p in REFS)
    sh = sum(should[(p, a)] for p in REFS)
    ws.append([a, round(po, 2), round(sh, 2), round(sh - po, 2)])
n = ws.max_row
ws.append(['Total', f'=SUM(B2:B{n})', f'=SUM(C2:C{n})', f'=SUM(D2:D{n})'])
for p in REFS:
    s = out.create_sheet(f'3700{p}')
    s.append(['Account', 'Posted', 'Should be', 'Change needed'])
    for a in accts:
        po, sh = gl[(p, a)], should[(p, a)]
        if abs(po) > 0.004 or abs(sh) > 0.004:
            s.append([a, round(po, 2), round(sh, 2), round(sh - po, 2)])
    m = s.max_row
    s.append(['Total', f'=SUM(B2:B{m})', f'=SUM(C2:C{m})', f'=SUM(D2:D{m})'])
d = out.create_sheet('EH detail')
d.append(['Pay run', 'Employee', 'Employment type', 'Job worked', 'Pay category', 'Gross', 'Correct account'])
for x in detail:
    d.append(list(x))
for s in out.worksheets:
    for c in s[1]:
        c.font = Font(bold=True)
    s.column_dimensions['A'].width = 36
    for col in 'BCDEFG':
        s.column_dimensions[col].width = 18
d.column_dimensions['B'].width = 24
d.column_dimensions['D'].width = 55
d.column_dimensions['G'].width = 32
out.save(OUT)
