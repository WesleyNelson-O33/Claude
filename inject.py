"""Rebuild the HR copy: clean master + HR data + README + hidden columns + HR formatting + tally preload."""
import sys, re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L
from copy import copy
src_path, clean_path, out_path = sys.argv[1:4]
tally_path = sys.argv[4] if len(sys.argv) > 4 else None
src=load_workbook(src_path); new=load_workbook(clean_path)
def copy_style(a,b):
    b.font=copy(a.font); b.fill=copy(a.fill); b.border=copy(a.border); b.alignment=copy(a.alignment); b.number_format=a.number_format
ss=src['Staff']; ns=new['Staff']
hs={ss.cell(row=4,column=c).value:c for c in range(1,ss.max_column+1) if ss.cell(row=4,column=c).value}
hn={ns.cell(row=4,column=c).value:c for c in range(1,ns.max_column+1) if ns.cell(row=4,column=c).value}
inputs=["Name","Surname","EmploymentType","DateOfBirth","StartDate","Casual Start","PTE Start","FTE Start","AnniversaryDate","PT Hrs/Week"]
for r in range(5,155):
    if ss.cell(row=r,column=hs["Name"]).value is None: continue
    for h in inputs:
        v=ss.cell(row=r,column=hs[h]).value
        if v is not None:
            c=ns.cell(row=r,column=hn[h]); c.value=v
            if hasattr(v,'year'): c.number_format='dd/mm/yyyy'
# formatting by header
for h,cs in hs.items():
    if h in hn and not h.startswith('h '):
        cn=hn[h]
        for r in range(3,155): copy_style(ss.cell(row=r,column=cs), ns.cell(row=r,column=cn))
        ns.column_dimensions[L(cn)].width=ss.column_dimensions[L(cs)].width
        ns.column_dimensions[L(cn)].hidden=ss.column_dimensions[L(cs)].hidden
for h in ("Earned this month (days)","Milestone this month","Days excluded (before scheme start)"):
    if h in hn and h not in hs:
        for r in range(3,155): copy_style(ss.cell(row=r,column=hs["Hours to Accrue"]), ns.cell(row=r,column=hn[h]))
        ns.cell(row=4,column=hn[h]).value=h
if "Milestone this month" in hn:
    for r in range(5,155): ns.cell(row=r,column=hn["Milestone this month"]).number_format='dd/mm/yyyy'
for r in range(1,5):
    ns.row_dimensions[r].height=ss.row_dimensions[r].height
    for c in range(1,3): copy_style(ss.cell(row=r,column=c), ns.cell(row=r,column=c))
ns.freeze_panes=ss.freeze_panes
ns['D1'].number_format='dd/mm/yyyy'
# Leave History data A:L
sl=src['Leave History']; nl=new['Leave History']
for r in range(1,sl.max_row+1):
    for c in range(1,13):
        v=sl.cell(row=r,column=c).value
        if v is not None: nl.cell(row=r,column=c).value=v
for name in ('Settings','Leave History'):
    a=src[name]; b=new[name]
    for row in a.iter_rows(min_row=1,max_row=(60 if name=='Settings' else 1)):
        for c in row:
            if c.has_style: copy_style(c, b.cell(row=c.row,column=c.column))
    for k,dim in a.column_dimensions.items(): b.column_dimensions[k].width=dim.width
for cell in ('C32','C33'): new['Settings'][cell].number_format='dd/mm/yyyy'
# Settings values the HR team set
for cell in ('C6','C7','C8','C9','C19','C31'):
    v=src['Settings'][cell].value
    if v is not None: new['Settings'][cell].value=v
# READ ME
if 'READ ME' in src.sheetnames:
    sr=src['READ ME']; nr=new.create_sheet('READ ME',0)
    for row in sr.iter_rows():
        for c in row:
            if c.value is not None:
                d=nr.cell(row=c.row,column=c.column,value=c.value); d.font=copy(c.font); d.alignment=copy(c.alignment)
    for k,dim in sr.column_dimensions.items(): nr.column_dimensions[k].width=dim.width
    nr.sheet_properties.tabColor="1F4E78"
for name in src.sheetnames:
    if name in new.sheetnames: new[name].sheet_properties.tabColor=src[name].sheet_properties.tabColor
# tally preload
if tally_path:
    names=["Christopher Lehmann","Daniel Sobkowski","Danielle Hurley","Duncan Lugstein","Jordan Sexty","Kyle Krishnappa","Milo Rankin","Patrick Lee","Simon Gruenefeld","Tyler Wood","Wade Tonna","Isaac Butterworth","Blake Crisford","Carlo Daru"]
    t=new['HR Tally Check']
    for p,nm in enumerate(names): t.cell(row=7,column=2+p).value=nm
    i=0
    for line in open(tally_path):
        parts=line.rstrip('\n').split('\t')
        if not re.match(r'[A-Z][a-z]{2}-\d{2}',parts[0]): continue
        t.cell(row=8+i,column=1).value=parts[0]
        for p in range(len(names)):
            v=parts[p+1].strip() if p+1<len(parts) else ''
            if v: t.cell(row=8+i,column=2+p).value=float(v)
        i+=1
# notes next to names
NOTES={
 ("Daniel","SOBKOWSKI"):"HR tally 44 vs tracker 43: HR back-dated 1 day to his Jul-17 anniversary, before the scheme started 27/11/17. 5-yr anniversary Jul-18 = 1 day under the 2017 scheme (tracker applies this).",
 ("Duncan","LUGSTEIN"):"HR tally runs a Feb/May/Aug/Nov cycle from 2017 (eligibility letter); tracker uses his 16/10 anniversary. 1.5 day variance = Aug-17 pre-scheme day + 0.5 in Nov-17 + cycle timing. Confirm cycle date with HR.",
 ("Jordan","SEXTY"):"HR moved him to the 5-day rate in Jan-23; his 5-yr anniversary was Apr-21. Tracker is 2 days above the HR tally - owed.",
 ("Tyler","WOOD"):"HR moved him to the 5-day rate in Jul-24; his 5-yr anniversary was Jul-22. Tracker is 2 days above the HR tally - owed.",
 ("Wade","TONNA"):"HR adjusted him to the 5-day rate in Jan-26; his 5-yr anniversary was Jun-23. Tracker is 1 day above the HR tally - owed.",
 ("Milo","RANKIN"):"HR tally cycle Feb/Jun/Oct implies a start around Feb-20; tracker start is 05/08/2019. Confirm start date with HR. Tracker is 3 days above the HR tally.",
 ("Kyle","Krishnappa"):"Part-time 19/04/22 to 05/09/22 counted as service. If that part-time contract was under 24 hrs, treat as casual (move the date to Casual Start). Tracker is 1 day above the HR tally.",
 ("Carlo","Daru"):"HR tally has nothing accrued; tracker 4 days from 19/09/25 (3 yrs from full-time start). HR eligibility date agrees (19/09/2025).",
 ("Danielle","HURLEY"):"Matches HR tally (8). 5-yr anniversary Apr-26 = 1 day under the 2017 scheme.",
 ("Isaac","BUTTERWORTH"):"Matches HR tally (9). Part-time days credited from Sep-24 per HR practice; part-time eligibility is formally from the 2026 policy.",
 ("Blake","Crisford"):"Matches HR tally (2).",
}
if "Notes" in hn:
    nc=hn["Notes"]
    for r in range(5,155):
        key=(ns[f'A{r}'].value, ns[f'B{r}'].value)
        for (fn,sn),txt in NOTES.items():
            if key[0]==fn and (key[1] or '').lower()==sn.lower(): ns.cell(row=r,column=nc).value=txt
    for r in range(3,155): copy_style(ss.cell(row=r,column=hs["Casual Start"]), ns.cell(row=r,column=nc))
    ns.cell(row=4,column=nc).value="Notes"; ns.cell(row=3,column=nc).value="free text - reconciliation notes, HR decisions"
    ns.column_dimensions[L(nc)].width=60
# Milestones default = Isaac
ms=new['Milestones']; key=ms['C5'].value
for r in range(5,155):
    if ns[f'A{r}'].value=='Isaac': ms['C5']=key[:-1]+str(r); break
new.calculation.fullCalcOnLoad=True; new.save(out_path); print('saved', out_path, new.sheetnames)
